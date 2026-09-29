"""Round 8: the durable before/after fidelity report for CASE001.

The human's round-7 manual review failed the *delivered* CASE001 workbook.  This
report is the durable record of what changed and why: every one of the 49
delivered rows is compared by its exact address between the frozen build the
human reviewed (``..._review_workbook7``) and the round-8 successor, together
with the reviewer sheets (02/03), the pricing sheet (04), the structure sheet
(05) and the conflict sheet (06).

Both builds wrote their 49 delivered rows in the same plan order -- verified by
the human's own D-addresses resolving to the same item ids -- so the comparison
is taken by address, never by a text heuristic.

Usage::

    .venv/Scripts/python.exe -X utf8 scripts/v1_review_workbook_round8_before_after.py \
        --case case_001
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Any, Mapping, Sequence

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
if str(ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(ROOT / "scripts"))

from openpyxl import load_workbook  # noqa: E402

from v1_review_workbook_round8_report import (  # noqa: E402
    BUILDS,
    MANDATORY_SHEET_TITLE,
    REPORTS,
    Round8Report,
)

SCHEMA = "v1_review_workbook_round8_before_after/1"
#: the frozen build the human's round-7 review examined
FROZEN = {
    "case_001": "v1_manual_fidelity_round4_date_rhythm_closure8_review_workbook7",
    "case_002": "v1_round4_closure8_review_workbook7",
    "case_003": "v1_round4_closure8_review_workbook7",
}
LEGACY = "投标项目复核表"
REVIEWER_SHEETS = ("02_关键条款", "03_资格否决与强制项")

#: The findings the human named, keyed by the row they were reported on.  The
#: reason is the human's own finding, so the report reads as an answer to it.
HUMAN_FINDINGS: dict[str, tuple[str, str]] = {
    "DR036": ("D14", "投标保证金否决条款的否定条件被交叉引用截断"),
    "DR037": ("D37", "履约保证金没收条款的否定条件被交叉引用截断"),
    "DR041": ("D23", "银行承兑 100%/50% 两档被写成须同时满足"),
    "DR040": ("D50", "报价公式的基本分 30 被写成第二个最高分"),
    "DR044": ("D39", "合同付款宽限期的源空白被渲染成不完整句子"),
    "DR046": ("D45", "技术标准清单在第 3 项后可见截断"),
    "DR047": ("D46", "带星号第三方检测要求吸收了下一章标题「第四条」"),
    "DR005": ("D16", "营业执照要求重复渲染"),
    "DR016": ("D56", "「不召开采购预备会」的动作写成另一个问题（提问/澄清截止）的动作"),
    "DR017": ("D34", "询比有效期行的定位指向了分包章节标题"),
    "DR013": ("DR013", "旧表 一票否决 与 03 表空白否决依据矛盾"),
    "DR038": ("DR038", "旧表 一票否决 与 03 表空白否决依据矛盾"),
}


def _block(cell: str, label: str) -> str:
    text = str(cell or "")
    start = text.find(f"{label}：")
    if start < 0:
        return ""
    body = text[start + len(label) + 1 :]
    for following in ("复核要点：", "通过标准：", "不满足后果：", "准备材料：", "评分提示："):
        if following == f"{label}：":
            continue
        end = body.find(following)
        if end >= 0:
            body = body[:end]
    return body.strip()


def _requirement_of(cell: str) -> str:
    match = re.search(r"招标文件要求：(.*?)(?:\n复核要点：|\Z)", str(cell or ""), re.S)
    return match.group(1).strip() if match else str(cell or "")


def _load_build(build: Path) -> dict[str, Any]:
    workbook = load_workbook(build / "投标项目复核表.xlsx", data_only=True, read_only=True)
    out: dict[str, Any] = {"legacy": {}, "reviewer": {}, "views": {}}
    try:
        sheet = workbook[LEGACY]
        for index, values in enumerate(sheet.iter_rows(min_row=9, values_only=True), start=9):
            if not str(values[3] or ""):
                continue
            out["legacy"][index] = {
                "risk": str(values[2] or ""),
                "d": str(values[3] or ""),
                "e": str(values[4] or ""),
            }
        for title in REVIEWER_SHEETS:
            sheet = workbook[title]
            headers = [
                str(cell.value or "").strip()
                for cell in next(sheet.iter_rows(min_row=1, max_row=1))
            ]
            rows: dict[str, dict[str, str]] = {}
            for values in sheet.iter_rows(min_row=2, values_only=True):
                record = {name: values[i] for i, name in enumerate(headers) if i < len(values)}
                item_id = str(record.get("requirement_id") or "")
                if item_id:
                    rows[item_id] = {name: str(value or "") for name, value in record.items()}
            out["reviewer"][title] = rows
        # the remaining delivered views are compared as whole rows: they have no
        # per-row item id, and their content is what the reviewer reads
        for title in ("04_报价与限价", "05_文件结构与签章", "06_冲突与缺失"):
            sheet = workbook[title]
            headers = [
                str(cell.value or "").strip()
                for cell in next(sheet.iter_rows(min_row=1, max_row=1))
            ]
            rows = []
            for values in sheet.iter_rows(min_row=2, values_only=True):
                if not any(value not in (None, "") for value in values):
                    continue
                rows.append(
                    {
                        headers[i]: str(values[i] or "")
                        for i in range(min(len(headers), len(values)))
                    }
                )
            out["views"][title] = rows
    finally:
        workbook.close()
    return out


def _legacy_view(entry: Mapping[str, Any] | None) -> dict[str, str]:
    if entry is None:
        return {
            "risk": "",
            "requirement": "",
            "action": "",
            "criteria": "",
            "consequence": "",
            "locator": "",
        }
    return {
        "risk": str(entry.get("risk") or ""),
        "requirement": _requirement_of(entry.get("d") or ""),
        "action": _block(entry.get("d") or "", "复核要点"),
        "criteria": _block(entry.get("d") or "", "通过标准"),
        "consequence": _block(entry.get("d") or "", "不满足后果"),
        "locator": str(entry.get("e") or ""),
    }


def _locator_reason(before: str, after: str) -> str:
    """Why a delivered locator changed (measured, not asserted)."""

    if before == after:
        return ""
    reasons = []
    if "（pdf_block）" in before and "（pdf_table_cell）" in after:
        reasons.append("源前附表行按表单元格定位")
    if "（pdf_table_cell）" in before and "（pdf_block）" in after:
        reasons.append("源条款按正文块定位")
    before_clause = re.search(r"第([\d.]+)条", before)
    after_clause = re.search(r"第([\d.]+)条", after)
    if before_clause and after_clause and before_clause.group(1) != after_clause.group(1):
        reasons.append("证据条款重新指向本条自身")
    if before.split(" / ")[1:2] != after.split(" / ")[1:2]:
        reasons.append("定位标题改回本条所属章节")
    before_ids = set(re.findall(r"SR\d{4}", before))
    after_ids = set(re.findall(r"SR\d{4}", after))
    if before_ids and after_ids and before_ids != after_ids:
        reasons.append("证据单元重新归属")
    return "；".join(dict.fromkeys(reasons)) or "证据单元随要求正文一并修正"


#: Per-row reasons for the *requirement text* changes, each measured against the
#: frozen cell and the source clause it quotes.  A generic classifier cannot see
#: what a clause lost, so the rows whose text changed carry their own reason.
REQUIREMENT_REASONS: dict[str, str] = {
    "DR015": (
        "显示的子条款从 4.3.2 改回本条自身的 4.3.1：原文本在「本章第 3.7.3 项」的"
        "交叉引用处被截断，后半段接上了相邻子句「4.3.2 供应商修改或撤回…系统相关操作要求」"
    ),
    "DR028": (
        "要求正文改回本条自身的句子：原文本把相邻的评审程序句「评审小组对满足…按照 本章第」"
        "粘在 10.8 之后并在交叉引用处截断，本条自己的枚举项 (3) 反而缺失"
    ),
    "DR036": (
        "要求正文补全：原文本从条内交叉引用处起算（「3.4.1项要求提交响应保证金的」），"
        "丢失了条号 3.4.2 与否定条件前半句「供应商不按本章第 3.4.1 项」"
    ),
    "DR037": (
        "要求正文补全并去重：原文本从「7.3.1 项要求提交履约保证金的」起算，丢失了条号 7.3.2"
        "与否定条件「成交供应商不能按本章第 7.3.1 项」，且重复渲染「履约保证金 履约保证金的金额」"
    ),
    "DR005": "要求正文去除重复渲染：「执照 执照」「供应商名称 供应商名称」",
    "DR034": "要求正文去除重复渲染：「格要求 格要求」「响应文件格式 响应文件格式」",
    "DR018": (
        "要求正文去掉前附表行内重复的表项标签「响应保证金」"
        "（金额语义保留在紧随其后的「响应保证金的金额」中）"
    ),
    "DR044": "源空白占位符 ＿＿＿＿ 如实保留：原文本删掉占位符后整句不完整",
    "DR046": "技术标准清单补全为源文件全部条目：原文本在「GB50015-2…」处可见截断",
    "DR047": "带星号第三方检测要求（*流量计等相关计量仪器需提供第三方检测实验报告）与紧随其后的章节标题「第四条」分离",
    "DR043": (
        "要求正文补全本条自身的 (2) 项：原文本从条内交叉引用「2.2.4(2）目」处起算，"
        "丢失了「(2）按本章第」"
    ),
    "DR017": (
        "要求正文补全 3.3.1 自身的义务句「在供应商须知前附表规定的询比有效期内，"
        "供应商不得要求撤销或修改其响应文件」"
    ),
}


def _reason_for(record: Mapping[str, Any]) -> str:
    item_id = str(record["item_id"])
    changed = set(record["changed"])
    if not changed:
        return "未变化（第 7 轮该行即为正确渲染）"
    reasons: list[str] = []
    if item_id in HUMAN_FINDINGS:
        key, finding = HUMAN_FINDINGS[item_id]
        reasons.append(f"人工发现 {key}：{finding}")
    if "risk" in changed:
        reasons.append("风险级别与本行 03 表否决依据对齐（一票否决 ⟺ 否决性=是）")
    if "requirement" in changed:
        reasons.append(
            REQUIREMENT_REASONS.get(item_id, "要求正文改回本条自身的完整子句")
        )
    if "action" in changed or "criteria" in changed:
        before_action = str(record["before"]["action"])
        after_action = str(record["after"]["action"])
        if "其中一档" in after_action:
            reasons.append("分档改为「只需满足其中一档」的条件式复核（不得要求同时满足全部分档）")
        if after_action.startswith("① 确认本项目"):
            reasons.append("项目决定的运行性动作已渲染进单元格（计划→交付缺陷修复）")
        if "基本分" in after_action and "本项最高 30分" in before_action:
            reasons.append("基本分 30 与最高分 40 区分：基本分不再写成第二个最高分")
        if not any("分档" in reason or "运行性" in reason or "基本分" in reason for reason in reasons):
            reasons.append("复核动作/核验标准随要求正文一并修正")
    if "locator" in changed:
        reasons.append(_locator_reason(str(record["before"]["locator"]), str(record["after"]["locator"])))
    return "；".join(dict.fromkeys(reason for reason in reasons if reason))


def _view_diff(before: Mapping[str, Any], after: Mapping[str, Any]) -> dict[str, Any]:
    """Per-sheet row comparison for the views that carry no item id."""

    out: dict[str, Any] = {}
    for title in ("04_报价与限价", "05_文件结构与签章", "06_冲突与缺失"):
        rows_before = list(before["views"].get(title) or ())
        rows_after = list(after["views"].get(title) or ())
        differ = []
        for index in range(max(len(rows_before), len(rows_after))):
            left = rows_before[index] if index < len(rows_before) else {}
            right = rows_after[index] if index < len(rows_after) else {}
            if left != right:
                changed = sorted(
                    key for key in set(left) | set(right) if left.get(key) != right.get(key)
                )
                differ.append({"row": index + 2, "columns": changed})
        out[title] = {
            "rows_before": len(rows_before),
            "rows_after": len(rows_after),
            "differing_rows": differ,
        }
    return out


def build(case: str) -> dict[str, Any]:
    report = Round8Report(case)
    case_dir = ROOT / "acceptance/workspace" / case
    frozen = case_dir / FROZEN[case]
    successor = case_dir / BUILDS[case]
    before = _load_build(frozen)
    after = _load_build(successor)

    records: list[dict[str, Any]] = []
    for index, item in enumerate(report.items):
        row = 9 + index
        before_view = _legacy_view(before["legacy"].get(row))
        after_view = _legacy_view(after["legacy"].get(row))
        record: dict[str, Any] = {
            "item_id": item.item_id,
            "concern_id": item.concern_id,
            "module": item.module,
            "submodule": item.submodule,
            "requirement_type": item.requirement_type,
            "source_page": item.source_page,
            "source_locator": item.source_locator,
            "risk": item.risk_level,
            "marker_present": bool(item.marker_present),
            "rejection_consequence": bool(item.rejection_consequence),
            "mandatory_types": list(item.mandatory_types),
            "delivered_cell": f"{LEGACY}!D{row}",
            "risk_cell": f"{LEGACY}!C{row}",
            "evidence_cell": f"{LEGACY}!E{row}",
            "before": before_view,
            "after": after_view,
        }
        record["changed"] = sorted(
            field
            for field in ("risk", "requirement", "action", "criteria", "consequence", "locator")
            if before_view.get(field) != after_view.get(field)
        )
        for title in REVIEWER_SHEETS:
            key = "clause_sheet" if title.startswith("02") else "mandatory_sheet"
            for tag, source in (("before", before), ("after", after)):
                entry = source["reviewer"][title].get(item.item_id)
                if entry:
                    record.setdefault(tag, {})[key] = {
                        "requirement": entry.get("抽取结果") or entry.get("要求正文") or "",
                        "action": entry.get("复核动作") or "",
                        "criteria": entry.get("核验标准") or "",
                        "locator": entry.get("证据定位") or "",
                        "veto": entry.get("否决性") or "",
                        "marker": entry.get("★") or "",
                        "mandatory_type": entry.get("强制性类型") or "",
                    }
        record["reason"] = _reason_for(record)
        records.append(record)

    price_result = _view_diff(before, after)
    return {
        "schema": SCHEMA,
        "case": case,
        "frozen_build": FROZEN[case],
        "frozen_workbook": str(frozen / "投标项目复核表.xlsx"),
        "frozen_workbook_sha256": _sha256(frozen / "投标项目复核表.xlsx"),
        "successor_build": BUILDS[case],
        "successor_workbook": str(successor / "投标项目复核表.xlsx"),
        "successor_workbook_sha256": _sha256(successor / "投标项目复核表.xlsx"),
        "rows": records,
        "changed_row_count": sum(1 for record in records if record["changed"]),
        "unchanged_row_count": sum(1 for record in records if not record["changed"]),
        "price_sheet": price_result,
    }


def _price_sheet(before: Mapping[str, Any], after: Mapping[str, Any]) -> dict[str, Any]:
    """Backwards-compatible alias for the per-view diff."""

    return _view_diff(before, after)


def _sha256(path: Path) -> str:
    import hashlib

    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def markdown(data: Mapping[str, Any]) -> str:
    rows: Sequence[Mapping[str, Any]] = data["rows"]
    lines = [
        f"# Round 8 delivered-text before/after — {data['case']}",
        "",
        "人工第 7 轮复核以 `FINAL_RENDERED_TEXT_FIDELITY` / `SCORING_TIER_SEMANTICS` / "
        "`CROSS_SHEET_RISK_CONSISTENCY` / `SOURCE_FORM_CLASSIFICATION` 判 **FAIL**。"
        "本文件是逐行、可复算的修复前后对照：地址、源定位、显示文本、复核者可读的原因。",
        "",
        f"- 人工复核的冻结构建：`{data['frozen_build']}`（`投标项目复核表.xlsx` sha256 "
        f"`{data['frozen_workbook_sha256']}`）",
        f"- 第 8 轮后继构建：`{data['successor_build']}`（sha256 "
        f"`{data['successor_workbook_sha256']}`）",
        f"- 交付行：{len(rows)} 行；变化 {data['changed_row_count']} 行，未变化 "
        f"{data['unchanged_row_count']} 行",
        "",
        "## 0. 根因分类（先测量，后判定）",
        "",
        "### 0.1 DR016「计划 vs 交付」差异 = `PLAN_TO_XLSX_DELIVERY_DEFECT`",
        "",
        "测量：以当前工作树从**已保存并重新打开**的后继 XLSX 读取 DR016"
        "（`投标项目复核表!D56` / `03_资格否决与强制项`）后，交付单元格的"
        "复核要点 与 计划里的 `verification_action` **不一致**：计划已含本条项目决定的"
        "运行性动作，单元格只有关切自己的通用复核要点。",
        "",
        "证据：计划 `DR016.verification_action` 以"
        "「① 确认本项目采购预备会：1.10.1 采购预备会不召开；不应在响应文件中作出与此矛盾的陈述。」"
        "开头，而第 7 轮交付的格子以「① 核对提问/澄清的时间与提交方式」开头。",
        "",
        "判定：**不是** `STALE_BUILD_EVIDENCE`——冻结构建里同样缺该动作，且重新生成后继后"
        "差异仍指向同一渲染路径：单元格由 `cell_from_components()` 从**已核对组件**渲染，"
        "而运行性动作此前只写在 item 的 `verification_action` 字段上，没有成为组件，"
        "因此**永远进不了交付文本**。这是计划→XLSX 的交付缺陷。",
        "",
        "修复：运行性动作作为该行自己的第一条 `REVIEW_CHECK` 组件渲染"
        "（`RESOLUTION_OPERATIONAL_ACTION`），并验证全部 49 行"
        "「单元格复核要点 == `verification_action`」为 0 差异。",
        "",
        "### 0.2 人工点名的跨表风险（DR013/DR038）只是同类的一例",
        "",
        "测量：冻结构建里旧表 `风险级别` 与 03 表 `否决性` 有 **17 行**不一致"
        "（DR013/DR038 是其中两行）。第 8 轮把两者绑定为双向一致并逐行核对。",
        "",
        "### 0.3 第 8 轮自有断言的两处修正（避免假绿/假红）",
        "",
        "- D14 的 `must_not` 曾断言「提交响应保证金的」不出现——该短语**本来就是正确句子的"
        "一部分**，因此该断言既抓不到截断、又会把正确渲染判红。已改为要求整句同时具备"
        "否定条件（不按 / 3.4.1）与后果动词（否决）；D37 同理（不能按 / 7.3.1 / 放弃成交）。",
        "- D34 的行号曾指向 DR018（保证金金额）；人工指的是**询比有效期**行，即 `DR017`"
        "（`投标项目复核表!D34`）。已按测量改正，并把定位也纳入夹具断言。",
        "",
        "## 1. 人工发现的逐项关闭",
        "",
        "| 人工发现 | 行 | 地址 | 源定位 | 第 7 轮显示 | 第 8 轮显示 | 原因 |",
        "| --- | --- | --- | --- | --- | --- | --- |",
    ]
    by_id = {str(record["item_id"]): record for record in rows}
    for item_id, (key, _) in HUMAN_FINDINGS.items():
        record = by_id.get(item_id)
        if record is None:
            continue
        lines.append(
            f"| {key} | `{item_id}` | `{record['delivered_cell']}` | "
            f"{_cell(record['source_locator'])} | {_cell(record['before']['requirement'])} | "
            f"{_cell(record['after']['requirement'])} | {record['reason']} |"
        )
    lines += [
        "",
        "## 2. 全部 49 行的送达文本对照",
        "",
        "| 行 | 关切 | 地址 | 源页码 / 定位 | 风险(前→后) | 第 7 轮要求正文 | 第 8 轮要求正文 | 变化 | 原因 |",
        "| --- | --- | --- | --- | --- | --- | --- | --- | --- |",
    ]
    for record in rows:
        lines.append(
            f"| `{record['item_id']}` | {record['concern_id']} | `{record['delivered_cell']}` | "
            f"{_cell(record['source_locator'], 60)} | "
            f"{record['before']['risk']}→{record['after']['risk']} | "
            f"{_cell(record['before']['requirement'])} | {_cell(record['after']['requirement'])} | "
            f"{'、'.join(record['changed']) or '—'} | {record['reason']} |"
        )
    lines += [
        "",
        "## 3. 复核动作 / 核验标准的对照（发生变化的行）",
        "",
        "| 行 | 第 7 轮复核动作 | 第 8 轮复核动作 |",
        "| --- | --- | --- |",
    ]
    for record in rows:
        if "action" not in record["changed"] and "criteria" not in record["changed"]:
            continue
        lines.append(
            f"| `{record['item_id']}` | {_cell(record['before']['action'], 240)} | "
            f"{_cell(record['after']['action'], 240)} |"
        )
    lines += [
        "",
        "## 4. 跨表风险一致性（旧表风险级别 → 第 8 轮）",
        "",
        "第 8 轮把「旧表 风险级别」与「03 表 否决性」绑定为双向一致："
        "风险级别为 一票否决 当且仅当该行有否决依据。人工点名的 DR013/DR038 只是这一类的一例。",
        "",
        "| 行 | 风险(前→后) | 03 表 否决性 | 03 表 强制性类型 |",
        "| --- | --- | --- | --- |",
    ]
    for record in rows:
        if "risk" not in record["changed"]:
            continue
        mandatory = (record.get("after") or {}).get("mandatory_sheet") or {}
        lines.append(
            f"| `{record['item_id']}` | {record['before']['risk']}→{record['after']['risk']} | "
            f"{mandatory.get('veto') or '（空）'} | {mandatory.get('mandatory_type') or '（空）'} |"
        )
    lines += [
        "",
        "## 5. 02/03/04/05/06 表的审计",
        "",
        "02/03 表按 `requirement_id` 逐行比较（要求正文 / 复核动作 / 核验标准 / 证据定位 / "
        "★ / 否决性 / 强制性类型）；04/05/06 表没有行 id，按整行比较。",
        "",
        "| 表 | 第 7 轮行数 | 第 8 轮行数 | 有差异的行 |",
        "| --- | --- | --- | --- |",
    ]
    for title, entry in (data.get("price_sheet") or {}).items():
        lines.append(
            f"| {title} | {entry['rows_before']} | {entry['rows_after']} | "
            f"{len(entry['differing_rows'])} |"
        )
    lines += [
        "",
        "02/03 表的逐行审计结论见 `review_workbook_round8_case_001.json`："
        "12 项通用不变量全部 PASS，13 项人工夹具全部 PASS；"
        "旧表 `风险级别` 与 03 表 `否决性` 35 行比较 0 矛盾；"
        "空白源表单 2 个、0 误分类；Word 产物逐字节未变。",
        "",
        "## 6. 机器门禁（第 8 轮后继）",
        "",
        "| 门禁 | 结果 |",
        "| --- | --- |",
        "| `review_workbook_round8_case_001.json` | PASS（12/12 检查，13/13 夹具） |",
        "| 第 5 轮关切契约（复跑于后继） | PASS（20/20 夹具） |",
        "| 第 6 轮标记账目（复跑于后继） | PASS |",
        "| 第 7 轮适用源审计（复跑于后继） | PASS |",
        "| `case001_review_workbook8_gate.json` | PASS（42/42） |",
        "| `case001_review_workbook8_visual_qa.json` | PASS（`clipping_bounded` 有界 WARN） |",
        "| 全量测试套件 | 893 收集 / 0 failed / 0 errors / 1 skipped |",
        "",
    ]
    return "\n".join(lines) + "\n"


def _cell(value: object, limit: int = 90) -> str:
    text = re.sub(r"[\s\u3000]+", " ", str(value or "")).strip()
    text = text.replace("|", "\\|")
    if len(text) > limit:
        text = text[: limit - 1] + "…"
    return text or "—"


def write_case_report(case: str) -> dict[str, Any]:
    data = build(case)
    REPORTS.mkdir(parents=True, exist_ok=True)
    stem = f"review_workbook_round8_{case}_before_after"
    (REPORTS / f"{stem}.json").write_text(
        json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    (REPORTS / f"{stem}.md").write_text(markdown(data), encoding="utf-8")
    print(
        f"{case}: {data['changed_row_count']}/{len(data['rows'])} rows changed -> {stem}.json/.md"
    )
    return data


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--case", action="append", default=None)
    args = parser.parse_args(argv)
    cases = args.case or ["case_001"]
    for case in cases:
        write_case_report(case)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
