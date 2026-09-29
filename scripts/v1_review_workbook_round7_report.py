"""Round 7: applicable-source resolution, rejection scope, evidence fidelity.

The human's round-6 manual review failed with
``SOURCE_APPLICABILITY_AND_EVIDENCE_FIDELITY``.  Round 7 closes four classes of
defect on the delivered workbook cells (never on the internal objects alone):

* a generic bidder-instruction clause is displayed where a project-specific
  schedule value (前附表 / 附表) governs the project;
* a consequence that is *not* a response-stage rejection (a scoring deduction, a
  post-award obligation, a contract liability, a procedural statement) is shown
  as 否决性 on an ordinary requirement;
* the displayed requirement and its evidence locator / section / excerpt describe
  different source semantics;
* a rendered source fragment is incomplete, duplicated or numerically corrupted,
  and different platform roles are collapsed into one fabricated conflict.

Usage::

    python scripts/v1_review_workbook_round7_report.py --three-case
    python scripts/v1_review_workbook_round7_report.py --case case_001
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
if str(ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(ROOT / "scripts"))

from openpyxl import load_workbook  # noqa: E402

from tender_basic.applicability_invariants import (  # noqa: E402
    check_applicable_source,
    check_evidence_locator,
    check_fragment_completeness,
    check_platform_roles,
    check_rejection_scope,
)
from tender_basic.evidence_unit import EvidenceUnitIndex, _flatten  # noqa: E402
from tender_basic.models import ProjectFacts  # noqa: E402
from tender_basic.platform_roles import (  # noqa: E402
    PLATFORM_ROLES,
    classify_platform_role,
)
from tender_basic.review_workbook_views import SHEET_TITLES  # noqa: E402
from v1_review_workbook_round4_report import Round4Report  # noqa: E402
from v1_review_workbook_round6_report import CASES, Round6Report  # noqa: E402

MANDATORY_SHEET_TITLE = "03_资格否决与强制项"
CLAUSE_SHEET_TITLE = "02_关键条款"
EXCEPTION_SHEET_TITLE = "06_冲突与缺失"
#: the delivered template sheet: every plan row is printed here as well
LEGACY_SHEET_TITLE = "投标项目复核表"

#: every delivered review row lives on one of these sheets
DELIVERED_SHEET_TITLES = (CLAUSE_SHEET_TITLE, MANDATORY_SHEET_TITLE)
#: the displayed-requirement and evidence columns differ per sheet
REQUIREMENT_COLUMN = {CLAUSE_SHEET_TITLE: "抽取结果", MANDATORY_SHEET_TITLE: "要求正文"}
LOCATOR_COLUMN = {CLAUSE_SHEET_TITLE: "证据定位", MANDATORY_SHEET_TITLE: "证据定位"}
EXCERPT_COLUMN = {CLAUSE_SHEET_TITLE: "证据摘要", MANDATORY_SHEET_TITLE: "证据摘要"}
PAGE_COLUMN = {CLAUSE_SHEET_TITLE: "源页码", MANDATORY_SHEET_TITLE: "证据页码"}

#: the legacy sheet prints locator + excerpt in one cell, locator first
_LEGACY_LOCATOR_RE = re.compile(r"^(.*?（pdf_(?:block|table_cell)）)")
_LEGACY_SOURCE_TAIL_RE = re.compile(r"（源条款")
#: the legacy cell labels its source block
_LEGACY_REQUIREMENT_RE = re.compile(r"招标文件要求：(.*?)(?:\n复核要点：|\Z)", re.S)

REPORTS = ROOT / "acceptance/reports/v1_generalization"

#: round-7 successor build per case
BUILDS: dict[str, str] = {
    "case_001": "v1_manual_fidelity_round4_date_rhythm_closure8_review_workbook7",
    "case_002": "v1_round4_closure8_review_workbook7",
    "case_003": "v1_round4_closure8_review_workbook7",
}

#: §24 round-5 fixtures: the value *relationships* that must never drift.  They
#: are CASE001's own project values (its retention money is 5% and its warranty is
#: 24 months); other cases state their own numbers, so the invariant asserted for
#: them is that a value is present and that the two relationships stay distinct.
VALUE_INVARIANTS: tuple[tuple[str, str, str], ...] = (
    ("PROJECT_WARRANTY", "24", "项目质保期 24 个月"),
    ("RETENTION_MONEY_RATIO", "5", "质保金比例 5%"),
    ("RETENTION_RELEASE_PERIOD", "12", "质保金释放期 12 个月"),
)
#: the case the invariants are the human's own fixtures for
VALUE_INVARIANT_CASE = "case_001"


def _clean(value: object) -> str:
    return re.sub(r"[\s\u3000]+", "", str(value or ""))


@dataclass
class Check:
    name: str
    ok: bool
    detail: str
    evidence: dict[str, Any] = field(default_factory=dict)

    def as_dict(self) -> dict[str, Any]:
        return {
            "check": self.name,
            "ok": bool(self.ok),
            "detail": self.detail,
            "evidence": self.evidence,
        }


class Round7Report:
    """Audit one case's round-7 successor workbook from its final cells."""

    def __init__(
        self,
        case: str,
        *,
        case_dir: Path | None = None,
        build_name: str | None = None,
    ) -> None:
        self.case = case
        spec = CASES[case]
        self.case_dir = case_dir or (ROOT / "acceptance/workspace" / case)
        #: The frozen round-7 successor is the default, so this module keeps
        #: reproducing the round-7 record.  A later round passes its own
        #: successor name to re-run the same gates against the newer build
        #: without rewriting the round-7 evidence on disk.
        self.build = self.case_dir / (build_name or BUILDS[case])
        self.workbook = self.build / "投标项目复核表.xlsx"
        self.r4 = Round4Report(self.build, case)
        self.plan = self.r4.plan
        self.document = self.r4.document
        self.facts: ProjectFacts = self.r4.facts
        self.items = list(self.r4.items)
        self.background = list(self.r4.background)
        self.by_item = {str(item.item_id): item for item in [*self.items, *self.background]}
        # the plan's atom stream, rebuilt exactly as the plan built it
        self.atoms = self._plan_atoms()
        self.units = EvidenceUnitIndex.build(self.document, self.atoms)
        self.checks: list[Check] = []
        self.fixtures: dict[str, dict[str, Any]] = {}
        self.counts: dict[str, int] = {}

    # -- cell access ------------------------------------------------------ #

    def _plan_atoms(self) -> list[Any]:
        """The plan's own atom stream (the source of every delivered row).

        Rebuilt the same way ``build_dynamic_review_plan`` builds it, so the audit
        can locate each row's evidence atom in the document.
        """

        from tender_basic.dynamic_requirements import build_requirement_index
        from tender_basic.review_concern import atomize_units
        from tender_basic.source_applicability import (
            apply_applicable_resolutions,
            discover_schedule_rows,
            resolve_applicable_sources,
        )

        atoms = atomize_units(build_requirement_index(self.document).units)
        schedule_rows = discover_schedule_rows(self.document)
        _applied, stream = apply_applicable_resolutions(
            atoms, resolve_applicable_sources(atoms, schedule_rows, self.document)
        )
        return stream

    def _rows(self, title: str) -> list[dict[str, Any]]:
        workbook = load_workbook(self.workbook, data_only=True, read_only=True)
        try:
            worksheet = workbook[title]
            headers = [
                str(cell.value or "").strip()
                for cell in next(worksheet.iter_rows(min_row=1, max_row=1))
            ]
            out: list[dict[str, Any]] = []
            for row in worksheet.iter_rows(min_row=2, values_only=True):
                if not any(value not in (None, "") for value in row):
                    continue
                out.append({headers[i]: row[i] for i in range(min(len(headers), len(row)))})
            return out
        finally:
            workbook.close()

    def delivered_rows(self) -> list[dict[str, Any]]:
        """Every delivered review row, read from the *final* workbook cells.

        Both reviewer sheets carry delivered rows (a row that is not mandatory
        lives on 02_关键条款), so the audit reads both and never trusts the plan
        alone.
        """

        out: list[dict[str, Any]] = []
        for title in DELIVERED_SHEET_TITLES:
            for row in self._rows(title):
                item_id = str(row.get("requirement_id") or "").strip()
                if not item_id:
                    continue
                item = self.by_item.get(item_id)
                out.append(
                    {
                        "sheet": title,
                        "item_id": item_id,
                        "item": item,
                        "requirement": str(row.get(REQUIREMENT_COLUMN[title]) or ""),
                        "action": str(row.get("复核动作") or ""),
                        "criteria": str(row.get("核验标准") or ""),
                        "marker": str(row.get("★") or ""),
                        "veto": str(row.get("否决性") or ""),
                        "mandatory_type": str(row.get("强制性类型") or ""),
                        "page": row.get(PAGE_COLUMN[title]),
                        "locator": str(row.get(LOCATOR_COLUMN[title]) or ""),
                        "excerpt": str(row.get(EXCERPT_COLUMN[title]) or ""),
                        "raw_marker": str(row.get("源标记") or ""),
                        "substantive_basis": str(row.get("实质性依据") or ""),
                        "rejection_basis": str(row.get("否决依据") or ""),
                    }
                )
        out.extend(self.legacy_rows())
        return out

    def legacy_rows(self) -> list[dict[str, Any]]:
        """The delivered template sheet rows (the human's own working sheet).

        A row that is neither a key clause nor mandatory is still delivered here,
        so round 7 audits this sheet too instead of trusting the reviewer views.
        The template prints no requirement id, so rows are bound to plan items by
        their rendered cell text (the same binding the round-4 report uses).
        """

        out: list[dict[str, Any]] = []
        for item in self.items:
            row = self.r4.rows_by_item.get(str(item.item_id))
            if row is None:
                continue
            requirement_cell = str(getattr(row, "d_text", "") or "")
            evidence_cell = str(getattr(row, "e_text", "") or "")
            requirement = requirement_cell
            found = _LEGACY_REQUIREMENT_RE.search(requirement_cell)
            if found:
                requirement = found.group(1).strip()
            locator = evidence_cell
            found = _LEGACY_LOCATOR_RE.match(evidence_cell)
            if found:
                locator = found.group(1).strip()
            excerpt = _LEGACY_LOCATOR_RE.sub("", evidence_cell).strip()
            excerpt = _LEGACY_SOURCE_TAIL_RE.split(excerpt)[0].strip()
            out.append(
                {
                    "sheet": LEGACY_SHEET_TITLE,
                    "item_id": str(item.item_id),
                    "item": item,
                    "requirement": requirement,
                    "action": "",
                    "criteria": "",
                    # the template prints the locator and a 110-character excerpt
                    # in one clipped cell, so its evidence cell is not a full
                    # locator: the strict unit-membership gate runs on the
                    # reviewer sheets, while this sheet is still audited for its
                    # displayed requirement, page and fragment integrity
                    "strict_locator": False,
                    "marker": "",
                    "veto": "",
                    "mandatory_type": "",
                    "page": item.source_page,
                    "locator": locator,
                    "excerpt": excerpt,
                    "raw_marker": "",
                    "substantive_basis": "",
                    "rejection_basis": "",
                }
            )
        return out

    def _sheet_values(self, title: str) -> list[tuple[Any, ...]]:
        workbook = load_workbook(self.workbook, data_only=True, read_only=True)
        try:
            worksheet = workbook[title]
            return [
                row
                for row in worksheet.iter_rows(values_only=True)
                if any(value not in (None, "") for value in row)
            ]
        finally:
            workbook.close()

    def check(self, name: str, ok: bool, detail: str, evidence: dict[str, Any] | None = None) -> None:
        self.checks.append(Check(name, bool(ok), detail, evidence or {}))

    def fixture(self, key: str, ok: bool, detail: str, evidence: str = "") -> None:
        self.fixtures[key] = {
            "status": "PASS" if ok else "FAIL",
            "detail": detail,
            "evidence": evidence[:400],
        }

    # -- checks ----------------------------------------------------------- #

    def audit(self) -> dict[str, Any]:
        rows = self.delivered_rows()
        self.check_applicable_sources(rows)
        self.check_rejection_scopes(rows)
        self.check_locators(rows)
        self.check_fragments(rows)
        self.check_platforms()
        self.check_values(rows)
        self.check_round6_regression()
        self.check_word_artifacts()
        failed = [check.name for check in self.checks if not check.ok]
        failed_fixtures = [key for key, value in self.fixtures.items() if value["status"] != "PASS"]
        verdict = "PASS" if not failed and not failed_fixtures else "FAIL"
        return {
            "schema": "v1_review_workbook_round7/1",
            "case": self.case,
            "build_id": self.build.name,
            "workbook": str(self.workbook.relative_to(ROOT).as_posix()),
            "workbook_sha256": _sha256(self.workbook),
            "verdict": verdict,
            "check_count": len(self.checks),
            "failed_checks": failed,
            "checks": [check.as_dict() for check in self.checks],
            "fixtures": self.fixtures,
            "fixture_summary": {
                "total": len(self.fixtures),
                "passed": len(self.fixtures) - len(failed_fixtures),
                "failed": failed_fixtures,
            },
            "counts": dict(self.counts),
            "delivered_row_count": len(rows),
        }

    # -- 1. applicable source --------------------------------------------- #

    def check_applicable_sources(self, rows: Sequence[Mapping[str, Any]]) -> None:
        result = check_applicable_source(self.build, self.document, self.items, rows)
        self.counts["GENERIC_REFERENCE_UNRESOLVED_COUNT"] = result["unresolved_count"]
        self.counts["SOURCE_UNREADABLE_REFERENCE_COUNT"] = result["source_unreadable_count"]
        self.check(
            "every generic schedule reference resolves to the project value",
            result["unresolved_count"] == 0,
            f"{result['resolved_count']} resolved, {result['unresolved_count']} unresolved, "
            f"{result['source_unreadable_count']} source-unreadable",
            {
                "unresolved": result["unresolved"][:8],
                "source_unreadable": result["source_unreadable"][:8],
            },
        )
        self.check(
            "no delivered row displays a superseded generic clause",
            result["superseded_displayed_count"] == 0,
            f"{result['superseded_displayed_count']} row(s) still display the generic template",
            {"rows": result["superseded_displayed"][:8]},
        )
        for key, fixture in result["fixtures"].items():
            self.fixture(key, fixture["ok"], fixture["detail"], fixture.get("evidence", ""))

    # -- 2. rejection scope ----------------------------------------------- #

    def check_rejection_scopes(self, rows: Sequence[Mapping[str, Any]]) -> None:
        result = check_rejection_scope(self.document, self.plan, rows)
        self.counts["FALSE_REJECTION_CLASSIFICATION_COUNT"] = result["false_count"]
        self.check(
            "no scoring / post-award / contract / procedural consequence is shown as 否决性",
            result["false_count"] == 0,
            f"{result['false_count']} false rejection classification(s) "
            f"across {result['vetoed_rows']} vetoed row(s)",
            {"false": result["false"][:8]},
        )
        for key, fixture in result["fixtures"].items():
            self.fixture(key, fixture["ok"], fixture["detail"], fixture.get("evidence", ""))

    # -- 3. evidence locator ---------------------------------------------- #

    def check_locators(self, rows: Sequence[Mapping[str, Any]]) -> None:
        result = check_evidence_locator(
            self.units, self.items, rows, self.atoms, case=self.case
        )
        self.counts["EVIDENCE_LOCATOR_SEMANTIC_MISMATCH_COUNT"] = result["mismatch_count"]
        self.check(
            "displayed requirement, page, section, clause and excerpt describe one unit",
            result["mismatch_count"] == 0,
            f"{result['checked']} row(s) checked, {result['mismatch_count']} mismatch(es)",
            {"mismatches": result["mismatches"][:8]},
        )
        for key, fixture in result["fixtures"].items():
            self.fixture(key, fixture["ok"], fixture["detail"], fixture.get("evidence", ""))

    # -- 4. fragment completeness ----------------------------------------- #

    def check_fragments(self, rows: Sequence[Mapping[str, Any]]) -> None:
        result = check_fragment_completeness(rows, case=self.case)
        self.counts["INCOMPLETE_RENDERED_SOURCE_FRAGMENT_COUNT"] = result["incomplete_count"]
        self.counts["CORRUPTED_RENDERED_NUMERIC_TEXT_COUNT"] = result["corrupted_count"]
        self.check(
            "no rendered source fragment is incomplete",
            result["incomplete_count"] == 0,
            f"{result['incomplete_count']} incomplete fragment(s) of {len(rows)} row(s)",
            {"incomplete": result["incomplete"][:8]},
        )
        self.check(
            "no rendered numeric text is duplicated or corrupted",
            result["corrupted_count"] == 0,
            f"{result['corrupted_count']} corrupted numeric fragment(s)",
            {"corrupted": result["corrupted"][:8]},
        )
        for key, fixture in result["fixtures"].items():
            self.fixture(key, fixture["ok"], fixture["detail"], fixture.get("evidence", ""))

    # -- 5. platform roles ------------------------------------------------ #

    def check_platforms(self) -> None:
        result = check_platform_roles(self.build, self.facts, self._rows(EXCEPTION_SHEET_TITLE))
        self.counts["FALSE_PLATFORM_CONFLICT_COUNT"] = result["false_conflict_count"]
        self.check(
            "different platform roles are not collapsed into one conflict",
            result["false_conflict_count"] == 0,
            f"{result['false_conflict_count']} false platform conflict(s); "
            f"roles={result['role_count']}",
            {
                "false_conflicts": result["false_conflicts"][:8],
                "roles": result["roles"],
                "platform_status": result["status"],
            },
        )
        for key, fixture in result["fixtures"].items():
            self.fixture(key, fixture["ok"], fixture["detail"], fixture.get("evidence", ""))

    # -- 6. round-5 value invariants -------------------------------------- #

    def check_values(self, rows: Sequence[Mapping[str, Any]]) -> None:
        """§24: the round-5/round-6 fact semantics must not drift.

        The invariants are the values the human fixed in round 5 for CASE001 (its
        project warranty is 24 months, its retention money is 5%, its retention is
        released after 12 months).  Every other case states its own numbers, so
        there the invariant is the *relationship*: a warranty period is not the
        retention release period, and a retention ratio is not the payment ratio.
        """

        problems: list[str] = []
        checked = 0
        if self.case == VALUE_INVARIANT_CASE:
            for concern_id, value, label in VALUE_INVARIANTS:
                items = [item for item in self.items if str(item.concern_id) == concern_id]
                if not items:
                    continue
                checked += 1
                rendered = " ".join(str(item.source_requirement) for item in items)
                if value not in rendered:
                    problems.append(f"{concern_id}: {value} missing from {label}")
        # the two relationships must stay distinct in every case where both exist
        warranty = [item for item in self.items if str(item.concern_id) == "PROJECT_WARRANTY"]
        release = [
            item for item in self.items if str(item.concern_id) == "RETENTION_RELEASE_PERIOD"
        ]
        if warranty and release:
            warranty_text = " ".join(str(item.source_requirement) for item in warranty)
            release_text = " ".join(str(item.source_requirement) for item in release)
            warranty_numbers = set(re.findall(r"\d+(?:\.\d+)?", warranty_text))
            release_numbers = set(re.findall(r"\d+(?:\.\d+)?", release_text))
            if warranty_numbers and release_numbers:
                checked += 1
                if "12" in warranty_numbers and "24" not in warranty_numbers:
                    problems.append("PROJECT_WARRANTY shows 12 months (retention release)")
                if "24" in release_numbers and "12" not in release_numbers:
                    problems.append("RETENTION_RELEASE_PERIOD shows 24 months (warranty)")
        self.counts["VALUE_INVARIANT_FAILURE_COUNT"] = len(problems)
        self.check(
            "round-5/round-6 fact invariants still hold",
            not problems,
            f"{checked} invariant(s) checked, {len(problems)} failure(s)",
            {"problems": problems},
        )

    # -- 7. banked round-6 accounting ------------------------------------- #

    def check_round6_regression(self) -> None:
        data = Round6Report(
            self.case, case_dir=self.case_dir, build_name=self.build.name
        ).run()
        counts = data["counts"]
        self.counts["ROUND6_MARKER_OCCURRENCE_COUNT"] = counts["source_marker_occurrence_count"]
        self.counts["ROUND6_MARKER_LOST_COUNT"] = counts["marker_lost_count"]
        self.check(
            "the banked round-6 marker accounting still holds",
            data["verdict"] == "PASS",
            f"round-6 accounting verdict={data['verdict']} "
            f"(failed={','.join(data['failed_checks']) or 'none'})",
            {
                "failed_checks": data["failed_checks"],
                "failed_fixtures": data["fixture_summary"]["failed"],
                "counts": {key: counts[key] for key in sorted(counts) if "marker" in key},
                "concern_contract_verdict": data["concern_contract_verdict"],
            },
        )
        self.counts["ROUND6_CONCERN_CONTRACT_FAILURES"] = len(
            data["concern_contract"]["failed_checks"]
        )

    # -- 8. word artifacts ------------------------------------------------- #

    def check_word_artifacts(self) -> None:
        manifest = json.loads((self.build / "build_manifest.json").read_text(encoding="utf-8"))
        identity = manifest.get("artifact_identity", {})
        ok = bool(identity) and all(entry.get("byte_identical") for entry in identity.values())
        self.check(
            "Word artifacts are byte-identical (not re-rendered)",
            ok,
            f"{len(identity)} artifact(s) compared against the accepted source build",
            {
                name: entry.get("successor_sha256")
                for name, entry in list(identity.items())[:3]
            },
        )


def _sha256(path: Path) -> str:
    import hashlib

    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def markdown(data: dict) -> str:
    lines = [
        f"# Round 7 review workbook audit — {data['case']}",
        "",
        f"- build: `{data['build_id']}`",
        f"- workbook: `{data['workbook']}`",
        f"- sha256: `{data['workbook_sha256']}`",
        f"- verdict: **{data['verdict']}**",
        f"- delivered rows audited: {data['delivered_row_count']}",
        "",
        "## Zero-count gates",
        "",
        "| gate | value |",
        "| --- | --- |",
    ]
    for key in sorted(data["counts"]):
        lines.append(f"| {key} | {data['counts'][key]} |")
    lines += ["", "## Checks", "", "| check | result | detail |", "| --- | --- | --- |"]
    for check in data["checks"]:
        lines.append(
            f"| {check['check']} | {'PASS' if check['ok'] else 'FAIL'} | {check['detail']} |"
        )
    lines += ["", "## Fixtures", "", "| fixture | result | detail |", "| --- | --- | --- |"]
    for key in sorted(data["fixtures"]):
        fixture = data["fixtures"][key]
        lines.append(f"| {key} | {fixture['status']} | {fixture['detail']} |")
    return "\n".join(lines) + "\n"


def write_case_report(
    case: str,
    *,
    case_dir: Path | None = None,
    build_name: str | None = None,
    stem: str | None = None,
) -> dict[str, Any]:
    report = Round7Report(case, case_dir=case_dir, build_name=build_name)
    data = report.audit()
    REPORTS.mkdir(parents=True, exist_ok=True)
    stem = stem or f"review_workbook_round7_{case}"
    (REPORTS / f"{stem}.json").write_text(
        json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    (REPORTS / f"{stem}.md").write_text(markdown(data), encoding="utf-8")
    print(f"{case}: {data['verdict']} ({len(data['failed_checks'])} failed check(s)) -> {stem}.json")
    for check in data["checks"]:
        if not check["ok"]:
            print(f"   FAIL {check['check']}: {check['detail']}")
    for key, value in data["fixtures"].items():
        if value["status"] != "PASS":
            print(f"   FIXTURE {key}: {value['status']} — {value['detail']}")
    return data


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--case", action="append", default=None)
    parser.add_argument("--three-case", action="store_true")
    args = parser.parse_args(argv)
    cases = list(BUILDS) if args.three_case or not args.case else args.case
    results = {case: write_case_report(case) for case in cases}
    if len(results) > 1:
        summary = {
            "schema": "v1_review_workbook_round7_generalization/1",
            "cases": {
                case: {
                    "verdict": data["verdict"],
                    "failed_checks": data["failed_checks"],
                    "failed_fixtures": data["fixture_summary"]["failed"],
                    "counts": data["counts"],
                }
                for case, data in results.items()
            },
            "verdict": (
                "PASS" if all(data["verdict"] == "PASS" for data in results.values()) else "FAIL"
            ),
        }
        path = REPORTS / "review_workbook_round7_generalization.json"
        path.write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(f"generalization: {summary['verdict']} -> {path.name}")
    return 0 if all(data["verdict"] == "PASS" for data in results.values()) else 1


if __name__ == "__main__":
    raise SystemExit(main())
