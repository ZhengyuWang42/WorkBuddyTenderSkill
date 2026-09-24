"""Round-6 source-visible criticality report (REVIEW WORKBOOK ROUND 6).

The round-5 review workbook was still losing the source's own emphasis markers,
and it treated three different source questions as one:

* which requirements does the *source* mark (``*`` / ``★`` / ...)?
* which requirements does the *source* call 实质性要求 (or 必备证明材料)?
* which requirements does the *source* attach a 否决/无效 consequence to?

Round 6 keeps those three dimensions apart.  ``tender_basic.source_criticality``
discovers, per document, the marker vocabulary, the marker's *meaning* as the
document states it, the governing substantive-requirement clause, the
consequence rules and the explicit references between clauses.  This report
audits the rendered workbook against that model and persists the per-row chain
(source atom -> marker evidence -> governing basis -> concern -> review point ->
final cell) for a human reviewer.

Usage::

    python -X utf8 scripts/v1_review_workbook_round6_report.py --case case_001
    python -X utf8 scripts/v1_review_workbook_round6_report.py --three-case
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
if str(ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(ROOT / "scripts"))

from openpyxl import load_workbook  # noqa: E402

from tender_basic.source_criticality import (  # noqa: E402
    BASIS_REFERENCE_PROPAGATION,
    BASIS_SOURCE_MARKER,
    COVERAGE_BACKGROUND_ROW,
    COVERAGE_DELIVERED_ROW,
    COVERAGE_REFERENCE_CHILD,
    COVERAGE_REFERENCE_PARENT,
    COVERAGE_UNCOVERED,
    MARKER_CHARS,
    MANDATORY_TYPE_STARRED,
)
from v1_review_workbook_round4_report import MANDATORY_SHEET, Round4Report  # noqa: E402
from v1_review_workbook_round5_report import Round5Report  # noqa: E402

CASES = {
    "case_001": {
        "build": "v1_manual_fidelity_round4_date_rhythm_closure8_review_workbook6",
        "source": "v1_manual_fidelity_round4_date_rhythm_closure8",
        "before": "v1_manual_fidelity_round4_date_rhythm_closure8_review_workbook5",
        "marker_fixtures": (
            "1.4.1",
            "1.4.4",
            "1.4.5",
            "1.4.6",
            "1.4.7",
            "1.4.8",
            "1.5.1",
            "1.5.2",
            "3.3.1",
            "3.4.1",
            "10.1",
            "流量计",
        ),
        "row_fixtures": (
            "DR001",
            "DR002",
            "DR003",
            "DR004",
            "DR012",
            "DR017",
            "DR018",
            "DR019",
            "DR020",
            "DR030",
            "DR047",
        ),
    },
    "case_002": {
        "build": "v1_round4_closure8_review_workbook6",
        "source": "v1_round4_closure8",
        "before": "v1_round4_closure8_review_workbook5",
    },
    "case_003": {
        "build": "v1_round4_closure8_review_workbook6",
        "source": "v1_round4_closure8",
        "before": "v1_round4_closure8_review_workbook5",
    },
}

CASES_SPECIAL = None

REPORTS = ROOT / "acceptance/reports/v1_generalization"


@dataclass
class Check:
    name: str
    ok: bool
    detail: str
    evidence: dict

    def as_dict(self) -> dict:
        return {
            "check": self.name,
            "ok": bool(self.ok),
            "detail": self.detail,
            "evidence": self.evidence,
        }


class Round6Report:
    def __init__(self, case: str, *, case_dir: Path | None = None) -> None:
        spec = CASES[case]
        self.case = case
        self.case_dir = case_dir or (ROOT / "acceptance/workspace" / case)
        self.build = self.case_dir / spec["build"]
        self.source_build = self.case_dir / spec["source"]
        self.before = self.case_dir / spec["before"]
        self.r4 = Round4Report(self.build, case)
        self.plan = self.r4.plan
        self.document = self.r4.document
        self.facts = self.r4.facts
        self.items = list(self.r4.items)
        self.background = list(self.r4.background)
        self.criticality = self.plan.criticality
        self.checks: list[Check] = []
        self.marker_records: list[dict] = []
        self.marker_false_positive_count = 0
        self.uncovered_substantive_count = 0
        self.row_records: list[dict] = []
        self.fixtures: dict[str, dict] = {}
        self.coverage_records: list[dict] = []

    # -- helpers ---------------------------------------------------------- #

    def check(self, name: str, ok: bool, detail: str, evidence: dict | None = None) -> None:
        self.checks.append(Check(name, bool(ok), detail, evidence or {}))

    def _sheet_rows(self, title: str) -> list[dict]:
        workbook = load_workbook(self.build / "投标项目复核表.xlsx", data_only=True, read_only=True)
        try:
            worksheet = workbook[title]
            headers = [
                str(cell.value or "").strip()
                for cell in next(worksheet.iter_rows(min_row=1, max_row=1))
            ]
            rows = []
            for row in worksheet.iter_rows(min_row=2, values_only=True):
                if not any(value not in (None, "") for value in row):
                    continue
                rows.append({headers[i]: row[i] for i in range(min(len(headers), len(row)))})
            return rows
        finally:
            workbook.close()

    def _item(self, item_id: str):
        return next((item for item in self.items if item.item_id == item_id), None)

    def marker_lost_count(self) -> int:
        """Source markers that never reached a review row.

        "Reached" means: a delivered row shows ``★`` for the atom carrying the
        marker, or the atom belongs to a non-bidder-facing clause that is kept
        for coverage only (round-4 background policy) -- in that case the marker
        is still recorded in this audit, it is deliberately not a delivery row.
        """

        lost_states = {COVERAGE_UNCOVERED, "LOST_DELIVERED_ROW_WITHOUT_STAR"}
        return sum(1 for record in self.marker_records if record["status"] in lost_states)

    # -- marker fidelity -------------------------------------------------- #

    def audit_markers(self) -> None:
        index = self.criticality
        # Concerns own atoms; map atom -> row through the plan's review points.
        atom_owner: dict[str, list[str]] = {}
        for item in [*self.items, *self.background]:
            for atom_id in getattr(item.review_point, "owned_atom_ids", ()) or ():
                atom_owner.setdefault(str(atom_id), []).append(item.item_id)

        star_by_item = {
            str(row.get("requirement_id")): str(row.get("★") or "").strip()
            for row in self._sheet_rows(MANDATORY_SHEET)
        }
        marker_atoms = [
            atom_id
            for atom_id, record in index.criticalities.items()
            if record.marker_present
        ]
        delivered_ids = {item.item_id for item in self.items}
        lost: list[dict] = []
        false_positive: list[dict] = []
        for atom_id in marker_atoms:
            record = index.criticalities[atom_id]
            owners = atom_owner.get(atom_id, [])
            delivered_owners = [owner for owner in owners if owner in delivered_ids]
            background_owners = [owner for owner in owners if owner not in delivered_ids]
            workbook_star = [
                star_by_item[owner] for owner in delivered_owners if star_by_item.get(owner) == "★"
            ]
            if not owners:
                status = COVERAGE_UNCOVERED
            elif delivered_owners and not workbook_star:
                status = "LOST_DELIVERED_ROW_WITHOUT_STAR"
            elif delivered_owners:
                status = COVERAGE_DELIVERED_ROW
            else:
                # owned only by a non-bidder-facing (background) clause: kept in
                # the plan and in this audit, deliberately not a delivery row
                status = COVERAGE_BACKGROUND_ROW
            self.marker_records.append(
                {
                    "source_atom_id": atom_id,
                    "atom_text": index.atom_text.get(atom_id, "")[:200],
                    "raw_source_marker": record.raw_source_marker,
                    "marker_source_location": record.marker_source_location,
                    "marker_source_text": record.marker_source_text[:200],
                    "marker_semantics": record.marker_semantics,
                    "marker_evidence_id": record.marker_evidence_id,
                    "owner_items": owners,
                    "delivered_owner_items": delivered_owners,
                    "background_owner_items": background_owners,
                    "workbook_star": workbook_star,
                    "status": status,
                }
            )
            if status in (COVERAGE_UNCOVERED, "LOST_DELIVERED_ROW_WITHOUT_STAR"):
                lost.append(
                    {
                        "source_atom_id": atom_id,
                        "raw_source_marker": record.raw_source_marker,
                        "locator": record.marker_source_location,
                        "owner_items": owners,
                        "status": status,
                    }
                )
            # a marker may only exist where the source evidence says so
            if record.raw_source_marker and record.raw_source_marker not in (
                record.marker_source_text or ""
            ) and record.raw_source_marker not in (index.atom_text.get(atom_id) or ""):
                false_positive.append(
                    {
                        "source_atom_id": atom_id,
                        "raw_source_marker": record.raw_source_marker,
                        "marker_source_text": record.marker_source_text[:120],
                    }
                )

        # every ★ in the sheet must be backed by a marker atom
        sheet_starred = {
            row.get("requirement_id")
            for row in self._sheet_rows(MANDATORY_SHEET)
            if str(row.get("★") or "").strip() == "★"
        }
        unbacked = sorted(
            item_id
            for item_id in sheet_starred
            if item_id and not (self._item(item_id) and self._item(item_id).marker_present)
        )
        false_positive.extend({"item_id": item_id, "reason": "★ without marker atom"} for item_id in unbacked)

        self.check(
            "marker preserved (no source marker is lost)",
            not lost,
            f"marker_lost_count={len(lost)}",
            {"marker_atom_count": len(marker_atoms), "lost": lost[:10]},
        )
        self.check(
            "marker false positives",
            not false_positive,
            f"marker_false_positive_count={len(false_positive)}",
            {"false_positives": false_positive[:10]},
        )
        self.marker_false_positive_count = len(false_positive)

    def check_marker_fixtures(self) -> None:
        rows = self._sheet_rows(MANDATORY_SHEET)
        star_by_item = {
            row.get("requirement_id"): str(row.get("★") or "").strip() for row in rows
        }
        index = self.criticality
        delivered_ids = {item.item_id for item in self.items}
        spec = CASES[self.case]
        clause_fixtures = tuple(spec.get("marker_fixtures") or ()) or tuple(
            dict.fromkeys(marker.clause for marker in index.markers if marker.clause)
        )[:12]
        row_fixtures = tuple(spec.get("row_fixtures") or ()) or tuple(
            item.item_id for item in self.items if item.marker_present
        )[:20]
        atom_owner: dict[str, list[str]] = {}
        for item in [*self.items, *self.background]:
            for atom_id in getattr(item.review_point, "owned_atom_ids", ()) or ():
                atom_owner.setdefault(str(atom_id), []).append(item.item_id)
        for fixture in clause_fixtures:
            evidence = [
                marker for marker in index.markers if marker.clause == fixture
            ] or [
                marker for marker in index.markers if fixture in marker.source_text
            ]
            evidence_ids = {marker.evidence_id for marker in evidence}
            is_clause = bool(re.fullmatch(r"\d+(?:\.\d+)*", fixture))
            atoms = [
                atom_id
                for atom_id, record in index.criticalities.items()
                if record.marker_present
                and (
                    record.marker_evidence_id in evidence_ids
                    or (not is_clause and fixture in index.atom_text.get(atom_id, ""))
                )
            ]
            owners = sorted(
                {owner for atom_id in atoms for owner in atom_owner.get(atom_id, [])}
            )
            delivered_rows = [owner for owner in owners if owner in delivered_ids]
            starred_rows = [owner for owner in delivered_rows if star_by_item.get(owner) == "★"]
            ok = bool(evidence) and bool(atoms) and (not delivered_rows or bool(starred_rows))
            self.fixtures[f"marker:{fixture}"] = {
                "expectation": f"source marker on {fixture} reaches the workbook",
                "status": "PASS" if ok else "FAIL",
                "evidence": {
                    "marker_evidence_ids": sorted(evidence_ids)[:4],
                    "marker_atoms": atoms[:6],
                    "owner_items": owners[:6],
                    "delivered_rows": delivered_rows[:6],
                    "starred_rows": starred_rows[:6],
                },
            }
        for fixture in row_fixtures:
            item = self._item(fixture)
            self.fixtures[f"row:{fixture}"] = {
                "expectation": f"{fixture} keeps its source marker in the workbook",
                "status": "PASS" if item is not None and star_by_item.get(fixture) == "★" else "FAIL",
                "evidence": {
                    "raw_source_markers": list(getattr(item, "raw_source_markers", ()) or ()),
                    "star": star_by_item.get(fixture, ""),
                },
            }
        failed = [key for key, value in self.fixtures.items() if value["status"] != "PASS"]
        self.check(
            "known source-marker fixtures",
            not failed,
            f"{len(self.fixtures) - len(failed)}/{len(self.fixtures)} fixture(s) marked",
            {"failed": failed},
        )

    # -- semantics -------------------------------------------------------- #

    def check_semantics(self) -> None:
        rules = [rule for rule in self.criticality.substantive_rules if rule.markers]
        vocabularies = sorted({marker for rule in rules for marker in rule.markers})
        semantics = sorted({rule.semantics for rule in rules})
        unknown = [
            rule.rule_id for rule in rules if not rule.semantics or rule.semantics == "UNESTABLISHED"
        ]
        marked_rows = [item for item in self.items if item.marker_present]
        unclassified = [
            item.item_id
            for item in marked_rows
            if not item.substantive_requirement
            and item.marker_semantics in ("", "UNESTABLISHED")
        ]
        self.check(
            "source marker vocabulary is discovered per document",
            bool(vocabularies),
            f"vocabulary={vocabularies}",
            {"rules": [rule.as_dict() for rule in rules]},
        )
        self.check(
            "marker meaning comes from the document's own statement",
            bool(semantics) and all(value and value != "UNESTABLISHED" for value in semantics),
            f"semantics={semantics}",
            {"unestablished_rules": unknown},
        )
        self.check(
            "no marker is promoted to a rejection consequence by itself",
            all(
                not (item.marker_present and item.rejection_consequence)
                or bool(item.rejection_basis_atom_ids)
                for item in self.items
            ),
            f"unclassified_marked_rows={unclassified}",
            {"marked_rows": [(item.item_id, item.marker_semantics) for item in marked_rows]},
        )

    def check_governing_clause(self) -> None:
        """SOURCE MARKER -> substantive rule -> consequence rule -> row."""

        rules = list(self.criticality.substantive_rules)
        marker_rules = [rule for rule in rules if rule.markers]
        governed = [rule for rule in marker_rules if rule.governing_clause]
        anchored = [
            rule
            for rule in marker_rules
            if rule.governing_clause or rule.semantics != "UNESTABLISHED"
        ]
        unanchored = [rule.rule_id for rule in marker_rules if rule not in anchored]
        unclassified_rows: list[dict] = []
        for item in self.items:
            if not item.marker_present:
                continue
            record = self.criticality.criticalities.get(
                next(
                    (
                        atom_id
                        for atom_id in item.review_point.owned_atom_ids
                        if self.criticality.criticalities.get(str(atom_id), None)
                        is not None
                        and self.criticality.criticalities[str(atom_id)].marker_present
                    ),
                    "",
                )
            )
            semantics = record.marker_semantics if record is not None else item.marker_semantics
            if not item.substantive_requirement and semantics == "UNESTABLISHED":
                unclassified_rows.append(
                    {"item_id": item.item_id, "marker": item.raw_source_markers}
                )
        rejection_rules = list(self.criticality.rejection_rules)
        consequence_backed = [
            item.item_id
            for item in self.items
            if item.rejection_consequence and item.rejection_basis_atom_ids
        ]
        unbacked_consequence = [
            item.item_id
            for item in self.items
            if item.rejection_consequence and not item.rejection_basis_atom_ids
        ]
        self.check(
            "the tender's own substantive-requirement rule is discovered and linked",
            bool(marker_rules) and not unanchored and not unclassified_rows,
            f"{len(anchored)}/{len(marker_rules)} marker rule(s) anchored in the source "
            f"({len(governed)} via a governing clause), "
            f"{len(unclassified_rows)} unclassified marked row(s)",
            {
                "governing_clauses": sorted({rule.governing_clause for rule in governed}),
                "semantics": sorted({rule.semantics for rule in marker_rules}),
                "unanchored": unanchored,
                "unclassified_rows": unclassified_rows,
                "rule_kinds": sorted({rule.kind for rule in rules}),
            },
        )
        self.check(
            "every rejection claim traces to a source consequence rule",
            bool(consequence_backed) and not unbacked_consequence,
            f"{len(consequence_backed)} consequence-backed row(s), "
            f"{len(rejection_rules)} source consequence rule(s)",
            {
                "unbacked": unbacked_consequence,
                "rules": [rule.rule_id for rule in rejection_rules][:10],
            },
        )

    # -- reference propagation -------------------------------------------- #

    def check_references(self) -> None:
        links = list(self.criticality.reference_links)
        resolved = [link for link in links if link.target_atom_ids]
        unresolved = [link for link in links if not link.target_atom_ids]
        chains = [
            {
                "parent_atom_id": link.parent_atom_id,
                "parent_clause": link.parent_clause,
                "marker": link.marker,
                "target": link.target,
                "target_kind": link.target_kind,
                "child_count": len(link.target_atom_ids),
                "child_atom_ids": list(link.target_atom_ids[:8]),
            }
            for link in links
        ]
        propagated_children = [
            item.item_id
            for item in self.items
            if item.substantive_basis_kind == BASIS_REFERENCE_PROPAGATION
        ]
        child_items = [
            item
            for item in self.items
            if item.substantive_basis_kind == BASIS_REFERENCE_PROPAGATION
        ]
        self.check(
            "explicit source references propagate the starred parent's meaning",
            not unresolved and (not links or bool(propagated_children)),
            f"{len(resolved)} resolved reference link(s), {len(unresolved)} unresolved, "
            f"{len(propagated_children)} propagated row(s)",
            {"chains": chains},
        )
        self.check(
            "reference children are marked as propagated, never as starred",
            all(
                MANDATORY_TYPE_STARRED not in item.mandatory_types
                or item.marker_present
                for item in self.items
            ),
            f"propagated_rows={propagated_children}",
            {
                "rows": [
                    {
                        "item_id": item.item_id,
                        "types": list(item.mandatory_types),
                        "star": item.marker_present,
                        "reference_parent_atom_id": item.reference_parent_atom_id,
                        "reference_target": item.reference_target,
                    }
                    for item in child_items
                ]
            },
        )

    # -- consequence rules ------------------------------------------------- #

    def check_consequences(self) -> None:
        index = self.criticality
        explicit = [item for item in self.items if item.rejection_kind == "EXPLICIT"]
        derived = [item for item in self.items if item.rejection_kind == "DERIVED"]
        unbacked = [
            item.item_id
            for item in [*explicit, *derived]
            if not item.rejection_basis_atom_ids
        ]
        rules = [rule.as_dict() for rule in index.rejection_rules]
        self.check(
            "rejection consequence is a separate, source-backed dimension",
            not unbacked,
            f"explicit={len(explicit)} derived={len(derived)} unbacked={len(unbacked)}",
            {"unbacked": unbacked, "rules": rules[:8]},
        )

    # -- workbook columns -------------------------------------------------- #

    def check_sheet_columns(self) -> None:
        rows = self._sheet_rows(MANDATORY_SHEET)
        marked = [
            row for row in rows if self._item(str(row.get("requirement_id") or "")) is not None
            and self._item(str(row.get("requirement_id"))).marker_present
        ]
        blank_star = [row.get("requirement_id") for row in marked if str(row.get("★") or "").strip() != "★"]
        starred = {row.get("requirement_id") for row in rows if str(row.get("★") or "").strip() == "★"}
        vetoed = {row.get("requirement_id") for row in rows if str(row.get("否决性") or "").strip() == "是"}
        types = {str(row.get("强制性类型") or "") for row in rows}
        self.check(
            "every source-marked row shows ★",
            not blank_star,
            f"{len(marked)} marked row(s), {len(blank_star)} blank",
            {"blank_marked_rows": blank_star[:10]},
        )
        self.check(
            "★, 否决性 and 强制性类型 are independent columns",
            starred != vetoed or not starred,
            f"starred={len(starred)} vetoed={len(vetoed)} symmetric_difference={len(starred ^ vetoed)}",
            {
                "starred_only": sorted(starred - vetoed)[:10],
                "vetoed_only": sorted(vetoed - starred)[:10],
                "mandatory_types": sorted(types),
            },
        )
        self.check(
            "强制性类型 states the source basis instead of the requirement type",
            any("SUBSTANTIVE" in value or value == "MANDATORY" for value in types),
            f"types={sorted(types)}",
            {"row_types": [(row.get("requirement_id"), row.get("强制性类型")) for row in rows[:8]]},
        )
        self.check(
            "raw marker and its basis are carried next to the normalised columns",
            all(
                (str(row.get("★") or "").strip() != "★" or str(row.get("源标记") or "").strip() != "")
                and (
                    "SUBSTANTIVE" not in str(row.get("强制性类型") or "")
                    or str(row.get("实质性依据") or "").strip() != ""
                )
                and (
                    str(row.get("否决性") or "").strip() != "是"
                    or str(row.get("否决依据") or "").strip() != ""
                )
                for row in rows
            ),
            "源标记/实质性依据/否决依据 columns carry the source basis",
            {
                "headers": list(rows[0].keys()) if rows else [],
                "sample": [
                    {
                        "requirement_id": row.get("requirement_id"),
                        "★": row.get("★"),
                        "源标记": row.get("源标记"),
                        "强制性类型": row.get("强制性类型"),
                        "实质性依据": row.get("实质性依据"),
                        "否决依据": row.get("否决依据"),
                    }
                    for row in rows
                    if str(row.get("★") or "").strip() == "★"
                ][:5],
            },
        )

    # -- dashboard --------------------------------------------------------- #

    def check_dashboard(self) -> None:
        workbook = load_workbook(self.build / "投标项目复核表.xlsx", data_only=False, read_only=True)
        try:
            sheet = workbook["00_复核总览"]
            labels: dict[str, str] = {}
            for row in sheet.iter_rows(values_only=True):
                if row and row[0] and row[1] and str(row[1]).startswith("="):
                    labels[str(row[0])] = str(row[1])
        finally:
            workbook.close()
        expected = {
            "源标记条款数（带★）": '"★"',
            "实质性要求数（源依据）": "SUBSTANTIVE_STARRED",
            "明示或可证明否决项数": '"是"',
        }
        missing = [label for label in expected if label not in labels]
        wrong = [
            label
            for label, token in expected.items()
            if label in labels and token not in labels[label]
        ]
        conflated = [label for label in labels if "★/一票否决" in label]
        self.check(
            "dashboard reports the three dimensions separately",
            not missing and not wrong and not conflated,
            f"missing={missing} conflated={conflated}",
            {"labels": {key: labels.get(key, "") for key in expected}},
        )

    # -- coverage ---------------------------------------------------------- #

    def check_coverage(self) -> None:
        index = self.criticality
        atom_owner: dict[str, str] = {}
        for item in [*self.items, *self.background]:
            for atom_id in getattr(item.review_point, "owned_atom_ids", ()) or ():
                atom_owner.setdefault(str(atom_id), item.item_id)
        reference_parents = {
            link.parent_atom_id for link in index.reference_links if link.target_atom_ids
        }
        reference_children = {
            child for link in index.reference_links for child in link.target_atom_ids
        }
        substantives = [
            (atom_id, record)
            for atom_id, record in index.criticalities.items()
            if record.substantive_requirement
        ]
        uncovered: list[dict] = []
        for atom_id, record in substantives:
            if atom_id in reference_children:
                coverage = COVERAGE_REFERENCE_CHILD
            elif atom_id in reference_parents:
                coverage = COVERAGE_REFERENCE_PARENT
            elif atom_id in atom_owner:
                owner = atom_owner[atom_id]
                coverage = (
                    COVERAGE_DELIVERED_ROW
                    if any(item.item_id == owner for item in self.items)
                    else COVERAGE_BACKGROUND_ROW
                )
            else:
                coverage = COVERAGE_UNCOVERED
                uncovered.append(
                    {
                        "source_atom_id": atom_id,
                        "atom_text": index.atom_text.get(atom_id, "")[:160],
                        "basis_kind": record.substantive_basis_kind,
                    }
                )
            self.coverage_records.append(
                {
                    "source_atom_id": atom_id,
                    "atom_text": index.atom_text.get(atom_id, "")[:200],
                    "basis_kind": record.substantive_basis_kind,
                    "basis_atom_ids": list(record.substantive_basis_atom_ids),
                    "owner_item": atom_owner.get(atom_id, ""),
                    "coverage": coverage,
                    "rejection_kind": record.rejection_kind,
                }
            )
        self.check(
            "every source substantive requirement is covered by the review",
            not uncovered,
            f"source_substantive_requirement_without_review_coverage_count={len(uncovered)}",
            {"uncovered": uncovered[:10]},
        )
        self.uncovered_substantive_count = len(uncovered)

    # -- contract / word artifacts ---------------------------------------- #

    def check_contract(self) -> dict:
        report = Round5Report(build=self.build, case=self.case)
        data = report.run()
        failed = [check["check"] for check in data["checks"] if not check["ok"]]
        contract_failed = failed
        self.check(
            "concern contracts still hold on the round-6 successor",
            not contract_failed and not data.get("fixtures_failed"),
            f"concern contract checks failed={len(contract_failed)}, "
            f"fixtures={data.get('fixtures_passed')}",
            {
                "failed_checks": contract_failed[:6],
                "fixtures_failed": data.get("fixtures_failed"),
                "contract_count": data.get("contract_count"),
            },
        )
        return data

    def check_word_artifacts(self) -> None:
        manifest = json.loads((self.build / "build_manifest.json").read_text(encoding="utf-8"))
        identity = manifest.get("artifact_identity", {})
        self.check(
            "Word artifacts are byte-identical (not re-rendered)",
            bool(identity) and all(entry.get("byte_identical") for entry in identity.values()),
            f"{len(identity)} artifact(s) unchanged",
            {"identity": identity},
        )

    # -- row audit --------------------------------------------------------- #

    def audit_rows(self) -> None:
        rows = {row.get("requirement_id"): row for row in self._sheet_rows(MANDATORY_SHEET)}
        for item in self.items:
            if not (
                item.marker_present or item.substantive_requirement or item.rejection_consequence
            ):
                continue
            row = rows.get(item.item_id)
            atoms = [
                str(atom_id) for atom_id in item.review_point.owned_atom_ids
            ]
            basis_records = [
                self.criticality.criticalities[atom_id].as_dict()
                for atom_id in atoms
                if atom_id in self.criticality.criticalities
                and (
                    self.criticality.criticalities[atom_id].marker_present
                    or self.criticality.criticalities[atom_id].substantive_requirement
                )
            ]
            self.row_records.append(
                {
                    "item_id": item.item_id,
                    "module": item.module,
                    "concern_id": item.concern_id,
                    "concern_label": item.concern_label,
                    "source_atom_ids": atoms,
                    "source_locator": item.source_locator,
                    "source_page": item.source_page,
                    "raw_source_markers": list(item.raw_source_markers),
                    "marker_present": item.marker_present,
                    "marker_semantics": item.marker_semantics,
                    "substantive_requirement": item.substantive_requirement,
                    "substantive_basis_kind": item.substantive_basis_kind,
                    "substantive_basis_atom_ids": list(item.substantive_basis_atom_ids),
                    "rejection_consequence": item.rejection_consequence,
                    "rejection_kind": item.rejection_kind,
                    "rejection_basis_atom_ids": list(item.rejection_basis_atom_ids),
                    "mandatory_types": list(item.mandatory_types),
                    "criticality_level": item.criticality_level,
                    "criticality_reason": item.criticality_reason,
                    "reference_parent_atom_id": item.reference_parent_atom_id,
                    "reference_target": item.reference_target,
                    "review_point_id": getattr(item.review_point, "point_id", ""),
                    "review_point_checks": list(getattr(item.review_point, "review_checks", ()) or ()),
                    "final_row": row.get("序号") if row else None,
                    "cell": {
                        "★": (row or {}).get("★", ""),
                        "否决性": (row or {}).get("否决性", ""),
                        "强制性类型": (row or {}).get("强制性类型", ""),
                        "源标记": (row or {}).get("源标记", ""),
                        "实质性依据": (row or {}).get("实质性依据", ""),
                        "否决依据": (row or {}).get("否决依据", ""),
                    },
                    "delivered_cell_prefix": str(row.get("要求正文") or "")[:40] if row else "",
                    "atom_criticalities": basis_records[:8],
                }
            )

    # -- public ------------------------------------------------------------ #

    def run(self) -> dict:
        self.audit_markers()
        self.check_marker_fixtures()
        self.check_semantics()
        self.check_governing_clause()
        self.check_references()
        self.check_consequences()
        self.check_sheet_columns()
        self.check_dashboard()
        self.check_coverage()
        self.check_word_artifacts()
        self.audit_rows()
        contract = self.check_contract()
        failed = [check.name for check in self.checks if not check.ok]
        failed_fixtures = [
            key for key, value in self.fixtures.items() if value["status"] != "PASS"
        ]
        verdict = "PASS" if not failed and not failed_fixtures else "FAIL"
        return {
            "schema": "v1_review_workbook_round6/1",
            "case": self.case,
            "build_id": self.build.name,
            "workbook": str(self.build / "投标项目复核表.xlsx"),
            "verdict": verdict,
            "check_count": len(self.checks),
            "failed_checks": failed,
            "checks": [check.as_dict() for check in self.checks],
            "fixtures": self.fixtures,
            "fixture_summary": {
                "total": len(self.fixtures),
                "passed": len(self.fixtures) - len(failed_fixtures),
                "failed": failed_fixtures,
            },
            "concern_contract_verdict": (
                "PASS"
                if not [check for check in contract["checks"] if not check["ok"]]
                and not contract.get("fixtures_failed")
                else "FAIL"
            ),
            "concern_contract": {
                "contract_source": contract.get("contract_source"),
                "contract_version": contract.get("contract_version"),
                "contract_count": contract.get("contract_count"),
                "review_row_count": contract.get("review_row_count"),
                "fixtures_passed": contract.get("fixtures_passed"),
                "failed_checks": [
                    check["check"] for check in contract["checks"] if not check["ok"]
                ],
            },
            "criticality_model": {
                "marker_count": len(self.criticality.markers),
                "marker_vocabulary": sorted(
                    {record.raw_marker for record in self.criticality.markers}
                ),
                "substantive_rules": [
                    rule.as_dict() for rule in self.criticality.substantive_rules
                ],
                "rejection_rule_count": len(self.criticality.rejection_rules),
                "reference_links": [
                    {
                        "parent_atom_id": link.parent_atom_id,
                        "parent_clause": link.parent_clause,
                        "target": link.target,
                        "target_kind": link.target_kind,
                        "child_count": len(link.target_atom_ids),
                    }
                    for link in self.criticality.reference_links
                ],
            },
            "counts": {
                "source_marker_atom_count": sum(
                    1 for record in self.criticality.criticalities.values() if record.marker_present
                ),
                "source_marked_row_count": sum(1 for item in self.items if item.marker_present),
                "substantive_requirement_row_count": sum(
                    1 for item in self.items if item.substantive_requirement
                ),
                "rejection_consequence_row_count": sum(
                    1 for item in self.items if item.rejection_consequence
                ),
                "explicit_rejection_row_count": sum(
                    1 for item in self.items if item.rejection_kind == "EXPLICIT"
                ),
                "derived_rejection_row_count": sum(
                    1 for item in self.items if item.rejection_kind == "DERIVED"
                ),
                "marker_lost_count": self.marker_lost_count(),
                "marker_false_positive_count": self.marker_false_positive_count,
                "source_substantive_requirement_without_review_coverage_count": (
                    self.uncovered_substantive_count
                ),
            },
            "marker_records": self.marker_records,
            "row_audit": self.row_records,
            "coverage_records": self.coverage_records,
        }


def write_case_report(case: str, *, case_dir: Path | None = None) -> dict:
    report = Round6Report(case, case_dir=case_dir)
    data = report.run()
    REPORTS.mkdir(parents=True, exist_ok=True)
    path = REPORTS / f"review_workbook_round6_criticality_audit_{case}.json"
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"{case}: {data['verdict']} ({len(data['failed_checks'])} failed check(s)) -> {path.name}")
    for check in data["checks"]:
        if not check["ok"]:
            print(f"   FAIL {check['check']}: {check['detail']}")
    for key, value in data["fixtures"].items():
        if value["status"] != "PASS":
            print(f"   FIXTURE {key}: {value['status']}")
    return data


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--case", action="append", default=None)
    parser.add_argument("--three-case", action="store_true")
    args = parser.parse_args(argv)
    cases = list(CASES) if args.three_case or not args.case else args.case
    results = {case: write_case_report(case) for case in cases}
    if len(results) > 1:
        REPORTS.mkdir(parents=True, exist_ok=True)
        summary = {
            "schema": "v1_review_workbook_round6_generalization/1",
            "cases": {
                case: {
                    "verdict": data["verdict"],
                    "failed_checks": data["failed_checks"],
                    "failed_fixtures": data["fixture_summary"]["failed"],
                    "marker_vocabulary": data["criticality_model"]["marker_vocabulary"],
                    "counts": data["counts"],
                }
                for case, data in results.items()
            },
            "verdict": (
                "PASS" if all(data["verdict"] == "PASS" for data in results.values()) else "FAIL"
            ),
        }
        path = REPORTS / "review_workbook_round6_generalization.json"
        path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"generalization: {summary['verdict']} -> {path.name}")
    return 0 if all(data["verdict"] == "PASS" for data in results.values()) else 1


if __name__ == "__main__":
    raise SystemExit(main())
