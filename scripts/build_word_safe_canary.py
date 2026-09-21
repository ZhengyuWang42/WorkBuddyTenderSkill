"""Small manual acceptance target exercising the production API subset."""
import argparse
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from docx.shared import Pt
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from tender_basic.word_safe_source_builder import new_word_safe_document, add_safe_run
from tender_basic.word_safe_scan import scan_word_safe_docx

def build(output):
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = new_word_safe_document()
    doc.sections[0].page_width, doc.sections[0].page_height = Pt(595), Pt(842)
    for section in doc.sections:
        section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Pt(36)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_safe_run(p, 'Word 安全构造验收').bold = True
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Pt(21)
    add_safe_run(p, '中文字体、换行测试\n第二行。').italic = True
    add_safe_run(p, '下划线填值', underline=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_safe_run(p, '右对齐')
    for columns in (2, 4):
        table = doc.add_table(4, columns)
        table.style = 'Table Grid'
        table.autofit = False
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for col in table.columns:
            col.width = Pt(480/columns)
        for row in table.rows:
            row.height, row.height_rule = Pt(24), WD_ROW_HEIGHT_RULE.AT_LEAST
            for cell in row.cells:
                cell.width = Pt(480/columns)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        table.cell(0,0).merge(table.cell(0,columns-1))
        table.cell(1,0).merge(table.cell(2,0))
        add_safe_run(table.cell(0,0).paragraphs[0], '可编辑合并表格')
        add_safe_run(table.cell(1,0).paragraphs[0], '纵向合并')
        add_safe_run(table.cell(3,1).paragraphs[0], '测试单元格')
    doc.add_page_break()
    add_safe_run(doc.add_paragraph(), '第二页：公开 API 分页。')
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.page_width, section.page_height = Pt(842), Pt(595)
    add_safe_run(doc.add_paragraph(), '第三页：公开 API 横向分节。')
    doc.save(output)
    report = scan_word_safe_docx(output)
    if report['result'] != 'PASS':
        raise ValueError(report)
    return output

if __name__ == '__main__':
    p = argparse.ArgumentParser()
    p.add_argument('--output', type=Path, default=Path('tmp/word-safe-canary.docx'))
    print(build(p.parse_args().output))
