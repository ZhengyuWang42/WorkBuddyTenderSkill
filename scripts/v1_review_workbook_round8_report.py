"""Round 8: delivered-text fidelity, conditional scoring, cross-sheet risk.

The human's round-7 manual review examined the *actual* CASE001 workbook and
failed it.  The round-7 machine gates had all passed, because a provenance ID and
a locator can both be right while the delivered sentence says the wrong thing.
Round 8 therefore audits the **final delivered text**:

* polarity -- a negated operative clause keeps its negation;
* conditional scoring -- tiers are alternatives, a base score is not a maximum;
* blank placeholders -- a source blank stays a blank;
* complete fragments -- no truncated enumeration, duplicated wording or foreign
  heading;
* cross-sheet risk -- the legacy 风险级别 cannot contradict the reviewer sheet's
  own 否决依据;
* source-form classification -- a history table is not a blank quotation form.

Usage::

    python scripts/v1_review_workbook_round8_report.py --three-case
    python scripts/v1_review_workbook_round8_report.py --case case_001
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Mapping, Sequence

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
if str(ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(ROOT / "scripts"))

from openpyxl import load_workbook  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from tender_basic.fidelity_invariants import (  # noqa: E402
    RISK_LEVELS,
    RISK_VETO,
    check_absorbed_headings,
    check_blank_placeholders,
    check_complete_fragments,
    check_conditional_scoring,
    check_cross_sheet_risk,
    check_duplicated_wording,
    check_form_classification,
    check_locator_headings,
    check_polarity,
    criticality_vocabulary_ok,
)
from tender_basic.review_workbook_views import (  # noqa: E402
    SHEET_TITLES,
    _price_table_rows,
)
from v1_review_workbook_round4_report import Round4Report  # noqa: E402
from v1_review_workbook_round6_report import CASES, Round6Report  # noqa: E402
from v1_review_workbook_round7_report import (  # noqa: E402
    BUILDS as ROUND7_BUILDS,
    Round7Report,
)

MANDATORY_SHEET_TITLE = "03_资格否决与强制项"
CLAUSE_SHEET_TITLE = "02_关键条款"
PRICE_SHEET_TITLE = "04_报价与限价"
LEGACY_SHEET_TITLE = "投标项目复核表"

DELIVERED_SHEET_TITLES = (CLAUSE_SHEET_TITLE, MANDATORY_SHEET_TITLE)
REQUIREMENT_COLUMN = {CLAUSE_SHEET_TITLE: "抽取结果", MANDATORY_SHEET_TITLE: "要求正文"}
LOCATOR_COLUMN = {CLAUSE_SHEET_TITLE: "证据定位", MANDATORY_SHEET_TITLE: "证据定位"}
EXCERPT_COLUMN = {CLAUSE_SHEET_TITLE: "证据摘要", MANDATORY_SHEET_TITLE: "证据摘要"}
PAGE_COLUMN = {CLAUSE_SHEET_TITLE: "源页码", MANDATORY_SHEET_TITLE: "证据页码"}

_LEGACY_LOCATOR_RE = re.compile(r"^(.*?（pdf_(?:block|table_cell)）)")
_LEGACY_REQUIREMENT_RE = re.compile(r"招标文件要求：(.*?)(?:\n复核要点：|\Z)", re.S)

REPORTS = ROOT / "acceptance/reports/v1_generalization"

#: round-8 successor build per case
BUILDS: dict[str, str] = {
    "case_001": "v1_manual_fidelity_round4_date_rhythm_closure8_review_workbook8",
    "case_002": "v1_round4_closure8_review_workbook8",
    "case_003": "v1_round4_closure8_review_workbook8",
}


def _clean(value: object) -> str:
    return re.sub(r"[\s\u3000]+", "", str(value or ""))


def _legacy_block(cell_text: str, label: str) -> str:
    """One labelled block of a legacy template cell ("复核要点：…")."""

    text = str(cell_text or "")
    start = text.find(f"{label}：")
    if start < 0:
        return ""
    body = text[start + len(label) + 1 :]
    for following in ("复核要点：", "通过标准：", "不满足后果：", "准备材料：", "评分提示："):
        if following == f"{label}：":
            continue
        end = body.find(following)
        if end >= 0:
            body = body[:end]
    return body.strip()


def _sha256(path: Path) -> str:
    import hashlib

    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


@dataclass
class Check:
    name: str
    ok: bool
    detail: str
    evidence: dict[str, Any] = field(default_factory=dict)

    def as_dict(self) -> dict[str, Any]:
        return {
            "check": self.name,
            "ok": bool(self.ok),
            "detail": self.detail,
            "evidence": self.evidence,
        }


class Round8Report:
    """Audit one case's round-8 successor workbook from its final cells."""

    def __init__(
        self, case: str, *, case_dir: Path | None = None, build_name: str | None = None
    ) -> None:
        self.case = case
        self.case_dir = case_dir or (ROOT / "acceptance/workspace" / case)
        # a later round passes its own successor so the same audit runs against the
        # newer build without rewriting the round-8 evidence
        self.build = self.case_dir / (build_name or BUILDS[case])
        self.workbook = self.build / "投标项目复核表.xlsx"
        # the round-7 report reads the *frozen* round-7 build for its plan; the
        # round-8 audit must read the delivered round-8 cells, so its own round-4
        # projection is bound to the successor build
        self.r7 = Round7Report(case, case_dir=self.case_dir)
        self.r4 = Round4Report(self.build, case)
        self.document = self.r7.document
        self.facts = self.r7.facts
        self.items = list(self.r7.items)
        self.background = list(self.r7.background)
        self.by_item = dict(self.r7.by_item)
        self.checks: list[Check] = []
        self.fixtures: dict[str, dict[str, Any]] = {}
        self.counts: dict[str, int] = {}

    # -- cells ------------------------------------------------------------ #

    def _rows(self, title: str) -> list[dict[str, Any]]:
        workbook = load_workbook(self.workbook, data_only=True, read_only=True)
        try:
            worksheet = workbook[title]
            headers = [
                str(cell.value or "").strip()
                for cell in next(worksheet.iter_rows(min_row=1, max_row=1))
            ]
            out: list[dict[str, Any]] = []
            for index, row in enumerate(
                worksheet.iter_rows(min_row=2, values_only=True), start=2
            ):
                if not any(value not in (None, "") for value in row):
                    continue
                out.append(
                    {
                        headers[i]: row[i] for i in range(min(len(headers), len(row)))
                    }
                    | {"__row__": index}
                )
            return out
        finally:
            workbook.close()

    def _address(self, title: str, column: str, row: Mapping[str, Any]) -> str:
        headers = list(row)
        position = headers.index(column) + 1 if column in headers else 0
        letter = get_column_letter(position) if position else "?"
        return f"{title}!{letter}{row.get('__row__')}"

    def reviewer_rows(self) -> list[dict[str, Any]]:
        out: list[dict[str, Any]] = []
        for title in DELIVERED_SHEET_TITLES:
            for row in self._rows(title):
                item_id = str(row.get("requirement_id") or "").strip()
                if not item_id:
                    continue
                out.append(
                    {
                        "sheet": title,
                        "item_id": item_id,
                        "item": self.by_item.get(item_id),
                        "requirement": str(row.get(REQUIREMENT_COLUMN[title]) or ""),
                        "action": str(row.get("复核动作") or ""),
                        "criteria": str(row.get("核验标准") or ""),
                        "marker": str(row.get("★") or ""),
                        "veto": str(row.get("否决性") or ""),
                        "mandatory_type": str(row.get("强制性类型") or ""),
                        "page": row.get(PAGE_COLUMN[title]),
                        "locator": str(row.get(LOCATOR_COLUMN[title]) or ""),
                        "excerpt": str(row.get(EXCERPT_COLUMN[title]) or ""),
                        "cell": self._address(
                            title, REQUIREMENT_COLUMN[title], row
                        ),
                    }
                )
        return out

    def legacy_rows(self) -> list[dict[str, Any]]:
        out: list[dict[str, Any]] = []
        for item in self.items:
            row = self.r4.rows_by_item.get(str(item.item_id))
            if row is None:
                continue
            requirement_cell = str(getattr(row, "d_text", "") or "")
            requirement = requirement_cell
            found = _LEGACY_REQUIREMENT_RE.search(requirement_cell)
            if found:
                requirement = found.group(1).strip()
            evidence_cell = str(getattr(row, "e_text", "") or "")
            locator = evidence_cell
            found = _LEGACY_LOCATOR_RE.match(evidence_cell)
            if found:
                locator = found.group(1).strip()
            # the delivered sheet's 风险级别 is column C: it is read from the
            # FINAL cells, never from a plan attribute, so a cross-sheet audit
            # compares two delivered values (round-8 fixtures DR013/DR038)
            risk = self.r4.cells.value(LEGACY_SHEET_TITLE, int(row.row), 3)
            out.append(
                {
                    "sheet": LEGACY_SHEET_TITLE,
                    "item_id": str(item.item_id),
                    "item": item,
                    "requirement": requirement,
                    "action": "",
                    "criteria": "",
                    "risk": risk,
                    "risk_cell": self.r4.cells.address(
                        LEGACY_SHEET_TITLE, int(row.row), 3
                    ),
                    "page": item.source_page,
                    "locator": locator,
                    "excerpt": "",
                    "cell_text": requirement_cell,
                    "cell": str(getattr(row, "d_address", "") or ""),
                }
            )
        return out

    def all_rows(self) -> list[dict[str, Any]]:
        return [*self.reviewer_rows(), *self.legacy_rows()]

    # -- checks ----------------------------------------------------------- #

    def check(self, name: str, ok: bool, detail: str, evidence: dict | None = None) -> None:
        self.checks.append(Check(name, bool(ok), detail, evidence or {}))

    def fixture(self, key: str, ok: bool, detail: str, evidence: str = "") -> None:
        self.fixtures[key] = {
            "status": "PASS" if ok else "FAIL",
            "detail": detail,
            "evidence": evidence[:400],
        }

    def audit(self) -> dict[str, Any]:
        rows = self.all_rows()
        reviewer = self.reviewer_rows()
        price_result = _price_table_rows(self.document)

        polarity = check_polarity(rows)
        self.counts["POLARITY_LOSS_COUNT"] = polarity["polarity_loss_count"]
        self.check(
            "no negated operative clause loses its negation",
            polarity["polarity_loss_count"] == 0,
            f"{polarity['checked']} negated clause(s) checked, "
            f"{polarity['polarity_loss_count']} polarity loss(es)",
            {"losses": polarity["polarity_losses"][:8]},
        )

        scoring = check_conditional_scoring(rows)
        self.counts["CONDITIONAL_SCORING_PROBLEM_COUNT"] = scoring[
            "conditional_scoring_problem_count"
        ]
        self.check(
            "scoring tiers are conditional alternatives and base/ceiling stay distinct",
            scoring["conditional_scoring_problem_count"] == 0,
            f"{scoring['tier_rows']} tier row(s), {scoring['base_score_rows']} base-score row(s), "
            f"{scoring['conditional_scoring_problem_count']} problem(s)",
            {"problems": scoring["conditional_scoring_problems"][:8]},
        )

        blanks = check_blank_placeholders(rows)
        self.counts["BLANK_PLACEHOLDER_PROBLEM_COUNT"] = blanks[
            "blank_placeholder_problem_count"
        ]
        self.check(
            "a source blank is delivered as a blank",
            blanks["blank_placeholder_problem_count"] == 0,
            f"{blanks['blank_marked_rows']} row(s) carry an explicit blank, "
            f"{blanks['blank_placeholder_problem_count']} problem(s)",
            {"problems": blanks["blank_placeholder_problems"][:8]},
        )

        fragments = check_complete_fragments(rows)
        self.counts["FRAGMENT_PROBLEM_COUNT"] = fragments["fragment_problem_count"]
        self.check(
            "no enumeration is truncated and no clause is rendered twice",
            fragments["fragment_problem_count"] == 0,
            f"{fragments['fragment_problem_count']} fragment problem(s)",
            {"problems": fragments["fragment_problems"][:8]},
        )

        headings = check_absorbed_headings(rows)
        self.counts["ABSORBED_HEADING_COUNT"] = headings["absorbed_heading_count"]
        self.check(
            "no requirement absorbed the next source heading",
            headings["absorbed_heading_count"] == 0,
            f"{headings['absorbed_heading_count']} absorbed heading(s)",
            {"absorbed": headings["absorbed_headings"][:8]},
        )

        locators = check_locator_headings(rows)
        self.counts["FOREIGN_LOCATOR_HEADING_COUNT"] = locators["foreign_locator_heading_count"]
        self.check(
            "no locator names a heading that governs a different clause",
            locators["foreign_locator_heading_count"] == 0,
            f"{locators['locator_heading_checked']} locator(s) checked, "
            f"{locators['foreign_locator_heading_count']} foreign heading(s)",
            {"problems": locators["foreign_locator_headings"][:8]},
        )

        duplicated = check_duplicated_wording(rows)
        self.counts["DUPLICATED_WORDING_COUNT"] = duplicated["duplicated_wording_count"]
        self.check(
            "no delivered requirement repeats a source phrase",
            duplicated["duplicated_wording_count"] == 0,
            f"{duplicated['duplicated_wording_count']} duplicated phrase(s)",
            {"problems": duplicated["duplicated_wording"][:8]},
        )

        risk = check_cross_sheet_risk(self._rows(MANDATORY_SHEET_TITLE), self.legacy_rows())
        self.counts["CROSS_SHEET_RISK_MISMATCH_COUNT"] = risk["cross_sheet_risk_mismatch_count"]
        self.check(
            "the legacy 风险级别 agrees with the reviewer sheet's 否决依据",
            risk["cross_sheet_risk_mismatch_count"] == 0,
            f"{risk['cross_sheet_risk_checked']} row(s) compared, "
            f"{risk['cross_sheet_risk_mismatch_count']} contradiction(s)",
            {"mismatches": risk["cross_sheet_risk_mismatches"][:8]},
        )

        forms = check_form_classification(price_result["blank_forms"])
        self.counts["FORM_CLASSIFICATION_PROBLEM_COUNT"] = forms[
            "form_classification_problem_count"
        ]
        self.check(
            "blank source forms are classified by heading and column schema",
            forms["form_classification_problem_count"] == 0,
            f"{forms['blank_form_checked']} blank form(s) checked, "
            f"{forms['form_classification_problem_count']} misclassified",
            {"problems": forms["form_classification_problems"][:8]},
        )

        self.check(
            "the criticality vocabulary matches the checks",
            criticality_vocabulary_ok(),
            "risk / criticality vocabularies are the ones the engine emits",
        )

        self.audit_fixtures(self.fixture_rows(), price_result)
        self.check_round_regressions()
        self.check_word_artifacts()

        failed = [check.name for check in self.checks if not check.ok]
        failed_fixtures = [key for key, value in self.fixtures.items() if value["status"] != "PASS"]
        verdict = "PASS" if not failed and not failed_fixtures else "FAIL"
        return {
            "schema": "v1_review_workbook_round8/1",
            "case": self.case,
            "build_id": self.build.name,
            "workbook": str(self.workbook.relative_to(ROOT).as_posix()),
            "workbook_sha256": _sha256(self.workbook),
            "verdict": verdict,
            "check_count": len(self.checks),
            "failed_checks": failed,
            "checks": [check.as_dict() for check in self.checks],
            "fixtures": self.fixtures,
            "fixture_summary": {
                "total": len(self.fixtures),
                "passed": len(self.fixtures) - len(failed_fixtures),
                "failed": failed_fixtures,
            },
            "counts": dict(self.counts),
            "delivered_row_count": len(self.items),
        }

    # -- the human's own fixtures ----------------------------------------- #

    def _row_for(self, rows: Sequence[Mapping[str, Any]], item_id: str) -> Mapping[str, Any] | None:
        """The fullest delivered row for ``item_id``.

        A row that is neither mandatory nor a key clause lives only on the legacy
        template, and the legacy cell prints the requirement and the action in the
        same cell.  The audit therefore prefers a reviewer-sheet row when one
        exists (it is the structured delivery) and otherwise takes the template
        row, whose requirement is recovered from the cell's own block labels.
        """

        candidates = [row for row in rows if str(row.get("item_id")) == item_id]
        if not candidates:
            return None
        for row in candidates:
            if str(row.get("sheet")) in DELIVERED_SHEET_TITLES:
                return row
        return candidates[0]

    def fixture_rows(self) -> list[dict[str, Any]]:
        """One row per delivered item, with the legacy cell's blocks recovered.

        The template prints "招标文件要求：…\\n复核要点：…\\n不满足后果：…" in one
        cell; the fixture checks need the requirement and the action separately,
        whichever sheet the item is delivered on.
        """

        reviewer = self.reviewer_rows()
        legacy = self.legacy_rows()
        out: list[dict[str, Any]] = []
        seen: set[str] = set()
        for row in [*reviewer, *legacy]:
            item_id = str(row.get("item_id") or "")
            if not item_id or item_id in seen:
                continue
            seen.add(item_id)
            chosen = self._row_for([*reviewer, *legacy], item_id)
            if chosen is None:
                continue
            merged = dict(chosen)
            if str(chosen.get("sheet")) == LEGACY_SHEET_TITLE:
                template = next(
                    (candidate for candidate in legacy if str(candidate.get("item_id")) == item_id),
                    None,
                )
                cell = str(template.get("cell_text") or "") if template else ""
                merged["action"] = merged.get("action") or _legacy_block(cell, "复核要点")
                merged["criteria"] = merged.get("criteria") or _legacy_block(cell, "核验标准")
                merged["consequence"] = _legacy_block(cell, "不满足后果")
            out.append(merged)
        return out

    def audit_fixtures(
        self, rows: Sequence[Mapping[str, Any]], price_result: Mapping[str, Any]
    ) -> None:
        """Every fixture the human named, checked independently by content.

        The fixtures are CASE001's own review findings; the checks above already
        run generically on all three cases.  These bind each finding to the row
        it was reported on, so a regression cannot hide behind an aggregate.
        """

        if self.case != "case_001":
            return
        # (key, item id, detail, tokens the row's *displayed requirement* must
        #  carry, tokens the row's own locator must carry, tokens the row's
        #  action/criteria must carry, tokens that must not appear anywhere in
        #  the row's delivered text)
        #
        # The negated-condition fixtures require the *whole* operative clause,
        # not just its first half: the round-7 defect was that "供应商不按本章第
        # 3.4.1项要求提交响应保证金的" was delivered without the rejection it
        # governs.  Its negation and its cross-reference were both present, so a
        # `must_not` on "提交响应保证金的" (which is legitimately part of the
        # correct sentence) asserted nothing and the fixture passed the broken
        # render.  Requiring the consequence verb is the polarity requirement.
        fixture_specs: tuple[
            tuple[
                str,
                str,
                str,
                tuple[str, ...],
                tuple[str, ...],
                tuple[str, ...],
                tuple[str, ...],
            ],
            ...,
        ] = (
            ("D14", "DR036", "the bid-bond rejection keeps its negated condition",
             ("不按", "3.4.1", "否决"), (), (), ()),
            ("D37", "DR037", "the performance-bond forfeiture keeps its negated condition",
             ("不能按", "7.3.1", "放弃成交"), (), (), ()),
            ("D23", "DR041", "bank-acceptance tiers are conditional alternatives",
             ("100%", "50%"), (), ("其中一档",), ("本项最高 4分", "本项最高 2分")),
            ("D50", "DR040", "the formula's base score is not a second maximum",
             ("基本分 30",), (), (), ("本项最高 30分",)),
            ("D39", "DR044", "the contract's blank grace period stays a blank",
             ("＿＿",), (), (), ()),
            ("D45", "DR046", "the technical-standard list is not visibly truncated",
             ("（11）",), (), (), ("…",)),
            ("D46", "DR047", "the starred test duty does not absorb the next heading",
             ("第三方检测",), (), (), ("第四条",)),
            ("D16", "DR005", "the licence wording is not duplicated",
             ("营业执照",), (), (), ("执照 执照", "供应商名称 供应商名称")),
            ("D34", "DR017", "the validity row's locator names its own clause",
             ("3.3.1",), ("3.3.1",), (), ("1.12 分包",)),
        )
        for key, item_id, detail, must, must_locator, must_action, must_not in fixture_specs:
            row = self._row_for(rows, item_id)
            if row is None:
                self.fixture(key, False, f"{detail} (row {item_id} is not delivered)")
                continue
            requirement = str(row.get("requirement") or "")
            locator = str(row.get("locator") or "")
            action = " ".join(
                str(row.get(field) or "") for field in ("action", "criteria")
            )
            ok = all(token in requirement for token in must)
            if ok and must_locator:
                ok = all(token in locator for token in must_locator)
            if ok and must_action:
                ok = all(token in action for token in must_action)
            if ok and must_not:
                whole = " ".join(
                    str(row.get(field) or "")
                    for field in ("requirement", "action", "criteria", "locator")
                )
                ok = not any(token in whole for token in must_not)
            self.fixture(
                key,
                ok,
                detail,
                f"build={self.build.name} | row={item_id} | sheet={row.get('sheet')} "
                f"| cell={row.get('cell')} | requirement={requirement[:150]} "
                f"| locator={locator[:90]}",
            )

        # D56: the action must name the decision's own subject, not another
        # question's subject
        row = self._row_for(rows, "DR016")
        if row is None:
            self.fixture("D56", False, "the no-pre-bid-meeting action names its own subject")
        else:
            action = str(row.get("action") or "")
            requirement = str(row.get("requirement") or "")
            ok = "采购预备会" in action and "不召开" in action
            if "提问与澄清截止" in action:
                ok = False
            self.fixture(
                "D56",
                ok,
                "the no-pre-bid-meeting row's action names its own decision",
                f"build={self.build.name} | row=DR016 | sheet={row.get('sheet')} "
                f"| cell={row.get('cell')} | action={action[:150]} "
                f"| requirement={requirement[:80]}",
            )

        # DR013 / DR038: the delivered risk must not claim a veto the row's own
        # 否决依据 does not support.  Both values are read from the final cells:
        # the legacy sheet's 风险级别 (column C) and the 03 sheet's 否决性.
        mandatory = {str(r.get("requirement_id")): r for r in self._rows(MANDATORY_SHEET_TITLE)}
        for key, item_id in (("DR013", "DR013"), ("DR038", "DR038")):
            row = self._row_for(self.legacy_rows(), item_id)
            reviewer = mandatory.get(item_id)
            if row is None or reviewer is None:
                self.fixture(key, False, f"{item_id} is not delivered on both sheets")
                continue
            risk = str(row.get("risk") or "")
            veto = str(reviewer.get("否决性") or "").strip()
            ok = risk in RISK_LEVELS and not (risk == RISK_VETO and veto != "是")
            self.fixture(
                key,
                ok,
                "the legacy 一票否决 does not contradict the row's own 否决依据",
                f"build={self.build.name} | row={item_id} | risk_cell={row.get('risk_cell')} "
                f"| risk={risk!r} | 否决性={veto!r} "
                f"| 强制性类型={reviewer.get('强制性类型')!r}",
            )

        # 04 sheet: the page-53 similar-project-history form is not a blank
        # quotation form
        forms = list(price_result["blank_forms"])
        history = [
            form
            for form in forms
            if "类似项目" in str(form.get("heading") or "")
            or "项目名称" in _clean(form.get("columns"))
        ]
        if not history:
            self.fixture(
                "FORM_P53",
                True,
                "no similar-project-history blank form in this case",
                "",
            )
        else:
            ok = all(
                "空白报价表单" not in str(form.get("form_class") or "") for form in history
            )
            self.fixture(
                "FORM_P53",
                ok,
                "the similar-project-history form is not classified as a blank quotation form",
                " | ".join(
                    f"p{form.get('page')}: {form.get('form_class')}" for form in history
                )[:300],
            )

    # -- banked rounds ---------------------------------------------------- #

    def check_round_regressions(self) -> None:
        # the banked round-5/round-6 gates must run against the *round-8
        # successor*: pointing them at the frozen round-7 build re-audited the
        # preserved evidence instead of the current delivery, which is how a
        # stale "37/49 rows audited" result was mistaken for a live regression.
        data = Round6Report(
            self.case, case_dir=self.case_dir, build_name=self.build.name
        ).run()
        self.check(
            "the banked round-6 marker accounting still holds",
            data["verdict"] == "PASS",
            f"round-6 accounting verdict={data['verdict']} "
            f"(failed={','.join(data['failed_checks']) or 'none'})",
            {
                "failed_checks": data["failed_checks"],
                "failed_fixtures": data["fixture_summary"]["failed"],
            },
        )

    def check_word_artifacts(self) -> None:
        manifest = json.loads((self.build / "build_manifest.json").read_text(encoding="utf-8"))
        identity = manifest.get("artifact_identity", {})
        ok = bool(identity) and all(entry.get("byte_identical") for entry in identity.values())
        self.check(
            "Word artifacts are byte-identical (not re-rendered)",
            ok,
            f"{len(identity)} artifact(s) compared against the accepted source build",
            {
                name: entry.get("successor_sha256")
                for name, entry in list(identity.items())[:3]
            },
        )


def markdown(data: dict) -> str:
    lines = [
        f"# Round 8 delivered-text audit — {data['case']}",
        "",
        f"- build: `{data['build_id']}`",
        f"- workbook: `{data['workbook']}`",
        f"- sha256: `{data['workbook_sha256']}`",
        f"- verdict: **{data['verdict']}**",
        f"- delivered rows audited: {data['delivered_row_count']}",
        "",
        "## Zero-count gates",
        "",
        "| gate | value |",
        "| --- | --- |",
    ]
    for key in sorted(data["counts"]):
        lines.append(f"| {key} | {data['counts'][key]} |")
    lines += ["", "## Checks", "", "| check | result | detail |", "| --- | --- | --- |"]
    for check in data["checks"]:
        lines.append(
            f"| {check['check']} | {'PASS' if check['ok'] else 'FAIL'} | {check['detail']} |"
        )
    lines += ["", "## Fixtures", "", "| fixture | result | detail |", "| --- | --- | --- |"]
    for key in sorted(data["fixtures"]):
        fixture = data["fixtures"][key]
        lines.append(f"| {key} | {fixture['status']} | {fixture['detail']} |")
    return "\n".join(lines) + "\n"


def write_case_report(case: str, *, case_dir: Path | None = None) -> dict[str, Any]:
    report = Round8Report(case, case_dir=case_dir)
    data = report.audit()
    REPORTS.mkdir(parents=True, exist_ok=True)
    stem = f"review_workbook_round8_{case}"
    (REPORTS / f"{stem}.json").write_text(
        json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    (REPORTS / f"{stem}.md").write_text(markdown(data), encoding="utf-8")
    print(f"{case}: {data['verdict']} ({len(data['failed_checks'])} failed check(s)) -> {stem}.json")
    for check in data["checks"]:
        if not check["ok"]:
            print(f"   FAIL {check['check']}: {check['detail']}")
    for key, value in data["fixtures"].items():
        if value["status"] != "PASS":
            print(f"   FIXTURE {key}: {value['status']} — {value['detail']}")
    return data


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--case", action="append", default=None)
    parser.add_argument("--three-case", action="store_true")
    args = parser.parse_args(argv)
    cases = list(BUILDS) if args.three_case or not args.case else args.case
    results = {case: write_case_report(case) for case in cases}
    if len(results) > 1:
        summary = {
            "schema": "v1_review_workbook_round8_generalization/1",
            "cases": {
                case: {
                    "verdict": data["verdict"],
                    "failed_checks": data["failed_checks"],
                    "failed_fixtures": data["fixture_summary"]["failed"],
                    "counts": data["counts"],
                }
                for case, data in results.items()
            },
            "verdict": (
                "PASS" if all(data["verdict"] == "PASS" for data in results.values()) else "FAIL"
            ),
        }
        path = REPORTS / "review_workbook_round8_generalization.json"
        path.write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(f"generalization: {summary['verdict']} -> {path.name}")
    return 0 if all(data["verdict"] == "PASS" for data in results.values()) else 1


__all__ = ["BUILDS", "Round8Report", "main", "markdown", "write_case_report"]


if __name__ == "__main__":
    raise SystemExit(main())
