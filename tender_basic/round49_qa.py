"""Round 4.9 alignment and destination-style QA."""

from __future__ import annotations

import re
from collections import Counter
from pathlib import Path

import pymupdf
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from .source_font_policy import normalize_pdf_font_name
from .source_format import SourceFormatTemplate, TypographyRole


def _run_style(run) -> dict[str, object]:
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    return {
        "east_asia_font": fonts.get(qn("w:eastAsia")),
        "latin_font": fonts.get(qn("w:ascii")),
        "font_size": run.font.size.pt if run.font.size is not None else None,
        "bold": bool(run.bold),
        "italic": bool(run.italic),
        "underline": bool(run.underline),
        "font_weight_role": "bold" if run.bold else "regular",
    }


def _alignment(paragraph) -> str:
    return {
        WD_ALIGN_PARAGRAPH.CENTER: "center",
        WD_ALIGN_PARAGRAPH.RIGHT: "right",
        WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
        WD_ALIGN_PARAGRAPH.LEFT: "left",
    }.get(paragraph.alignment, "left")


def build_filled_slot_style_audit(generation_report: dict) -> dict:
    rows = []
    family = size = weight = alignment = 0
    for item in generation_report.get("filled_slots", []):
        expected = item.get("destination_expected_style", {})
        actual = item.get("generated_value_style", {})
        family_ok = expected.get("east_asia_font") == actual.get("east_asia_font")
        size_ok = expected.get("font_size") is not None and actual.get("font_size") is not None and abs(
            float(expected["font_size"]) - float(actual["font_size"])
        ) <= 0.01
        weight_ok = expected.get("font_weight_role") == actual.get("font_weight_role")
        alignment_ok = expected.get("paragraph_alignment") == actual.get("paragraph_alignment")
        family += not family_ok
        size += not size_ok
        weight += not weight_ok
        alignment += not alignment_ok
        rows.append({
            "field": item.get("field", []), "page": item.get("source_page"),
            "slot_role": item.get("slot_role"), "slot_id": item.get("slot_id"),
            "value": item.get("value"),
            "destination_expected_font": expected.get("east_asia_font"),
            "actual_font": actual.get("east_asia_font"),
            "expected_size": expected.get("font_size"), "actual_size": actual.get("font_size"),
            "expected_weight": expected.get("font_weight_role"),
            "actual_weight": actual.get("font_weight_role"),
            "expected_alignment": expected.get("paragraph_alignment"),
            "alignment": actual.get("paragraph_alignment"),
            "anchor_source": expected.get("anchor_source"),
            "style_match": family_ok and size_ok and weight_ok and alignment_ok,
        })
    return {
        "total": len(rows),
        "filled_slot_font_family_mismatch": int(family),
        "filled_slot_font_size_mismatch": int(size),
        "filled_slot_weight_mismatch": int(weight),
        "filled_slot_alignment_mismatch": int(alignment),
        "filled_slots": rows,
        "result": "PASS" if not (family or size or weight or alignment) else "FAIL",
    }


def build_project_number_alignment_qa(
    docx_path: Path, rendered_pdf: Path, generation_report: dict, *, source_page: int
) -> dict:
    logical = next(item for item in generation_report.get("logical_paragraph_records", [])
                   if item.get("source_page") == source_page and "项目编号" in item.get("source_text", ""))
    document = Document(docx_path)
    paragraph = next(item for item in document.paragraphs if "项目编号" in item.text)
    with pymupdf.open(rendered_pdf) as pdf:
        page = pdf[0]
        rectangles = page.search_for("项目编号")
        if not rectangles:
            raise AssertionError("Rendered project-number anchor not found")
        label = rectangles[0]
        line = next((line for block in page.get_text("dict").get("blocks", [])
                     for line in block.get("lines", [])
                     if any("项目编号" in span.get("text", "") for span in line.get("spans", []))), None)
        boxes = [span["bbox"] for span in line.get("spans", []) if span.get("text", "").strip()] if line else [label]
        x0 = min(float(box[0]) for box in boxes)
        x1 = max(float(box[2]) for box in boxes)
        generated_center = (x0 + x1) / 2.0
    source_center = float(logical["source_container_center_x"])
    return {
        "source_alignment": logical["alignment_role"],
        "generated_alignment": _alignment(paragraph).upper(),
        "source_line_center_x": logical.get("source_line_center_x"),
        "source_center_x": source_center,
        "generated_center_x": generated_center,
        "center_error_pt": abs(generated_center - source_center),
        "left_indent_pt": paragraph.paragraph_format.left_indent.pt if paragraph.paragraph_format.left_indent else 0.0,
        "right_indent_pt": paragraph.paragraph_format.right_indent.pt if paragraph.paragraph_format.right_indent else 0.0,
        "first_line_indent_pt": paragraph.paragraph_format.first_line_indent.pt if paragraph.paragraph_format.first_line_indent else 0.0,
        "total_source_width": logical.get("total_source_width"),
        "slot_width": logical.get("slot_width"),
        "result": "PASS" if _alignment(paragraph) == "center" and abs(generated_center-source_center) <= 6.0
        and not any((paragraph.paragraph_format.left_indent, paragraph.paragraph_format.right_indent,
                     paragraph.paragraph_format.first_line_indent)) else "FAIL",
    }


def build_case001_table_project_name_qa(docx_path: Path) -> dict:
    document = Document(docx_path)
    for table in document.tables:
        for row in table.rows:
            for index, cell in enumerate(row.cells[:-1]):
                if "项目名称、标段" not in cell.text:
                    continue
                label_run = next(run for paragraph in cell.paragraphs for run in paragraph.runs if run.text.strip())
                value_cell = row.cells[index + 1]
                value_run = next(run for paragraph in value_cell.paragraphs for run in paragraph.runs if run.text.strip())
                label, value = _run_style(label_run), _run_style(value_run)
                match = label["east_asia_font"] == value["east_asia_font"] and abs(
                    float(label["font_size"]) - float(value["font_size"])
                ) <= .01 and label["font_weight_role"] == value["font_weight_role"]
                return {"label": label, "value": value, "style_match": match,
                        "label_text": cell.text, "value_text": value_cell.text,
                        "result": "PASS" if match else "FAIL"}
    raise AssertionError("case001 project-name table row not found")


def build_cover_typography_audit(template: SourceFormatTemplate) -> dict:
    page = template.source_pages[0]
    role_map = {
        TypographyRole.COVER_CHAPTER: "chapter",
        TypographyRole.COVER_PROJECT_TITLE: "project_title",
        TypographyRole.COVER_DOCUMENT_TITLE: "document_title",
        TypographyRole.FORM_LABEL: "project_number_label",
    }
    grouped: dict[str, list[dict[str, object]]] = {value: [] for value in role_map.values()}
    for paragraph in page.paragraphs:
        key = role_map.get(paragraph.typography_role)
        if key is None or (key == "project_number_label" and "项目编号" not in paragraph.text):
            continue
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            normalized, _ = normalize_pdf_font_name(run.font_name)
            grouped[key].append({
                "text": run.text, "source_pdf_font_name": run.raw_font_name or run.font_name,
                "normalized_font_name": normalized, "font_size": run.font_size,
                "font_flags": run.flags, "font_weight_role": run.font_weight_role,
                "bold": run.bold, "italic": run.italic,
                "bbox_height": run.bbox[3] - run.bbox[1], "bbox": list(run.bbox),
            })
    return {key: {"spans": values, "span_count": len(values)} for key, values in grouped.items()}


def build_cover_render_qa(template: SourceFormatTemplate, rendered_pdf: Path) -> dict:
    page_model = template.source_pages[0]
    role_map = {
        TypographyRole.COVER_CHAPTER: "chapter",
        TypographyRole.COVER_PROJECT_TITLE: "project_title",
        TypographyRole.COVER_DOCUMENT_TITLE: "document_title",
        TypographyRole.FORM_LABEL: "project_number_label",
    }
    result: dict[str, object] = {}
    with pymupdf.open(rendered_pdf) as pdf:
        page = pdf[0]
        generated_spans = [span for block in page.get_text("dict").get("blocks", [])
                           for line in block.get("lines", []) for span in line.get("spans", [])]
        for role, key in role_map.items():
            candidates = [paragraph for paragraph in page_model.paragraphs if paragraph.typography_role == role]
            if key == "project_number_label":
                candidates = [paragraph for paragraph in candidates if "项目编号" in paragraph.text]
            source_runs = [run for paragraph in candidates for run in paragraph.runs if run.text.strip()]
            matched = []
            for run in source_runs:
                token = run.text.strip()
                candidates = [span for span in generated_spans if token and token == span.get("text", "").strip()]
                if not candidates:
                    candidates = [span for span in generated_spans if token and token in span.get("text", "")]
                found = min(candidates, key=lambda span: abs(float(span.get("size", 0.0))-run.font_size), default=None)
                if found is not None and found not in matched:
                    matched.append(found)
            if not matched:
                result[key] = {"matched": False}
                continue
            x0=min(float(span["bbox"][0]) for span in matched)
            x1=max(float(span["bbox"][2]) for span in matched)
            sizes=[float(span.get("size",0.0)) for span in matched]
            bold=any(bool(int(span.get("flags",0)) & 16) or "bold" in span.get("font","").lower()
                     for span in matched)
            result[key]={
                "matched": True, "x_center": (x0+x1)/2.0,
                "font_size_class": {"min": min(sizes), "max": max(sizes)},
                "visual_weight_class": "bold" if bold else "regular",
                "y_anchor": min(float(span["bbox"][1]) for span in matched),
                "generated_fonts": sorted({span.get("font","") for span in matched}),
            }
    return result


__all__ = [
    "build_case001_table_project_name_qa", "build_cover_typography_audit",
    "build_cover_render_qa", "build_filled_slot_style_audit", "build_project_number_alignment_qa",
]
