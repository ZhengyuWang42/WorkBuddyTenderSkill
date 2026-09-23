"""Round 5.0 style-first Word-safe source renderer."""

from __future__ import annotations

import json
import re
from pathlib import Path
from types import SimpleNamespace

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from .page_layout import (
    AlignmentRole,
    EditableBlankRenderStyle,
    infer_semantic_line_spacing,
    classify_rule_relation,
    source_rule_geometry_intent,
    VALUE_REPLACING_POLICIES,
    RULE_TEXT_OCCUPANCY_GAP,
    RULE_TEXT_OCCUPANCY_UNDERLINE,
    FormBlock,
    LogicalParagraph,
    ParagraphLayout,
    ParagraphLayoutRole,
    bounded_scale,
    build_page_layout,
    join_visual_lines,
    single_visual_row_join,
    paragraph_coordinate_frame,
    restore_inline_rule_compositions,
    restore_inline_rule_blanks,
    source_visual_rows,
    plan_source_visual_line,
    underline_owner_runs,
)
from .source_fill_policy import recorded_slot_presentation, slot_replacement
from .source_form_execution import (
    execution_owner_metrics,
    plan_application_metrics,
)
from .source_fill_patterns import infer_fill_patterns, pattern_index
from .source_page_frame import derive_source_section_page_frames
from .source_visual_typography import bold_overrides, infer_visual_typography
from .source_format import DestinationStyleProfile, SourceLine
from .source_font_policy import font_name
from .word_safe_xml import set_east_asian_font
from .style_architecture import (
    FixtureStylePack,
    TenderStyleProfile,
    _toc_entries,
    classify_paragraph,
    fixture_integrity_report,
    infer_tender_style_profile,
    marker_info,
)
from .word_safe_scan import scan_word_safe_docx
from .word_safe_source_builder import WordSafeSourceDocumentBuilder, _key, set_run_spacing


def _row_join_remap(offset, removed):
    """Translate a payload offset across the same-row separator join."""

    if not removed or offset is None:
        return offset
    return int(offset) - sum(1 for index in removed if index < offset)


def _run_character_text_offsets(run):
    """Each recorded source character's own offset inside the run's text.

    The source's per-character boxes are the finest geometry the extraction
    carries, so a rule that covers only part of a run can be resolved to the exact
    glyphs it was drawn beneath.  The run's text may carry separators the
    extraction inserted, so the characters are matched in order instead of by
    index.  Returns ``[]`` when the source recorded no per-character geometry.
    """

    characters = list(getattr(run, "characters", ()) or ())
    text = str(getattr(run, "text", "") or "")
    if not characters or not text:
        return []
    offsets = []
    cursor = 0
    for char in characters:
        glyph = str(getattr(char, "character", "") or "")
        if not glyph:
            return []
        at = text.find(glyph, cursor)
        if at < 0:
            return []
        offsets.append(at)
        cursor = at + len(glyph)
    return offsets


def _character_majority_under_rule(char, rule_x0, rule_x1) -> bool:
    """Whether a source character is mostly beneath a drawn rule's extent.

    A rule ends where the source stopped drawing it, which is rarely exactly on a
    glyph boundary; requiring only an overlap would underline a whole trailing
    space as well.  A majority of the character's own box is the line's own
    coverage of it.
    """

    x0, x1 = float(char.x0), float(char.x1)
    width = x1 - x0
    if width <= 0:
        return False
    overlap = min(x1, rule_x1) - max(x0, rule_x0)
    return overlap >= width / 2.0


def _cell_rule_text_occupancy(rule, runs):
    """How much of one drawn cell rule the cell's own glyphs sit on.

    The same occupancy the page-level classifier uses, measured against the
    cell's runs because a table cell's text is not part of the page's own span
    list.  A rule the glyphs barely touch is a blank; a rule they cover is an
    underline of them.
    """

    from .page_layout import rule_text_occupancy

    return rule_text_occupancy(rule, list(runs))


class StyleFirstSourceDocumentBuilder(WordSafeSourceDocumentBuilder):
    """Build a source-first document with fixture-derived editable styles.

    The source geometry and fact-slot machinery remain Round 4.9's public
    implementation. This subclass changes the document architecture at the
    point where semantic blocks become Word paragraphs/cells. Source-owned
    numbering stays literal; generated numbering is a separate opt-in path.
    """

    def __init__(
        self,
        facts,
        template,
        *,
        page_scales=None,
        fixture_path: str | Path | None = None,
        style_profile: TenderStyleProfile | None = None,
        source_text_qa: dict | None = None,
        source_path: str | Path | None = None,
    ):
        super().__init__(facts, template, page_scales=page_scales)
        # One Word table per logical table: the PDF page fragments that continue
        # an open table are folded into it instead of becoming separate tables.
        self.logical_plan = self._build_logical_plan()
        self.style_pack = FixtureStylePack(fixture_path)
        self.source_path = source_path
        # Round 5.6: rendered source glyph weight decides the visual role weight,
        # and the source fill patterns decide how a fact value occupies a source
        # field.  Both are measured from real source evidence, never hard-coded.
        self.visual_typography = infer_visual_typography(source_path, template, self.layouts)
        self.fill_patterns = infer_fill_patterns(template, self.layouts)
        self.fill_pattern_index = pattern_index(self.fill_patterns)
        self.section_geometries: list[dict] = []
        #: V1 manual-layout closure: one stable ``SourceSectionPageFrame`` per
        #: source page group.  Every page of a group resolves to the *same*
        #: section frame, so no page's own bounding box can move a margin.
        self.section_frames: dict = {}
        #: Classified source rule object id -> stable ``P<page>-R<index>`` id, so
        #: a composition knows which source rule it represents.
        self._rule_ids_by_object: dict[int, str] = {}
        #: Source runs whose underline semantics a SOURCE_RULE_COMPOSITION has
        #: taken over, so the legacy whole-run underline does not re-paint the
        #: characters outside the physical rule.
        self._composition_owned_source_run_ids: set[int] = set()
        #: Production transformation policy per source rule, derived from source
        #: evidence during rendering (see ``source_rule_transformation_policy``).
        self._rule_transformation_policy: dict = {}
        #: Authoritative slot binding per source rule, from template-level
        #: ``SourceFillSlot`` evidence.
        self._rule_slot_binding: dict = {}
        self.rule_compositions_created = 0
        self.style_profile = style_profile or infer_tender_style_profile(
            template,
            self.layouts,
            visual_weights=bold_overrides(self.visual_typography),
        )
        self.body_list_num_id: int | None = None
        self.fixture_report: dict = {}
        self._toc_pages = {
            page.page for page in template.source_pages
            if any("".join(paragraph.text.split()) in {"目录", "目 录"} for paragraph in page.paragraphs)
        }
        self._toc_entries = _toc_entries(template)
        self.numbering_exception_reason: str | None = None
        self.source_text_qa = source_text_qa or {"source_text_disagreement_count": 0, "source_glyph_mapping_warning_count": 0, "repairs": []}

    def _cell_source_text(self, source_cell, slots):
        """A table cell's text with the rules the source drew inside it restored.

        A cell's own drawn rules are content, not table geometry: a rule occupied
        by the cell's glyphs is the underline *of that value* - so the run that
        owns it is marked underlined instead of receiving a blank - and a rule
        clear of them is a fixed blank the source left for an editable field,
        which becomes an inline blank at the rule's own measured span.
        """

        rules = self._cell_content_rules(source_cell)
        if not rules:
            return source_cell.text, slots, source_cell.runs
        runs = list(getattr(source_cell, "runs", ()) or ())
        covered = [
            covered_range
            for covered_range in (
                self._cell_rule_covered_range(rule, runs, source_cell.text)
                for rule in rules
            )
            if covered_range is not None
        ]
        for text_start, text_end in covered:
            runs = self._split_runs_for_covered_range(
                source_cell.text, runs, text_start, text_end,
            )
        text, remapped, restored = restore_inline_rule_blanks(
            source_cell.text, runs, rules, slots,
        )
        self.vector_blanks_recovered += restored
        return text, remapped, runs

    @staticmethod
    def _cell_rule_covered_range(rule, runs, text):
        """The cell-text character range a drawn rule underlines, if any.

        A rule whose extent is occupied by the cell's own glyphs is the underline
        *of that text*: the covered range is read from the run's own box, because
        a delivered Word run is underlined all-or-nothing and the source may draw
        the line under only part of it (``…截止之日起 90 日历天`` underlines the
        figure alone).  A rule the glyphs leave clear is a blank, not an
        underline, and carries no covered range here.
        """

        if not runs or not text:
            return None
        occupied = _cell_rule_text_occupancy(rule, runs)
        if occupied < RULE_TEXT_OCCUPANCY_UNDERLINE:
            return None
        rule_x0, rule_x1 = float(rule.bbox[0]), float(rule.bbox[2])
        for run in runs:
            run_text = str(getattr(run, "text", "") or "")
            if not run_text.strip():
                continue
            x0, _y0, x1, _y1 = (float(value) for value in run.bbox)
            if x1 <= rule_x0 or x0 >= rule_x1 or x1 <= x0:
                continue
            at = text.find(run_text)
            if at < 0:
                continue
            offsets = _run_character_text_offsets(run)
            if offsets:
                covered_offsets = [
                    offset
                    for char, offset in zip(run.characters, offsets)
                    if _character_majority_under_rule(char, rule_x0, rule_x1)
                ]
                if covered_offsets:
                    start = at + min(covered_offsets)
                    end = at + max(covered_offsets) + 1
                    if end > start:
                        return (start, end)
                continue
            width = x1 - x0
            span = len(run_text)
            start_fraction = (max(rule_x0, x0) - x0) / width
            end_fraction = (min(rule_x1, x1) - x0) / width
            start = at + int(round(start_fraction * span))
            end = at + int(round(end_fraction * span))
            start = max(at, min(start, at + span))
            end = max(start, min(end, at + span))
            if end > start:
                return (start, end)
        return None

    @staticmethod
    def _split_runs_for_covered_range(text, runs, text_start, text_end):
        """Split the run covering a character range into three source runs.

        The middle copy carries the source's own underline, so ``_add_run``
        underlines exactly the glyphs the source drew the rule beneath and leaves
        the rest of the run plain - a Word run cannot be underlined in part.
        """

        result = []
        cursor = 0
        for run in runs:
            run_text = str(getattr(run, "text", "") or "")
            at = text.find(run_text, cursor) if run_text else -1
            if at < 0 or at >= text_end or at + len(run_text) <= text_start:
                result.append(run)
                if at >= 0:
                    cursor = at + len(run_text)
                continue
            cursor = at + len(run_text)
            local_start = max(text_start - at, 0)
            local_end = min(text_end - at, len(run_text))
            x0, y0, x1, y1 = (float(value) for value in run.bbox)
            span = max(1, len(run_text))
            pieces = (
                (run_text[:local_start], False, 0, local_start),
                (run_text[local_start:local_end], True, local_start, local_end),
                (run_text[local_end:], False, local_end, len(run_text)),
            )
            for piece, is_covered, offset_start, offset_end in pieces:
                if not piece:
                    continue
                piece_x0 = x0 + (x1 - x0) * offset_start / span
                piece_x1 = x0 + (x1 - x0) * offset_end / span
                result.append(
                    run.model_copy(
                        update={
                            "text": piece,
                            "bbox": (piece_x0, y0, piece_x1, y1),
                            "underline": True if is_covered else run.underline,
                        }
                    )
                )
        return result

    @staticmethod
    def _cell_content_rules(source_cell):
        """The cell's own drawn rules as rule geometry the restorers understand."""
        rules = []
        for entry in getattr(source_cell, "content_rules", ()) or ():
            try:
                x0, x1, y = (float(value) for value in entry)
            except (TypeError, ValueError):
                continue
            rules.append(
                SourceLine(
                    bbox=(min(x0, x1), y - 0.5, max(x0, x1), y + 0.5),
                    width=0.5,
                    color=None,
                    orientation="horizontal",
                )
            )
        return rules

    def _add_run(self, paragraph, text, source=None, underline=None, spacing_pt=None):
        """Apply the Named Style baseline, then explicit source-owned overrides."""

        run = paragraph.add_run()
        if text:
            for index, line in enumerate(str(text).split("\n")):
                if index:
                    run.add_break(WD_BREAK.LINE)
                run.add_text(line)
        if spacing_pt:
            set_run_spacing(run, spacing_pt)
        baseline_role = str(getattr(source, "baseline_role", "NORMAL") or "NORMAL").upper()
        if baseline_role == "SUPERSCRIPT":
            run.font.superscript = True
        elif baseline_role == "SUBSCRIPT":
            run.font.subscript = True
        if underline is not None:
            run.underline = underline
        elif bool(getattr(source, "underline", False)):
            run.underline = True
        if source is not None:
            if isinstance(source, DestinationStyleProfile):
                east_asia = source.east_asia_font
                latin = source.latin_font
                size = source.font_size
            else:
                east_asia, _changed = font_name(getattr(source, "font_name", None))
                latin = "Times New Roman"
                size = float(getattr(source, "font_size", 0.0) or 0.0)
            if east_asia:
                run.font.name = latin or east_asia
                set_east_asian_font(run, east_asia)
            if size > 0:
                run.font.size = Pt(size)
            # Round 5.6: the delivered weight is owned by the semantic role.
            # A run only *adds* emphasis when the source span itself is bold;
            # it never forces ``w:b val=0``, which would erase the visual weight
            # the style profile measured from the rendered source page.
            if bool(getattr(source, "bold", False)):
                run.bold = True
            run.italic = bool(getattr(source, "italic", False))
            # Round 5.8: a source rule whose extent is occupied by this span's
            # glyphs is an underline *of that text*, not a separate blank.  The
            # owning run carries the underline, so no invented fill tab and no
            # extra blank appear after the text.
            if run.underline is not True and self._source_run_is_underlined(source):
                run.underline = True
        return run

    def _resolved_source_values(self) -> tuple:
        """Resolved ProjectFacts values, for VALUE_UNDERLINE classification."""

        values = []
        fields = getattr(self.facts, "fields", None)
        for name in dir(fields) if fields is not None else []:
            if name.startswith("_"):
                continue
            try:
                field = getattr(fields, name)
            except Exception:
                continue
            value = getattr(field, "resolved_value", None)
            status = getattr(getattr(field, "status", None), "value", None)
            if status == "RESOLVED" and value:
                text = str(value).strip()
                if len(text) >= 2:
                    values.append(text)
        return tuple(values)

    def _resolved_fact_fields(self) -> frozenset:
        """Fact fields ProjectFacts has actually resolved, for slot binding."""

        names = set()
        fields = getattr(self.facts, "fields", None)
        for name in dir(fields) if fields is not None else []:
            if name.startswith("_"):
                continue
            try:
                field = getattr(fields, name)
            except Exception:
                continue
            value = getattr(field, "resolved_value", None)
            status = getattr(getattr(field, "status", None), "value", None)
            if status == "RESOLVED" and value:
                names.add(str(name).upper())
        return frozenset(names)

    def _resolved_fact_values(self) -> dict:
        """Resolved fact field -> resolved value, the one value authority."""

        values = {}
        fields = getattr(self.facts, "fields", None)
        for name in dir(fields) if fields is not None else []:
            if name.startswith("_"):
                continue
            try:
                field = getattr(fields, name)
            except Exception:
                continue
            value = getattr(field, "resolved_value", None)
            status = getattr(getattr(field, "status", None), "value", None)
            if status == "RESOLVED" and value:
                text = str(value).strip()
                if text:
                    values[str(name).upper()] = text
        return values

    def _annotated_source_rule_registry(self) -> list:
        """Registry records carrying their derived policy and geometry intent."""

        annotated = []
        for entry in self.source_rule_registry:
            policy = self._rule_transformation_policy.get(entry["source_rule_id"])
            binding = self._rule_slot_binding.get(entry["source_rule_id"]) or {}
            record = dict(entry)
            record["transformation_policy"] = policy
            record["geometry_intent"] = source_rule_geometry_intent(policy)
            record["authoritative_slot_id"] = binding.get("slot_id")
            record["authoritative_slot_binding"] = binding.get("reason")
            record["allowed_fact_fields"] = binding.get("allowed_fact_fields") or []
            record["resolved_fact_fields"] = binding.get("resolved_fact_fields") or []
            annotated.append(record)
        return annotated

    def _rule_relation_for(self, rule):
        """The source-local relation record already built for this rule."""

        rule_id = self._rule_ids_by_object.get(id(rule))
        if rule_id is None:
            return None
        for entry in self._page_rule_relations:
            if entry.get("source_rule_id") == rule_id:
                return entry
        return None

    def _composition_form_line_rules(self, item, layout):
        """Rules that need their own line context because of partial occupation.

        A rule whose extent is only partly occupied by source glyphs has interior
        empty segments, so it has to paint from its own source anchor.  On a
        source visual line that is not this element's first line, continuous flow
        may already have carried the cursor past that anchor, and a tab cannot
        move backwards.  Evidence is source visual-line geometry plus the rule's
        own glyph occupancy - never a page number, rule id or literal text.
        """

        lines = list(item.source_lines)
        if len(lines) < 2:
            return []
        found = []
        for rule in layout.horizontal_rules:
            if getattr(rule, "orientation", None) != "horizontal":
                continue
            relation = self._rule_relation_for(rule)
            if relation is None:
                continue
            occupancy = float(relation.get("text_occupancy") or 0.0)
            if not RULE_TEXT_OCCUPANCY_GAP <= occupancy < RULE_TEXT_OCCUPANCY_UNDERLINE:
                continue
            rule_y = (float(rule.bbox[1]) + float(rule.bbox[3])) / 2.0
            if any(
                float(line.bbox[1]) - 1.0 <= rule_y <= float(line.bbox[3]) + 2.0
                for line in lines[1:]
            ):
                found.append(rule)
        found.sort(key=lambda rule: (float(rule.bbox[1]), float(rule.bbox[0])))
        return found

    def _source_run_is_underlined(self, source) -> bool:
        """Whether this exact source run owns a classified underline rule.

        A source run that a SOURCE_RULE_COMPOSITION has split and consumed no
        longer owns a whole-run underline: the composition itself carries the
        underline for its own TEXT segments, and the characters the split left
        outside the physical rule stay plain.
        """

        if id(source) in self._composition_owned_source_run_ids:
            return False
        return id(source) in self._underlined_source_run_ids
    def _semantic(self, item, page, layout) -> object:
        return classify_paragraph(
            item,
            page,
            layout_classification=layout.classification,
            toc_pages=self._toc_pages,
            toc_entries=self._toc_entries,
            neighboring_items=[candidate for candidate in layout.elements if hasattr(candidate, "source")],
        )

    @staticmethod
    def _trim_marker(raw: str, fragment_slots: list, *, remove: bool):
        if not remove:
            return raw, fragment_slots, False
        marker, _remainder = marker_info(raw)
        if not marker:
            return raw, fragment_slots, False
        # marker_info strips leading whitespace; recompute the exact source
        # prefix so slot offsets remain tied to the original source container.
        match = re.match(r"^\s*" + re.escape(marker), raw)
        if match is None:
            return raw, fragment_slots, False
        skip = match.end()
        adjusted = []
        for slot in fragment_slots:
            start, end = slot.text_start, slot.text_end
            if start is None or end is None:
                adjusted.append(slot)
            elif start >= skip:
                adjusted.append(slot.model_copy(update={"text_start": start - skip, "text_end": end - skip}))
            elif end <= skip:
                continue
            else:
                adjusted.append(slot.model_copy(update={"text_start": 0, "text_end": max(0, end - skip)}))
        return raw[skip:], adjusted, True

    def _record_fill(self, slot, value, method, run, paragraph):
        # Ensure filled values are attached to a named destination style before
        # the inherited Round 4.9 audit record is produced.
        if slot.container_type == "table_cell":
            paragraph.style = "Tender Table Value"
        elif paragraph.style.name in {"Normal", "正文（首行缩进两字）"}:
            paragraph.style = "Tender Signature Block"
        super()._record_fill(slot, value, method, run, paragraph)
        # The inherited audit helper reads run fonts through get_or_add_rFonts,
        # which would leave an empty direct rFonts node on an otherwise
        # style-owned run. Remove that no-op override before serialization.
        rpr = run._r.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is not None and not rfonts.attrib:
            rpr.remove(rfonts)
        record = self.filled[-1]
        profile = slot.destination_style
        record["generated_paragraph_index"] = len(paragraph.part.document.paragraphs) - 1
        presentation = recorded_slot_presentation(slot)
        if presentation is not None:
            # The value Word received is composed of several presentation
            # components.  Each one is recorded separately, so the resolved fact
            # is distinguishable from the source's own separator and from the
            # source form's not-applicable marker - and the marker can never be
            # read back as a resolved value.  The composite occupies one source
            # form slot, so the slot's own underline intent is inherited by the
            # components that sit in it; the prose before and after the slot is
            # written by its own runs and keeps only its own decoration.
            components = [component["kind"] for component in presentation.get("components", ())]
            record["value_presentation"] = {
                **presentation,
                "source_slot_underline_inherited": bool(run.underline),
                "underline_components": components if run.underline else [],
                "underline_provenance": (
                    "the source drew one rule across this form slot, so every "
                    "component occupying the slot inherits that decoration"
                    if run.underline
                    else "the source drew no rule on this form slot"
                ),
            }
        record["generated_value_style"] = {
            "east_asia_font": profile.east_asia_font,
            "latin_font": profile.latin_font,
            "font_size": profile.font_size,
            "bold": profile.bold,
            "italic": profile.italic,
            "underline": bool(run.underline),
            "color": "000000",
            "font_weight_role": profile.font_weight_role,
            "paragraph_alignment": profile.paragraph_alignment,
            "word_style": paragraph.style.name,
        }

    def form_block(self, doc, block, *, page_content_x0=18.0, page_content_x1=None, scale=1.0, initial_space_before=0.0):
        first = len(doc.paragraphs)
        last = super().form_block(
            doc,
            block,
            page_content_x0=page_content_x0,
            page_content_x1=page_content_x1,
            scale=scale,
            initial_space_before=initial_space_before,
        )
        for paragraph in doc.paragraphs[first:]:
            paragraph.style = "Tender Signature Block"
        for record in self.logical_records:
            if record.get("paragraph_index", -1) >= first and record.get("kind") == "FormRow":
                record.update({
                    "word_style": "Tender Signature Block",
                    "numbering_family": None,
                    "heading_level": None,
                    "body_list_level": None,
                })
        return last

    def _register_page_rules(self, page, layout) -> None:
        """Register this source page's classified horizontal rules.

        Every classified source rule gets a stable id sorted by its position on
        its page, so a generated rule's provenance can be resolved document-wide
        instead of only within one QA page.  Registration is a property of the
        source page, not of whether that page owns an emitted section, so a page
        whose table fragment was folded into another section still contributes.
        """

        page_spans = [
            span
            for paragraph in page.paragraphs
            for line in paragraph.lines
            for span in line.spans
        ]
        self._underline_rule_spans = []
        self._underlined_source_run_ids = set()
        self._page_rule_relations = []
        page_visual_lines = []
        for paragraph in page.paragraphs:
            for line in paragraph.lines:
                bands = [
                    span.bbox
                    for span in line.spans
                    if getattr(span, "text", "").strip()
                ]
                if not bands:
                    continue
                page_visual_lines.append(
                    {
                        "y0": min(float(band[1]) for band in bands),
                        "y1": max(float(band[3]) for band in bands),
                        "x0": min(float(band[0]) for band in bands),
                        "x1": max(float(band[2]) for band in bands),
                    }
                )
        page_rules = sorted(
            (
                rule
                for rule in layout.horizontal_rules
                if rule.orientation == "horizontal"
            ),
            key=lambda rule: (
                (float(rule.bbox[1]) + float(rule.bbox[3])) / 2.0,
                float(rule.bbox[0]),
            ),
        )
        for rule_index, rule in enumerate(page_rules, start=1):
            relation = classify_rule_relation(
                rule, page_spans, resolved_values=self._resolved_source_values()
            )
            rule_id = "P%d-R%d" % (page.page, rule_index)
            # The source visual form line a rule belongs to, so an anchored value
            # can name the form line it was positioned into.
            rule_y = (float(rule.bbox[1]) + float(rule.bbox[3])) / 2.0
            for line_index, visual_line in enumerate(page_visual_lines, start=1):
                if visual_line["y0"] - 1.0 <= rule_y <= visual_line["y1"] + 2.0:
                    self._rule_source_visual_line[rule_id] = "S%d-L%d" % (
                        page.page,
                        line_index,
                    )
                    break
            self._page_rule_relations.append(
                {
                    "source_rule_id": rule_id,
                    "source_page": page.page,
                    "rule_x0": round(float(rule.bbox[0]), 2),
                    "rule_x1": round(float(rule.bbox[2]), 2),
                    "rule_y": round((float(rule.bbox[1]) + float(rule.bbox[3])) / 2, 2),
                    **relation,
                }
            )
            self._rule_ids_by_object[id(rule)] = rule_id
            self.source_rule_registry.append(
                {
                    "source_rule_id": rule_id,
                    "source_page": page.page,
                    "x0": round(float(rule.bbox[0]), 2),
                    "x1": round(float(rule.bbox[2]), 2),
                    "y": round((float(rule.bbox[1]) + float(rule.bbox[3])) / 2, 2),
                    "relation_type": relation["relation_type"],
                    "text_occupancy": relation["text_occupancy"],
                    "bearing": relation["relation_type"]
                    in (
                        "TEXT_UNDERLINE",
                        "PLACEHOLDER_UNDERLINE",
                        "VALUE_UNDERLINE",
                        "GAP_FILL_RULE",
                    ),
                }
            )
            if relation["relation_type"] in (
                "TEXT_UNDERLINE",
                "PLACEHOLDER_UNDERLINE",
                "VALUE_UNDERLINE",
            ):
                self._underline_rule_spans.append(
                    (
                        float(rule.bbox[0]),
                        float(rule.bbox[2]),
                        (float(rule.bbox[1]) + float(rule.bbox[3])) / 2,
                    )
                )

    def table(self, doc, source, usable, *, page_content_x0=None):
        before = len(doc.tables)
        result = super().table(doc, source, usable, page_content_x0=page_content_x0)
        if len(doc.tables) <= before:
            return result
        table = doc.tables[-1]
        slots_by_cell = {
            (slot.row_index, slot.column_index): slot
            for slot in self.template.fill_slots
            if slot.container_type == "table_cell"
            and slot.source_page == source.page
            and slot.table_index == source.table_index
        }
        for row in source.rows:
            for source_cell in row.cells:
                if source_cell.row_index >= len(table.rows) or source_cell.column_index >= len(table.columns):
                    continue
                cell = table.cell(source_cell.row_index, source_cell.column_index)
                slot = slots_by_cell.get((source_cell.row_index, source_cell.column_index))
                role = getattr(source_cell, "cell_role", "BODY_TEXT")
                style_name = {
                    "HEADER": "Tender Table Header",
                    "NUMERIC": "Tender Table Numeric",
                    "UNIT": "Tender Table Unit",
                    "BODY_TEXT": "Tender Table Value" if slot or source_cell.column_index else "Tender Table Label",
                }.get(role, "Table Text")
                for paragraph in cell.paragraphs:
                    paragraph.style = style_name
                    paragraph.alignment = {
                        "left": 0, "center": 1, "right": 2, "justify": 3,
                    }.get(getattr(source_cell, "horizontal_alignment", "left"), 0)
        return result

    def _page_has_emittable_content(self, layout) -> bool:
        """Whether this source page still owes the document any content.

        A source page whose only elements are table fragments that continue an
        open logical table contributes nothing of its own: those rows render in
        the head fragment's section, and the page's non-table elements - if any -
        are what remain for it to contribute.  The evidence is the page's own
        layout model resolved against the logical plan, never a page number, a
        table index or a case identity.
        """

        for item in layout.elements or ():
            page_number = getattr(item, "page", None)
            table_index = getattr(item, "table_index", None)
            if page_number is None or table_index is None:
                return True
            if self.logical_plan.head_page_for(int(page_number), int(table_index)) is None:
                return True
        return False

    def _prepare_logical_paragraph(self, doc, item, page, layout, gap, page_content_x0, scale,
                                   *, page_content_x1=None):
        semantic = self._semantic(item, page, layout)
        paragraph = doc.add_paragraph(style=semantic.style_name)
        paragraph.alignment = {
            "center": 1, "right": 2, "justify": 3, "left": 0,
        }.get(item.alignment_hint, 0)
        self._positioned_paragraph_left = getattr(self, "_positioned_paragraph_left", {})
        pf = paragraph.paragraph_format
        pf.space_before = Pt(gap)
        pf.space_after = Pt(0)
        # SOURCE PARAGRAPH INDENT.  The paragraph's own source visual rows decide
        # whether its first row is indented, hanging or aligned with the rows it
        # wraps into, and the body boundary they return to is the paragraph's left
        # indent.  A list item is not a special case: writing its first row's x as
        # a whole-paragraph left indent moves every wrapped row off the body
        # boundary the source returns it to, and forcing its first-line indent to
        # zero drops the source's own first-line offset.
        self._apply_source_paragraph_indent(pf, item, page_content_x0)
        # The left edge a tab stop is measured from.  Tab stops are measured from
        # the *section text margin*
        # (``WORD_TAB_REFERENCE_MODEL``), which a first-line or hanging indent
        # does not move.  Only the paragraph's own left indent shifts the ruler.
        self._positioned_paragraph_left[id(paragraph)] = (
            float(page_content_x0) + float(pf.left_indent.pt or 0.0)
        )
        # A source line that runs past the stable body frame is element-level
        # geometry, not a reason to widen the section: the paragraph keeps its
        # source extent through a Word-native negative right indent, which
        # extends the paragraph into the right margin exactly as far as the
        # source does.  Without it the frame change would wrap a source line the
        # source itself keeps on one line.
        pf.right_indent = Pt(self._frame_overhang_right_indent(item, page_content_x1))
        pf.line_spacing = infer_semantic_line_spacing(
            item.font_size_pt, item.line_pitch_pt * scale,
            source_locator=item.source.locator,
        ).word_value
        pf.widow_control = False
        return paragraph, semantic

    def _frame_overhang_right_indent(self, item, page_content_x1) -> float:
        """Negative right indent needed to keep a source line's own extent.

        The stable body frame is a *bound*: most source lines end inside it, and
        a few genuinely end past it.  Reproducing those few must not move the
        section margin for every page, so the difference rides on the element:
        ``w:ind/@w:right`` is negative by exactly the source overhang.  A
        paragraph that stays inside the frame keeps a right indent of zero.
        """

        if page_content_x1 is None:
            return 0.0
        extents = [
            max((float(span.bbox[2]) for span in line.spans
                 if str(getattr(span, "text", "") or "").strip()),
                default=None)
            for line in (getattr(item, "source_lines", None) or ())
        ]
        extents = [value for value in extents if value is not None]
        if not extents:
            extents = [float(item.bbox[2])]
        overhang = max(extents) - float(page_content_x1)
        return round(min(0.0, -overhang), 2)

    def _form_line_rule_spans(self, item, layout):
        """Source rules that fall on this element's own source visual lines.

        The rule has to lie on a source line this element owns, and its glyph
        band must be clear of that line's text, so it is a real fill gap rather
        than an underline that belongs to a text run.
        """

        spans = [
            span
            for line in item.source_lines
            for span in getattr(line, "spans", [])
            if getattr(span, "text", "").strip()
        ]
        if not spans:
            return []
        found = []
        for rule in layout.horizontal_rules:
            if getattr(rule, "orientation", None) != "horizontal":
                continue
            rule_x0, rule_x1 = float(rule.bbox[0]), float(rule.bbox[2])
            rule_y = (float(rule.bbox[1]) + float(rule.bbox[3])) / 2.0
            on_line = [
                span
                for span in spans
                if float(span.bbox[1]) - 1.0 <= rule_y <= float(span.bbox[3]) + 2.0
            ]
            if not on_line:
                continue
            occupied = any(
                min(float(span.bbox[2]), rule_x1) - max(float(span.bbox[0]), rule_x0) > 0.05
                for span in on_line
            )
            if occupied:
                continue
            found.append(rule)
        found.sort(key=lambda rule: (float(rule.bbox[1]), float(rule.bbox[0])))
        return found

    def _later_source_row_rule_spans(self, item, layout):
        """Gap rules that sit on a source visual row after the element's first.

        Such a rule has its own row origin, which continuous text flow cannot
        guarantee, so that whole row is assembled into its own paragraph context
        before any of it is emitted.  The evidence is source geometry only.
        """

        rows = source_visual_rows(item.source_lines)
        if len(rows) < 2:
            return []
        later = rows[1:]
        found = []
        for rule in self._form_line_rule_spans(item, layout):
            rule_y = (float(rule.bbox[1]) + float(rule.bbox[3])) / 2.0
            if any(
                min(float(line.bbox[1]) for line in row) - 1.0
                <= rule_y
                <= max(float(line.bbox[3]) for line in row) + 2.0
                for row in later
            ):
                found.append(rule)
        return found

    def _activate_source_form_line_paragraphs(self, doc, paragraph, item, semantic, rules):
        """Turn on form-line paragraphisation for this element."""

        style_name = semantic.style_name
        alignment = paragraph.alignment
        source_lines = list(item.source_lines)
        state = {"closed": []}

        def factory(previous, record):
            new_paragraph = doc.add_paragraph(style=style_name)
            new_paragraph.alignment = alignment
            formatting = new_paragraph.paragraph_format
            # Explicit, conservative and source-derived: no inherited spacing
            # that would invent vertical space between form lines.
            formatting.space_before = Pt(0)
            formatting.space_after = Pt(0)
            formatting.left_indent = Pt(0)
            formatting.right_indent = Pt(0)
            formatting.first_line_indent = Pt(0)
            formatting.line_spacing = paragraph.paragraph_format.line_spacing
            record["generated_paragraph_index"] = len(doc.paragraphs) - 1
            record["paragraph_format"] = {
                "left_indent_pt": 0.0,
                "right_indent_pt": 0.0,
                "first_line_indent_pt": 0.0,
                "space_before_pt": 0.0,
                "space_after_pt": 0.0,
                "line_spacing": formatting.line_spacing,
            }
            record["source_visual_line_ids"] = [
                "S%d-L%d" % (self._active_source_page, index + 1)
                for index in range(len(source_lines))
            ]
            record["source_line_y"] = (
                round(float(source_lines[0].bbox[1]), 2) if source_lines else None
            )
            record["source_line_x0"] = (
                round(float(source_lines[0].bbox[0]), 2) if source_lines else None
            )
            record["source_line_x1"] = (
                round(float(source_lines[0].bbox[2]), 2) if source_lines else None
            )
            state["closed"].append(record)
            # The rules on this new form line, mapped by the factory caller.
            return new_paragraph

        self._source_form_line_mode = True
        self._form_line_paragraph_factory = factory
        self._form_line_has_content = False
        self._form_line_previous = None
        self._form_line_composition_split = False
        self._form_line_source_rows = set()
        self._pending_form_line_records = state

    def _deactivate_source_form_line_paragraphs(self):
        state = getattr(self, "_pending_form_line_records", None)
        if state is not None:
            for record in state["closed"]:
                self.source_form_line_paragraphs.append(record)
        self._source_form_line_mode = False
        self._form_line_paragraph_factory = None
        self._form_line_has_content = False
        self._form_line_composition_split = False
        self._form_line_source_rows = set()
        self._pending_form_line_records = None
        self._row_context_paragraph = None
        self._current_emission_paragraph = None

    def _covered_text_range_for_rule(self, rule, runs, text):
        """The character range of ``text`` a rule's own glyphs occupy.

        The same authoritative mapping the restorers use: visible characters are
        matched to source runs in reading order, so the rule's extent resolves to
        the offsets of the characters it was drawn beneath.  This is the evidence
        that lets the emitter hand a source decoration to the semantic slot that
        occupies the rule, instead of to whichever run happens to contain it.
        """

        positions = [i for i, character in enumerate(text) if not character.isspace()]
        compact = "".join(text[i] for i in positions)
        mapped, cursor = [], 0
        for run in runs:
            key = "".join(str(getattr(run, "text", "") or "").split())
            if not key:
                continue
            at = compact.find(key, cursor)
            if at >= 0:
                mapped.append((run, positions[at], positions[at + len(key) - 1] + 1))
                cursor = at + len(key)
        x0, x1 = float(rule.bbox[0]), float(rule.bbox[2])
        covered = [
            (start, end)
            for run, start, end in mapped
            if float(run.bbox[2]) > x0 + 0.5 and float(run.bbox[0]) < x1 - 0.5
        ]
        if not covered:
            return None
        return min(start for start, _ in covered), max(end for _, end in covered)

    def _decoration_transfers_to_replacement_value(self, rule, runs, text, slots) -> bool:
        """Whether a rule's decoration belongs to a replacement value.

        The test is evidence the renderer actually holds: a rule transfers its
        decoration when the source characters it was drawn beneath are a
        placeholder that a resolved value replaces.  The decoration then belongs
        to that value and must not also be painted onto the surrounding source
        run, which is what inverts the underline when a placeholder and its
        prose share one source run.  A rule covering no replaced slot keeps the
        accepted whole-run underline.
        """

        covered = self._covered_text_range_for_rule(rule, runs, text)
        if covered is None:
            return False
        for slot in slots:
            start = getattr(slot, "text_start", None)
            end = getattr(slot, "text_end", None)
            if start is None or end is None:
                continue
            if not (covered[0] <= start and end <= covered[1]):
                continue
            if slot_replacement(slot, self.facts) is not None:
                return True
        return False

    def _render_logical_content(self, paragraph, item, layout, semantic):
        item_slots = [
            s for fragment, _start, _end in item.fragments
            for s in self.template.fill_slots
            if _key(s.source_locator) == _key(fragment.locator)
        ]
        has_project_and_lot_slot = any(s.slot_type == "PROJECT_AND_LOT_SLOT" for s in item_slots)
        # Every paragraph reconstructed from the tender source is
        # SOURCE_LITERAL by default. A Named Style supplies navigation and
        # editability; it must not replace the source-visible prefix.
        remove_marker = semantic.numbering_ownership.value == "GENERATED_AUTO"
        removed_marker = False
        # SOURCE-LINE-AWARE EMISSION: the source's own physical rows are the
        # assembly unit.  One reconstructed element can own several physical
        # rows, and a row that carries a positioned rule must be assembled - its
        # leading source text, its rules and its trailing text together - before
        # any of it is emitted, because continuous flow cannot guarantee the
        # row's own origin.
        rows = source_visual_rows(item.source_lines)
        row_indexes = self._fragment_visual_row_indexes(item, rows)
        rule_spans = self._source_rule_spans_for_emission_plan()
        self._current_emission_paragraph = paragraph
        # SOURCE VISUAL ROW COUNT: the source's own wrapping is not a line break.
        # A reconstructed element that owns several physical rows is one flowing
        # source paragraph whose rows wrap, so none of those rows may open a Word
        # paragraph of its own; only an element that IS one physical row is a
        # source form line, and only that may be isolated for its own geometry.
        self._current_source_visual_row_count = len(rows)
        for fragment_index, (fragment, start, end) in enumerate(item.fragments):
            raw = fragment.text[start:end]
            fragment_slots = [
                s.model_copy(update={"text_start": s.text_start - start, "text_end": s.text_end - start})
                for s in self.template.fill_slots
                if _key(s.source_locator) == _key(fragment.locator)
                and s.text_start is not None and s.text_end is not None
                and start <= s.text_start <= s.text_end <= end
            ]
            # A source row drawn as several PDF text objects arrives with an
            # extraction newline between them.  That newline is presentation, not
            # a source line break: joining it here keeps a one-row source element
            # on one rendered row instead of splitting it into two paragraphs.
            # An extraction newline a fill slot addresses is the opposite - it is
            # the blank the source drew between two text objects - so those
            # offsets are protected, and every remaining offset the join removes
            # is translated in the slots below instead of shifting them.
            protected = {
                offset
                for slot in fragment_slots
                for offset in range(int(slot.text_start), int(slot.text_end) + 1)
            }
            raw, row_join_offsets = single_visual_row_join(
                raw, getattr(fragment, "lines", None) or (), protected=protected)
            if row_join_offsets:
                fragment_slots = [
                    s.model_copy(update={
                        "text_start": _row_join_remap(s.text_start, row_join_offsets),
                        "text_end": _row_join_remap(s.text_end, row_join_offsets),
                    })
                    for s in fragment_slots
                ]
            if fragment_index == 0:
                raw, fragment_slots, removed_marker = self._trim_marker(
                    raw, fragment_slots, remove=remove_marker,
                )
            raw, fragment_slots, restored = restore_inline_rule_blanks(
                raw, fragment.runs, layout.horizontal_rules, fragment_slots,
                spans=[s for line in item.source_lines for s in line.spans]
                if all(hasattr(line, "spans") for line in item.source_lines)
                else None,
            )
            self.vector_blanks_recovered += restored
            # SOURCE_RULE_COMPOSITION: a rule that spans whitespace and text is
            # one logical representation built from ordered segments, so its own
            # visible characters are consumed here and re-emitted positioned.
            raw, fragment_slots, composed = restore_inline_rule_compositions(
                raw, fragment.runs, layout.horizontal_rules, fragment_slots,
                # The container text the covered ranges and the character
                # evidence are both expressed in.
                container_text=raw,
                spans=[s for line in item.source_lines for s in line.spans]
                if all(hasattr(line, "spans") for line in item.source_lines)
                else None,
                resolved_values=self._resolved_source_values(),
                rule_ids=self._rule_ids_by_object,
                owned_runs=self._composition_owned_source_run_ids,
                policy_out=self._rule_transformation_policy,
                # Template-level SourceFillSlot objects are the single source of
                # truth for slot geometry, allowed fact fields and placeholder
                # semantics; a glyph-free rule can only be bound by their own
                # source geometry, since fragment slots are filtered by text
                # range and drop exactly those slots.
                authoritative_slots=self.template.fill_slots,
                resolved_fact_fields=self._resolved_fact_fields(),
                binding_out=self._rule_slot_binding,
                # Source page per rule, so a slot can only be claimed by a rule
                # cut out of the same source page.
                rule_pages={
                    entry["source_rule_id"]: entry.get("source_page")
                    for entry in self.source_rule_registry
                },
                # Compiled once per document, before rendering: the source glyph
                # evidence each physical rule owns, and the per-rule field plan
                # that execution resolves against at the emission event.
                rule_glyph_evidence=self._source_rule_glyph_evidence,
                rule_field_plans=self._rule_field_plans,
            )
            self.rule_compositions_created += composed
            # A rule classified as a text/placeholder/value underline belongs to
            # one owning source run, which then carries the underline itself -
            # unless the rule decorates a placeholder that a resolved value
            # replaces.  In that case the source decoration belongs to the
            # semantic slot and is inherited by the replacement value, so
            # painting the owning run as well would underline prose the rule
            # never covered and leave the value undecorated.
            inherited_decoration_ranges = []
            # SOURCE_RULE_DECORATION.  A rule the compiler bound to a slot whose
            # placeholder a resolved value replaces is decoration evidence in its
            # own right: the binding names the exact slot the source drew the rule
            # on.  Such a rule's extent commonly carries a blank, the placeholder
            # and more blank, so the occupancy heuristic labels it a form-layout
            # rule rather than an underline and the label-based eligibility would
            # never see the decoration the source drew on the slot.  The binding
            # is the stronger evidence, so those rules are eligible as well - and
            # only those, so a rule whose slot keeps its placeholder or stays
            # empty still paints nothing onto a value.
            bound_decoration_rule_ids = {
                rule_id
                for rule_id, binding in (self._rule_slot_binding or {}).items()
                if (binding or {}).get("resolved_fact_fields")
                and (self._rule_transformation_policy or {}).get(rule_id)
                in VALUE_REPLACING_POLICIES
            }
            for owner_run, relation, rule in underline_owner_runs(
                layout.horizontal_rules,
                fragment.runs,
                resolved_values=self._resolved_source_values(),
                bound_decoration_rule_ids=bound_decoration_rule_ids,
            ):
                if self._decoration_transfers_to_replacement_value(
                    rule, fragment.runs, raw, fragment_slots
                ):
                    covered = self._covered_text_range_for_rule(rule, fragment.runs, raw)
                    if covered is not None:
                        inherited_decoration_ranges.append(covered)
                    self._page_rule_relations.append(
                        {
                            "source_page": self._active_source_page,
                            "emission_mechanism": "VALUE_INHERITS_SLOT_DECORATION",
                            "relation_type": relation["relation_type"],
                            "source_rule_id": (self._rule_ids_by_object or {}).get(id(rule)),
                            "covered_text_range": None if covered is None else list(covered),
                            "owner_text": str(getattr(owner_run, "text", ""))[:60],
                            "owner_bbox": [round(float(v), 2) for v in owner_run.bbox],
                            "reason": (
                                "rule decorates a placeholder bound to a resolved fact, so the "
                                "decoration is inherited by the replacement value instead"
                            ),
                        }
                    )
                    continue
                self._underlined_source_run_ids.add(id(owner_run))
                self._page_rule_relations.append(
                    {
                        "source_page": self._active_source_page,
                        "emission_mechanism": "RUN_UNDERLINE",
                        "relation_type": relation["relation_type"],
                        "source_rule_id": (self._rule_ids_by_object or {}).get(id(rule)),
                        "owner_text": str(getattr(owner_run, "text", ""))[:60],
                        "owner_bbox": [round(float(v), 2) for v in owner_run.bbox],
                    }
                )
            if fragment_index and has_project_and_lot_slot:
                for field_name in ("project_number", "tender_number"):
                    fact = getattr(self.facts.fields, field_name)
                    if fact.status.value != "RESOLVED" or fact.resolved_value is None:
                        continue
                    incompatible = str(fact.resolved_value)
                    if raw.startswith(incompatible) and re.match(r"(?:询比|招标|采购)文件", raw[len(incompatible):]):
                        raw = raw[len(incompatible):]
                        self.semantic_slot_corrections.append({
                            "slot_id": next(s.slot_id for s in item_slots if s.slot_type == "PROJECT_AND_LOT_SLOT"),
                            "removed_incompatible_value": incompatible,
                            "reason": "identifier in PROJECT_AND_LOT_SLOT across PDF fragments",
                        })
                        break
            if fragment_index and paragraph.text and raw and paragraph.text[-1].isascii() and paragraph.text[-1].isalnum() and raw[0].isascii() and raw[0].isalnum():
                self._add_run(paragraph, " ", fragment.runs[0] if fragment.runs else None)
            # The row's emission atoms are known before the row's first glyph is
            # written, so the whole source visual line is assembled into one Word
            # paragraph context and emitted in source-x order.
            self._row_context_paragraph = None
            row_index = row_indexes[fragment_index] if fragment_index < len(row_indexes) else None
            plan = plan_source_visual_line(
                raw,
                source_page=self._active_source_page,
                source_row_index=0 if row_index is None else row_index,
                source_lines=(
                    item.source_lines
                    if row_index is None or row_index >= len(rows)
                    else rows[row_index]
                ),
                owning_element_kind=item.kind,
                rule_spans=rule_spans,
                source_locator=fragment.locator,
            )
            # A row that is not the element's first row cannot trust continuous
            # flow for its own origin, so it is assembled inside its own
            # paragraph context - opened *before* its leading text is emitted -
            # and the row's leading text and every atom of the row end up in one
            # paragraph as one rendered row.
            #
            # A row that IS the element's first row keeps the element's own
            # paragraph context.  Its positioned atoms are reached forward from
            # that paragraph's origin by the source-derived tab the atom's own
            # representation emits, so isolating it would open a second Word
            # paragraph for a source visual line that never needed one and would
            # drop the element's own style indent on the continuation.  The
            # clause ``4、本项目询比有效期为提交响应文件截止之日起90日历天…``
            # was split exactly that way: its fixed leading text ends at the R7
            # anchor, so the atom is forward-reachable and the row must stay one
            # paragraph.
            if not plan.needs_own_line_context and plan.positioned_atoms:
                plan.forward_reachable = True
            if plan.needs_own_line_context and self._row_is_its_own_source_line(
                plan, len(rows)
            ):
                plan.line_context_available = (
                    getattr(self, "_form_line_paragraph_factory", None) is not None
                )
                if plan.line_context_available:
                    paragraph = self._assemble_source_visual_line(
                        paragraph, plan, fragment
                    )
                else:
                    # The element has no source-form-line paragraph machinery, so
                    # the row keeps the accepted inline emission.  The gap is
                    # recorded rather than silently assumed away, and the frozen
                    # horizontal gates still measure what it actually painted.
                    self.source_visual_line_assembly_gaps.append(
                        {
                            "source_page": plan.source_page,
                            "source_line_identity": plan.source_line_identity,
                            "source_row_index": plan.source_row_index,
                            "source_y0": round(float(plan.source_y0), 2),
                            "source_rule_ids": plan.source_rule_ids,
                            "reason": (
                                "element has no source-form-line paragraph "
                                "activated; row keeps the accepted inline emission"
                            ),
                        }
                    )
            elif plan.needs_own_line_context:
                # The row is a wrapped row of a flowing paragraph, so it is not a
                # line the source drew: the row keeps the element's own paragraph
                # and its positioned atoms are placed by flow.  Recorded, not
                # assumed away, so the horizontal gates still measure the result.
                plan.line_context_available = False
                plan.wrapped_row_stays_in_paragraph = True
                self.source_visual_line_assembly_gaps.append(
                    {
                        "source_page": plan.source_page,
                        "source_line_identity": plan.source_line_identity,
                        "source_row_index": plan.source_row_index,
                        "source_y0": round(float(plan.source_y0), 2),
                        "source_rule_ids": plan.source_rule_ids,
                        "reason": (
                            "row is one of the element's own wrapped rows, not a "
                            "source line: the element stays one Word paragraph "
                            "and the row is placed by flow"
                        ),
                    }
                )
            # The row's own line context, when it was granted one, is the paragraph
            # the row's rules are positioned in: the emitter's rule geometry may
            # then use the source-derived tab stops, and the paragraph's own
            # justification cannot move them.
            self._row_owns_line_context = bool(
                getattr(plan, "owns_line_context", False)
            )
            self._inherited_value_decoration_ranges = tuple(inherited_decoration_ranges)
            self.text(paragraph, raw, fragment.runs, fragment_slots, flow=item.flow)
            self._inherited_value_decoration_ranges = ()
            self._row_owns_line_context = False
            self.source_visual_line_emission_plans.append(plan.as_dict())
            paragraph = self._current_emission_paragraph or paragraph
        return removed_marker

    def _row_is_its_own_source_line(self, plan, element_row_count) -> bool:
        """Whether a source visual row must own its own Word paragraph.

        The row assembly exists so a row that carries positioned atoms can be
        built from its own origin instead of trusting continuous flow.  That is
        honest for a row the source *drew as its own line* - and it is also what
        the frozen horizontal contract requires of a wrapped row that carries a
        rule whose accepted geometry intent is ``EXACT_SOURCE_SPAN``.  Such a rule
        owes both of its source endpoints inside the frozen 2.0 pt tolerance, and
        continuous flow cannot reach them: the authority order puts the rule's own
        geometry above Word paragraph continuity, and a paragraph boundary with
        zero before/after spacing is the authorised remedy (never a ``<w:br/>``).

        A wrapped row that carries no exact-span rule stays in the element's single
        paragraph and is placed by flow, exactly as the accepted single-row form
        lines already are: turning the source's own wrapping into paragraph
        boundaries is visible, and the row's rules only owe their start anchor,
        which flow keeps.

        The evidence is the element's own source visual row count and the rule
        registry's accepted geometry intent - source geometry and the frozen
        contract, never a page, a rule id or a literal string.
        """

        if int(element_row_count or 1) <= 1:
            return True
        return self._row_owes_exact_source_span(plan)

    def _row_owes_exact_source_span(self, plan) -> bool:
        """Whether one of this row's own rules owes both of its source endpoints.

        The accepted geometry intent is derived from the rule's own transformation
        policy through the same function the registry annotation uses, so the row
        decision and the delivered registry can never disagree about which rules
        owe an exact source span.
        """

        for rule_id in getattr(plan, "source_rule_ids", None) or ():
            policy = self._rule_transformation_policy.get(str(rule_id))
            if source_rule_geometry_intent(policy) == "EXACT_SOURCE_SPAN":
                return True
        return False

    def _source_form_field_slot_ids(self, runs, text, slots) -> frozenset:
        """Slot ids whose own glyphs the source's rule is drawn beneath.

        ``PLACEHOLDER_UNDERLINE`` with full text occupancy is the source's own
        statement that the rule decorates the *placeholder's* characters, not a
        wider blank: at full occupancy the placeholder's glyphs are the whole
        span the source ruled, so its brackets are the source's form text and the
        value is substituted inside that frame.  ``TEXT_UNDERLINE`` with partial
        occupancy is the opposite - the rule is wider than the text it crosses,
        which is a blank to fill, and the accepted substitution replaces the
        whole placeholder.

        The rule must also *own* a slot: the registry's ``authoritative_slot_id``
        names the slot a rule governs, and a rule that governs no slot
        (``NO_AUTHORITATIVE_SLOT``) cannot be the reason to preserve any
        placeholder's frame, however its x-band happens to fall across one.  A
        rule that does own a slot decorates the whole span it covers, so every
        placeholder inside that span inherits the frame, not only the one the
        binding names.  Both the relation and the occupancy are read from the
        registered source rule, and the character range is resolved with the same
        mapping the decoration restorers use.  Nothing here is keyed to a page, a
        rule id, a case or a literal value.
        """

        if not slots or not runs or not text:
            return frozenset()
        bindings = getattr(self, "_rule_slot_binding", {}) or {}
        covered_ranges = []
        for entry in self.source_rule_registry:
            if str(entry.get("relation_type")) != "PLACEHOLDER_UNDERLINE":
                continue
            authoritative = (bindings.get(entry.get("source_rule_id")) or {}).get(
                "slot_id"
            )
            if authoritative is None:
                continue
            try:
                occupancy = float(entry.get("text_occupancy"))
            except (TypeError, ValueError):
                continue
            if occupancy < 1.0 - 1e-6:
                continue
            x0, x1 = entry.get("x0"), entry.get("x1")
            if x0 is None or x1 is None:
                continue
            page = entry.get("source_page")
            if page is None:
                continue
            y = float(entry.get("y") or 0.0)
            span = SimpleNamespace(bbox=(float(x0), y, float(x1), y))
            covered = self._covered_text_range_for_rule(span, runs, text)
            if covered is not None:
                # The rule's x-band is only meaningful on the page it was read
                # from: the same x-band on another page crosses different glyphs,
                # so a rule can never decorate a placeholder on a different page.
                covered_ranges.append((int(page), covered[0], covered[1]))
        if not covered_ranges:
            return frozenset()
        return frozenset(
            str(slot.slot_id)
            for slot in slots
            if slot.text_start is not None
            and slot.text_end is not None
            and getattr(slot, "source_page", None) is not None
            and any(
                page == int(slot.source_page)
                and start <= slot.text_start
                and slot.text_end <= end
                for page, start, end in covered_ranges
            )
        )

    def _source_rule_spans_for_emission_plan(self):
        """``(x0, x1, rule_id)`` of every registered rule, for atom provenance."""

        return [
            (entry.get("x0"), entry.get("x1"), entry.get("source_rule_id"))
            for entry in self.source_rule_registry
        ]

    def _fragment_visual_row_indexes(self, item, rows):
        """The source visual row each fragment of this element belongs to.

        A fragment's own run geometry is the key: the extracted lines are grouped
        by the same physical-row rule the source itself is grouped by, so the
        mapping never depends on a page, a rule id or literal text.
        """

        from tender_basic.page_layout import SOURCE_VISUAL_ROW_TOLERANCE

        if not rows:
            return [None] * len(item.fragments)
        indexes = []
        for fragment, _start, _end in item.fragments:
            top = min(
                (
                    float(run.bbox[1])
                    for run in fragment.runs
                    if getattr(run, "bbox", None)
                ),
                default=None,
            )
            if top is None:
                indexes.append(None)
                continue
            index = None
            for row_index, row in enumerate(rows):
                if abs(float(row[0].bbox[1]) - top) <= SOURCE_VISUAL_ROW_TOLERANCE:
                    index = row_index
                    break
            if index is None:
                index = min(
                    range(len(rows)),
                    key=lambda i: abs(float(rows[i][0].bbox[1]) - top),
                )
            indexes.append(index)
        return indexes

    def _assemble_source_visual_line(self, paragraph, plan, fragment):
        """Give one source visual line its own Word paragraph context.

        The row is opened *before* its first glyph is emitted, so the row's
        leading source text, its positioned rules and its trailing text are all
        written into the same paragraph, from the row's own origin.  The row then
        owns that paragraph: no atom on it opens a second one.
        """

        if getattr(self, "_form_line_paragraph_factory", None) is None:
            return paragraph
        positioned = plan.positioned_atoms
        if not positioned:
            return paragraph
        first = positioned[0]
        opened = self._start_source_form_line_paragraph(
            paragraph,
            fragment.runs[0] if fragment.runs else None,
            span=(first.source_x0, first.source_x1),
        )
        self._form_line_managed_paragraphs.add(id(opened))
        self._row_context_paragraph = opened
        plan.owns_line_context = True
        record = self._last_form_line_record()
        if record is not None:
            plan.generated_paragraph_index = record.get("generated_paragraph_index")
            record["source_rule_ids"] = plan.source_rule_ids
            record["source_visual_line_emission"] = plan.as_dict()
        return opened

    def _last_form_line_record(self):
        """The form-line paragraph record most recently opened for this element."""

        state = getattr(self, "_pending_form_line_records", None) or {}
        closed = state.get("closed") or []
        if closed:
            return closed[-1]
        if self.source_form_line_paragraphs:
            return self.source_form_line_paragraphs[-1]
        return None

    def build(self, output, generation_report_path=None):
        # SourceFormPlan compiler, glyph dimension: compile every physical rule's
        # covered source text once, from the parsed pages, before any rendering.
        # Rendering then consumes this plan instead of re-deriving rule evidence
        # from fragment-local state.
        from tender_basic.page_layout import (
            compile_source_form_field_plans,
            compile_source_rule_glyph_evidence,
        )
        from tender_basic.source_format import _slot_contract

        self._source_rule_glyph_evidence = compile_source_rule_glyph_evidence(
            self.template.source_pages
        )
        # Field-binding dimension of the same plan: glyph-free physical rules are
        # bound to their source form field through the existing deterministic
        # hint contract and real ProjectFacts status, before any rendering.
        self._rule_field_plans = compile_source_form_field_plans(
            source_pages=self.template.source_pages,
            glyph_evidence=self._source_rule_glyph_evidence,
            resolved_fact_fields=self._resolved_fact_fields(),
            slot_contract=_slot_contract,
        )
        if not self.template.source_pages:            raise ValueError("Source format has no pages")
        # V1 manual-layout closure: derive the document/section page frame from
        # repeated *source* geometry before any page renders.  The frame is a
        # property of the source section, not of a page's content bounding box.
        self.section_frames = derive_source_section_page_frames(
            self.template.source_pages, self.layouts
        )
        # The source page frame is now known, so the shared vertical rhythm is
        # re-sampled from the frame's own body top.  Sampling before this point
        # would model the page's first element from a different origin than the
        # emitter uses.
        self._prepare_source_vertical_rhythm()
        doc = self.style_pack.new_document()
        self.fixture_report = self.style_pack.apply_profile(doc, self.style_profile)
        # The generated list family is retained in the package for explicit
        # GENERATED_AUTO callers, but source paragraphs never reference it.
        self.body_list_num_id = self.style_pack.generated_body_list_num_id
        previous = None
        for page, layout in zip(self.template.source_pages, self.layouts):
            folded_into = self.logical_plan.head_page_for(page.page, 0)
            if folded_into not in (None, page.page) and not self._page_has_emittable_content(
                layout
            ):
                # This source page's content was folded into the logical table
                # emitted at the seam: its table fragment is a continuation whose
                # rows already render in the head fragment's section, and it
                # carries nothing else.  Opening a section for it would add a
                # source boundary Word has to fill with an empty page, so the
                # page contributes no section at all.  The test is the source
                # page's own emittable content, not whether it happens to carry a
                # table fragment - a continuation page always does.  The page's
                # rules stay in the provenance registry, because the registry
                # describes the source document rather than the emission
                # schedule.
                self._register_page_rules(page, layout)
                previous = (page.width, page.height)
                continue
            section = doc.sections[-1]
            if previous is not None:
                section = doc.add_section(WD_SECTION.NEW_PAGE)
            previous = (page.width, page.height)
            section.page_width, section.page_height = Pt(page.width), Pt(page.height)
            # V1 manual-layout closure: the section carries the *stable source
            # page frame*, never a bounding box measured on this one page.
            # Round 5.6 derived the margin from the page's own min/max content
            # x, which made a sparse title-only page acquire enormous margins
            # (the clamp at 30% of the page width) and made a table wider than
            # the body contract the text column.  Both are element-level
            # geometry problems: a source x that differs from the frame is
            # represented on the element (indent, tab stop, table indent,
            # positioned blank) and never by moving the section margin.
            frame = self.section_frames[page.page]
            self.section_geometries.append(frame.page_geometry(page.page))
            top = frame.top_body_frame
            section.top_margin = Pt(top)
            section.bottom_margin = Pt(frame.bottom_body_frame)
            section.left_margin = Pt(frame.left_margin)
            section.right_margin = Pt(frame.right_margin)
            self._active_source_page = page.page
            page_content_x0 = frame.left_margin
            page_content_x1 = frame.body_x1
            usable = frame.usable_text_width
            scale = self.page_scales.get(page.page, 1.0)
            # Round 5.8: classify this page's rules once.  The ones whose extent
            # is occupied by source text are underlines of that text and are
            # applied to the owning run; only gaps become positioned blanks.
            self._register_page_rules(page, layout)
            cursor_y = top
            last_paragraph = None
            for position, item in enumerate(layout.elements):
                # SHARED VERTICAL RHYTHM.  The gap above a source element is a
                # real source block boundary, so it is expressed as the source's
                # own shared spacing level rather than as this element's private
                # y offset.  The cursor is the bottom of the furthest content
                # already emitted: a form block whose rows end above its own
                # declared extent must not move it backwards, because that would
                # silently shrink the next real source boundary into a clamp.
                gap = self._page_vertical_space(
                    max(0, item.bbox[1] - cursor_y) * scale,
                    page=page.page,
                    position=position,
                )
                if isinstance(item, FormBlock):
                    last_paragraph = self.form_block(
                        doc, item, page_content_x0=page_content_x0,
                        page_content_x1=page_content_x1,
                        scale=scale, initial_space_before=gap,
                    )
                    cursor_y = max(cursor_y, item.bbox[3])
                    continue
                if not isinstance(item, LogicalParagraph):
                    source = item
                    if last_paragraph is not None:
                        last_paragraph.paragraph_format.space_after = Pt(gap)
                    if (source.page, source.table_index) in self.logical_plan.continuation_map:
                        # A PDF page fragment that continues an open logical
                        # table is not its own Word table; its rows are already
                        # folded into the head table emitted at the seam.
                        cursor_y = max(cursor_y, source.bbox[3])
                        last_paragraph = None
                        continue
                    logical = self.logical_plan.table_for(source.page, source.table_index)
                    # A table object has no ``space_before`` of its own in Word,
                    # so the source gap above it is carried by the preceding
                    # paragraph's ``space_after`` (set on the generic
                    # ``not isinstance(item, LogicalParagraph)`` path above).
                    # Applying the same gap again inside the table's first cell
                    # double-counts it and pushes the table up the page.
                    self.table(
                        doc,
                        logical if logical is not None else source,
                        usable,
                        page_content_x0=page_content_x0,
                    )
                    cursor_y = max(cursor_y, source.bbox[3])
                    last_paragraph = None
                    continue

                source = item.source
                paragraph, semantic = self._prepare_logical_paragraph(
                    doc, item, page, layout, gap, page_content_x0, scale,
                    page_content_x1=page_content_x1,
                )
                # SOURCE_FORM_LINE_PARAGRAPH activation (narrow): only for a
                # reconstructed paragraph that carries more than one source fill
                # rule, where each rule's own source visual form line needs
                # independent horizontal geometry.  Single-rule paragraphs such
                # as the proven P42-R4 path keep flowing as one paragraph.  A
                # paragraph whose later source visual rows carry positioned rules
                # is included too: those rows are assembled into their own
                # paragraph context, so the factory must exist for them.
                form_line_rules = self._form_line_rule_spans(item, layout)
                composition_line_rules = self._composition_form_line_rules(item, layout)
                if (
                    len(form_line_rules) > 1
                    or composition_line_rules
                    or self._later_source_row_rule_spans(item, layout)
                ):
                    self._activate_source_form_line_paragraphs(
                        doc, paragraph, item, semantic,
                        form_line_rules if len(form_line_rules) > 1 else composition_line_rules,
                    )
                removed_marker = self._render_logical_content(paragraph, item, layout, semantic)
                self._deactivate_source_form_line_paragraphs()
                if semantic.role == "BODY_LIST":
                    role = ParagraphLayoutRole.LIST_ITEM
                elif semantic.role == "HEADING":
                    role = ParagraphLayoutRole.COVER_TITLE if layout.classification in ("COVER_PAGE", "CHAPTER_DIVIDER") else ParagraphLayoutRole.HEADING
                elif item.kind == "Caption" or len(item.logical_text) <= 40 and item.logical_text.strip().startswith(("致", "收件人")):
                    role = ParagraphLayoutRole.ADDRESSEE
                else:
                    role = ParagraphLayoutRole.FLOW_BODY
                source_anchor = (
                    float(item.source_lines[0].bbox[0])
                    if semantic.role == "BODY_LIST" and item.source_lines
                    else item.bbox[0]
                )
                paragraph_layout = ParagraphLayout(
                    role=role,
                    container_x0=float(item.container.container_x0 if item.container else page_content_x0),
                    container_x1=float(item.container.container_x1 if item.container else page.width - page_content_x0),
                    source_x0=float(source_anchor), source_y0=float(item.bbox[1]),
                    alignment=item.alignment_hint,
                    left_indent_pt=float(paragraph.paragraph_format.left_indent.pt or 0.0),
                    right_indent_pt=float(paragraph.paragraph_format.right_indent.pt or 0.0),
                    first_line_indent_pt=float(paragraph.paragraph_format.first_line_indent.pt or 0.0),
                    space_before_pt=float(paragraph.paragraph_format.space_before.pt or 0.0),
                    space_after_pt=float(paragraph.paragraph_format.space_after.pt or 0.0),
                    line_spacing=float(paragraph.paragraph_format.line_spacing or 1.0),
                    tab_stops=[], runs=list(paragraph.runs), source_locator=source.locator,
                    source_page=page.page,
                )
                self.paragraph_layouts.append(paragraph_layout)
                # The delivered line origin.  A paragraph's *first* line carries
                # its own first-line offset, so the line that starts at the
                # source's first row is ``left_indent + first_line_indent``.
                # Measuring the left indent alone would report a first-line
                # indent as a horizontal error exactly as wide as itself.
                generated_x = page_content_x0 + (
                    float(paragraph.paragraph_format.left_indent.pt or 0.0)
                    + float(paragraph.paragraph_format.first_line_indent.pt or 0.0)
                )
                source_first_row_x = (
                    float(item.source_lines[0].bbox[0])
                    if item.source_lines
                    else float(item.bbox[0])
                )
                self.paragraph_x_errors.append(
                    abs(generated_x - source_first_row_x)
                    if item.alignment_hint != "center" else 0.0
                )
                self.paragraph_y_errors.append(0.0)
                self.logical_records.append({
                    "kind": "Heading" if semantic.role == "HEADING" else ("List" if semantic.role == "BODY_LIST" else item.kind),
                    "role": role.value,
                    "source_page": page.page,
                    "paragraph_index": len(doc.paragraphs) - 1,
                    "text": paragraph.text,
                    "source_text": item.logical_text,
                    "text_bbox": item.text_bbox,
                    "container_bbox": item.paragraph_container_bbox,
                    "source_lines": len(item.source_lines),
                    "fragment_count": len(item.fragments),
                    "source_baselines": sorted({round(line.bbox[1], 1) for line in item.source_lines}),
                    "alignment": item.alignment_hint,
                    "alignment_role": item.alignment_role.value,
                    "list_level": item.list_level,
                    "list_number_start_x": source_anchor if semantic.role == "BODY_LIST" else None,
                    "list_text_start_x": source_anchor if semantic.role == "BODY_LIST" else None,
                    "list_container_right": item.container.container_x1 if semantic.role == "BODY_LIST" and item.container else None,
                    "list_source_line_xs": [line.bbox[0] for line in item.source_lines] if semantic.role == "BODY_LIST" else [],
                    "list_group_key": (
                        page.page, item.list_level, item.list_group_id,
                        round(item.left_indent_pt, 1), round(item.first_line_indent_pt, 1),
                    ) if semantic.role == "BODY_LIST" else None,
                    "left_indent_pt": paragraph.paragraph_format.left_indent.pt or 0.0,
                    "first_line_indent_pt": paragraph.paragraph_format.first_line_indent.pt or 0.0,
                    "source_indent": (
                        item.source_indent.as_dict()
                        if getattr(item, "source_indent", None) is not None
                        else None
                    ),
                    "layout_role": role.value,
                    "container_x0": paragraph_layout.container_x0,
                    "container_x1": paragraph_layout.container_x1,
                    "source_x": source_anchor,
                    "generated_left_indent": paragraph_layout.left_indent_pt,
                    "source_y": item.bbox[1],
                    "space_before": paragraph_layout.space_before_pt,
                    "tab_stops": [],
                    "x_error_pt": self.paragraph_x_errors[-1],
                    "source_first_row_x": round(source_first_row_x, 4),
                    "generated_first_line_x": round(generated_x, 4),
                    "y_error_pt": 0.0,
                    "word_style": semantic.style_name,
                    "numbering_family": semantic.numbering_family,
                    "numbering_ownership": semantic.numbering_ownership.value,
                    "source_number_token": semantic.source_number_token.as_dict() if semantic.source_number_token else None,
                    "source_number_prefix": semantic.source_number_token.raw_prefix if semantic.source_number_token else "",
                    "source_number_separator_after_prefix": semantic.source_number_token.separator_after_prefix if semantic.source_number_token else "",
                    "source_number_confidence": semantic.source_number_token.confidence if semantic.source_number_token else None,
                    "classification_status": semantic.classification_status,
                    "heading_level": semantic.heading_level,
                    "body_list_level": semantic.list_level,
                    "literal_number_removed": removed_marker,
                    "semantic_evidence": semantic.evidence,
                })
                visual_rows = len({round(line.bbox[1], 0) for line in item.source_lines}) or 1
                cursor_y = item.bbox[1] + visual_rows * item.line_pitch_pt * scale
                last_paragraph = paragraph

        doc.core_properties.keywords = "format_source=SOURCE_DOCUMENT;compatibility_mode=WORD_SAFE;style_architecture=FIXTURE_BASED"
        doc.core_properties.subject = "招标文件格式继承;format_source=SOURCE_DOCUMENT;fixture_style_baseline"
        # A few paragraph-native form spacers are intentionally empty. They
        # remain real paragraphs for source flow, but must still participate in
        # the named style architecture rather than silently inheriting Normal.
        for paragraph in doc.paragraphs:
            if paragraph._p.pPr is None or paragraph._p.pPr.pStyle is None:
                paragraph.style = "Tender Body No Indent"
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        if paragraph._p.pPr is None or paragraph._p.pPr.pStyle is None:
                            paragraph.style = "Table Text"
        # Merged-cell continuation paragraphs are not exposed as distinct
        # python-docx cell proxies. Walk the story XML once so even those
        # empty editable table paragraphs receive the baseline table style.
        table_style_id = doc.styles["Table Text"].style_id
        for p_element in doc._element.body.iter(qn("w:p")):
            if p_element.getparent() is None or p_element.getparent().tag != qn("w:tc"):
                continue
            ppr = p_element.get_or_add_pPr()
            if ppr.find(qn("w:pStyle")) is None:
                pstyle = OxmlElement("w:pStyle")
                pstyle.set(qn("w:val"), table_style_id)
                ppr.insert(0, pstyle)
        if len(doc.tables) != len(self.logical_plan.source_tables):
            raise ValueError("Logical source table count changed before serialization")
        path = Path(output)
        path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(path)
        scan = scan_word_safe_docx(path)
        if scan["result"] != "PASS":
            raise ValueError(f"Word-safe package validation failed: {scan}")
        all_document_text = "".join(paragraph.text for paragraph in doc.paragraphs)
        all_document_text += "".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
        nbsp_only = int("\u00a0" in all_document_text and all_document_text.replace("\u00a0", "").strip() == "")
        blank_errors = sorted(self.blank_width_errors)
        # A composite source placeholder filled by adjacent slots is one logical
        # write event, so its applications collapse before serialization and
        # before any plan/application comparison runs.
        self._merge_adjacent_logical_applications()
        owner_metrics = execution_owner_metrics(self)
        application_metrics = plan_application_metrics(
            planned=self._planned_resolved_fields(),
            applications=getattr(self, "source_fill_applications", []),
        )
        report = {
            "architecture": "StyleFirstSourceDocumentBuilder",
            "compatibility_mode": "WORD_SAFE",
            "format_source": "SOURCE_DOCUMENT",
            "style_architecture": "SOURCE_PROFILE_NAMED_STYLE_BASELINE",
            "previous_docx_loaded": False,
            "settings_policy": "FRESH_PYTHON_DOCX_MINIMAL_DEPENDENCY_COMPLETE",
            "source_text_qa": self.source_text_qa,
            "source_heading": self.template.source_heading,
            "source_page_count": len(self.template.source_pages),
            "source_table_count": self.template.source_table_count,
            "source_real_tables": self.template.source_table_count,
            "pdf_table_fragment_count": len(self.logical_plan.head_keys) + len(self.logical_plan.continuation_map),
            "logical_table_count": len(self.logical_plan.source_tables),
            "logical_tables": [
                {
                    "table_id": logical.table_id,
                    "head_page": self.logical_plan.head_keys[index][0],
                    "head_table_index": self.logical_plan.head_keys[index][1],
                    "source_pages": logical.source_pages,
                    "fragment_count": len(logical.fragments),
                    "row_count": len(self.logical_plan.source_tables[index].rows),
                    "column_count": self.logical_plan.source_tables[index].columns,
                }
                for index, logical in enumerate(self.logical_plan.logical_tables)
            ],
            "continuation_fragments_merged": self.logical_plan.report.continuation_fragments_merged,
            "orphan_continuation_fragments": self.logical_plan.report.orphan_continuation_fragments,
            "false_continuation_merges": self.logical_plan.report.false_continuation_merges,
            "logical_table_decisions": self.logical_plan.report.decisions,
            "generated_table_count": len(doc.tables),
            "generated_textbox_count": 0,
            "generated_real_tables": len(doc.tables),
            "synthetic_layout_tables": 0,
            "form_layout_table_count": 0,
            "form_block_count": self.form_block_count,
            "paragraph_form_count": self.paragraph_form_count,
            "vector_blanks_recovered": self.vector_blanks_recovered,
            "editable_blanks": self.blank_records,
            "detected_editable_blanks": len(self.blank_records),
            "visible_rule_blanks": self.visible_rule_blank_count,
            "plain_empty_blanks": self.plain_empty_blank_count,
            "nbsp_only_placeholder_count": nbsp_only,
            "invisible_required_blank_count": 0,
            "generic_role_overrode_source_blank_kind_count": len(
                self.generic_role_overrode_source_blank_kind
            ),
            "generic_role_overrode_source_blank_kind": self.generic_role_overrode_source_blank_kind,
            "blank_kind_authority": self.blank_kind_authority,
            "blank_width_error_max": max(blank_errors, default=0.0),
            "blank_width_error_median": blank_errors[len(blank_errors) // 2] if blank_errors else 0.0,
            "inline_blank_count": self.inline_blank_count,
            "paragraph_layout_count": len(self.paragraph_layouts),
            "paragraph_layout_records": self.logical_records,
            "source_vertical_rhythm": self._source_vertical_rhythm_report(),
            "slot_value_presentations": self._slot_value_presentation_report(),
            "forward_reachable_emission_count": sum(
                1
                for plan in self.source_visual_line_emission_plans
                if plan.get("forward_reachable")
            ),
            "forward_reachable_emission_plans": [
                {
                    "source_page": plan.get("source_page"),
                    "source_line_identity": plan.get("source_line_identity"),
                    "source_row_index": plan.get("source_row_index"),
                    "source_rule_ids": plan.get("source_rule_ids"),
                    "source_y0": plan.get("source_y0"),
                    "reachability": "FORWARD_REACHABLE_SAME_ROW",
                    "evidence": (
                        "the row's positioned atoms are reached forward from the "
                        "row's own paragraph origin by their own source-derived "
                        "anchor tabs, so the source visual line stays one Word "
                        "paragraph"
                    ),
                }
                for plan in self.source_visual_line_emission_plans
                if plan.get("forward_reachable")
            ],
            "structural_isolation_rows": [
                {
                    "source_page": record.get("source_page"),
                    "source_rule_id": record.get("source_rule_id"),
                    "source_rule_ids": record.get("source_rule_ids"),
                    "generated_paragraph_index": record.get("generated_paragraph_index"),
                    "isolation_reason": record.get("isolation_reason"),
                    "space_before_pt": (
                        record.get("paragraph_format", {}) or {}
                    ).get("space_before_pt"),
                    "space_after_pt": (
                        record.get("paragraph_format", {}) or {}
                    ).get("space_after_pt"),
                }
                for record in self.source_form_line_paragraphs
            ],
            "structural_isolation_reason": (
                "STRUCTURAL_ISOLATION_REQUIRED_BY_RESOLVED_REFLOW"
            ),
            "paragraph_x_error_max": max(self.paragraph_x_errors, default=0.0),
            "paragraph_x_error_median": sorted(self.paragraph_x_errors)[len(self.paragraph_x_errors) // 2] if self.paragraph_x_errors else 0.0,
            "paragraph_y_error_max": max(self.paragraph_y_errors, default=0.0),
            "paragraph_y_error_median": sorted(self.paragraph_y_errors)[len(self.paragraph_y_errors) // 2] if self.paragraph_y_errors else 0.0,
            "tab_stop_usage": self.tab_stop_usage,
            "positioned_blank_records": self.positioned_blank_records,
            "positioned_blank_count": len(self.positioned_blank_records),
            "unreachable_positioned_blank_count": self.unreachable_positioned_blank_count,
            "flow_inline_blank_count": self.flow_inline_blank_count,
            "positioned_form_layout_breaks": self.positioned_form_layout_breaks,
            "positioned_coordinate_frames": self._positioned_coordinate_frames,
            "source_form_line_paragraphs": self.source_form_line_paragraphs,            "source_form_line_paragraph_count": len(self.source_form_line_paragraphs),
            "source_visual_line_emission_plans": self.source_visual_line_emission_plans,
            "source_visual_line_emission_plan_count": len(self.source_visual_line_emission_plans),
            "source_visual_line_context_count": sum(
                1 for plan in self.source_visual_line_emission_plans
                if plan.get("owns_line_context")
            ),
            "source_visual_line_assembly_gaps": self.source_visual_line_assembly_gaps,
            "source_anchored_value_runs": self.source_anchored_value_runs,
            "source_rule_compositions": self.source_rule_compositions,
            "rule_compositions_created": self.rule_compositions_created,
            "source_anchored_value_run_count": len(self.source_anchored_value_runs),
            "rule_source_visual_lines": self._rule_source_visual_line,
            "source_rule_registry": self._annotated_source_rule_registry(),
            "source_form_layout_line_count": len(self.positioned_form_layout_breaks),
            "form_layout_break_count": len(self.positioned_form_layout_breaks),
            "fill_slots_detected": len(self.template.fill_slots),
            "fill_slots_filled": len(self.filled),
            "fill_slots_left_blank": len(self.template.fill_slots) - len(self.filled),
            "source_fill_applications": [
                        item.as_dict() if hasattr(item, "as_dict") else item
                        for item in getattr(self, "source_fill_applications", [])
                    ],
                    "source_fill_application_count": len(
                        getattr(self, "source_fill_applications", [])
                    ),
                    "source_form_execution_owners": [
                        item.as_dict() if hasattr(item, "as_dict") else item
                        for item in getattr(self, "source_form_execution_owners", [])
                    ],
                    "source_form_execution_owner_count": len(
                        getattr(self, "source_form_execution_owners", [])
                    ),
                    "source_form_execution_owner_metrics": owner_metrics,
                    "eligible_value_owner_count": owner_metrics[
                        "eligible_value_owner_count"
                    ],
                    "multiple_eligible_owner_conflicts": owner_metrics[
                        "multiple_eligible_owner_conflicts"
                    ],
                    "source_form_line_value_runs": self.source_form_line_value_runs,
                    "source_form_line_value_run_count": len(
                        self.source_form_line_value_runs
                    ),
                    "plan_application_metrics": application_metrics,
                    "planned_resolved_application_count": application_metrics[
                        "planned_resolved_application_count"
                    ],
                    "actual_resolved_application_count": application_metrics[
                        "actual_resolved_application_count"
                    ],
                    "logical_application_record_count": application_metrics[
                        "logical_application_record_count"
                    ],
                    "missing_planned_application_count": application_metrics[
                        "missing_planned_application_count"
                    ],
                    "unexpected_application_count": application_metrics[
                        "unexpected_application_count"
                    ],
                    "application_fact_mismatch_count": application_metrics[
                        "application_fact_mismatch_count"
                    ],
                    "application_value_mismatch_count": application_metrics[
                        "application_value_mismatch_count"
                    ],
                    "application_field_mismatch_count": application_metrics[
                        "application_field_mismatch_count"
                    ],
                    "planned_resolved_fact_count": application_metrics[
                        "planned_resolved_fact_count"
                    ],
                    "actual_covered_resolved_fact_count": application_metrics[
                        "actual_covered_resolved_fact_count"
                    ],
                    "logical_application_merges": self.logical_application_merges,
                    "owner_driven_value_emission_count": self.tab_stop_usage.get(
                        "owner_driven_value_emissions", 0
                    ),
                    # The compiled SourceFormPlan, serialized for QA/debug only.
                    # It is compiled before rendering and read back by no
                    # production code path.
                    "source_form_field_plans": self._compiled_field_plan_records(),
                    "source_form_field_plan_count": len(
                        getattr(self, "_rule_field_plans", None) or {}
                    ),
                    "filled_slots": self.filled,
            "table_style_outliers": self.table_style_outliers,
            "table_style_outlier_count": len(self.table_style_outliers),
            "same_role_style_discontinuities": sum(1 for item in self.table_style_outliers if item.get("resolution") == "PRESERVED_UNEXPLAINED_VARIATION"),
            "semantic_slot_corrections": self.semantic_slot_corrections,
            "raised_glyph_count": sum(len(cell.raised_glyphs) for page in self.template.source_pages for table in page.tables for row in table.rows for cell in row.cells),
            "layout_mode": "GEOMETRY_AWARE_STYLE_FIRST",
            "logical_paragraph_records": self.logical_records,
            "page_layouts": [{"source_page": p.source_page, "classification": p.classification, "content_box": p.content_box} for p in self.layouts],
            "artifact_filter_report": self.template.artifact_filter_report,
            "page_scales": self.page_scales,
            "layout_warnings": self.warnings,
            "font_substitutions": self.font_substitutions,
            "font_repairs": [],
            "word_open_status": "PENDING_MANUAL_CONFIRMATION",
            "word_safe_scan": scan,
            "table_geometry_records": self.table_geometry_records,
            "style_profile": self.style_profile.as_dict(),
            "fixture_integrity": fixture_integrity_report(self.style_pack, self.style_profile),
            "fixture_body_list_num_id": self.body_list_num_id,
            "generated_heading_num_id": self.style_pack.generated_heading_num_id,
            "generated_body_list_num_id": self.style_pack.generated_body_list_num_id,
            "numbering_strategy": "SOURCE_FORMAT_FIRST_STYLE_SYSTEM_SECOND_AUTOMATION_THIRD",
            "source_literal_count": sum(
                1 for record in self.logical_records
                if record.get("numbering_ownership") == "SOURCE_LITERAL"
            ),
            "generated_auto_count": sum(
                1 for record in self.logical_records
                if record.get("numbering_ownership") == "GENERATED_AUTO"
            ),
            "source_literal_records_have_no_auto_numbering": True,
            "numbering_exception_reason": self.numbering_exception_reason,
            # Round 5.6: the source-visual models the emitter actually consumed.
            "source_visual_typography": {
                "status": self.visual_typography.get("status"),
                "reason": self.visual_typography.get("reason", ""),
                "baseline_weight": self.visual_typography.get("baseline_weight"),
                "roles": self.visual_typography.get("roles", {}),
                "samples": self.visual_typography.get("samples", []),
                "source_pdf": self.visual_typography.get("source_pdf", ""),
            },
            "source_fill_patterns": [pattern.as_dict() for pattern in self.fill_patterns],
            "source_fill_pattern_index": {
                key: value.as_dict() for key, value in self.fill_pattern_index.items()
            },
            "fill_pattern_usage": self.fill_pattern_usage,
            "section_geometry": self.section_geometries,
            "blank_representation_usage": dict(self.tab_stop_usage),
            "blank_representation_records": self.blank_records,
            # Per-logical-table provenance: which Word table each source logical
            # table became, and the cell-level account of what the emission lost,
            # duplicated or invented.  The pairing is read back from the saved
            # package's ``w:body`` document order, never from emission intent.
            "logical_table_provenance": self._logical_table_provenance(doc),
        }
        if generation_report_path:
            Path(generation_report_path).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        return path, report


    def _logical_table_provenance(self, doc):
        """Per-logical-table provenance and cell accounting for this build."""

        from .logical_table_provenance import provenance_for_plan, source_rows_for_plan

        return provenance_for_plan(
            doc,
            self.logical_plan,
            source_rows=source_rows_for_plan(self.logical_plan),
            fill_applications=getattr(self, "source_fill_applications", ()),
        )


__all__ = ["StyleFirstSourceDocumentBuilder"]
