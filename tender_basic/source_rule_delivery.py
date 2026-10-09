"""Source-rule delivery ledger: a source-visible rule is first-class evidence.

SOURCE TEXT COMPLETENESS AND SOURCE BLANK COMPLETENESS ARE DISTINCT.  A document
can carry every source character and still lose every rule the source drew: a form
rule may own no text at all, and a text-completeness ledger cannot see it.  This
module tracks source-visible rules instead, from the source's own drawing
primitives to the saved Word artifact.

SOURCE RULES HAVE FIRST-CLASS GEOMETRIC IDENTITY.  A rule's identity is its own
stroke geometry on its own page - never a label, and never an extractor row
number.  The generator's rule registry is evidence about physical strokes, not a
list of physical strokes: the same drawn line can be described more than once
(``EXTRACTION_ALIAS``), a degenerate entry can describe no line at all
(``ZERO_WIDTH_NON_RULE``), and a rule the source uses as a separator is not a
fillable field (``DECORATIVE_RULE``).  Normalization comes first, because a raw
registry count is not a physical rule count.

RULE EXTRACTION ALIASES DO NOT REQUIRE DUPLICATE DELIVERY.  Two registry entries
that describe one physical line need one delivered line between them; two
genuinely distinct lines may share coordinates in x and still be distinct in y,
and they need two.

PARAGRAPH FORM RULE DISCOVERY MUST NOT DEPEND ON LABEL REGEX.  Label recognition
supplies a field's *meaning* and its expected editing behaviour; it is never the
gate that decides whether the source drew a rule, because a rule whose label the
vocabulary has never seen would then be dropped without a record.

Nothing here is case-, page- or label-specific: every classification is derived
from measured source geometry and the artifact's own records.
"""

from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from pathlib import Path

SCHEMA = "source_rule_delivery/1"

# --------------------------------------------------------------------------- #
# Rule occurrence classes (one raw registry entry)
# --------------------------------------------------------------------------- #

#: A real drawn line that the delivery scope can carry as a field.
PHYSICAL_VISIBLE_RULE = "PHYSICAL_VISIBLE_RULE"

#: Another description of a line already described by a canonical entry.
EXTRACTION_ALIAS = "EXTRACTION_ALIAS"

#: A degenerate entry that describes no line (collapsed span).
ZERO_WIDTH_NON_RULE = "ZERO_WIDTH_NON_RULE"

#: A line the source uses as a separator, not as a fillable field.
DECORATIVE_RULE = "DECORATIVE_RULE"

#: A line drawn under source text: that text's own underline.
TEXT_UNDERLINE = "TEXT_UNDERLINE"

#: A line whose own class cannot be established from the available evidence.
UNRESOLVED_CLASSIFICATION = "UNRESOLVED_CLASSIFICATION"

# --------------------------------------------------------------------------- #
# Delivery states (one canonical physical rule)
# --------------------------------------------------------------------------- #

DELIVERED_VISIBLE_EDITABLE = "DELIVERED_VISIBLE_EDITABLE"
DELIVERED_STRUCTURALLY_BUT_NOT_VISIBLE = "DELIVERED_STRUCTURALLY_BUT_NOT_VISIBLE"
NOT_DELIVERED = "NOT_DELIVERED"
NOT_APPLICABLE_WITH_SOURCE_EVIDENCE = "NOT_APPLICABLE_WITH_SOURCE_EVIDENCE"
EVIDENCE_UNAVAILABLE = "EVIDENCE_UNAVAILABLE"

#: The stages every canonical rule is tracked through.
STAGES = (
    "STAGE_0_SOURCE_RULE",
    "STAGE_1_SOURCE_FORMAT",
    "STAGE_2_PAGE_LAYOUT",
    "STAGE_3_WORD_EMISSION_PLAN",
    "STAGE_4_IN_MEMORY_DOCX",
    "STAGE_5_SAVED_REOPENED_DOCX",
    "STAGE_6_RENDERER_OBSERVATION",
)

#: Below this a registry entry describes no line at all.
MIN_RULE_WIDTH_PT = 1.0

#: Stroke geometry agreement between a registry entry and a PDF drawing.
STROKE_Y_TOLERANCE_PT = 1.6
STROKE_X_TOLERANCE_PT = 2.5

#: A rule spanning this fraction of its page is a separator, not a field.
DECORATIVE_PAGE_WIDTH_RATIO = 0.5

#: Rule relation types that are a *field* rather than text decoration.
FILL_RELATIONS = frozenset({"GAP_FILL_RULE"})

#: Relation types that are decoration of source text.
TEXT_UNDERLINE_RELATIONS = frozenset(
    {"TEXT_UNDERLINE", "VALUE_UNDERLINE", "PLACEHOLDER_UNDERLINE"}
)

#: A visible rule must cover at least this fraction of its own source width with
#: painted glyph advances.  A single space whose width comes from character
#: spacing carries the underline property but paints a stub.
MIN_VISIBLE_COVERAGE_RATIO = 0.9


@dataclass(frozen=True)
class SourceRuleAtom:
    """One canonical physical source rule, with its alias provenance."""

    canonical_rule_id: str
    source_rule_id: str
    raw_rule_aliases: tuple[str, ...]
    source_page: int
    source_x0: float
    source_x1: float
    source_y: float
    source_width_pt: float
    occurrence_class: str
    relation_type: str
    transformation_policy: str
    stroke_evidence: str
    source_container_id: str | None = None
    source_label: str | None = None
    authoritative_slot_id: str | None = None
    authoritative_slot_binding: str | None = None
    resolved_fact_fields: tuple[str, ...] = ()
    allowed_fact_fields: tuple[str, ...] = ()

    @property
    def value_owned(self) -> bool:
        """Whether the source slot this rule belongs to resolves to a fact value.

        RULE OWNERSHIP MUST SURVIVE SOURCE-FILL FACT RESOLUTION.  A rule the
        registry binds to a source slot that carries resolved ProjectFacts fields
        is not delivered as an empty blank: the *resolved value* occupies it, and
        that value's own underlined run is the rule's Word owner.  The binding is
        declared by the registry itself (slot id plus fact fields), so this is
        provenance, not a guess from geometry or from a label.
        """

        return bool(self.resolved_fact_fields) and bool(self.authoritative_slot_id)

    @property
    def owner_kind(self) -> str:
        if not self.applicable:
            return "NOT_APPLICABLE_SOURCE_BACKED"
        if self.value_owned:
            return "SOURCE_RULE_OWNED_BY_RESOLVED_VALUE"
        return "INLINE_GAP_BLANK"

    @property
    def applicable(self) -> bool:
        """Whether this rule is a field the delivery must carry."""

        return self.occurrence_class == PHYSICAL_VISIBLE_RULE

    @property
    def requires_visible_rule(self) -> bool:
        return self.applicable

    @property
    def requires_editable_blank(self) -> bool:
        return self.applicable

    def as_dict(self) -> dict:
        return {
            "canonical_rule_id": self.canonical_rule_id,
            "source_rule_id": self.source_rule_id,
            "raw_rule_aliases": list(self.raw_rule_aliases),
            "source_page": self.source_page,
            "source_container_id": self.source_container_id,
            "source_x0": round(self.source_x0, 2),
            "source_x1": round(self.source_x1, 2),
            "source_y": round(self.source_y, 2),
            "source_width_pt": round(self.source_width_pt, 2),
            "occurrence_class": self.occurrence_class,
            "relation_type": self.relation_type,
            "transformation_policy": self.transformation_policy,
            "stroke_evidence": self.stroke_evidence,
            "source_label": self.source_label,
            "authoritative_slot_id": self.authoritative_slot_id,
            "authoritative_slot_binding": self.authoritative_slot_binding,
            "resolved_fact_fields": list(self.resolved_fact_fields),
            "allowed_fact_fields": list(self.allowed_fact_fields),
            "owner_kind": self.owner_kind,
            "value_owned": self.value_owned,
            "applicable": self.applicable,
            "requires_visible_rule": self.requires_visible_rule,
            "requires_editable_blank": self.requires_editable_blank,
        }


@dataclass
class RuleStageRecord:
    """What one stage of the delivery dataflow knows about one rule."""

    stage: str
    present: bool
    owner: str | None = None
    construct_kind: str | None = None
    mechanism: str | None = None
    generated_width_pt: float | None = None
    reason: str | None = None

    def as_dict(self) -> dict:
        return {
            "stage": self.stage,
            "present": self.present,
            "owner": self.owner,
            "construct_kind": self.construct_kind,
            "mechanism": self.mechanism,
            "generated_width_pt": self.generated_width_pt,
            "reason": self.reason,
        }


@dataclass
class SourceRuleDeliveryRecord:
    """One canonical rule's full journey and its verdict."""

    atom: SourceRuleAtom
    stages: list[RuleStageRecord] = field(default_factory=list)
    verdict: str = EVIDENCE_UNAVAILABLE
    first_missing_stage: str | None = None
    word_owner: str | None = None
    docx_evidence: list[str] = field(default_factory=list)
    generated_width_pt: float | None = None
    visible_coverage_ratio: float | None = None
    editable: bool | None = None
    owner_paragraph: int | None = None
    matched_advance_pt: float | None = None
    reasons: list[str] = field(default_factory=list)

    def as_dict(self) -> dict:
        return {
            **self.atom.as_dict(),
            "stages": [stage.as_dict() for stage in self.stages],
            "verdict": self.verdict,
            "first_missing_stage": self.first_missing_stage,
            "word_owner": self.word_owner,
            "docx_evidence": list(self.docx_evidence),
            "generated_width_pt": (
                None
                if self.generated_width_pt is None
                else round(float(self.generated_width_pt), 3)
            ),
            "visible_coverage_ratio": (
                None
                if self.visible_coverage_ratio is None
                else round(float(self.visible_coverage_ratio), 4)
            ),
            "editable": self.editable,
            "reasons": list(self.reasons),
        }


@dataclass
class SourceRuleDeliveryLedger:
    """Every canonical rule of one build, with its counters and verdict."""

    records: list[SourceRuleDeliveryRecord] = field(default_factory=list)
    source_pdf: str | None = None
    source_pdf_available: bool = False
    raw_rule_count: int = 0
    extraneous_rules: list[dict] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)

    # -- counters ---------------------------------------------------------- #

    @property
    def raw_source_rule_count(self) -> int:
        return self.raw_rule_count

    @property
    def canonical_physical_rule_count(self) -> int:
        return len(
            [
                record
                for record in self.records
                if record.atom.occurrence_class == PHYSICAL_VISIBLE_RULE
            ]
        )

    @property
    def extraction_alias_count(self) -> int:
        return sum(len(record.atom.raw_rule_aliases) for record in self.records)

    @property
    def zero_width_noise_count(self) -> int:
        return len(
            [
                record
                for record in self.records
                if record.atom.occurrence_class == ZERO_WIDTH_NON_RULE
            ]
        )

    @property
    def decorative_rule_count(self) -> int:
        return len(
            [
                record
                for record in self.records
                if record.atom.occurrence_class == DECORATIVE_RULE
            ]
        )

    @property
    def text_underline_count(self) -> int:
        return len(
            [
                record
                for record in self.records
                if record.atom.occurrence_class == TEXT_UNDERLINE
            ]
        )

    @property
    def applicable_rule_count(self) -> int:
        return len([record for record in self.records if record.atom.applicable])

    @property
    def delivered_rule_count(self) -> int:
        return len(
            [
                record
                for record in self.records
                if record.atom.applicable
                and record.verdict
                in (DELIVERED_VISIBLE_EDITABLE, DELIVERED_STRUCTURALLY_BUT_NOT_VISIBLE)
            ]
        )

    @property
    def visible_rule_count(self) -> int:
        return len(
            [
                record
                for record in self.records
                if record.atom.applicable and record.verdict == DELIVERED_VISIBLE_EDITABLE
            ]
        )

    @property
    def missing_rule_count(self) -> int:
        return len(
            [
                record
                for record in self.records
                if record.atom.applicable and record.verdict == NOT_DELIVERED
            ]
        )

    @property
    def invisible_rule_count(self) -> int:
        return len(
            [
                record
                for record in self.records
                if record.atom.applicable
                and record.verdict == DELIVERED_STRUCTURALLY_BUT_NOT_VISIBLE
            ]
        )

    @property
    def unresolved_rule_count(self) -> int:
        """Rules whose own *class* could not be established from source evidence."""

        return len(
            [
                record
                for record in self.records
                if record.atom.occurrence_class == UNRESOLVED_CLASSIFICATION
            ]
        )

    @property
    def evidence_unavailable_rule_count(self) -> int:
        """Rules whose class is known but whose delivery could not be measured."""

        return len(
            [
                record
                for record in self.records
                if record.atom.occurrence_class != UNRESOLVED_CLASSIFICATION
                and record.verdict == EVIDENCE_UNAVAILABLE
            ]
        )

    @property
    def duplicate_delivery_count(self) -> int:
        return sum(
            1
            for record in self.records
            if "DUPLICATE_DELIVERY" in record.reasons
            or "DUPLICATE_DELIVERY_OWNER" in record.reasons
        )

    @property
    def slot_ownership_violation_count(self) -> int:
        """Bidder-owned slots carrying content the source never authorised."""

        return sum(
            1 for record in self.records if "SOURCE_SLOT_OWNERSHIP" in record.reasons
        )

    @property
    def wrong_owner_count(self) -> int:
        return sum(1 for record in self.records if "WRONG_OWNER" in record.reasons)

    @property
    def not_editable_rule_count(self) -> int:
        """Delivered lines that are not native editable Word fields."""

        return sum(
            1 for record in self.records if "SOURCE_FORM_EDITABILITY" in record.reasons
        )

    @property
    def extraneous_rule_count(self) -> int:
        """Painted lines no canonical applicable source rule accounts for."""

        return len(self.extraneous_rules)

    def counters(self) -> dict:
        return {
            "raw_source_rule_count": self.raw_source_rule_count,
            "canonical_physical_rule_count": self.canonical_physical_rule_count,
            "extraction_alias_count": self.extraction_alias_count,
            "zero_width_noise_count": self.zero_width_noise_count,
            "decorative_rule_count": self.decorative_rule_count,
            "text_underline_count": self.text_underline_count,
            "applicable_rule_count": self.applicable_rule_count,
            "delivered_rule_count": self.delivered_rule_count,
            "visible_rule_count": self.visible_rule_count,
            "missing_rule_count": self.missing_rule_count,
            "invisible_rule_count": self.invisible_rule_count,
            "duplicate_delivery_count": self.duplicate_delivery_count,
            "wrong_owner_count": self.wrong_owner_count,
            "slot_ownership_violation_count": self.slot_ownership_violation_count,
            "not_editable_rule_count": self.not_editable_rule_count,
            "extraneous_rule_count": self.extraneous_rule_count,
            "unresolved_rule_count": self.unresolved_rule_count,
            "evidence_unavailable_rule_count": self.evidence_unavailable_rule_count,
        }

    @property
    def verdict(self) -> str:
        if not self.source_pdf_available:
            return "SOURCE_BLANK_COMPLETENESS = EVIDENCE_UNAVAILABLE"
        if self.unresolved_rule_count:
            return "SOURCE_BLANK_COMPLETENESS = FAIL_UNRESOLVED_RULES"
        if self.evidence_unavailable_rule_count:
            return "SOURCE_BLANK_COMPLETENESS = FAIL_EVIDENCE_UNAVAILABLE"
        if (
            self.missing_rule_count
            or self.invisible_rule_count
            or self.duplicate_delivery_count
            or self.wrong_owner_count
            or self.not_editable_rule_count
            or self.extraneous_rule_count
            or self.slot_ownership_violation_count
        ):
            return "SOURCE_BLANK_COMPLETENESS = FAIL"
        return "SOURCE_BLANK_COMPLETENESS = PASS"

    def as_dict(self) -> dict:
        return {
            "schema": SCHEMA,
            "source_pdf": self.source_pdf,
            "source_pdf_available": self.source_pdf_available,
            "verdict": self.verdict,
            "counters": self.counters(),
            "records": [record.as_dict() for record in self.records],
            "notes": list(self.notes),
        }


# --------------------------------------------------------------------------- #
# Source-side evidence
# --------------------------------------------------------------------------- #


def _rule_key(page: int, x0: float, x1: float, y: float) -> tuple:
    """A physical stroke's own identity: page plus its measured geometry."""

    return (int(page), round(float(y), 2), round(float(x0), 2), round(float(x1), 2))


def source_page_strokes(document, page_number: int) -> list[tuple[float, float, float]]:
    """Every horizontal stroke on one source page as ``(y, x0, x1)``.

    The stroke list is the source's own drawing evidence, so a registry entry can
    be checked against a real primitive rather than being trusted as one.
    """

    strokes: list[tuple[float, float, float]] = []
    page = document[page_number - 1]
    for drawing in page.get_drawings():
        for item in drawing["items"]:
            if item[0] == "l":
                first, second = item[1], item[2]
                if abs(first.y - second.y) <= 1.0 and abs(second.x - first.x) >= 3.0:
                    strokes.append(
                        (float(first.y), min(first.x, second.x), max(first.x, second.x))
                    )
            elif item[0] == "re":
                rect = item[1]
                if rect.height <= 2.5 and rect.width >= 3.0:
                    strokes.append(
                        (float((rect.y0 + rect.y1) / 2.0), float(rect.x0), float(rect.x1))
                    )
    return strokes


def stroke_matches(
    strokes, x0: float, x1: float, y: float
) -> list[tuple[float, float, float]]:
    return [
        stroke
        for stroke in strokes
        if abs(stroke[0] - y) <= STROKE_Y_TOLERANCE_PT
        and abs(stroke[1] - x0) <= STROKE_X_TOLERANCE_PT
        and abs(stroke[2] - x1) <= STROKE_X_TOLERANCE_PT
    ]


def source_label_for_rule(document, page_number: int, x0: float, y: float) -> str | None:
    """The source text immediately left of a rule on its own visual row.

    The label is read from the source's own glyphs so the ledger can name a field
    without a label vocabulary - and, critically, without using the vocabulary to
    decide whether the field exists.
    """

    page = document[page_number - 1]
    best: tuple[float, str] | None = None
    for block in page.get_text("dict")["blocks"]:
        for line in block.get("lines", []):
            bbox = line["bbox"]
            if not (bbox[1] - 3.0 <= y <= bbox[3] + 4.0):
                continue
            if bbox[2] > x0 + 1.5:
                continue
            text = "".join(span["text"] for span in line["spans"]).strip()
            if not text:
                continue
            if best is None or bbox[2] > best[0]:
                best = (float(bbox[2]), text)
    return None if best is None else best[1]


# --------------------------------------------------------------------------- #
# The ledger
# --------------------------------------------------------------------------- #


def normalize_source_rules(report: dict, document=None) -> list[SourceRuleAtom]:
    """Normalize the generator's rule registry into canonical physical rules."""

    entries = list(report.get("source_rule_registry") or [])
    page_widths: dict[int, float] = {}
    strokes_by_page: dict[int, list] = {}
    if document is not None:
        for entry in entries:
            page_number = int(entry["source_page"])
            if page_number not in strokes_by_page:
                strokes_by_page[page_number] = source_page_strokes(document, page_number)
                page_widths[page_number] = float(document[page_number - 1].rect.width)

    groups: dict[tuple, list[dict]] = {}
    order: list[tuple] = []
    for entry in sorted(
        entries, key=lambda item: (item["source_page"], item["y"], item["x0"])
    ):
        key = _rule_key(entry["source_page"], entry["x0"], entry["x1"], entry["y"])
        if key not in groups:
            groups[key] = []
            order.append(key)
        groups[key].append(entry)

    atoms: list[SourceRuleAtom] = []
    for key in order:
        members = groups[key]
        primary = members[0]
        page_number, y, x0, x1 = key
        width = float(x1) - float(x0)
        relation = str(primary.get("relation_type") or "")
        aliases = tuple(
            str(member.get("source_rule_id")) for member in members[1:]
        )
        # Physical stroke evidence: does the source actually draw this line?
        if document is None:
            occurrence = UNRESOLVED_CLASSIFICATION
            evidence = "no source document was available, so the stroke could not be verified"
        else:
            matches = stroke_matches(
                strokes_by_page.get(page_number, []), x0, x1, y
            )
            if width < MIN_RULE_WIDTH_PT:
                occurrence = ZERO_WIDTH_NON_RULE
                evidence = (
                    f"the entry's own span is {width:.2f} pt wide, which describes no line"
                )
            elif relation in TEXT_UNDERLINE_RELATIONS:
                # A line drawn under source text is that text's own decoration.  Its
                # evidence is the text it sits beneath, not the extractor's stroke
                # list - which is why it is classified from the relation the source
                # owner established rather than from a geometry re-match.
                occurrence = TEXT_UNDERLINE
                evidence = (
                    f"relation={relation} with source text occupancy "
                    f"{float(primary.get('text_occupancy') or 0.0):.3f}"
                )
            elif relation == "FORM_LAYOUT_RULE" and (
                width >= page_widths.get(page_number, 0.0) * DECORATIVE_PAGE_WIDTH_RATIO
                or float(primary.get("text_occupancy") or 0.0) >= 0.3
            ):
                occurrence = DECORATIVE_RULE
                evidence = (
                    f"a {width:.2f} pt layout rule carrying "
                    f"{float(primary.get('text_occupancy') or 0.0):.3f} source text "
                    "occupancy: the source uses it as a separator, not a field"
                )
            elif relation in FILL_RELATIONS:
                if matches:
                    occurrence = PHYSICAL_VISIBLE_RULE
                    evidence = (
                        f"{len(matches)} source stroke(s) at y={matches[0][0]:.2f} "
                        f"x={matches[0][1]:.2f}..{matches[0][2]:.2f}"
                    )
                else:
                    # A fill rule with no drawn primitive behind it is exactly the
                    # ambiguity that must block: the registry claims a field the
                    # source's own drawings cannot confirm.
                    occurrence = UNRESOLVED_CLASSIFICATION
                    evidence = (
                        "the registry records a gap-fill rule but no source drawing "
                        "primitive matches its geometry"
                    )
            else:
                occurrence = UNRESOLVED_CLASSIFICATION
                evidence = (
                    f"relation type {relation!r} has no established class for this "
                    "evidence set"
                )
        atoms.append(
            SourceRuleAtom(
                canonical_rule_id=f"R{page_number}-{y:.2f}-{x0:.2f}-{x1:.2f}",
                source_rule_id=str(primary.get("source_rule_id")),
                raw_rule_aliases=aliases,
                source_page=page_number,
                source_x0=float(x0),
                source_x1=float(x1),
                source_y=float(y),
                source_width_pt=width,
                occurrence_class=occurrence,
                relation_type=relation,
                transformation_policy=str(primary.get("transformation_policy") or ""),
                stroke_evidence=evidence,
                source_label=(
                    source_label_for_rule(document, page_number, float(x0), float(y))
                    if document is not None
                    else None
                ),
                authoritative_slot_id=(
                    None
                    if primary.get("authoritative_slot_id") in (None, "")
                    else str(primary.get("authoritative_slot_id"))
                ),
                authoritative_slot_binding=(
                    None
                    if primary.get("authoritative_slot_binding") in (None, "")
                    else str(primary.get("authoritative_slot_binding"))
                ),
                resolved_fact_fields=tuple(
                    str(field) for field in (primary.get("resolved_fact_fields") or ())
                ),
                allowed_fact_fields=tuple(
                    str(field) for field in (primary.get("allowed_fact_fields") or ())
                ),
            )
        )
    return atoms


def _delivered_index(report: dict) -> dict:
    """Delivered blanks and positioned emissions keyed by page."""

    blanks: dict[int, list[dict]] = {}
    for record in report.get("editable_blanks") or []:
        locator = record.get("source_locator") or {}
        page = locator.get("page")
        if page is None:
            continue
        blanks.setdefault(int(page), []).append(
            {
                "source": "editable_blanks",
                "x0": float(record.get("source_x0") or 0.0),
                "x1": float(record.get("source_x1") or 0.0),
                "representation_kind": record.get("representation_kind"),
                "render_style": record.get("render_style"),
                "rendered_width_pt": record.get("rendered_width_pt"),
                "visible": bool(record.get("visible")),
                "block_index": locator.get("block_index"),
                "semantic_slot": record.get("semantic_slot"),
                "mechanism": None,
                "consumed": False,
            }
        )
    for record in report.get("positioned_blank_records") or []:
        page = record.get("source_page")
        if page is None:
            continue
        blanks.setdefault(int(page), []).append(
            {
                "source": "positioned_blank_records",
                "x0": float(record.get("source_x0") or 0.0),
                "x1": float(record.get("source_x1") or 0.0),
                "representation_kind": record.get("representation_kind"),
                "render_style": None,
                "rendered_width_pt": record.get("source_span_pt"),
                "visible": True,
                "block_index": None,
                "semantic_slot": record.get("semantic_slot"),
                "mechanism": record.get("emission_mechanism"),
                "consumed": False,
            }
        )
    return blanks


def _span_tolerance(width: float) -> float:
    return max(1.5, min(3.0, width * 0.02))


def _consume_match(candidates: list[dict], x0: float, x1: float, block_index):
    """Take exactly one delivered construct for this rule, or none.

    ONE-TO-ONE OWNERSHIP.  A delivered rule satisfies exactly one canonical rule.
    Matching prefers the rule's own source row (``block_index``) so two rules that
    share an x span but sit on different rows cannot be served by one emission.
    """

    tolerance = _span_tolerance(x1 - x0)
    ordered = sorted(
        (candidate for candidate in candidates if not candidate["consumed"]),
        key=lambda candidate: (
            0 if candidate["block_index"] is None or block_index is None
            or candidate["block_index"] == block_index
            else 1,
            abs(candidate["x0"] - x0) + abs(candidate["x1"] - x1),
        ),
    )
    for candidate in ordered:
        if candidate["block_index"] is not None and block_index is not None:
            if candidate["block_index"] != block_index:
                continue
        if (
            abs(candidate["x0"] - x0) <= tolerance
            and abs(candidate["x1"] - x1) <= tolerance
        ):
            candidate["consumed"] = True
            return candidate
    return None


def build_source_rule_ledger(
    build_dir: str | Path,
    *,
    source_pdf: str | Path | None = None,
    report: dict | None = None,
) -> SourceRuleDeliveryLedger:
    """Track every canonical source rule of one build to the saved artifact."""

    build = Path(build_dir)
    if report is None:
        report = json.loads((build / "generation_report.json").read_text(encoding="utf-8"))
    ledger = SourceRuleDeliveryLedger()
    if source_pdf is None:
        pointer = report.get("source_document") or {}
        candidate = pointer.get("path") if isinstance(pointer, dict) else None
        source_pdf = candidate
    import pymupdf  # local import: the ledger is usable without a source document

    document = None
    if source_pdf and Path(str(source_pdf)).is_file():
        document = pymupdf.open(str(source_pdf))
        ledger.source_pdf = str(source_pdf)
        ledger.source_pdf_available = True
    else:
        ledger.notes.append(
            "the source document was not readable, so physical stroke identity and "
            "label evidence are unavailable; every rule is EVIDENCE_UNAVAILABLE"
        )
    try:
        atoms = normalize_source_rules(report, document)
        ledger.raw_rule_count = len(report.get("source_rule_registry") or [])
        delivered = _delivered_index(report)
        _assign_delivery(atoms, delivered, report, ledger, build)
    finally:
        if document is not None:
            document.close()
    return ledger


def _row_block_for_rule(document, page_number: int, y: float) -> int | None:
    """The source text block whose own band contains this rule's stroke."""

    if document is None:
        return None
    page = document[page_number - 1]
    best: tuple[float, int] | None = None
    for index, block in enumerate(page.get_text("dict")["blocks"]):
        bbox = block.get("bbox")
        if not bbox:
            continue
        if bbox[1] - 2.0 <= y <= bbox[3] + 5.0:
            distance = abs(y - bbox[3])
            if best is None or distance < best[0]:
                best = (distance, index)
    return None if best is None else best[1]


def _glyph_advance(character: str, size_pt: float) -> float:
    """A painted glyph's own advance, WITHOUT any character spacing.

    This is the whole point of the measurement: ``w:spacing`` moves the cursor but
    paints nothing, so a rule whose width comes from character spacing has almost
    no painted coverage however wide the paragraph believes it is.
    """

    if character == "\u2007":
        return size_pt * 0.5
    if character == "\u3000":
        return size_pt
    if character.isspace():
        return size_pt * 0.25
    if ord(character) > 0x2E80:
        return size_pt
    return size_pt * 0.5


class DocxPainterIndex:
    """The saved DOCX's own painted coverage, paragraph by paragraph.

    VISIBILITY IS MEASURED ON THE ARTIFACT.  A paragraph's record states the
    advance its underlined glyphs actually paint, and whether it carries a tab
    whose declared leader paints the span.  An underline *property* on a spacer
    that paints nothing is not a visible rule, and this index is what tells the
    two apart.
    """

    def __init__(self, docx_path: Path):
        from docx import Document
        from docx.oxml.ns import qn

        self._qn = qn
        self.paragraphs: list[dict] = []
        document = Document(str(docx_path))
        for index, paragraph in enumerate(document.paragraphs):
            record = self._measure(paragraph)
            record["index"] = index
            self.paragraphs.append(record)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        record = self._measure(paragraph)
                        # A cell paragraph has no body index; it is addressed by
                        # its own table/row/cell identity instead.
                        record["index"] = None
                        self.paragraphs.append(record)

    def _measure(self, paragraph) -> dict:
        qn = self._qn
        text = paragraph.text
        painted = 0.0
        underlined = 0
        leader = False
        run_advances: list[float] = []
        run_texts: list[str] = []
        for run in paragraph.runs:
            if not run.underline:
                continue
            underlined += 1
            size = float(getattr(run.font.size, "pt", 0.0) or 10.5)
            advance = 0.0
            for character in run.text or "":
                advance += _glyph_advance(character, size)
            painted += advance
            run_advances.append(advance)
            run_texts.append(run.text or "")
            if run._r.findall(qn("w:tab")):
                # A tab whose own declared stop carries an underline leader paints
                # its leader glyphs across the span, so it covers the rule.
                leader = True
        return {
            "text": text,
            "painted_advance_pt": painted,
            "run_advances": run_advances,
            "run_texts": run_texts,
            "underlined_runs": underlined,
            "leader_painted": leader,
            "native_bottom_border": self._bottom_border(paragraph),
            # A rule painted by a floating shape or an image is not an *editable*
            # Word field, however visible it is.
            "has_drawing": bool(
                paragraph._p.findall(".//" + qn("w:drawing"))
                or paragraph._p.findall(".//" + qn("w:pict"))
            ),
        }

    def _bottom_border(self, paragraph) -> bool:
        """Whether the paragraph paints its own native bottom rule."""

        qn = self._qn
        properties = paragraph._p.find(qn("w:pPr"))
        if properties is None:
            return False
        borders = properties.find(qn("w:pBdr"))
        if borders is None:
            return False
        bottom = borders.find(qn("w:bottom"))
        if bottom is None:
            return False
        return str(bottom.get(qn("w:val")) or "").lower() not in ("", "none", "nil")

    def underlined_run_carrying(self, value: str):
        """The paragraph that carries ``value`` as a resolved source value.

        RULE OWNERSHIP MUST SURVIVE SOURCE-FILL FACT RESOLUTION.  The value's own
        presentation is the rule's Word owner, so it is located by the value the
        build resolved - not by the rule's *empty* span, whose width the value does
        not reproduce.  The search is by the paragraph's own text, because a long
        resolved value is not guaranteed to sit in a single run: Word splits it
        where the source's own prose resumes.
        """

        compact = re.sub(r"\s+", "", value or "")
        if not compact:
            return None
        for record in self.paragraphs:
            if compact not in re.sub(r"\s+", "", record["text"]):
                continue
            for run_text in record.get("run_texts") or ():
                if not re.sub(r"\s+", "", run_text):
                    continue
                return {
                    "paragraph_index": record["index"],
                    "run_text": run_text,
                    "coverage": 1.0,
                    "evidence": (
                        f"paragraph {record['index']}: the paragraph carries the "
                        f"resolved value and has underlined run {run_text[:24]!r}"
                    ),
                }
        return None

    def native_rule_coverage(self, width_pt: float, *, paragraph_index=None):
        """Coverage of a rule painted by the paragraph's own native border.

        A MECHANISM MUST BE MEASURED BY WHAT IT PAINTS.  A native paragraph bottom
        border is drawn across the paragraph's text column independently of any
        glyph, so the underlined-glyph formula reports a false deficit for it.  The
        border's own owner paragraph is the evidence, and its recorded source width
        is what the emitter measured; a border is a continuous line by construction.
        """

        if paragraph_index is None:
            return None, []
        for record in self.paragraphs:
            if record["index"] != paragraph_index:
                continue
            if record.get("native_bottom_border"):
                return 1.0, [
                    f"paragraph {paragraph_index}: native bottom border paints a "
                    f"continuous rule across the {width_pt:.2f} pt source span"
                ]
        return None, []

    def painted_coverage(
        self,
        width_pt: float,
        *,
        paragraph_index=None,
        label: str | None = None,
    ) -> tuple[float | None, list[str]]:
        """The painted coverage of the run that owns this source span.

        OWNERSHIP IS PER RULE, NOT PER PARAGRAPH.  Several source rules can share
        one delivered paragraph (a source visual row holds a label and two fields),
        so summing a paragraph's underlined advance and dividing it by one rule's
        width reports coverage above one for the first rule and far below one for
        the second - neither of which is a measurement of either rule.  The run
        whose own painted advance equals this rule's own source width is that
        rule's owner, because a source-rectified rule is painted to the source's
        measured span.
        """

        if width_pt <= 0:
            return None, []
        tolerance = max(2.0, width_pt * 0.06)
        compact_label = re.sub(r"\s+", "", label or "")
        best: float | None = None
        evidence: list[str] = []
        for record in self.paragraphs:
            if paragraph_index is not None:
                if record["index"] != paragraph_index:
                    continue
            elif compact_label:
                if compact_label not in re.sub(r"\s+", "", record["text"]):
                    continue
            elif not record["underlined_runs"]:
                continue
            if record["leader_painted"]:
                return 1.0, [
                    f"paragraph {record['index']}: underline leader tab paints the span"
                ]
            for advance in record["run_advances"]:
                if abs(advance - width_pt) <= tolerance:
                    return 1.0, [
                        f"paragraph {record['index']}: a painted run of "
                        f"{advance:.2f} pt owns this {width_pt:.2f} pt source span"
                    ]
            if record["underlined_runs"]:
                coverage = record["painted_advance_pt"] / width_pt
                if best is None or coverage > best:
                    best = coverage
                    evidence = [
                        f"paragraph {record['index']}: {record['underlined_runs']} "
                        f"underlined run(s) paint {record['painted_advance_pt']:.2f} pt "
                        f"of the {width_pt:.2f} pt source span"
                    ]
        return best, evidence

    def run_widths_in(self, paragraph_index) -> list[float]:
        """The painted advances of one paragraph's underlined runs."""

        for record in self.paragraphs:
            if record["index"] == paragraph_index:
                return list(record["run_advances"])
        return []

    def paragraph_text(self, paragraph_index) -> str:
        for record in self.paragraphs:
            if record["index"] == paragraph_index:
                return record["text"]
        return ""

    def paragraphs_with_run_width(self, width_pt, *, exclude=None, tolerance_pt=2.5):
        """Every paragraph whose own run paints this rule's width."""

        found: list[int] = []
        for record in self.paragraphs:
            if record["index"] == exclude or record["index"] is None:
                continue
            if record["leader_painted"]:
                continue
            for advance in record["run_advances"]:
                if abs(advance - width_pt) <= max(tolerance_pt, width_pt * 0.06):
                    found.append(record["index"])
                    break
        return found

    def editability(self, paragraph_index):
        """Whether the rule in this paragraph is a native, editable Word field.

        A FORM RULE MUST BE VISIBLE **AND EDITABLE**.  A line painted by a floating
        shape or an image is visible and is not a field the bidder can fill in, so
        it is rejected here however well it covers the source span.  An underline
        *property* is equally insufficient on its own: the construct has to be a
        run the bidder can type into.
        """

        if paragraph_index is None:
            return None, "the rule's Word owner is unknown"
        for record in self.paragraphs:
            if record["index"] != paragraph_index:
                continue
            if record["underlined_runs"]:
                return True, "native underlined run"
            if record.get("native_bottom_border"):
                # Visible, and not a field: the bidder cannot type into a border.
                return False, "the rule is a paragraph decoration, not an editable field"
            if record.get("has_drawing"):
                return False, "the rule is painted by a floating shape, not an editable run"
            return False, "the owner paragraph carries no editable rule"
        return None, "the rule's Word owner paragraph was not found"

    def extraneous_figure_space_runs(self, applicable_widths, *, tolerance_pt=2.5):
        """Painted rules the source never drew.

        A PLAIN GAP MUST NOT GAIN AN INVENTED RULE.  Every underlined figure-space
        run is a line the delivery painted; if no canonical applicable source rule
        has that width, the run has no source behind it.  Matching against *any*
        applicable width is deliberate - this establishes that a run is extraneous,
        and it must never be narrowed by width equality alone.  A tab leader is not
        an independently painted span, and a run whose text is not figure spaces is
        source text (or its own underline), not a fill rule.
        """

        found: list[dict] = []
        for record in self.paragraphs:
            if record["leader_painted"]:
                continue
            for advance, text in zip(record["run_advances"], record["run_texts"]):
                if advance <= 0:
                    continue
                if not text or text.strip("\u2007"):
                    # Only a figure-space run is this emitter's own painted rule;
                    # source text and its underline are never a fill rule.
                    continue
                if any(
                    abs(advance - width) <= max(tolerance_pt, width * 0.06)
                    for width in applicable_widths
                ):
                    continue
                found.append(
                    {
                        "paragraph_index": record["index"],
                        "painted_width_pt": round(advance, 3),
                        "text": record["text"][:40],
                        "reason": "EXTRANEOUS_SOURCE_RULE",
                    }
                )
        return found


def rule_owner_paragraphs(report: dict) -> dict[str, int]:
    """The Word paragraph the emitter itself recorded as each rule's owner.

    DETERMINISTIC ONE-TO-ONE IDENTITY, WITHOUT A LABEL.  The emitter already
    records which paragraph it built a source row into - a source form line's own
    paragraph, an emission plan's ``generated_paragraph_index``, a composition's
    ``paragraph_index`` - and those records name the rule ids they carry.  Reading
    the owner from them is exact; reading it from a label is a heuristic that
    fails outright when the source wrote no label to the left of a field.
    """

    owners: dict[str, int] = {}
    for record in report.get("source_form_line_paragraphs") or []:
        index = record.get("generated_paragraph_index")
        if index is None:
            continue
        for rule_id in record.get("source_rule_ids") or []:
            owners.setdefault(str(rule_id), int(index))
    for record in report.get("source_rule_compositions") or []:
        index = record.get("paragraph_index")
        if index is None:
            continue
        rule_id = record.get("source_rule_id")
        if rule_id:
            owners.setdefault(str(rule_id), int(index))
    for plan in report.get("source_visual_line_emission_plans") or []:
        index = plan.get("generated_paragraph_index")
        if index is None:
            continue
        for rule_id in plan.get("source_rule_ids") or []:
            owners.setdefault(str(rule_id), int(index))
    return owners


def resolved_values_by_field(build: Path) -> dict[str, str]:
    """The build's resolved fact values, keyed by the registry's own field names."""

    path = build / "project_facts.json"
    if not path.is_file():
        return {}
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except Exception:  # pragma: no cover - defensive artifact boundary
        return {}
    values: dict[str, str] = {}
    for name, fact in (payload.get("fields") or {}).items():
        if not isinstance(fact, dict):
            continue
        if str(fact.get("status") or "").upper() != "RESOLVED":
            continue
        value = fact.get("resolved_value")
        if isinstance(value, str) and value.strip():
            values[str(name).upper()] = value
    return values


def audit_rule_ownership(ledger, painter) -> None:
    """One source rule, one delivered owner - and the owner has to be the right one.

    A delivered owner must be linked by SOURCE PROVENANCE, not by width equality.
    Two rules that share a width are still two rules; a rule recorded against a
    paragraph that does not paint it was delivered somewhere other than its own
    source row; and a bidder-owned slot must not be carrying a value the source
    never authorised.
    """

    claimed: dict[tuple, SourceRuleDeliveryRecord] = {}
    for record in ledger.records:
        atom = record.atom
        if not atom.applicable:
            continue
        owner = record.owner_paragraph
        label = atom.source_label
        if owner is not None and painter is not None and record.verdict in (
            DELIVERED_STRUCTURALLY_BUT_NOT_VISIBLE,
            EVIDENCE_UNAVAILABLE,
        ):
            owner_text = painter.paragraph_text(owner)
            owner_paints = bool(painter.run_widths_in(owner))
            elsewhere = painter.paragraphs_with_run_width(
                atom.source_width_pt, exclude=owner
            )
            if not owner_paints and elsewhere:
                if label and label.replace(" ", "") in owner_text.replace(" ", ""):
                    # The rule's own row text is here, but the line it should carry
                    # was painted into a different paragraph: the rule and its row
                    # were separated.
                    record.reasons.append("SOURCE_FORM_STRUCTURE")
                    record.first_missing_stage = "STAGE_5_SAVED_REOPENED_DOCX"
                    record.verdict = DELIVERED_STRUCTURALLY_BUT_NOT_VISIBLE
                else:
                    record.reasons.append("WRONG_OWNER")
                    record.first_missing_stage = "STAGE_5_SAVED_REOPENED_DOCX"
                    record.verdict = DELIVERED_STRUCTURALLY_BUT_NOT_VISIBLE
            elif not owner_paints and not atom.value_owned and label:
                # A bidder-owned slot that carries text but paints no rule of its
                # own has been filled with something the source never authorised.
                residue = owner_text.replace(label, "", 1).strip()
                if residue:
                    record.reasons.append("SOURCE_SLOT_OWNERSHIP")
                    record.first_missing_stage = "STAGE_5_SAVED_REOPENED_DOCX"
                    record.verdict = NOT_DELIVERED
        # MULTIPLE SOURCE OWNERS: one delivered construct cannot serve two rules.
        key = (
            record.owner_paragraph,
            None if record.matched_advance_pt is None else round(record.matched_advance_pt, 2),
        )
        if record.verdict == DELIVERED_VISIBLE_EDITABLE and None not in key:
            if key in claimed:
                record.reasons.append("DUPLICATE_DELIVERY_OWNER")
                record.first_missing_stage = "STAGE_5_SAVED_REOPENED_DOCX"
                record.verdict = DELIVERED_STRUCTURALLY_BUT_NOT_VISIBLE
            else:
                claimed[key] = record


def _assign_delivery(atoms, delivered, report, ledger, build: Path) -> None:
    document = None
    try:
        import pymupdf

        if ledger.source_pdf:
            document = pymupdf.open(ledger.source_pdf)
    except Exception:  # pragma: no cover - defensive source boundary
        document = None

    by_id = {}
    for record in report.get("positioned_blank_records") or []:
        if record.get("source_rule_id"):
            by_id.setdefault(str(record["source_rule_id"]), record)

    painter = None
    docx_path = build / "基础投标文件.docx"
    if docx_path.is_file():
        try:
            painter = DocxPainterIndex(docx_path)
        except Exception as exc:  # pragma: no cover - defensive artifact boundary
            ledger.notes.append(f"the saved DOCX could not be measured: {exc}")
    owners = rule_owner_paragraphs(report)
    resolved = resolved_values_by_field(build)

    try:
        for atom in atoms:
            entry = SourceRuleDeliveryRecord(atom=atom)
            page_candidates = delivered.get(atom.source_page, [])
            block_index = _row_block_for_rule(document, atom.source_page, atom.source_y)
            atom_block = block_index
            match = None
            if atom.occurrence_class == PHYSICAL_VISIBLE_RULE and not atom.value_owned:
                match = _consume_match(
                    page_candidates, atom.source_x0, atom.source_x1, atom_block
                )
            entry.word_owner = (
                None
                if match is None
                else f"source_block:{match['block_index']} via {match['source']}"
            )
            entry.stages = _stage_records(atom, match, by_id)
            _decide(
                entry,
                match,
                by_id,
                painter,
                owner_paragraph=owners.get(atom.source_rule_id),
                resolved_values=[
                    resolved.get(str(field).upper())
                    for field in atom.resolved_fact_fields
                    if resolved.get(str(field).upper())
                ],
            )
            ledger.records.append(entry)
        applicable_widths = [
            record.atom.source_width_pt
            for record in ledger.records
            if record.atom.applicable
        ]
        if painter is not None:
            ledger.extraneous_rules = painter.extraneous_figure_space_runs(
                applicable_widths
            )
        audit_rule_ownership(ledger, painter)
    finally:
        if document is not None:
            document.close()


def _stage_records(atom, match, by_id) -> list[RuleStageRecord]:
    emission = by_id.get(atom.source_rule_id)
    delivered = match is not None
    stages = [
        RuleStageRecord(
            stage="STAGE_0_SOURCE_RULE",
            present=atom.occurrence_class != UNRESOLVED_CLASSIFICATION,
            owner=atom.source_rule_id,
            construct_kind=atom.occurrence_class,
            reason=atom.stroke_evidence,
        ),
        RuleStageRecord(
            stage="STAGE_1_SOURCE_FORMAT",
            present=atom.applicable,
            owner=atom.canonical_rule_id,
            construct_kind=atom.relation_type,
            reason=(
                "the source rule is in the delivery scope"
                if atom.applicable
                else f"not an applicable field rule ({atom.occurrence_class})"
            ),
        ),
        RuleStageRecord(
            stage="STAGE_2_PAGE_LAYOUT",
            present=delivered,
            owner=None if match is None else str(match.get("block_index")),
            construct_kind=None if match is None else match.get("render_style"),
            reason=(
                None
                if delivered
                else "no layout construct owns this rule's span and row"
            ),
        ),
        RuleStageRecord(
            stage="STAGE_3_WORD_EMISSION_PLAN",
            present=bool(emission),
            owner=None if emission is None else emission.get("emission_id"),
            construct_kind=None if emission is None else emission.get("representation_kind"),
            mechanism=None if emission is None else emission.get("emission_mechanism"),
            reason=None if emission else "no Word emission plan references this rule",
        ),
        RuleStageRecord(
            stage="STAGE_4_IN_MEMORY_DOCX",
            present=delivered,
            owner=None if match is None else match.get("source"),
            construct_kind=None if match is None else match.get("representation_kind"),
            generated_width_pt=None if match is None else match.get("rendered_width_pt"),
        ),
        RuleStageRecord(
            stage="STAGE_5_SAVED_REOPENED_DOCX",
            present=delivered,
            owner=None if match is None else match.get("source"),
            construct_kind=None if match is None else match.get("representation_kind"),
            generated_width_pt=None if match is None else match.get("rendered_width_pt"),
        ),
        RuleStageRecord(
            stage="STAGE_6_RENDERER_OBSERVATION",
            present=False,
            reason=(
                "renderer observation is diagnostic and is not a machine closure "
                "authority; Microsoft Word desktop remains the final authority"
            ),
        ),
    ]
    return stages


def _decide(
    entry: SourceRuleDeliveryRecord,
    match,
    by_id,
    painter,
    *,
    owner_paragraph=None,
    resolved_values=(),
) -> None:
    atom = entry.atom
    if atom.occurrence_class == UNRESOLVED_CLASSIFICATION:
        entry.verdict = EVIDENCE_UNAVAILABLE
        entry.reasons.append("UNRESOLVED_CLASSIFICATION")
        entry.first_missing_stage = "STAGE_0_SOURCE_RULE"
        return
    if not atom.applicable:
        entry.verdict = NOT_APPLICABLE_WITH_SOURCE_EVIDENCE
        entry.reasons.append(f"NOT_APPLICABLE:{atom.occurrence_class}")
        return
    # SOURCE_RULE_OWNED_BY_RESOLVED_VALUE.  The slot this rule belongs to resolved
    # to a fact value, so the rule's Word owner is that value's own run - located by
    # the value the build resolved, never by the rule's empty span, whose width the
    # value does not reproduce.
    if atom.value_owned:
        carried = None
        if painter is not None:
            for value in resolved_values:
                carried = painter.underlined_run_carrying(value)
                if carried is not None:
                    break
        if carried is not None:
            entry.docx_evidence = [
                f"owner_kind={atom.owner_kind}",
                f"authoritative_slot_id={atom.authoritative_slot_id}",
                f"resolved_fact_fields={list(atom.resolved_fact_fields)}",
                carried["evidence"],
            ]
            entry.word_owner = (
                None
                if carried["paragraph_index"] is None
                else f"word_paragraph:{carried['paragraph_index']}"
            )
            entry.visible_coverage_ratio = carried["coverage"]
            entry.editable = True
            entry.verdict = DELIVERED_VISIBLE_EDITABLE
            return
        entry.verdict = NOT_DELIVERED
        entry.first_missing_stage = "STAGE_5_SAVED_REOPENED_DOCX"
        entry.reasons.append("RESOLVED_VALUE_NOT_PRESENT_AS_AN_UNDERLINED_RUN")
        return
    if match is None and owner_paragraph is None:
        entry.verdict = NOT_DELIVERED
        entry.first_missing_stage = "STAGE_2_PAGE_LAYOUT"
        entry.reasons.append("NO_DELIVERED_CONSTRUCT_OWNS_THIS_SOURCE_RULE")
        return
    if match is not None:
        entry.generated_width_pt = (
            None if match.get("rendered_width_pt") is None
            else float(match["rendered_width_pt"])
        )
        entry.docx_evidence = [
            f"{match['source']}",
            f"representation_kind={match.get('representation_kind')}",
            f"render_style={match.get('render_style')}",
            f"recorded_width_pt={match.get('rendered_width_pt')}",
        ]
    else:
        entry.docx_evidence = [
            f"the emitter recorded paragraph {owner_paragraph} as this rule's own "
            "source-form-line owner",
        ]
    # The owner is a *Word paragraph index* when the emitter recorded one in a
    # rule-id-bearing record; a blank's own ``block_index`` is a source block, not
    # a Word paragraph, so it is never used as one.
    owner = owner_paragraph
    # VISIBILITY IS MEASURED ON THE SAVED DOCX.  The emitter's own recorded width
    # states what it *intended* to paint; the painted glyph advance states what the
    # artifact actually covers, and only the second can prove a visible rule.
    coverage = None
    if painter is not None:
        # A NATIVE RULE IS MEASURED BY WHAT IT PAINTS.  A paragraph bottom border is
        # drawn across the paragraph's own column independently of any glyph, so the
        # underlined-glyph formula would report a false deficit for it.
        coverage, native_evidence = painter.native_rule_coverage(
            atom.source_width_pt, paragraph_index=(owner if isinstance(owner, int) else None)
        )
        if coverage is not None:
            entry.docx_evidence.extend(native_evidence)
        else:
            coverage, painter_evidence = painter.painted_coverage(
                atom.source_width_pt,
                paragraph_index=(owner if isinstance(owner, int) else None),
                label=(
                    atom.source_label
                    or (match.get("semantic_slot") if match is not None else None)
                ),
            )
            entry.docx_evidence.extend(painter_evidence)
    entry.visible_coverage_ratio = coverage
    entry.editable = True
    entry.owner_paragraph = owner if isinstance(owner, int) else None
    entry.matched_advance_pt = (
        atom.source_width_pt
        if coverage is not None and coverage >= MIN_VISIBLE_COVERAGE_RATIO
        else None
    )
    if coverage is None:
        entry.verdict = EVIDENCE_UNAVAILABLE
        entry.reasons.append("VISIBILITY_NOT_MEASURABLE_FROM_THE_SAVED_DOCX")
        entry.first_missing_stage = "STAGE_5_SAVED_REOPENED_DOCX"
        return
    if coverage >= MIN_VISIBLE_COVERAGE_RATIO:
        # A FORM RULE MUST BE VISIBLE **AND EDITABLE**.  Coverage alone cannot tell a
        # native underlined field from a floating shape that happens to be the right
        # width, so the owner's own construct decides it.
        if painter is not None and isinstance(owner, int):
            editable_ok, editable_reason = painter.editability(owner)
            if editable_ok is False:
                entry.verdict = DELIVERED_STRUCTURALLY_BUT_NOT_VISIBLE
                entry.reasons.append("SOURCE_FORM_EDITABILITY")
                entry.first_missing_stage = "STAGE_5_SAVED_REOPENED_DOCX"
                entry.editable = False
                entry.docx_evidence.append(f"editability: {editable_reason}")
                return
        entry.verdict = DELIVERED_VISIBLE_EDITABLE
    else:
        entry.verdict = DELIVERED_STRUCTURALLY_BUT_NOT_VISIBLE
        entry.reasons.append("INSUFFICIENT_PAINTED_COVERAGE")
        entry.first_missing_stage = "STAGE_5_SAVED_REOPENED_DOCX"


__all__ = [
    "DELIVERED_STRUCTURALLY_BUT_NOT_VISIBLE",
    "DELIVERED_VISIBLE_EDITABLE",
    "DECORATIVE_RULE",
    "EVIDENCE_UNAVAILABLE",
    "EXTRACTION_ALIAS",
    "MIN_VISIBLE_COVERAGE_RATIO",
    "NOT_APPLICABLE_WITH_SOURCE_EVIDENCE",
    "NOT_DELIVERED",
    "PHYSICAL_VISIBLE_RULE",
    "SCHEMA",
    "STAGES",
    "SourceRuleAtom",
    "SourceRuleDeliveryLedger",
    "SourceRuleDeliveryRecord",
    "TEXT_UNDERLINE",
    "UNRESOLVED_CLASSIFICATION",
    "ZERO_WIDTH_NON_RULE",
    "build_source_rule_ledger",
    "normalize_source_rules",
    "source_label_for_rule",
    "source_page_strokes",
    "stroke_matches",
]
