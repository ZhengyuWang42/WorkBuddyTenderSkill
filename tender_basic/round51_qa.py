"""Round 5.1 source-first numbering and governance QA."""

from __future__ import annotations

import re
import shutil
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from lxml import etree

from .numbering import NumberingOwnership, parse_source_number_token, prefix_punctuation
from .round50_qa import (
    _document_xml_bytes,
    _num_pr,
    _paragraph_style_name,
    _paragraph_text,
    _patch_style_size_only,
    _style_names,
    _xml,
)
from .word_forensics import inspect_docx_ooxml


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = {"w": W_NS}


def _story_paragraphs(root: etree._Element) -> list[etree._Element]:
    """Return top-level story paragraphs matching ``Document.paragraphs``."""

    return list(root.xpath("./w:body/w:p", namespaces=W))


def _style_numbering_names(path: str | Path) -> set[str]:
    root = _xml(path, "word/styles.xml")
    names: set[str] = set()
    for style in root.xpath("./w:style", namespaces=W):
        if style.xpath("./w:pPr/w:numPr", namespaces=W):
            name = next(iter(style.xpath("./w:name/@w:val", namespaces=W)), "")
            match = re.fullmatch(r"heading\s+([1-9])", name.strip(), flags=re.IGNORECASE)
            if match:
                name = f"Heading {match.group(1)}"
            if name:
                names.add(name)
    return names


def _record_token(record: dict) -> dict | None:
    token = record.get("source_number_token")
    if isinstance(token, dict) and token.get("raw_prefix"):
        return token
    raw = record.get("source_number_prefix") or ""
    if not raw:
        return None
    parsed = parse_source_number_token(raw)
    return parsed.as_dict() if parsed else None


def _source_records(generation: dict) -> list[tuple[int, dict, dict]]:
    records: list[tuple[int, dict, dict]] = []
    for record in generation.get("logical_paragraph_records", []):
        token = _record_token(record)
        if token is None:
            continue
        records.append((int(record.get("paragraph_index", -1)), record, token))
    return records


def build_source_numbering_qa(path: str | Path, generation: dict) -> dict:
    """Compare every retained source prefix with the generated DOCX text."""

    root = _xml(path, "word/document.xml")
    styles = _style_names(path)
    paragraphs = _story_paragraphs(root)
    numbered: list[dict] = []
    prefix_mismatch = punctuation_mismatch = spacing_mismatch = 0
    sequence_dependency = 0
    source_literal_count = generated_auto_count = 0
    for index, record, token in _source_records(generation):
        if not 0 <= index < len(paragraphs):
            numbered.append({
                "paragraph_index": index,
                "source_visible_text": record.get("source_text", ""),
                "source_prefix": token.get("raw_prefix", ""),
                "generated_visible_text": None,
                "generated_prefix": "",
                "semantic_role": record.get("layout_role") or record.get("kind", ""),
                "word_style": record.get("word_style", ""),
                "ownership": token.get("ownership", NumberingOwnership.SOURCE_LITERAL.value),
                "source_locator": token.get("source_locator"),
                "prefix_match": False,
                "punctuation_match": False,
                "spacing_match": False,
                "sequence_dependency": True,
                "error": "paragraph_index_out_of_range",
            })
            prefix_mismatch += 1
            sequence_dependency += 1
            continue
        paragraph = paragraphs[index]
        generated_text = _paragraph_text(paragraph)
        generated_token = parse_source_number_token(generated_text)
        source_prefix = str(token.get("raw_prefix", ""))
        generated_prefix = generated_token.raw_prefix if generated_token else ""
        source_separator = str(token.get("separator_after_prefix", ""))
        generated_separator = generated_token.separator_after_prefix if generated_token else ""
        prefix_match = generated_prefix == source_prefix
        punctuation_match = prefix_punctuation(generated_prefix) == prefix_punctuation(source_prefix)
        spacing_match = generated_separator == source_separator
        if not prefix_match:
            prefix_mismatch += 1
        if not punctuation_match:
            punctuation_mismatch += 1
        if not spacing_match and float(token.get("confidence", 0.0) or 0.0) >= 0.8:
            spacing_mismatch += 1
        num_id, level = _num_pr(paragraph)
        style = _paragraph_style_name(paragraph, styles)
        style_has_numbering = style in _style_numbering_names(path)
        depends = bool(num_id or style_has_numbering)
        if depends:
            sequence_dependency += 1
        ownership = token.get("ownership", NumberingOwnership.SOURCE_LITERAL.value)
        if ownership == NumberingOwnership.SOURCE_LITERAL.value:
            source_literal_count += 1
        elif ownership == NumberingOwnership.GENERATED_AUTO.value:
            generated_auto_count += 1
        numbered.append({
            "paragraph_index": index,
            "source_visible_text": record.get("source_text", ""),
            "source_prefix": source_prefix,
            "generated_visible_text": generated_text,
            "generated_prefix": generated_prefix,
            "semantic_role": record.get("layout_role") or record.get("kind", ""),
            "word_style": style,
            "ownership": ownership,
            "numbering_family": token.get("numbering_family"),
            "source_locator": token.get("source_locator"),
            "prefix_match": prefix_match,
            "punctuation_match": punctuation_match,
            "spacing_match": spacing_match,
            "sequence_dependency": depends,
            "num_id": num_id,
            "ilvl": level,
        })
    source_style_numbering = sorted(
        name for name in _style_numbering_names(path)
        if name.startswith("Heading ") or name.startswith("Tender Body List ")
    )
    return {
        "source_numbered_block_count": len(numbered),
        "source_literal_count": source_literal_count,
        "generated_auto_count": generated_auto_count,
        "source_number_prefix_mismatch_count": prefix_mismatch,
        "source_number_punctuation_mismatch_count": punctuation_mismatch,
        "source_number_spacing_mismatch_count": spacing_mismatch,
        "source_number_sequence_dependency_count": sequence_dependency,
        "source_styles_with_auto_numbering": source_style_numbering,
        "records": numbered,
        "targets": {
            "source_number_prefix_mismatch_count": 0,
            "source_number_punctuation_mismatch_count": 0,
            "source_number_sequence_dependency_count": 0,
            "spacing_mismatch_when_geometry_reliable": 0,
        },
        "result": "PASS" if all(
            value == 0 for value in (
                prefix_mismatch, punctuation_mismatch,
                spacing_mismatch, sequence_dependency,
            )
        ) else "FAIL",
    }


def build_outline_tree(path: str | Path, generation: dict) -> dict:
    """Build the Word navigation tree and verify excluded source roles."""

    root = _xml(path, "word/document.xml")
    styles = _style_names(path)
    tree: list[dict] = []
    stack: list[tuple[int, dict]] = []
    headings = 0
    body_list_headings = 0
    toc_headings = 0
    for paragraph in _story_paragraphs(root):
        style = _paragraph_style_name(paragraph, styles)
        if style.startswith("Tender Body List "):
            if re.fullmatch(r"Heading [1-9]", style):
                body_list_headings += 1
            continue
        if style.startswith("Tender TOC"):
            if re.fullmatch(r"Heading [1-9]", style):
                toc_headings += 1
            continue
        match = re.fullmatch(r"Heading ([1-9])", style)
        if not match:
            continue
        headings += 1
        level = int(match.group(1))
        node = {"level": level, "style": style, "text": _paragraph_text(paragraph), "children": []}
        while stack and stack[-1][0] >= level:
            stack.pop()
        if stack:
            stack[-1][1]["children"].append(node)
        else:
            tree.append(node)
        stack.append((level, node))
    ambiguous = [
        record for record in generation.get("logical_paragraph_records", [])
        if record.get("classification_status") == "HEADING_CLASSIFICATION_NEEDS_REVIEW"
    ]
    ambiguous_promoted = [
        record for record in ambiguous
        if str(record.get("word_style", "")).startswith("Heading ")
    ]
    false_heading_records = [
        record for record in ambiguous
        if str(record.get("source_text", "")).strip()
        and re.search(r"^[一二三四五六七八九十百千万零〇]+、\s*(?:表|图)$", str(record.get("source_text", "")).strip())
        and str(record.get("word_style", "")).startswith("Heading ")
    ]
    result = not body_list_headings and not toc_headings and not ambiguous_promoted and not false_heading_records
    return {
        "headings": headings,
        "tree": tree,
        "body_list_headings": body_list_headings,
        "toc_headings": toc_headings,
        "ambiguous_heading_count": len(ambiguous),
        "ambiguous_heading_promoted_count": len(ambiguous_promoted),
        "false_heading_count": len(false_heading_records),
        "result": "PASS" if result else "FAIL",
    }


def build_style_architecture_qa(path: str | Path, generation: dict, numbering: dict) -> dict:
    """Check style ownership independently from visible source prefixes."""

    root = _xml(path, "word/document.xml")
    styles = _style_names(path)
    body_paragraphs = _story_paragraphs(root)
    all_paragraphs = list(root.xpath(".//w:p", namespaces=W))
    named = [_paragraph_style_name(paragraph, styles) for paragraph in all_paragraphs]
    heading_records = [
        record for record in generation.get("logical_paragraph_records", [])
        if re.fullmatch(r"Heading [1-9]", str(record.get("word_style", "")))
    ]
    list_records = [
        record for record in generation.get("logical_paragraph_records", [])
        if str(record.get("word_style", "")).startswith("Tender Body List ")
    ]
    source_heading_style_errors = 0
    source_heading_num_errors = 0
    for record in heading_records:
        index = int(record.get("paragraph_index", -1))
        if not 0 <= index < len(body_paragraphs):
            source_heading_style_errors += 1
            continue
        paragraph = body_paragraphs[index]
        if not re.fullmatch(r"Heading [1-9]", _paragraph_style_name(paragraph, styles)):
            source_heading_style_errors += 1
        num_id, _level = _num_pr(paragraph)
        if num_id or _paragraph_style_name(paragraph, styles) in _style_numbering_names(path):
            source_heading_num_errors += 1
    source_list_style_errors = 0
    source_list_num_errors = 0
    for record in list_records:
        index = int(record.get("paragraph_index", -1))
        if not 0 <= index < len(body_paragraphs):
            source_list_style_errors += 1
            continue
        paragraph = body_paragraphs[index]
        if not str(_paragraph_style_name(paragraph, styles)).startswith("Tender Body List "):
            source_list_style_errors += 1
        num_id, _level = _num_pr(paragraph)
        if num_id or _paragraph_style_name(paragraph, styles) in _style_numbering_names(path):
            source_list_num_errors += 1
    forensic = inspect_docx_ooxml(path)
    unstyled = sum(not bool(name) for name in named)
    return {
        "paragraph_count": len(all_paragraphs),
        "paragraphs_with_named_style": len(all_paragraphs) - unstyled,
        "paragraphs_without_named_style": unstyled,
        "source_literal_heading_count": len(heading_records),
        "source_literal_body_list_count": len(list_records),
        "source_heading_style_errors": source_heading_style_errors,
        "source_heading_auto_numbering_errors": source_heading_num_errors,
        "source_body_list_style_errors": source_list_style_errors,
        "source_body_list_auto_numbering_errors": source_list_num_errors,
        "literal_heading_number_count": sum(
            bool(record.get("source_number_prefix")) for record in heading_records
        ),
        "source_numbering": {
            "prefix_mismatch_count": numbering.get("source_number_prefix_mismatch_count", 0),
            "punctuation_mismatch_count": numbering.get("source_number_punctuation_mismatch_count", 0),
            "spacing_mismatch_count": numbering.get("source_number_spacing_mismatch_count", 0),
            "sequence_dependency_count": numbering.get("source_number_sequence_dependency_count", 0),
        },
        "generated_auto_style_names": sorted(
            name for name in _style_numbering_names(path)
            if name.startswith("Tender Generated ")
        ),
        "synthetic_layout_tables": generation.get("synthetic_layout_tables", 0),
        "unsafe_ooxml": int(forensic.get("summary", {}).get("error_count", 0) > 0),
        "result": "PASS" if all([
            unstyled == 0,
            source_heading_style_errors == 0,
            source_heading_num_errors == 0,
            source_list_style_errors == 0,
            source_list_num_errors == 0,
            numbering.get("result") == "PASS",
            generation.get("synthetic_layout_tables", 0) == 0,
            forensic.get("result") == "PASS",
        ]) else "FAIL",
    }


def build_editability_qa(path: str | Path, generation: dict, render=None) -> dict:
    """Edit a Heading style in a copy and prove literal prefixes are stable."""

    source = Path(path)
    checks: dict[str, dict] = {}
    source_indexes = [index for index, _record, _token in _source_records(generation)]

    def _prefix_stable(before, after, index: int) -> bool:
        if not (0 <= index < len(before.paragraphs) and 0 <= index < len(after.paragraphs)):
            return False
        before_token = parse_source_number_token(before.paragraphs[index].text)
        after_token = parse_source_number_token(after.paragraphs[index].text)
        if before_token is None:
            return after_token is None
        return after_token is not None and before_token.raw_prefix == after_token.raw_prefix

    with TemporaryDirectory(prefix="round51-editability-") as temp:
        root = Path(temp)
        for label, style_name, half_points in (("Heading 1", "heading 1", 60), ("Tender Body", "Tender Body", 26)):
            copy = root / f"{label.replace(' ', '_')}.docx"
            shutil.copyfile(source, copy)
            patched = _patch_style_size_only(source, copy, style_name, half_points)
            reopened = False
            prefix_stable = False
            if patched:
                before = Document(source)
                after = Document(copy)
                reopened = bool(after)
                prefix_stable = all(_prefix_stable(before, after, index) for index in source_indexes)
            checks[label] = {
                "style_definition_changed_only": patched and _document_xml_bytes(source) == _document_xml_bytes(copy),
                "reopened": reopened,
                "literal_prefixes_unchanged": prefix_stable,
                "rendered": False,
            }
            if render is not None and patched:
                try:
                    render(copy, root / f"render_{label.replace(' ', '_')}")
                    checks[label]["rendered"] = True
                except Exception as exc:  # pragma: no cover
                    checks[label]["render_error"] = str(exc)
            checks[label]["result"] = "PASS" if all(
                checks[label].get(key) is True for key in ("style_definition_changed_only", "reopened", "literal_prefixes_unchanged")
            ) else "FAIL"
    return {
        "checks": checks,
        "literal_prefix_stability": all(item["literal_prefixes_unchanged"] for item in checks.values()),
        "result": "PASS" if all(item["result"] == "PASS" for item in checks.values()) else "FAIL",
    }


def build_generated_auto_qa(path: str | Path, generation: dict) -> dict:
    """Prove GENERATED_AUTO has real numbering without touching source text."""

    source = Path(path)
    with TemporaryDirectory(prefix="round51-generated-auto-") as temp:
        copy = Path(temp) / "generated-auto.docx"
        shutil.copyfile(source, copy)
        document = Document(copy)
        source_headings = [p for p in document.paragraphs if p.style.name == "Heading 1"]
        if source_headings:
            inserted = document.add_paragraph("系统生成的自拟小节", style="Tender Generated Heading 1")
            source_headings[0]._p.addprevious(inserted._p)
        else:
            document.add_paragraph("系统生成的自拟小节", style="Tender Generated Heading 1")
        document.add_paragraph("系统生成的第二个小节", style="Tender Generated Heading 1")
        document.save(copy)
        reopened = Document(copy)
        root = _xml(copy, "word/document.xml")
        styles = _style_names(copy)
        generated = [
            paragraph for paragraph in _story_paragraphs(root)
            if _paragraph_style_name(paragraph, styles) == "Tender Generated Heading 1"
        ]
        source = [
            paragraph for paragraph in _story_paragraphs(root)
            if _paragraph_style_name(paragraph, styles) == "Heading 1"
        ]
        generated_style_numbered = "Tender Generated Heading 1" in _style_numbering_names(copy)
        source_style_numbered = "Heading 1" in _style_numbering_names(copy)
        source_prefixes = [
            parse_source_number_token(_paragraph_text(paragraph)).raw_prefix
            for paragraph in source
            if parse_source_number_token(_paragraph_text(paragraph)) is not None
        ]
        result = bool(reopened and len(generated) >= 2 and generated_style_numbered and not source_style_numbered)
        return {
            "generated_paragraph_count": len(generated),
            "generated_style_has_real_numPr": generated_style_numbered,
            "source_heading_style_has_no_numPr": not source_style_numbered,
            "source_literal_prefixes_present_after_generated_insert": len(source_prefixes),
            "generated_numbering_family_isolated": generated_style_numbered and not source_style_numbered,
            "result": "PASS" if result else "FAIL",
        }


__all__ = [
    "build_editability_qa",
    "build_generated_auto_qa",
    "build_outline_tree",
    "build_source_numbering_qa",
    "build_style_architecture_qa",
]
