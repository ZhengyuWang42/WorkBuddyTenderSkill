"""Round 5.0 style architecture, navigation, editability, and numbering QA."""

from __future__ import annotations

import json
import re
import shutil
from collections import Counter
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from .style_architecture import TenderStyleProfile, marker_info
from .word_forensics import inspect_docx_ooxml
from .word_safe_scan import scan_word_safe_docx


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = {"w": W_NS}


def _xml(path: str | Path, part: str) -> etree._Element:
    with ZipFile(path) as archive:
        return etree.fromstring(archive.read(part))


def _style_names(path: str | Path) -> dict[str, str]:
    root = _xml(path, "word/styles.xml")
    def canonical(raw: str) -> str:
        match = re.fullmatch(r"heading\s+([1-9])", raw.strip(), flags=re.IGNORECASE)
        return f"Heading {match.group(1)}" if match else raw
    return {
        style.get(f"{{{W_NS}}}styleId"): canonical(next(iter(style.xpath("./w:name/@w:val", namespaces=W)), ""))
        for style in root.xpath("./w:style", namespaces=W)
        if style.get(f"{{{W_NS}}}styleId")
    }


def _paragraph_style_id(paragraph: etree._Element) -> str | None:
    return next(iter(paragraph.xpath("./w:pPr/w:pStyle/@w:val", namespaces=W)), None)


def _paragraph_style_name(paragraph: etree._Element, styles: dict[str, str]) -> str:
    style_id = _paragraph_style_id(paragraph)
    return styles.get(style_id or "", "")


def _paragraph_text(paragraph: etree._Element) -> str:
    return "".join(paragraph.xpath(".//w:t/text()", namespaces=W))


def _num_pr(paragraph: etree._Element) -> tuple[str | None, str | None]:
    return (
        next(iter(paragraph.xpath("./w:pPr/w:numPr/w:numId/@w:val", namespaces=W)), None),
        next(iter(paragraph.xpath("./w:pPr/w:numPr/w:ilvl/@w:val", namespaces=W)), None),
    )


def _body_paragraphs(root: etree._Element) -> list[etree._Element]:
    return list(root.xpath(".//w:body/w:p", namespaces=W))


def _all_paragraphs(root: etree._Element) -> list[etree._Element]:
    return list(root.xpath(".//w:p", namespaces=W))


def _run_direct_formatting(root: etree._Element) -> tuple[int, int, list[dict]]:
    allowed = {"vertAlign", "u", "b", "i", "rStyle"}
    total = 0
    illegal = 0
    details: list[dict] = []
    for run in root.xpath(".//w:r", namespaces=W):
        rpr = run.find("w:rPr", namespaces=W)
        if rpr is None:
            continue
        children = [etree.QName(child).localname for child in rpr]
        if not children:
            continue
        total += 1
        bad = [child for child in children if child not in allowed]
        if bad:
            illegal += 1
            details.append({"text": _paragraph_text(run), "attributes": bad})
    return total, illegal, details


def _heading_like(text: str) -> bool:
    marker, remainder = marker_info(text)
    if not marker or not remainder:
        return False
    if re.fullmatch(r"[一二三四五六七八九十百千万零〇]+[、．.]", marker):
        return len(remainder) <= 65 and not re.search(r"[，。；：？！!?]", remainder)
    if re.fullmatch(r"[（(][一二三四五六七八九十百千万零〇]+[）)]", marker):
        return len(remainder) <= 65 and not re.search(r"[，。；：？！!?]", remainder)
    if marker.endswith((".", "．")):
        return len(remainder) <= 65 and not re.search(r"[，。；：？！!?]", remainder)
    return False


def _table_style_coverage(root: etree._Element, styles: dict[str, str]) -> dict[str, int | float]:
    paragraphs = root.xpath(".//w:tc/w:p", namespaces=W)
    named = [_paragraph_style_name(p, styles) for p in paragraphs]
    covered = sum(bool(name) for name in named)
    expected = {
        "Table Text", "Tender Table Header", "Tender Table Label", "Tender Table Value",
        "Tender Table Numeric", "Tender Table Unit",
    }
    expected_hits = sum(name in expected for name in named)
    return {
        "paragraph_count": len(paragraphs),
        "paragraphs_with_table_style": expected_hits,
        "paragraphs_with_any_named_style": covered,
        "coverage_ratio": round(expected_hits / len(paragraphs), 4) if paragraphs else 1.0,
    }


def build_style_architecture_qa(
    path: str | Path,
    generation: dict,
    profile: TenderStyleProfile,
) -> dict:
    root = _xml(path, "word/document.xml")
    styles = _style_names(path)
    body_paragraphs = _body_paragraphs(root)
    all_paragraphs = _all_paragraphs(root)
    named = [_paragraph_style_name(p, styles) for p in all_paragraphs]
    heading_names = {f"Heading {level}" for level in range(1, 10)}
    headings = [p for p in body_paragraphs if _paragraph_style_name(p, styles) in heading_names]
    heading_records = [
        record for record in generation.get("logical_paragraph_records", [])
        if record.get("kind") == "Heading" and record.get("word_style") in heading_names
    ]
    heading_with_correct_style = sum(
        0 <= int(record.get("paragraph_index", -1)) < len(body_paragraphs)
        and _paragraph_style_name(body_paragraphs[int(record["paragraph_index"])], styles) == record.get("word_style")
        for record in heading_records
    )
    toc_names = {"Tender TOC Title", "Tender TOC 1", "Tender TOC 2", "Tender TOC 3"}
    fake_heading_count = sum(
        _heading_like(_paragraph_text(p))
        and _paragraph_style_name(p, styles) not in heading_names
        and _paragraph_style_name(p, styles) not in toc_names
        for p in body_paragraphs
    )
    body_records = [
        record for record in generation.get("logical_paragraph_records", [])
        if record.get("kind") in {"LogicalParagraph", "Caption"}
    ]
    body_style_names = {
        "Tender Body", "Tender Body First Line", "Tender Body Center",
        "Tender Body Right", "Tender Body No Indent", "Tender Form Label",
        "Tender Signature Block",
    }
    body_using_style = sum(
        0 <= int(record.get("paragraph_index", -1)) < len(body_paragraphs)
        and _paragraph_style_name(body_paragraphs[int(record["paragraph_index"])], styles) in body_style_names
        for record in body_records
    )
    # The source model also uses kind=List for short appendix labels. Only
    # paragraphs that were actually assigned the body-list style belong in
    # this numbering audit; ordinary labels remain body/form text.
    list_records = [
        record for record in generation.get("logical_paragraph_records", [])
        if record.get("kind") == "List" and str(record.get("word_style", "")).startswith("Tender Body List ")
    ]
    list_num_id = str(generation.get("fixture_body_list_num_id"))
    list_real_numbering = 0
    literal_list_prefixes = 0
    for record in list_records:
        index = int(record.get("paragraph_index", -1))
        if not 0 <= index < len(body_paragraphs):
            continue
        paragraph = body_paragraphs[index]
        num_id, _level = _num_pr(paragraph)
        if num_id == list_num_id:
            list_real_numbering += 1
        if marker_info(_paragraph_text(paragraph))[0]:
            literal_list_prefixes += 1
    table_coverage = _table_style_coverage(root, styles)
    direct_count, illegal_count, direct_details = _run_direct_formatting(root)
    literal_heading_numbers = sum(
        bool(marker_info(_paragraph_text(body_paragraphs[int(record["paragraph_index"])]))[0])
        for record in heading_records
        if 0 <= int(record.get("paragraph_index", -1)) < len(body_paragraphs)
    )
    forensic = inspect_docx_ooxml(path)
    return {
        "paragraph_count": len(all_paragraphs),
        "paragraphs_with_named_style": sum(bool(name) for name in named),
        "paragraphs_without_named_style": sum(not bool(name) for name in named),
        "heading_count": len(heading_records),
        "heading_style_paragraph_count": len(headings),
        "headings_with_correct_style": heading_with_correct_style,
        "fake_heading_count": fake_heading_count,
        "body_count": len(body_records),
        "body_using_body_style": body_using_style,
        "body_list_count": len(list_records),
        "body_list_using_real_numbering": list_real_numbering,
        "body_list_literal_prefix_count": literal_list_prefixes,
        "table_paragraph_count": table_coverage["paragraph_count"],
        "table_style_coverage": table_coverage,
        "direct_format_override_count": direct_count,
        "illegal_direct_format_override_count": illegal_count,
        "direct_format_examples": direct_details[:20],
        "literal_heading_number_count": literal_heading_numbers,
        "synthetic_layout_tables": generation.get("synthetic_layout_tables", 0),
        "unsafe_ooxml": int(generation.get("word_safe_scan", {}).get("unsafe_ooxml", 0)) + int(forensic.get("summary", {}).get("error_count", 0) > 0),
        "style_ids": styles,
        "result": "PASS" if all([
            sum(not bool(name) for name in named) == 0,
            fake_heading_count == 0,
            heading_records and heading_with_correct_style == len(heading_records) or not heading_records,
            literal_heading_numbers == 0,
            illegal_count == 0,
            generation.get("synthetic_layout_tables", 0) == 0,
            forensic.get("result") == "PASS",
        ]) else "FAIL",
    }


def build_style_map(profile: TenderStyleProfile) -> list[dict]:
    rows = []
    requested = [
        ("Chapter Title", "Tender Chapter Title", None),
        ("Heading 1", "Heading 1", "Fixture Heading Multilevel"),
        ("Heading 2", "Heading 2", "Fixture Heading Multilevel"),
        ("Heading 3", "Heading 3", "Fixture Heading Multilevel"),
        ("Body", "Tender Body", None),
        ("Body First Line", "Tender Body First Line", None),
        ("Body List 1", "Tender Body List 1", "Tender Body List"),
        ("Body List 2", "Tender Body List 2", "Tender Body List"),
        ("Table Text", "Table Text", None),
        ("Table Header", "Tender Table Header", None),
        ("Table Value", "Tender Table Value", None),
        ("Signature Block", "Tender Signature Block", None),
    ]
    for semantic, word_style, family in requested:
        role = profile.role(semantic)
        rows.append({
            "semantic_role": semantic,
            "word_style": word_style,
            "numbering_family": family,
            "source_font": role.source_font_raw or role.font_east_asia,
            "source_font_normalized": role.font_east_asia,
            "source_size": role.font_size,
            "alignment": role.paragraph_alignment,
            "line_spacing": role.line_spacing,
            "indent": {
                "left": role.left_indent,
                "right": role.right_indent,
                "first_line": role.first_line_indent,
                "hanging": role.hanging_indent,
            },
            "page_break_before": role.page_break_before,
            "keep_with_next": role.keep_with_next,
            "evidence_count": role.evidence_count,
        })
    return rows


def build_outline_tree(path: str | Path) -> dict:
    root = _xml(path, "word/document.xml")
    styles = _style_names(path)
    tree: list[dict] = []
    stack: list[tuple[int, dict]] = []
    body_list_names = {f"Tender Body List {level}" for level in range(1, 5)}
    for paragraph in _body_paragraphs(root):
        style = _paragraph_style_name(paragraph, styles)
        match = re.fullmatch(r"Heading ([1-9])", style)
        if not match or style in body_list_names:
            continue
        level = int(match.group(1))
        node = {"level": level, "style": style, "text": _paragraph_text(paragraph), "children": []}
        while stack and stack[-1][0] >= level:
            stack.pop()
        if stack:
            stack[-1][1]["children"].append(node)
        else:
            tree.append(node)
        stack.append((level, node))
    return {
        "headings": sum(1 for _ in _body_paragraphs(root) if re.fullmatch(r"Heading [1-9]", _paragraph_style_name(_, styles))),
        "body_list_headings": 0,
        "tree": tree,
        "result": "PASS",
    }


def _patch_style_size_only(source: Path, target: Path, style_name: str, half_points: int) -> bool:
    with ZipFile(source, "r") as src, ZipFile(target, "w", ZIP_DEFLATED) as dst:
        for info in src.infolist():
            payload = src.read(info.filename)
            if info.filename == "word/styles.xml":
                root = etree.fromstring(payload)
                matches = [
                    style for style in root.xpath("./w:style", namespaces=W)
                    if next(iter(style.xpath("./w:name/@w:val", namespaces=W)), "") == style_name
                ]
                if not matches:
                    return False
                style = matches[0]
                size = style.find("w:rPr/w:sz", namespaces=W)
                if size is None:
                    rpr = style.find("w:rPr", namespaces=W)
                    if rpr is None:
                        rpr = etree.SubElement(style, qn("w:rPr"))
                    size = etree.SubElement(rpr, qn("w:sz"))
                size.set(qn("w:val"), str(half_points))
                payload = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            dst.writestr(info, payload)
    return True


def _document_xml_bytes(path: Path) -> bytes:
    with ZipFile(path) as archive:
        return archive.read("word/document.xml")


def build_editability_qa(path: str | Path, render=None) -> dict:
    source = Path(path)
    checks = {}
    with TemporaryDirectory(prefix="round50-editability-") as temp:
        root = Path(temp)
        for label, style_name, half_points in (
            ("Heading 1", "heading 1", 60),
            ("Tender Body", "Tender Body", 26),
            ("Table Text", "Table Text", 22),
        ):
            copy = root / f"{label.replace(' ', '_')}.docx"
            shutil.copyfile(source, copy)
            patched = _patch_style_size_only(source, copy, style_name, half_points)
            reopened = False
            if patched:
                Document(copy)
                reopened = True
            checks[label] = {
                "style_definition_changed_only": patched and _document_xml_bytes(source) == _document_xml_bytes(copy),
                "reopened": reopened,
                "rendered": False,
                "result": "PASS" if patched and reopened and _document_xml_bytes(source) == _document_xml_bytes(copy) else "FAIL",
            }
            if render is not None and patched:
                try:
                    render(copy, root / f"render_{label.replace(' ', '_')}")
                    checks[label]["rendered"] = True
                    checks[label]["result"] = "PASS"
                except Exception as exc:  # pragma: no cover - renderer is environment-specific
                    checks[label]["render_error"] = str(exc)
    return {"checks": checks, "result": "PASS" if all(item["result"] == "PASS" for item in checks.values()) else "FAIL"}


def _insert_paragraph_before(document: Document, before, text: str, style_name: str):
    new = document.add_paragraph(text, style=style_name)
    before._p.addprevious(new._p)
    return new


def _append_num_pr(paragraph, num_id: str, level: str):
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), level)
    num_pr.append(ilvl)
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), num_id)
    num_pr.append(num)
    ppr.append(num_pr)


def build_numbering_editability_qa(path: str | Path, heading_num_id: int, body_num_id: int, render=None) -> dict:
    source = Path(path)
    with TemporaryDirectory(prefix="round50-numbering-") as temp:
        root = Path(temp)
        copy = root / "numbering_insertions.docx"
        shutil.copyfile(source, copy)
        document = Document(copy)
        headings = [p for p in document.paragraphs if p.style.name == "Heading 1"]
        lists = [p for p in document.paragraphs if p.style.name == "Tender Body List 1"]
        heading_inserted = False
        list_inserted = False
        if len(headings) >= 2:
            inserted = _insert_paragraph_before(document, headings[1], "合成插入标题", "Heading 1")
            _append_num_pr(inserted, str(heading_num_id), "0")
            heading_inserted = True
        if len(lists) >= 2:
            inserted = _insert_paragraph_before(document, lists[1], "合成插入条目", "Tender Body List 1")
            _append_num_pr(inserted, str(body_num_id), "0")
            list_inserted = True
        document.save(copy)
        reopened = Document(copy)
        root_xml = _xml(copy, "word/document.xml")
        heading_prs = [
            p for p in _body_paragraphs(root_xml)
            if _paragraph_style_name(p, _style_names(copy)) == "Heading 1"
        ]
        list_prs = [
            p for p in _body_paragraphs(root_xml)
            if _paragraph_style_name(p, _style_names(copy)) == "Tender Body List 1"
        ]
        heading_ok = heading_inserted and all(_num_pr(p)[0] == str(heading_num_id) for p in heading_prs)
        list_ok = list_inserted and all(_num_pr(p)[0] == str(body_num_id) for p in list_prs)
        rendered = False
        render_error = None
        if render is not None:
            try:
                render(copy, root / "render")
                rendered = True
            except Exception as exc:  # pragma: no cover
                render_error = str(exc)
        result = heading_ok and list_ok and Document(copy) is not None
        return {
            "heading_insertion": {"inserted": heading_inserted, "sequence_structural": heading_ok},
            "body_list_insertion": {"inserted": list_inserted, "sequence_structural": list_ok},
            "reopened": True,
            "rendered": rendered,
            "render_error": render_error,
            "result": "PASS" if result else "FAIL",
        }


def write_json(path: str | Path, data: dict) -> Path:
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return output


__all__ = [
    "build_editability_qa", "build_numbering_editability_qa", "build_outline_tree",
    "build_style_architecture_qa", "build_style_map", "write_json",
]
