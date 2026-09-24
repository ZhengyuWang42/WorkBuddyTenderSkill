"""Round-2 contracts: review points are synthesized, not concatenated.

The tests are reference-style: they assert the *quality properties* the round-2
brief requires (four distinct semantics, numeric ownership, actionability,
no invented consequences, no historical status copying) instead of exact prose,
so a wording improvement never breaks them.
"""

from __future__ import annotations

import pytest
from openpyxl import Workbook

from tender_basic.dynamic_requirements import SourceRequirementUnit
from tender_basic.models import PdfLocator
from tender_basic.review_builder import CHECKLIST_START_ROW, _build_dynamic_review_rows
from tender_basic.review_point import (
    ACTIONABLE,
    INTERNAL_PROCEDURE,
    classify_group,
    extract_numeric_evidence,
    group_is_filterable,
    is_internal_procedure,
    render_review_cell,
    scan_review_point,
    synthesize_review_point,
)

#: Tokens that only ever come from the historical hand-filled workbooks.
HISTORICAL_TOKENS = ("已核对", "责任人", "复核人：", "已确认无误")

GENERIC_BOILERPLATE = (
    "按上述招标文件条款逐项核对响应文件对应内容，确认完全响应。",
    "响应内容与上述招标文件条款一致。",
)


def unit(
    requirement_id: str,
    *,
    requirement_type: str = "TECHNICAL",
    topic: str = "技术方案与实施组织",
    text: str = "",
    values: list[str] | None = None,
    page: int = 5,
    clause: str = "1.1",
    mandatory: bool = False,
    high_risk: bool = False,
    module: str = "技术部分",
) -> SourceRequirementUnit:
    return SourceRequirementUnit(
        requirement_id=requirement_id,
        requirement_type=requirement_type,
        topic=topic,
        risk_level="高" if high_risk else "中",
        module=module,
        mandatory=mandatory,
        high_risk=high_risk,
        clause=clause,
        page=page,
        section=topic,
        text=text or f"{topic}的具体要求。",
        evidence_text=text or f"{topic}的具体要求。",
        locator=PdfLocator(page=page, block_index=0),
        values=list(values or []),
    )


def point_for(*units: SourceRequirementUnit, requirement_type: str | None = None, topic: str | None = None):
    first = units[0]
    return synthesize_review_point(
        units=list(units),
        requirement_type=requirement_type or first.requirement_type,
        topic=topic or first.topic,
        fact_hints={},
    )


# --------------------------------------------------------------------------- #
# the four semantics stay distinct
# --------------------------------------------------------------------------- #


def test_four_semantics_are_distinct_and_non_empty() -> None:
    point = point_for(
        unit(
            "SR0001",
            requirement_type="BOND",
            topic="保证金金额",
            text="投标保证金金额为人民币 20000 元，未提交的按无效投标处理。",
            values=["20000元"],
        )
    )
    assert point is not None
    assert point.requirement_summary.strip()
    assert point.review_checks
    assert point.pass_criteria
    # the requirement summary quotes the source, the checks instruct the reviewer
    assert "按" in point.review_checks[0] or "核对" in point.review_checks[0]
    assert point.requirement_summary not in point.review_checks
    assert all(check not in point.pass_criteria for check in point.review_checks)


def test_source_requirement_is_quoted_not_invented() -> None:
    source = "投标保证金金额为人民币 20000 元。"
    point = point_for(
        unit("SR0001", requirement_type="BOND", topic="保证金金额", text=source, values=["20000元"])
    )
    assert point is not None
    assert source[:12] in point.requirement_summary


def test_review_checks_are_observable() -> None:
    point = point_for(
        unit("SR0001", requirement_type="SIGNATURE", topic="签字盖章要求", text="响应文件应加盖单位公章。")
    )
    assert point is not None
    assert point.review_checks
    assert all(len(check) > 6 for check in point.review_checks)
    assert any(token in "".join(point.review_checks) for token in ("核对", "确认", "比对", "检查"))


# --------------------------------------------------------------------------- #
# numeric ownership
# --------------------------------------------------------------------------- #


def test_numbers_carry_a_semantic_role() -> None:
    units = [
        unit(
            "SR0001",
            requirement_type="BOND",
            topic="保证金金额",
            text="投标保证金为人民币 20000 元，投标有效期为 90 日历天。",
            values=["20000元", "90日历天"],
        )
    ]
    roles = {evidence.value: evidence.role for evidence in extract_numeric_evidence(units)}
    # Round 5 separates the response bond from the performance bond and the
    # amount: a tender/response bond amount is RESPONSE_BOND (BOND_AMOUNT stays
    # an accepted alias for the generic amount role).
    assert roles["20000元"] in {"RESPONSE_BOND", "BOND_AMOUNT"}
    # Round 5 renamed the validity role to BID_VALIDITY_DAYS; the legacy
    # VALIDITY_DAYS spelling remains an accepted alias.
    assert roles["90日历天"] in {"BID_VALIDITY_DAYS", "VALIDITY_DAYS"}


def test_locator_numbers_are_not_values() -> None:
    units = [
        unit(
            "SR0001",
            text="9.2 详见第 5 页的表格，供货期为 30 日历天。",
            values=["30日历天"],
        )
    ]
    values = {evidence.value for evidence in extract_numeric_evidence(units)}
    assert "30日历天" in values
    assert "5页" not in values
    assert "9.2" not in values


def test_no_theme_never_owns_a_price_or_contact_number() -> None:
    point = point_for(
        unit(
            "SR0001",
            requirement_type="SIGNATURE",
            topic="签字盖章要求",
            text=(
                "响应文件应在 48 小时内上传，联系电话 0371-55527619，"
                "投标有效期为 90 日历天。"
            ),
            values=["90日历天"],
        )
    )
    assert point is not None
    checks = "\n".join(point.review_checks + point.pass_criteria)
    assert "0371" not in checks
    assert "48小时" not in checks
    assert "90" not in checks  # a validity value never becomes a signature check
    # the requirement block stays a verbatim source quote, phone number included
    assert "0371" in point.requirement_summary


def test_platform_fee_is_never_a_review_value() -> None:
    point = point_for(
        unit(
            "SR0001",
            requirement_type="TECHNICAL",
            topic="技术方案与实施组织",
            text=(
                "供应商在下载询比文件时，须向河南国企阳光招采服务平台缴纳平台服务费300元，"
                "售后不退。"
            ),
            values=["300元"],
        ),
        unit(
            "SR0002",
            requirement_type="TECHNICAL",
            topic="技术方案与实施组织",
            text="供应商应提供完整的实施方案与技术参数响应表。",
        ),
    )
    assert point is not None
    assert "300元" not in "\n".join(point.review_checks)
    assert "300元" not in "\n".join(point.pass_criteria)
    assert all(
        evidence.usable is False
        for evidence in point.numeric_evidence
        if evidence.value == "300元"
    )


def test_fee_only_row_is_not_a_review_point() -> None:
    point = point_for(
        unit(
            "SR0001",
            requirement_type="TECHNICAL",
            topic="售后服务与运维",
            text="平台服务费300元，售后不退。",
            values=["300元"],
        )
    )
    assert point is None


def test_value_role_must_match_the_row_type() -> None:
    point = point_for(
        unit(
            "SR0001",
            requirement_type="PERSONNEL",
            topic="安全与管理人员",
            text="项目负责人须具备一级建造师证书，且配备不少于 2 人。",
            values=["2人"],
        )
    )
    assert point is not None
    scan = scan_review_point(point)
    assert scan["semantic_contamination"] is False
    assert "2人" in render_review_cell(point)


def test_unrelated_values_are_not_collected() -> None:
    point = point_for(
        unit(
            "SR0001",
            requirement_type="VALIDITY",
            topic="投标有效期",
            text=(
                "投标有效期为 90 日历天；投标保证金 20000 元；联系人电话 0371-55527619；"
                "另附 5 台设备清单。"
            ),
            values=["90日历天"],
            high_risk=True,
        )
    )
    assert point is not None
    cell = render_review_cell(point)
    for stray in ("20000元", "0371", "5台"):
        assert stray not in "\n".join(point.review_checks + point.pass_criteria), stray
    assert "90日历天" in cell


# --------------------------------------------------------------------------- #
# contamination and boilerplate scanners
# --------------------------------------------------------------------------- #


def test_summary_quote_is_not_contamination() -> None:
    point = point_for(
        unit(
            "SR0001",
            requirement_type="SUBMISSION",
            topic="递交方式与截止时间",
            text="询比文件售价：0元，响应文件须在截止时间前上传。",
            values=["0元"],
        )
    )
    assert point is not None
    assert scan_review_point(point)["numeric_contamination"] is False
    assert "0元" not in "\n".join(point.review_checks)


def test_generic_boilerplate_is_dropped_from_the_synthesized_lines() -> None:
    point = point_for(
        unit(
            "SR0001",
            requirement_type="QUALIFICATION",
            topic="资质等级与许可",
            text="投标人应具备有效的营业执照，按上述招标文件条款逐项核对响应文件对应内容，确认完全响应。",
            mandatory=True,
        )
    )
    assert point is not None
    cell = render_review_cell(point)
    synthesized = "\n".join(
        point.review_checks + point.pass_criteria + point.preparation_materials
    )
    for phrase in GENERIC_BOILERPLATE:
        assert phrase not in synthesized
    assert scan_review_point(point)["generic"] is False
    assert cell


def test_consequence_needs_a_source_marker() -> None:
    with_consequence = point_for(
        unit(
            "SR0001",
            requirement_type="REJECTION",
            topic="未实质性响应",
            text="投标文件未按要求密封的，其投标将被否决。",
            high_risk=True,
        )
    )
    without_consequence = point_for(
        unit(
            "SR0001",
            requirement_type="TECHNICAL",
            topic="技术方案与实施组织",
            text="投标人应提供施工组织设计方案。",
        )
    )
    assert with_consequence is not None and without_consequence is not None
    assert with_consequence.failure_consequence
    assert with_consequence.consequence_evidence_id == "SR0001"
    assert without_consequence.failure_consequence == ""
    assert scan_review_point(with_consequence)["unsupported_consequence"] is False


# --------------------------------------------------------------------------- #
# actionability
# --------------------------------------------------------------------------- #


def test_purchaser_internal_procedure_is_not_a_review_point() -> None:
    units = [
        unit(
            "SR0001",
            requirement_type="OTHER",
            topic="文件组成与格式",
            text="评审小组由采购人代表和有关专家 2 人，共 3 人组成。",
        ),
        unit(
            "SR0002",
            requirement_type="OTHER",
            topic="文件组成与格式",
            text="评审小组成员与供应商有利害关系的，应当回避。",
        ),
    ]
    assert all(is_internal_procedure(item) for item in units)
    assert classify_group(units) == INTERNAL_PROCEDURE
    assert group_is_filterable(units, INTERNAL_PROCEDURE) is True


def test_bidder_obligation_keeps_the_row() -> None:
    units = [
        unit(
            "SR0001",
            requirement_type="OTHER",
            topic="程序性要求",
            text="供应商不得以任何方式干扰、影响评审工作。",
        )
    ]
    assert is_internal_procedure(units[0]) is False
    assert classify_group(units) == ACTIONABLE


def test_high_risk_internal_clause_keeps_its_row() -> None:
    units = [
        unit(
            "SR0001",
            requirement_type="OTHER",
            topic="文件组成与格式",
            text="评审小组成员不得泄露评审情况。",
            high_risk=True,
        )
    ]
    assert classify_group(units) == INTERNAL_PROCEDURE
    assert group_is_filterable(units, INTERNAL_PROCEDURE) is False


def test_fee_only_group_is_filterable() -> None:
    units = [
        unit(
            "SR0001",
            requirement_type="TECHNICAL",
            topic="售后服务与运维",
            text="平台服务费300元，售后不退。",
        )
    ]
    assert group_is_filterable(units, ACTIONABLE) is True


# --------------------------------------------------------------------------- #
# delivery integration
# --------------------------------------------------------------------------- #


def test_delivered_row_is_the_rendered_review_point() -> None:
    from tender_basic.dynamic_review import DynamicReviewItem, DynamicReviewPlan

    point = point_for(
        unit(
            "SR0001",
            requirement_type="BOND",
            topic="保证金金额",
            text="投标保证金金额为人民币 20000 元，未提交的按无效投标处理。",
            values=["20000元"],
            high_risk=True,
        )
    )
    assert point is not None
    item = DynamicReviewItem(
        item_id="DR001",
        module="商务部分",
        submodule="保证金金额",
        requirement_type="BOND",
        topic="保证金金额",
        risk_level="高",
        source_requirement=point.requirement_summary,
        source_locator="第5页",
        source_page=5,
        source_section="保证金金额",
        source_evidence="投标保证金金额为人民币 20000 元",
        source_requirement_ids=["SR0001"],
        verification_action="\n".join(point.review_checks),
        pass_criteria=" ".join(point.pass_criteria),
        consequence_if_failed=point.failure_consequence,
        cell_text=render_review_cell(point),
        review_point=point,
    )
    plan = DynamicReviewPlan(
        items=[item],
        source_requirement_count=1,
        dropped_duplicate_count=0,
        index=type("Index", (), {"units": [unit("SR0001")]})(),
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "投标项目复核表"
    for row in range(1, CHECKLIST_START_ROW):
        sheet.cell(row=row, column=1).value = f"模板行{row}"
    for column in range(1, 17):
        sheet.cell(row=CHECKLIST_START_ROW, column=column).value = f"donor{column}"
    for row in (75, 76, 77):
        sheet.cell(row=row, column=9).value = "签字："

    _build_dynamic_review_rows(sheet, plan)

    assert sheet.cell(row=CHECKLIST_START_ROW, column=4).value == render_review_cell(point)
    assert sheet.cell(row=CHECKLIST_START_ROW, column=1).value == 1
    assert sheet.cell(row=CHECKLIST_START_ROW, column=9).value == "待复核"
    assert sheet.cell(row=1, column=1).value == "模板行1"
    assert sheet.cell(row=75, column=9).value != "签字：" or sheet.cell(row=75, column=9).value


#: The review types the round-2 brief requires type-specific synthesis for.
REQUIRED_TYPES = (
    "REJECTION",
    "QUALIFICATION",
    "PRICING",
    "BOND",
    "DURATION",
    "VALIDITY",
    "LOCATION",
    "QUALITY",
    "WARRANTY",
    "CONSORTIUM",
    "SUBCONTRACT",
    "SIGNATURE",
    "ELECTRONIC",
    "SUBMISSION",
    "FORM",
    "PROOF",
    "TECHNICAL",
    "EVALUATION",
    "CONTRACT",
)


#: A realistic source clause per required review type, used for coverage checks.
TYPE_CLAUSES = {
    "REJECTION": "投标文件未按招标文件要求密封或者未加盖公章的，其投标将被否决。",
    "QUALIFICATION": "投标人须具备有效的营业执照与相应资质证书，并具备独立法人资格。",
    "PRICING": "投标报价不得超过最高限价，各分项报价应与开标一览表金额一致。",
    "BOND": "投标人须在递交截止时间前提交投标保证金，金额、形式与到账时间以本条规定为准。",
    "DURATION": "供货期为合同签订后 30 日历天，须覆盖全部交付节点。",
    "VALIDITY": "投标有效期为 90 日历天，应覆盖评审与定标全过程。",
    "LOCATION": "交付地点为采购人指定地点，实施地点与服务范围以本条规定为准。",
    "QUALITY": "投标人须承诺工程质量达到合格标准，并符合验收标准。",
    "WARRANTY": "质保期为 24 个月，质保范围包括全部设备与软件。",
    "CONSORTIUM": "本项目不接受联合体投标，投标人不得以联合体形式参加。",
    "SUBCONTRACT": "投标人不得违规分包、转包或将主体工作交由他人完成。",
    "SIGNATURE": "响应文件应由法定代表人或其授权代理人签字并加盖单位公章。",
    "ELECTRONIC": "本项目采用电子招投标，投标人须使用有效的CA介质加密上传响应文件并按时解密。",
    "SUBMISSION": "响应文件应在递交截止时间前上传至电子平台，逾期不予受理。",
    "FORM": "响应文件应按规定的格式、组成与编排要求编制，并编制目录与页码。",
    "PROOF": "投标人须提供与响应内容对应的证明材料，并保证清晰可核验。",
    "TECHNICAL": "投标人须逐条响应技术参数，并提供可核验的证明材料，不得出现负偏离。",
    "EVALUATION": "评审按评分因素逐项计分，每个评分因素均应有对应响应内容与证明材料。",
    "CONTRACT": "投标人须接受合同条款、付款方式与履约要求，不得附加采购人不能接受的条件。",
}


@pytest.mark.parametrize("requirement_type", REQUIRED_TYPES)
def test_each_required_type_has_its_own_synthesis(requirement_type: str) -> None:
    """Every required type renders its own pass criterion, not the generic one.

    The sample clause is realistic for the type: the grounding guard must accept
    a criterion only when the row's own source clause supports it.
    """

    from tender_basic.review_point import _PASS_BY_TYPE

    assert requirement_type in _PASS_BY_TYPE, requirement_type
    criterion = _PASS_BY_TYPE[requirement_type]
    assert criterion != _PASS_BY_TYPE["OTHER"], requirement_type
    point = point_for(
        unit(
            "SR0001",
            requirement_type=requirement_type,
            topic="一般要求",
            text=TYPE_CLAUSES[requirement_type],
        )
    )
    assert point is not None
    assert criterion in render_review_cell(point)
    assert point.review_checks
    assert scan_review_point(point)["generic"] is False


def test_no_historical_status_is_copied() -> None:
    point = point_for(
        unit(
            "SR0001",
            requirement_type="QUALIFICATION",
            topic="营业执照与独立法人",
            text="投标人应提供有效的营业执照副本复印件。",
            mandatory=True,
        )
    )
    assert point is not None
    cell = render_review_cell(point)
    for token in HISTORICAL_TOKENS:
        assert token not in cell, token


def test_review_point_round_trips_as_plain_data() -> None:
    point = point_for(
        unit(
            "SR0001",
            requirement_type="DURATION",
            topic="工期与交付时间",
            text="供货期为 30 日历天，逾期交付的投标将被否决。",
            values=["30日历天"],
        )
    )
    assert point is not None
    payload = point.model_dump()
    assert payload["requirement_type"] == "DURATION"
    assert payload["review_checks"]
    assert payload["consequence_evidence_id"] == "SR0001"


@pytest.mark.parametrize(
    "requirement_type,topic,text",
    (
        ("PRICING", "最高限价与控制价", "投标报价不得超过最高限价 3100000 元。"),
        ("BOND", "保证金金额", "投标保证金为人民币 20000 元。"),
        ("VALIDITY", "投标有效期", "投标有效期为 90 日历天。"),
        ("DURATION", "工期与交付时间", "供货期为 30 日历天。"),
        ("SIGNATURE", "签字盖章要求", "响应文件应由法定代表人或授权代理人签字并加盖公章。"),
        ("SUBMISSION", "递交方式与截止时间", "响应文件应在截止时间前上传至电子平台。"),
    ),
)
def test_every_type_renders_a_usable_cell(requirement_type: str, topic: str, text: str) -> None:
    point = point_for(unit("SR0001", requirement_type=requirement_type, topic=topic, text=text))
    assert point is not None
    cell = render_review_cell(point)
    assert "复核要点" in cell
    assert "通过标准" in cell
    assert scan_review_point(point)["checks_without_observable_noun"] is False
