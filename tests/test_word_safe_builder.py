"""Production API and corruption regressions for the Word-safe boundary."""
import ast
import inspect
from pathlib import Path
from zipfile import ZipFile
import pytest
from docx import Document
from lxml import etree
from tender_basic import word_safe_source_builder as safe
from tender_basic.word_safe_scan import scan_word_safe_docx
from tender_basic.models import FieldName
from tender_basic.source_format import _paragraph_slots
from test_review_builder import make_project_facts
from test_source_format_fidelity import _paragraph, _paragraph_template, _resolved
from test_source_format_fidelity import _table_template
from scripts.build_word_safe_canary import build as canary


def test_clean_document_bootstrap_and_no_prior_docx_load(tmp_path, monkeypatch):
    calls=[]
    real=safe.Document
    def fresh(*args,**kwargs):
        calls.append(args)
        assert args == () and not kwargs
        return real()
    monkeypatch.setattr(safe,'Document',fresh)
    template=_paragraph_template(_paragraph('源格式\n第二行'))
    out,_=safe.build_source_format_docx(make_project_facts(),template,tmp_path/'safe.docx')
    assert calls==[()]
    assert Document(out).paragraphs[0].text=='源格式\n第二行'
    assert scan_word_safe_docx(out)['result']=='PASS'


def test_production_has_no_arbitrary_xml_or_relationship_authoring():
    root=Path(__file__).resolve().parents[1]/'tender_basic'
    for name in ('word_safe_source_builder.py','bid_document_builder.py','source_fill_policy.py'):
        code=(root/name).read_text(encoding='utf-8')
        tree=ast.parse(code)
        assert 'source_docx_builder' not in code
        for node in ast.walk(tree):
            if isinstance(node,ast.Attribute):
                assert node.attr not in {'_element','_tbl','_p','_r','relate_to','load_rel','SubElement','OxmlElement','parse_xml'}
    assert 'Document()' in inspect.getsource(safe.new_word_safe_document)


def test_canary_all_production_constructs(tmp_path):
    out=canary(tmp_path/'canary.docx')
    result=scan_word_safe_docx(out)
    assert result['result']=='PASS'
    assert result['forbidden_elements']==result['forbidden_relationships']==[]
    doc=Document(out)
    assert len(doc.sections)==2 and len(doc.tables)==2
    assert doc.tables[0].cell(0,0).text == doc.tables[0].cell(0,1).text
    with ZipFile(out) as z:
        xml=etree.fromstring(z.read('word/document.xml'))
        ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        assert xml.xpath('.//w:gridSpan',namespaces=ns)
        assert xml.xpath('.//w:vMerge',namespaces=ns)
        assert xml.xpath('.//w:r/w:br[@w:type="page"]',namespaces=ns)
        assert not xml.xpath('.//w:p/w:br',namespaces=ns)
        assert not any(n.startswith('customXml/') for n in z.namelist())


@pytest.mark.parametrize('text,value,expected',[
    ('项目名称：________','脱敏工程','项目名称：脱敏工程'),
    ('（项目名称）','脱敏工程','脱敏工程'),
    ('供货期        日历天','90日历天','供货期90日历天'),
])
def test_safe_filling_preserves_source_slot(tmp_path,text,value,expected):
    field=FieldName.DURATION if '供货期' in text else FieldName.PROJECT_NAME
    facts=_resolved(make_project_facts(),field,value)
    p=_paragraph(text)
    slots=_paragraph_slots(p,0)
    template=_paragraph_template(p,slots=slots)
    out,report=safe.build_source_format_docx(facts,template,tmp_path/'filled.docx')
    actual=Document(out).paragraphs[0].text
    assert ''.join(actual.split())==''.join(expected.split())
    assert report['fill_slots_filled']==1
    from tender_basic.source_format_qa import build_source_format_qa
    qa=build_source_format_qa(template,facts,out,generation_report=report)
    assert qa['result']=='PASS_WITH_DEGRADATIONS'
    assert qa['page_count_generated'] is None
    assert qa['word_normal_open']=='PENDING_MANUAL_CONFIRMATION'
    if '_' in text:
        assert any(r.text==value and r.underline for r in Document(out).paragraphs[0].runs)


def test_unknown_and_fixed_and_bidder_slots_untouched(tmp_path):
    from tender_basic.models import ResolvedFact, FactStatus
    facts=make_project_facts()
    facts.fields.project_name=ResolvedFact(field=FieldName.PROJECT_NAME, status=FactStatus.NOT_FOUND, confidence=0, resolution_reason='Absent in synthetic source')
    text='项目名称：________\n投标人：________\n供货期：合同签订后30天\n年    月    日'
    p=_paragraph(text)
    out,_=safe.build_source_format_docx(facts,_paragraph_template(p,slots=_paragraph_slots(p,0)),tmp_path/'blank.docx')
    actual='\n'.join(paragraph.text for paragraph in Document(out).paragraphs)
    assert '项目名称：' in actual
    assert '投标人：' in actual
    assert '供货期：合同签订后30天' in actual
    # The packed synthetic paragraph exposes only the project-name span to
    # the slot model; the unclassified bidder marker remains source text.
    assert '项目名称：________' not in actual
    assert '\u00a0' not in actual
    assert any(run.underline and '\u2007' in run.text for paragraph in Document(out).paragraphs for run in paragraph.runs)


def test_table_slot_fixed_value_widths_and_heights(tmp_path):
    from docx.enum.table import WD_ROW_HEIGHT_RULE
    template,_=_table_template()
    out,report=safe.build_source_format_docx(make_project_facts(),template,tmp_path/'table.docx')
    table=Document(out).tables[0]
    assert table.cell(0,1).text=='测试工程'
    assert table.cell(1,1).text=='原模板固定质量标准'
    assert [round(c.width.pt) for c in table.columns]==[100,200]
    assert table.rows[0].height.pt==28
    assert table.rows[0].height_rule==WD_ROW_HEIGHT_RULE.AT_LEAST
    assert report['fill_slots_filled']==1


def test_merge_duplicate_rectangles_once_but_distinct_content_degrades(tmp_path):
    template,_=_table_template()
    template.fill_slots=[]
    source=template.source_pages[0].tables[0]
    source.merged_cells=[(0,0,0,1)]
    first,second=source.rows[0].cells
    second.text=first.text
    second.bbox=first.bbox
    out,report=safe.build_source_format_docx(make_project_facts(),template,tmp_path/'merged.docx')
    with ZipFile(out) as z:
        tree=etree.fromstring(z.read('word/document.xml'))
        assert len(tree.xpath('.//w:t[text()="项目名称"]',namespaces={'w':safe_scan_ns()}))==1
    second.text='独立固定条款'
    out,report=safe.build_source_format_docx(make_project_facts(),template,tmp_path/'rectangular.docx')
    assert Document(out).tables[0].cell(0,1).text=='独立固定条款'
    assert any('TABLE_MERGE_DEGRADED' in w for w in report['layout_warnings'])


def safe_scan_ns():
    return 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


def test_watermark_text_does_not_remove_same_text_from_cover():
    from types import SimpleNamespace
    from tender_basic.document_models import DocumentBlock, DocumentPage, NormalizedDocument, DocumentStatus
    from tender_basic.models import PdfLocator, SourceType
    from tender_basic.source_format import build_source_format_model
    pages=[]
    for n in (1,2,3):
        blocks=[DocumentBlock(block_index=0,text='脱敏项目标题',bbox=(170,300,420,560),locator=PdfLocator(page=n,block_index=0)),
                DocumentBlock(block_index=2,text='2025-01-02 12:34:56',bbox=(200,350,390,480),locator=PdfLocator(page=n,block_index=2))]
        if n==1:
            blocks.append(DocumentBlock(block_index=1,text='脱敏项目标题',bbox=(130,220,470,245),locator=PdfLocator(page=n,block_index=1)))
        pages.append(DocumentPage(page_number=n,width=595,height=842,blocks=blocks))
    document=NormalizedDocument(source_file='synthetic.pdf',source_type=SourceType.PDF,status=DocumentStatus.PARSED,
                                page_count=3,text_length=24,pages=pages)
    fmt=SimpleNamespace(source_locator=PdfLocator(page=1,block_index=1),source_heading='投标文件格式',
                        elements=[SimpleNamespace(locator=PdfLocator(page=3,block_index=0))])
    model=build_source_format_model(document,fmt)
    assert [p.text for pg in model.source_pages for p in pg.paragraphs]==['脱敏项目标题']


@pytest.mark.parametrize('fragment',[
    '<w:tblpPr/>','<w:br/>','<w:p><w:fldSimple w:instr="TEST"/></w:p>',
    '<w:p><w:r><w:pict/></w:r></w:p>',
])
def test_forbidden_construct_is_rejected(tmp_path,fragment):
    good=canary(tmp_path/'good.docx')
    bad=tmp_path/'bad.docx'
    with ZipFile(good) as src,ZipFile(bad,'w') as dest:
        for name in src.namelist():
            data=src.read(name)
            if name=='word/document.xml':
                data=data.replace(b'<w:body>',b'<w:body>'+fragment.encode())
            dest.writestr(name,data)
    assert scan_word_safe_docx(bad)['result']=='FAIL'
