from __future__ import annotations

import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from tender_basic.models import PdfTableLocator
from tender_basic.source_format import (
    SourceCell,
    SourceRow,
    SourceRun,
    SourceTable,
    apply_source_glyph_repairs,
    build_table_typography_profile,
)
from tender_basic.word_ooxml import scan_docx_package


def _rewrite_part(source: Path, target: Path, part: str, transform) -> None:
    with zipfile.ZipFile(source) as incoming, zipfile.ZipFile(target, "w") as outgoing:
        for item in incoming.infolist():
            payload = incoming.read(item.filename)
            if item.filename == part:
                payload = transform(payload)
            outgoing.writestr(item, payload)


def _fresh(tmp_path: Path) -> Path:
    path = tmp_path / "fresh.docx"
    document = Document()
    document.add_paragraph("valid")
    document.save(path)
    return path


def test_no_note_document_has_no_invalid_note_settings(tmp_path: Path) -> None:
    report = scan_docx_package(_fresh(tmp_path))
    assert report["invalid_note_settings"] == []
    assert report["python_docx_reopen"] is True


def test_footnote_settings_without_part_fail(tmp_path: Path) -> None:
    source = _fresh(tmp_path)
    target = tmp_path / "bad-footnote.docx"

    def add_note(payload: bytes) -> bytes:
        root = ET.fromstring(payload)
        note = ET.SubElement(root, qn("w:footnotePr"))
        child = ET.SubElement(note, qn("w:footnote"))
        child.set(qn("w:id"), "0")
        return ET.tostring(root, encoding="utf-8", xml_declaration=True)

    _rewrite_part(source, target, "word/settings.xml", add_note)
    report = scan_docx_package(target)
    assert report["invalid_note_settings"]
    assert report["result"] == "FAIL"


def test_endnote_settings_without_part_fail(tmp_path: Path) -> None:
    source = _fresh(tmp_path)
    target = tmp_path / "bad-endnote.docx"

    def add_note(payload: bytes) -> bytes:
        root = ET.fromstring(payload)
        note = ET.SubElement(root, qn("w:endnotePr"))
        child = ET.SubElement(note, qn("w:endnote"))
        child.set(qn("w:id"), "0")
        return ET.tostring(root, encoding="utf-8", xml_declaration=True)

    _rewrite_part(source, target, "word/settings.xml", add_note)
    assert scan_docx_package(target)["invalid_note_settings"]


def test_missing_relationship_and_content_type_targets_fail(tmp_path: Path) -> None:
    source = _fresh(tmp_path)
    rel_bad = tmp_path / "bad-rel.docx"
    ct_bad = tmp_path / "bad-ct.docx"

    def add_rel(payload: bytes) -> bytes:
        root = ET.fromstring(payload)
        child = ET.SubElement(root, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship")
        child.attrib.update({"Id": "rIdMissing", "Type": "urn:round52:test", "Target": "missing.xml"})
        return ET.tostring(root, encoding="utf-8", xml_declaration=True)

    def add_override(payload: bytes) -> bytes:
        root = ET.fromstring(payload)
        child = ET.SubElement(root, "{http://schemas.openxmlformats.org/package/2006/content-types}Override")
        child.attrib.update({"PartName": "/word/missing.xml", "ContentType": "application/xml"})
        return ET.tostring(root, encoding="utf-8", xml_declaration=True)

    _rewrite_part(source, rel_bad, "word/_rels/document.xml.rels", add_rel)
    _rewrite_part(source, ct_bad, "[Content_Types].xml", add_override)
    assert scan_docx_package(rel_bad)["relationship_targets_missing"]
    assert scan_docx_package(ct_bad)["content_type_targets_missing"]


def test_missing_style_and_numbering_references_fail(tmp_path: Path) -> None:
    style_path = tmp_path / "bad-style.docx"
    document = Document()
    paragraph = document.add_paragraph("style")
    paragraph._p.get_or_add_pPr().append(OxmlElement("w:pStyle"))
    paragraph._p.pPr.pStyle.set(qn("w:val"), "MissingRound52Style")
    document.save(style_path)
    assert scan_docx_package(style_path)["missing_style_refs"]

    num_path = tmp_path / "bad-numbering.docx"
    document = Document()
    paragraph = document.add_paragraph("number")
    num_pr = OxmlElement("w:numPr")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "9999")
    num_pr.append(num_id)
    paragraph._p.get_or_add_pPr().append(num_pr)
    document.save(num_path)
    assert scan_docx_package(num_path)["missing_numbering_refs"]


def test_same_semantic_table_role_preserves_source_variants() -> None:
    table = SourceTable(
        page=1, table_index=0, bbox=(0, 0, 100, 50), columns=1,
        column_widths=[100], rows=[SourceRow(row_index=0, cells=[
            SourceCell(row_index=0, column_index=0, text="variant-a", bbox=(0, 0, 100, 20),
                       runs=[SourceRun(text="variant-a", bbox=(0, 0, 80, 20), font_name="FangSong", font_size=12)],
                       locator=PdfTableLocator(page=1, table_index=0, row_index=0, column_index=0)),
        ]), SourceRow(row_index=1, cells=[
            SourceCell(row_index=1, column_index=0, text="variant-b", bbox=(0, 20, 100, 40),
                       runs=[SourceRun(text="variant-b", bbox=(0, 20, 40, 40), font_name="SimSun", font_size=10.45)],
                       locator=PdfTableLocator(page=1, table_index=0, row_index=1, column_index=0)),
        ])],
    )
    preserved = build_table_typography_profile(table)
    runs = [cell.runs[0] for row in preserved.rows for cell in row.cells]
    assert [(run.font_name, run.font_size) for run in runs] == [("FangSong", 12), ("SimSun", 10.45)]
