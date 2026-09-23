"""Round 5.9 regression: source character pairing and row-scoped label evidence.

Two generalization defects found on a cross-page quotation-table case are
regression-locked here in generic terms:

* a rule composition split must pair each run's *visible* character records with
  the run's visible character offsets by ordinal position, because a character
  record's own ``index`` may address the PDF span stream rather than the run
  text.  Index-keyed pairing silently drops a real source character when the two
  streams disagree (a whitespace record present in one and not the other);
* a glyph-free rule's field binding may only use a neighbouring container's label
  when that container shares the rule's own source visual row.  A container that
  merely precedes the rule in document order belongs to another visual row of the
  same paragraph block - frequently another clause's heading - and binding its
  label prints a value the field never asked for.

One further channel is regression-locked too, because a source sentence that
wraps ends the row above the field's own row: a label sitting at the *end* of
that row is the wrapped field's label (``RESOLVED_VALUE_IN_FIXED_SLOT`` when the
field is resolved).  The channel is deliberately narrow - the row above must end
the label, must not end a sentence, must not end a label terminator, and must be
one bounded row pitch away - and the tests below pin each of those guards as
well as the binding itself.
"""

from __future__ import annotations

from tender_basic.logical_table_provenance import (
    audit_logical_tables,
    normalize_cell_text,
)
from tender_basic.page_layout import character_composition_segments

_WORD_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


class _Character:
    def __init__(self, character, x0, x1, index, kind="EXACT_PDF_CHAR"):
        self.character = character
        self.x0 = x0
        self.x1 = x1
        self.index = index
        self.evidence_kind = kind


class _Run:
    def __init__(self, text, characters, bbox=(85.08, 192.42, 367.08, 204.42)):
        self.text = text
        self.characters = characters
        self.bbox = bbox
        self.font_size = 10.5
        self.origin = None


class _Rule:
    def __init__(self, bbox):
        self.bbox = bbox
        self.orientation = "horizontal"


def _real_document(rows_by_table):
    from docx import Document

    document = Document()
    for rows in rows_by_table:
        columns = max(len(row) for row in rows)
        table = document.add_table(rows=len(rows), cols=columns)
        for row_index, row in enumerate(rows):
            for column_index, text in enumerate(row):
                table.cell(row_index, column_index).text = text
    return document


def _minimal_report(**overrides):
    entry = {
        "table_id": "logical_0",
        "head_page": 1,
        "head_table_index": 0,
        "source_pages": [1],
        "fragment_count": 1,
        "row_count": 2,
        "column_count": 2,
    }
    entry.update(overrides)
    return {"logical_tables": [entry]}


def _span_stream_shifted_run():
    """A run whose character records are indexed one earlier than its text.

    The run text holds a literal space before the digits while the recorded
    character stream does not, which is exactly the stream disagreement that made
    an index-keyed pairing lose a glyph.
    """

    text = "为 90"
    # The recorded character stream holds only the visible glyphs, so the digits
    # are recorded at indices 1 and 2 while the run text holds them at indices 2
    # and 3 - exactly the stream disagreement that made an index-keyed pairing
    # drop a real glyph.  ``为`` is 12pt wide, the space 2pt and the digits 6pt.
    characters = [
        _Character("为", 0.00, 12.00, 0),
        _Character("9", 17.00, 23.00, 1),
        _Character("0", 23.00, 29.00, 2),
    ]
    return _Run(text, characters, bbox=(0.0, 192.42, 29.00, 204.42))


def test_composition_split_keeps_a_glyph_the_character_stream_index_misplaces():
    run = _span_stream_shifted_run()
    # The rule spans the digits only; both of them are real source glyphs.
    rule = _Rule((12.0, 204.2, 29.0, 204.2))
    segments = character_composition_segments(rule, [(run, 0, len(run.text))])[0]
    text_segments = [s for s in segments if s["segment_type"] == "TEXT_RULE_SEGMENT"]
    assert "".join(s["source_text"] for s in text_segments) == "90"
    assert segments[0]["segment_type"] == "EMPTY_RULE_SEGMENT"


def test_composition_split_without_the_stream_shift_keeps_the_same_glyphs():
    """The same rule over an agreeing stream must produce the same visible text."""

    run = _Run(
        "为 90",
        [
            _Character("为", 0.00, 12.00, 0),
            _Character(" ", 12.00, 14.00, 1),
            _Character("9", 17.00, 23.00, 2),
            _Character("0", 23.00, 29.00, 3),
        ],
        bbox=(0.0, 192.42, 29.00, 204.42),
    )
    rule = _Rule((12.0, 204.2, 29.0, 204.2))
    segments = character_composition_segments(rule, [(run, 0, len(run.text))])[0]
    text_segments = [s for s in segments if s["segment_type"] == "TEXT_RULE_SEGMENT"]
    assert "".join(s["source_text"] for s in text_segments) == "90"


def test_composition_split_fails_closed_when_the_streams_disagree_on_length():
    """A mismatch that cannot be paired by position authorises no split."""

    run = _Run("有效期为 90", [_Character("有", 100.0, 112.0, 0)])
    rule = _Rule((100.0, 204.2, 112.0, 204.2))
    assert character_composition_segments(rule, [(run, 0, len(run.text))]) is None


def _box(text, y0, *, is_leaf, x0=0.0, x1=90.0, y1=None):
    return {
        "x0": x0,
        "y0": y0,
        "x1": x1,
        "y1": y0 + 12.0 if y1 is None else y1,
        "text": text,
        "kind": "SourceParagraph",
        "ownership": (),
        "container_path": "",
        "is_leaf": is_leaf,
    }


def test_preceding_container_label_is_row_scoped(monkeypatch):
    """Only a container on the rule's own visual row is label evidence."""

    from tender_basic import page_layout

    boxes = [
        _box("（招标人名称）：", 100.0, is_leaf=True, x1=40.0),
        _box("注册于\n（工商行政管理局名称）之", 130.0, is_leaf=False, y1=166.0),
    ]
    monkeypatch.setattr(
        page_layout, "compile_source_visual_text_boxes", lambda page: (boxes, {})
    )
    candidates, evidence = page_layout.source_rule_structural_context(
        object(), 70.0, 134.0, 94.0, 134.0
    )
    assert evidence["preceding_container_texts"] == []
    assert all(kind != "preceding_container" for kind, _text in candidates)


def test_preceding_container_on_the_same_row_is_still_evidence(monkeypatch):
    from tender_basic import page_layout

    boxes = [
        _box("（招标人名称）：", 130.0, is_leaf=True, x1=60.0),
        _box("（招标人名称）：________", 130.0, is_leaf=False, y1=166.0),
    ]
    monkeypatch.setattr(
        page_layout, "compile_source_visual_text_boxes", lambda page: (boxes, {})
    )
    candidates, evidence = page_layout.source_rule_structural_context(
        object(), 70.0, 134.0, 94.0, 134.0
    )
    assert evidence["preceding_container_texts"] == ["（招标人名称）："]
    assert ("preceding_container", "（招标人名称）：") in candidates


def test_cell_text_normalization_ignores_wrapping_only():
    assert normalize_cell_text(" 设备\n清单 ") == "设备清单"


class _Page:
    def __init__(self, page, rules):
        self.page = page
        self.lines = rules


def _wrapped_row_plan(monkeypatch, *, row_above, resolved=("QUALITY_TARGET",)):
    """Compile the plan of a glyph-free rule whose label ends the row above."""

    from tender_basic import page_layout
    from tender_basic.source_format import _slot_contract

    boxes = [
        _box(row_above, 186.78, is_leaf=True, x0=70.80, x1=522.83),
        _box("达到\n。", 206.70, is_leaf=False, x0=70.80, x1=163.08, y1=218.98),
        _box("达到", 206.70, is_leaf=True, x0=70.80, x1=95.04, y1=218.98),
    ]
    monkeypatch.setattr(
        page_layout, "compile_source_visual_text_boxes", lambda page: (boxes, {})
    )
    rule = _Rule((95.25, 218.40, 151.05, 218.40))
    key = (42, round(95.25, 1), round(151.05, 1), round(218.40, 1))
    plans = page_layout.compile_source_form_field_plans(
        source_pages=[_Page(42, [rule])],
        glyph_evidence={key: ("", False)},
        resolved_fact_fields=frozenset(resolved),
        slot_contract=_slot_contract,
    )
    return plans[key]


def test_a_label_that_ends_the_wrapped_row_above_is_offered_as_evidence(monkeypatch):
    """The row above a wrapped slot offers the label it ends with."""

    from tender_basic import page_layout

    boxes = [
        _box("价，供货期，按合同约定实施并完成本项目规定的所有工作内容，供货质量", 186.78, is_leaf=True, x0=70.8, x1=522.83),
        _box("达到\n。", 206.70, is_leaf=False, x0=70.8, x1=163.08, y1=218.98),
    ]
    monkeypatch.setattr(
        page_layout, "compile_source_visual_text_boxes", lambda page: (boxes, {})
    )
    candidates, evidence = page_layout.source_rule_structural_context(
        object(), 95.25, 218.40, 151.05, 218.40
    )
    assert evidence["wrap_continuation_row_text"].endswith("供货质量")
    assert ("wrap_continuation_row", evidence["wrap_continuation_row_text"]) in candidates


def test_a_label_that_ends_its_own_sentence_is_not_carried_across_the_wrap(monkeypatch):
    """A sentence that already ended owns its label; nothing crosses the wrap."""

    from tender_basic import page_layout

    boxes = [
        _box("供货质量。", 186.78, is_leaf=True, x0=70.8, x1=120.0),
        _box("达到\n。", 206.70, is_leaf=False, x0=70.8, x1=163.08, y1=218.98),
    ]
    monkeypatch.setattr(
        page_layout, "compile_source_visual_text_boxes", lambda page: (boxes, {})
    )
    candidates, evidence = page_layout.source_rule_structural_context(
        object(), 95.25, 218.40, 151.05, 218.40
    )
    assert evidence["wrap_continuation_row_text"] == ""
    assert all(kind != "wrap_continuation_row" for kind, _text in candidates)


def test_a_label_terminator_row_is_not_carried_across_the_wrap(monkeypatch):
    """A trailing ``：`` means the label expects its own value in place."""

    from tender_basic import page_layout

    boxes = [
        _box("（招标人名称）：", 186.78, is_leaf=True, x0=70.8, x1=120.0),
        _box("注册于\n（工商行政管理局名称）之", 206.70, is_leaf=False, x0=70.8, x1=163.08, y1=218.98),
    ]
    monkeypatch.setattr(
        page_layout, "compile_source_visual_text_boxes", lambda page: (boxes, {})
    )
    candidates, evidence = page_layout.source_rule_structural_context(
        object(), 95.25, 218.40, 151.05, 218.40
    )
    assert evidence["wrap_continuation_row_text"] == ""
    assert all(kind != "wrap_continuation_row" for kind, _text in candidates)


def test_a_far_row_above_is_not_carried_across_the_wrap(monkeypatch):
    """One bounded row pitch: a blank gap is never bridged by a label."""

    from tender_basic import page_layout

    boxes = [
        _box("供货质量", 100.0, is_leaf=True, x0=70.8, x1=120.0),
        _box("达到\n。", 206.70, is_leaf=False, x0=70.8, x1=163.08, y1=218.98),
    ]
    monkeypatch.setattr(
        page_layout, "compile_source_visual_text_boxes", lambda page: (boxes, {})
    )
    candidates, evidence = page_layout.source_rule_structural_context(
        object(), 95.25, 218.40, 151.05, 218.40
    )
    assert evidence["wrap_continuation_row_text"] == ""
    assert all(kind != "wrap_continuation_row" for kind, _text in candidates)


def test_a_wrapped_label_binds_a_resolved_value_into_the_fixed_slot(monkeypatch):
    """The wrapped field plans ``RESOLVED_VALUE_IN_FIXED_SLOT``."""

    from tender_basic.page_layout import TRANSFORMATION_RESOLVED_VALUE_IN_FIXED_SLOT

    plan = _wrapped_row_plan(
        monkeypatch,
        row_above="价，供货期，按合同约定实施并完成本项目规定的所有工作内容，供货质量",
    )
    assert plan[0] == TRANSFORMATION_RESOLVED_VALUE_IN_FIXED_SLOT
    assert [str(getattr(field, "value", field)) for field in plan[1]] == ["quality_target"]


def test_a_wrapped_label_for_an_unresolved_field_stays_a_fixed_empty_slot(monkeypatch):
    """No resolved fact means no value: the slot stays empty, never invented."""

    from tender_basic.page_layout import TRANSFORMATION_FIXED_EMPTY_SLOT

    plan = _wrapped_row_plan(
        monkeypatch,
        row_above="价，供货期，按合同约定实施并完成本项目规定的所有工作内容，供货质量",
        resolved=("DURATION",),
    )
    assert plan[0] == TRANSFORMATION_FIXED_EMPTY_SLOT
    # The field the label names is recorded as provenance, but an unresolved fact
    # plans no insertion: the slot is painted empty exactly as the source drew it.
    assert [str(getattr(field, "value", field)) for field in plan[1]] == ["quality_target"]


def test_a_wrapped_label_ending_its_own_sentence_does_not_bind(monkeypatch):
    from tender_basic.page_layout import TRANSFORMATION_FIXED_EMPTY_SLOT

    plan = _wrapped_row_plan(monkeypatch, row_above="供货质量。")
    assert plan[0] == TRANSFORMATION_FIXED_EMPTY_SLOT
    assert plan[1] == ()


def test_a_wrapped_label_terminator_does_not_bind(monkeypatch):
    from tender_basic.page_layout import TRANSFORMATION_FIXED_EMPTY_SLOT

    plan = _wrapped_row_plan(monkeypatch, row_above="（招标人名称）：")
    assert plan[0] == TRANSFORMATION_FIXED_EMPTY_SLOT
    assert plan[1] == ()
    assert normalize_cell_text("设备清单") != normalize_cell_text("设备清单表")


def test_audit_pairs_logical_tables_by_document_order_and_counts_cells():
    audit = audit_logical_tables(
        _minimal_report(source_pages=[2, 3], fragment_count=2),
        _real_document([[["序号", "名称"], ["1", "泵站"]]]),
        {"logical_0": [["序号", "名称"], ["1", "泵站"]]},
    )
    assert audit["logical_table_count"] == 1
    record = audit["tables"][0]
    assert record["generated_word_table_index"] == 0
    assert record["editable_word_table"] is True
    assert record["continuation_fragments"] == 1
    assert record["lost_cell_count"] == 0
    assert record["duplicated_cell_count"] == 0
    assert record["invented_cell_count"] == 0
    assert record["cell_coverage"] == 1.0
    assert audit["total_source_cell_count"] == 4


def test_audit_reports_lost_and_invented_cells():
    audit = audit_logical_tables(
        _minimal_report(),
        _real_document([[["序号", "名称"], ["1", "错值"]]]),
        {"logical_0": [["序号", "名称"], ["1", "泵站"]]},
    )
    record = audit["tables"][0]
    assert record["lost_cell_count"] == 1
    assert record["invented_cell_count"] == 1
    assert record["lost_cells"][0]["cell_text"] == "泵站"
    assert record["invented_cells"][0]["cell_text"] == "错值"
    assert audit["total_lost_cell_count"] == 1
    assert audit["total_invented_cell_count"] == 1


def test_audit_reports_a_duplicated_cell():
    audit = audit_logical_tables(
        _minimal_report(),
        _real_document([[["序号", "名称"], ["1", "泵站"], ["1", "泵站"]]]),
        {"logical_0": [["序号", "名称"], ["1", "泵站"]]},
    )
    record = audit["tables"][0]
    assert record["lost_cell_count"] == 0
    assert record["duplicated_cell_count"] == 2
    assert audit["total_duplicated_cell_count"] == 2


def test_generated_cell_carrying_whole_source_cells_is_aggregated_not_invented():
    """A template cell that writes label + value is not an invented value."""

    audit = audit_logical_tables(
        _minimal_report(),
        _real_document([[["项目名称：某某工程", ""]]]),
        {"logical_0": [["项目名称：", "某某工程"]]},
    )
    record = audit["tables"][0]
    assert record["lost_cell_count"] == 0
    assert record["invented_cell_count"] == 0
    assert audit["total_invented_cell_count"] == 0
    assert record["aggregated_cell_count"] == 1
    aggregate = record["aggregated_cells"][0]
    assert aggregate["aggregated_source_cell_count"] == 2
    assert set(aggregate["aggregated_source_cells"]) == {"项目名称：", "某某工程"}
    assert record["cell_coverage"] == 1.0


def test_generated_cell_with_no_source_content_is_invented():
    audit = audit_logical_tables(
        _minimal_report(),
        _real_document([[["序号", "名称"], ["1", "泵站"], ["", "凭空写入的值"]]]),
        {"logical_0": [["序号", "名称"], ["1", "泵站"]]},
    )
    record = audit["tables"][0]
    assert record["lost_cell_count"] == 0
    assert record["invented_cell_count"] == 1
    assert record["invented_cells"][0]["cell_text"] == "凭空写入的值"


def test_recorded_resolved_fill_is_not_reported_as_invented():
    """A cell filled from a recorded SourceFillApplication is source-evidenced."""

    report = _minimal_report()
    report["source_fill_applications"] = [
        {
            "application_id": "SFA1",
            "fact_fields": ["project_name"],
            "resolved_values": ["某某工程"],
        }
    ]
    audit = audit_logical_tables(
        report,
        _real_document([[["项目名称", "某某工程"], ["序号", "1"]]]),
        {"logical_0": [["项目名称", ""], ["序号", "1"]]},
    )
    record = audit["tables"][0]
    assert record["invented_cell_count"] == 0
    assert record["filled_cell_count"] == 1
    assert record["filled_cells"][0]["fact_fields"] == ["project_name"]
    assert record["lost_cell_count"] == 0
    assert audit["total_invented_cell_count"] == 0
    assert audit["recorded_fill_field_count"] == 1


def test_a_value_that_is_neither_source_nor_recorded_fill_is_invented():
    report = _minimal_report()
    report["source_fill_applications"] = [
        {
            "application_id": "SFA1",
            "fact_fields": ["project_name"],
            "resolved_values": ["某某工程"],
        }
    ]
    audit = audit_logical_tables(
        report,
        _real_document([[["项目名称", "某某工程"], ["序号", "另一个凭空值"]]]),
        {"logical_0": [["项目名称", ""], ["序号", "1"]]},
    )
    record = audit["tables"][0]
    assert record["filled_cell_count"] == 1
    assert record["invented_cell_count"] == 1
    assert record["invented_cells"][0]["cell_text"] == "另一个凭空值"


def test_short_source_marker_is_lost_when_no_generated_cell_carries_it():
    """A one-character marker needs an exact generated cell, not a coincidence."""

    audit = audit_logical_tables(
        _minimal_report(),
        _real_document([[["序号表头", "名称"]]]),
        {"logical_0": [["1", "名称"]]},
    )
    record = audit["tables"][0]
    assert record["lost_cell_count"] == 1
    assert record["lost_cells"][0]["cell_text"] == "1"


def test_audit_pairs_each_logical_table_with_its_own_document_order_table():
    report = {
        "logical_tables": [
            {"table_id": "logical_0", "head_page": 1, "head_table_index": 0,
             "source_pages": [1], "fragment_count": 1, "row_count": 1,
             "column_count": 2},
            {"table_id": "logical_1", "head_page": 2, "head_table_index": 0,
             "source_pages": [2, 3], "fragment_count": 2, "row_count": 1,
             "column_count": 2},
        ]
    }
    audit = audit_logical_tables(
        report,
        _real_document([[["甲", "乙"]], [["丙", "丁"]]]),
        {"logical_0": [["甲", "乙"]], "logical_1": [["丙", "丁"]]},
    )
    assert [r["generated_word_table_index"] for r in audit["tables"]] == [0, 1]
    assert audit["unpaired_generated_word_table_count"] == 0
    assert audit["tables"][1]["source_fragment_count"] == 2
    assert audit["tables"][1]["continuation_fragments"] == 1
    assert audit["total_lost_cell_count"] == 0
    assert audit["total_invented_cell_count"] == 0
