"""Round-64: review-concern ownership, atomization and human review semantics.

Thirty-two coverage areas guard the round-3 contract:

    SourceRequirementAtom -> ReviewConcern -> ReviewPoint -> workbook views

with the invariant that every rendered component is source-backed *and* owned by
the same human review concern, plus the CASE001 warranty/retention semantics
(24-month project warranty vs 5% retention money vs 12-month release period).
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

from tender_basic.document_models import NormalizedDocument  # noqa: E402
from tender_basic.dynamic_review import (  # noqa: E402
    build_dynamic_review_plan,
    dynamic_review_qa,
    order_review_items,
)
from tender_basic.dynamic_requirements import build_requirement_index  # noqa: E402
from tender_basic.models import ProjectFacts  # noqa: E402
from tender_basic.review_concern import (  # noqa: E402
    CONCERNS,
    actionable_concerns,
    atomize_units,
    build_concerns,
    concern_spec,
    ownership_violations,
)
from tender_basic.review_point import (  # noqa: E402
    render_review_cell,
    scan_review_point,
    synthesize_concern_point,
)

import v1_review_workbook_round3_report as report  # noqa: E402

CASE = ROOT / "acceptance/workspace/case_001"
BUILD3 = CASE / "v1_manual_fidelity_round4_date_rhythm_closure8_review_workbook3"
BUILD2 = CASE / "v1_manual_fidelity_round4_date_rhythm_closure8_review_workbook2"
WORD = CASE / "v1_manual_fidelity_round4_date_rhythm_closure8"

pytestmark = pytest.mark.skipif(
    not (BUILD3 / "project_facts.json").is_file(),
    reason="round-3 CASE001 successor build is not present",
)


@pytest.fixture(scope="module")
def ctx() -> dict:
    facts = ProjectFacts.model_validate(
        json.loads((BUILD3 / "project_facts.json").read_text(encoding="utf-8"))
    )
    document = NormalizedDocument.model_validate(
        json.loads((BUILD3 / "normalized_document.json").read_text(encoding="utf-8"))
    )
    index = build_requirement_index(document)
    atoms = atomize_units(index.units)
    concerns = build_concerns(atoms)
    kept, filtered = actionable_concerns(concerns)
    points = [
        (concern, point)
        for concern, point in (
            (concern, synthesize_concern_point(concern, fact_hints={})) for concern in kept
        )
        if point is not None
    ]
    plan = build_dynamic_review_plan(document, facts)
    return {
        "facts": facts,
        "document": document,
        "index": index,
        "atoms": atoms,
        "concerns": concerns,
        "kept": kept,
        "filtered": filtered,
        "points": points,
        "plan": plan,
    }


def _violations(concern, point) -> dict:
    return ownership_violations(
        concern,
        rendered_numbers=point.numeric_evidence,
        rendered_materials=point.preparation_materials,
        rendered_consequence=point.failure_consequence,
        rendered_evidence=point.evidence_ids,
    )


# 1 ------------------------------------------------------------------------- #
def test_atom_inventory_is_built(ctx):
    assert len(ctx["atoms"]) > 0
    assert all(re.fullmatch(r"SRA\d{4}(\.\w+)?", atom.atom_id) for atom in ctx["atoms"])


# 2 ------------------------------------------------------------------------- #
def test_every_atom_is_source_located(ctx):
    assert all(atom.source_text.strip() for atom in ctx["atoms"])
    assert all(atom.source_page or atom.source_locator for atom in ctx["atoms"])
    assert all(atom.source_clause_id for atom in ctx["atoms"])


# 3 ------------------------------------------------------------------------- #
def test_every_atom_has_a_registered_owner(ctx):
    assert all(atom.owner_concern_id in CONCERNS for atom in ctx["atoms"])
    assert all(atom.owner_reason for atom in ctx["atoms"])


# 4 ------------------------------------------------------------------------- #
def test_no_unclassified_actionable_concern(ctx):
    assert not [concern for concern in ctx["kept"] if concern.concern_id == "UNCLASSIFIED"]


# 5 ------------------------------------------------------------------------- #
def test_concerns_are_unique_per_concern_id(ctx):
    ids = [concern.concern_id for concern in ctx["concerns"]]
    assert len(ids) == len(set(ids))


# 6 ------------------------------------------------------------------------- #
def test_review_points_carry_their_concern(ctx):
    assert all(point.concern_id and point.concern_label for _c, point in ctx["points"])
    assert all(
        point.owned_atom_ids == list(concern.atom_ids)
        for concern, point in ctx["points"]
    )


# 7 ------------------------------------------------------------------------- #
def test_point_concern_ids_are_registered(ctx):
    assert all(point.concern_id in CONCERNS for _c, point in ctx["points"])


# 8 ------------------------------------------------------------------------- #
def test_owned_numbers_use_the_concern_roles(ctx):
    for concern, point in ctx["points"]:
        roles = concern_spec(concern.concern_id).numeric_roles
        assert all(value.role in roles for value in point.numeric_evidence)


# 9 ------------------------------------------------------------------------- #
def test_no_source_concern_mismatch(ctx):
    assert all(not _violations(c, p)["source_concern_mismatch"] for c, p in ctx["points"])


# 10 ------------------------------------------------------------------------ #
def test_no_numeric_concern_mismatch(ctx):
    assert all(not _violations(c, p)["numeric_concern_mismatch"] for c, p in ctx["points"])


# 11 ------------------------------------------------------------------------ #
def test_no_material_concern_mismatch(ctx):
    assert all(not _violations(c, p)["material_concern_mismatch"] for c, p in ctx["points"])


# 12 ------------------------------------------------------------------------ #
def test_no_consequence_concern_mismatch(ctx):
    assert all(not _violations(c, p)["consequence_concern_mismatch"] for c, p in ctx["points"])


# 13 ------------------------------------------------------------------------ #
def test_no_evidence_concern_mismatch(ctx):
    assert all(not _violations(c, p)["evidence_concern_mismatch"] for c, p in ctx["points"])


# 14 ------------------------------------------------------------------------ #
def test_no_authority_scope_mismatch(ctx):
    assert all(not _violations(c, p)["authority_scope_mismatch"] for c, p in ctx["points"])


# 15 ------------------------------------------------------------------------ #
def test_no_foreign_role_numeric_component(ctx):
    assert all(not _violations(c, p)["foreign_role_numeric"] for c, p in ctx["points"])


# 16 ------------------------------------------------------------------------ #
def test_component_ownership_is_coherent_for_every_point(ctx):
    counters = report.ownership_counters(ctx["kept"], ctx["points"], ctx["atoms"])
    for key in (
        "source_concern_mismatch",
        "numeric_concern_mismatch",
        "material_concern_mismatch",
        "consequence_concern_mismatch",
        "evidence_concern_mismatch",
        "authority_scope_mismatch",
        "foreign_role_numeric",
    ):
        assert counters[key] == 0, key


# 17 ------------------------------------------------------------------------ #
def test_project_warranty_owns_the_project_warranty_period(ctx):
    match = [p for c, p in ctx["points"] if c.concern_id == "PROJECT_WARRANTY"]
    assert match
    roles = {value.role for value in match[0].numeric_evidence}
    assert roles == {"WARRANTY_MONTHS"}
    assert any("24" in value.value for value in match[0].numeric_evidence)


# 18 ------------------------------------------------------------------------ #
def test_retention_money_ratio_owns_the_ratio(ctx):
    match = [p for c, p in ctx["points"] if c.concern_id == "RETENTION_MONEY_RATIO"]
    assert match
    assert any(
        value.role == "PAYMENT_RATIO" and "5%" in value.value
        for value in match[0].numeric_evidence
    )


# 19 ------------------------------------------------------------------------ #
def test_retention_release_period_owns_the_release_period(ctx):
    match = [p for c, p in ctx["points"] if c.concern_id == "RETENTION_RELEASE_PERIOD"]
    assert match
    assert any("12" in value.value for value in match[0].numeric_evidence)


# 20 ------------------------------------------------------------------------ #
def test_warranty_and_retention_are_distinct_concerns(ctx):
    ids = {
        concern.concern_id
        for concern in ctx["concerns"]
        if concern.concern_id.startswith(("PROJECT_WARRANTY", "RETENTION"))
    }
    assert {"PROJECT_WARRANTY", "RETENTION_RELEASE_PERIOD"} <= ids
    warranty = [p for c, p in ctx["points"] if c.concern_id == "PROJECT_WARRANTY"][0]
    release = [p for c, p in ctx["points"] if c.concern_id == "RETENTION_RELEASE_PERIOD"][0]
    assert not ({v.value for v in warranty.numeric_evidence} & {v.value for v in release.numeric_evidence})


# 21 ------------------------------------------------------------------------ #
def test_warranty_retention_fixture_passes(ctx):
    fixture = report.warranty_retention_fixture(ctx["points"], ctx["atoms"], "case_001")
    assert fixture["result"] == "PASS", fixture
    assert fixture["generic_separation_ok"] is True
    assert fixture["reported_as_conflict"] is False


# 22 ------------------------------------------------------------------------ #
def test_no_free_text_score_consequence_outside_scoring(ctx):
    for concern, point in ctx["points"]:
        if concern.concern_id == "EVALUATION_SCORING":
            continue
        assert point.failure_consequence not in {"0分", "扣分", "不得分"}


# 23 ------------------------------------------------------------------------ #
def test_no_tender_fee_fragment_leaks_into_a_review_point(ctx):
    for _concern, point in ctx["points"]:
        blob = render_review_cell(point)
        assert not re.search(r"(询比文件售价|标书费|工本费|平台服务费)", blob), point.concern_id


# 24 ------------------------------------------------------------------------ #
def test_every_point_has_checks_and_pass_criteria(ctx):
    for concern, point in ctx["points"]:
        assert point.review_checks, concern.concern_id
        assert point.pass_criteria, concern.concern_id


# 25 ------------------------------------------------------------------------ #
def test_every_point_keeps_its_own_evidence(ctx):
    for concern, point in ctx["points"]:
        assert point.owned_atom_ids
        assert point.concern_id == concern.concern_id


# 26 ------------------------------------------------------------------------ #
def test_scan_review_point_reports_a_decision_for_every_row(ctx):
    for _concern, point in ctx["points"]:
        scan = scan_review_point(point)
        assert "generic" in scan


# 27 ------------------------------------------------------------------------ #
def test_rendered_cell_names_the_requirement_and_concern(ctx):
    for _concern, point in ctx["points"]:
        cell = render_review_cell(point)
        assert "招标文件要求" in cell
        assert "复核要点" in cell
        assert point.concern_label and point.concern_question


# 28 ------------------------------------------------------------------------ #
def test_human_blocker_fixtures_a_to_n_pass(ctx):
    fixtures = report.evaluate_fixtures(ctx["points"], ctx["atoms"], ctx["filtered"])
    failed = {key: item for key, item in fixtures.items() if item["result"] != "PASS"}
    assert not failed, failed
    assert len(fixtures) == 14


# 29 ------------------------------------------------------------------------ #
def test_manual_style_audit_is_fully_coherent(ctx):
    audit = report.human_style_audit(ctx["points"])
    assert audit["coherent_count"] == audit["sample_size"] == 20


# 30 ------------------------------------------------------------------------ #
def test_before_after_examples_cover_at_least_fifteen_concerns(ctx):
    before = report.legacy_rows(BUILD2 / "投标项目复核表.xlsx")
    assert before
    produced = report.examples(ctx["points"], before)
    assert len(produced) >= 15
    assert all(item["after_source_atoms"] for item in produced)


# 31 ------------------------------------------------------------------------ #
def test_false_conflict_count_is_zero(ctx):
    conflicts = report.conflict_report(ctx["points"], ctx["atoms"])
    assert conflicts["false_conflict_count"] == 0
    assert conflicts["false_conflicts"] == []


# 32 ------------------------------------------------------------------------ #
def test_successor_workbook_matches_the_plan_and_keeps_the_word_artifacts(ctx):
    plan = ctx["plan"]
    qa = dynamic_review_qa(plan, ctx["document"], ctx["facts"])
    assert qa["result"] == "PASS", qa["hard_gate_failures"]

    workbook = load_workbook(BUILD3 / "投标项目复核表.xlsx", data_only=True, read_only=True)
    assert workbook.sheetnames[0] == "投标项目复核表"
    sheet = workbook["投标项目复核表"]
    written = [
        str(row[0].value) if row[0].value is not None else ""
        for row in sheet.iter_rows(min_row=9, min_col=4, max_col=4)
    ]
    workbook.close()
    while written and not written[-1]:
        written.pop()
    expected = [item.cell_text for item in order_review_items(plan.items)]
    assert written == expected

    identities = json.loads(
        (ROOT / "acceptance/reports/v1_generalization/case_001_review_workbook_build_round3.json").read_text(
            encoding="utf-8"
        )
    )["artifact_identity"]
    assert all(entry["byte_identical"] for entry in identities.values())
