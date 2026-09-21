from __future__ import annotations

import json

import pytest
from openpyxl import load_workbook

from tender_basic.document_models import DocumentStatus
from tender_basic.models import (
    CandidateFact,
    DocxParagraphLocator,
    DocxTableLocator,
    FactStatus,
    FieldName,
    ProjectFacts,
    ProjectFields,
    ResolvedFact,
    SourceDocument,
    SourceType,
)
from tender_basic.review_builder import (
    DEFAULT_TEMPLATE_PATH,
    TEMPLATE_MANUAL_CELLS,
    TEMPLATE_SHEET_NAME,
    build_review_workbook,
)


def _candidate(
    value: str,
    *,
    paragraph_index: int | None = None,
    table_index: int | None = None,
    row_index: int = 0,
    cell_index: int = 0,
    method: str = "label_value_same_line",
    confidence: float = 0.95,
) -> CandidateFact:
    if paragraph_index is not None:
        locator = DocxParagraphLocator(paragraph_index=paragraph_index)
    else:
        locator = DocxTableLocator(
            table_index=table_index or 0,
            row_index=row_index,
            cell_index=cell_index,
        )
    return CandidateFact(
        value=value,
        normalized_value=value,
        confidence=confidence,
        method=method,
        source_file="sample.docx",
        source_type=SourceType.DOCX,
        locator=locator,
        section="招标公告",
        evidence_text=f"字段证据：{value}",
    )


def _resolved(field: FieldName, value: str, candidates: list[CandidateFact]) -> ResolvedFact:
    return ResolvedFact(
        field=field,
        resolved_value=value,
        status=FactStatus.RESOLVED,
        confidence=max(candidate.confidence for candidate in candidates),
        candidates=candidates,
        resolution_reason="Test resolved fact.",
    )


def _review(field: FieldName, values: list[str]) -> ResolvedFact:
    candidates = [
        _candidate(value, paragraph_index=20 + index, confidence=0.8)
        for index, value in enumerate(values)
    ]
    return ResolvedFact(
        field=field,
        resolved_value=None,
        status=FactStatus.NEEDS_REVIEW,
        confidence=0.8,
        candidates=candidates,
        resolution_reason="Conflicting test candidates.",
    )


def _not_found(field: FieldName) -> ResolvedFact:
    return ResolvedFact(
        field=field,
        resolved_value=None,
        status=FactStatus.NOT_FOUND,
        confidence=0.0,
        candidates=[],
        resolution_reason="No test candidate.",
    )


def make_project_facts() -> ProjectFacts:
    fields = {}
    for field in FieldName:
        fields[field.value] = _not_found(field)

    fields[FieldName.PROJECT_NAME.value] = _resolved(
        FieldName.PROJECT_NAME,
        "测试工程",
        [
            _candidate("测试工程", paragraph_index=0),
            _candidate("测试工程", paragraph_index=1),
            _candidate("测试工程", table_index=0, row_index=0, cell_index=1, method="table_label_exact", confidence=0.98),
        ],
    )
    fields[FieldName.PROJECT_NUMBER.value] = _resolved(
        FieldName.PROJECT_NUMBER,
        "PRJ-123",
        [_candidate("PRJ-123", paragraph_index=2)],
    )
    fields[FieldName.TENDER_NUMBER.value] = _resolved(
        FieldName.TENDER_NUMBER,
        "BID-456",
        [_candidate("BID-456", paragraph_index=3)],
    )
    fields[FieldName.PURCHASER.value] = _resolved(
        FieldName.PURCHASER,
        "采购人甲",
        [_candidate("采购人甲", paragraph_index=4)],
    )
    fields[FieldName.BUDGET.value] = _resolved(
        FieldName.BUDGET,
        "80万元",
        [_candidate("80万元", paragraph_index=5)],
    )
    fields[FieldName.MAX_PRICE.value] = _resolved(
        FieldName.MAX_PRICE,
        "100万元",
        [_candidate("100万元", paragraph_index=6)],
    )
    fields[FieldName.PROJECT_LOCATION.value] = _review(
        FieldName.PROJECT_LOCATION,
        ["地点A", "地点B"],
    )

    project_fields = ProjectFields(**fields)
    return ProjectFacts.from_fields(
        source_document=SourceDocument(
            source_file="sample.docx",
            source_type=SourceType.DOCX,
        ),
        fields=project_fields,
    )


def test_review_workbook_uses_the_supplied_template_as_the_primary_sheet(tmp_path) -> None:
    output = tmp_path / "投标项目复核表.xlsx"
    build_review_workbook(make_project_facts(), output)

    workbook = load_workbook(output, read_only=False, data_only=False)
    try:
        assert workbook.sheetnames == [TEMPLATE_SHEET_NAME]
        worksheet = workbook[TEMPLATE_SHEET_NAME]
        assert worksheet.max_row >= 77
        assert worksheet.max_column >= 13
        assert "A1:N1" in {str(value) for value in worksheet.merged_cells.ranges}
        assert "A75:C75" in {str(value) for value in worksheet.merged_cells.ranges}
        assert worksheet["B2"].value == "测试工程"
        assert worksheet["F2"].value == "项目编号：PRJ-123；招标编号：BID-456"
        assert worksheet["B4"].value == "预算：80万元；最高限价：100万元"
        assert worksheet["L4"].value == "【待补充】"
        assert worksheet["E9"].value == "【待核对】"
        assert [
            worksheet.cell(row=row, column=1).value
            for row in range(9, 74)
            if worksheet.cell(row=row, column=1).value is not None
        ] == list(range(1, 65))
    finally:
        workbook.close()


def test_review_workbook_keeps_numbers_independent_and_manual_fields_unfilled(tmp_path) -> None:
    output = tmp_path / "review.xlsx"
    build_review_workbook(make_project_facts(), output)
    workbook = load_workbook(output, read_only=False, data_only=False)
    try:
        worksheet = workbook[TEMPLATE_SHEET_NAME]
        assert worksheet["F2"].value != "BID-456"
        assert "PRJ-123" in worksheet["F2"].value
        assert "BID-456" in worksheet["F2"].value
        assert worksheet["B5"].value == "【填写】"
        assert worksheet["F5"].value == "【填写】"
        assert worksheet["I5"].value == "【填写】"
        assert worksheet["L5"].value == "【填写】"
        assert worksheet["F4"].value == "【待补充】"
        assert worksheet["L4"].value != "公开招标"
        assert all(
            "自动填充" not in str(cell.value)
            for row in worksheet.iter_rows()
            for cell in row
            if cell.value is not None
        )
        for address, expected in TEMPLATE_MANUAL_CELLS.items():
            assert worksheet[address].value == expected
    finally:
        workbook.close()


def test_review_workbook_preserves_template_validation_and_structure(tmp_path) -> None:
    output = tmp_path / "review.xlsx"
    build_review_workbook(make_project_facts(), output)
    workbook = load_workbook(output, read_only=False, data_only=False)
    try:
        worksheet = workbook[TEMPLATE_SHEET_NAME]
        assert len(worksheet.data_validations.dataValidation) >= 3
        assert len(worksheet.conditional_formatting) >= 3
        assert len(worksheet.merged_cells.ranges) >= 30
        assert worksheet["A75"].value == "投标小组评审意见："
        assert worksheet["A76"].value == "分总评审意见："
        assert worksheet["A77"].value == "大区投标负责人评审意见："
    finally:
        workbook.close()


def test_review_workbook_template_is_present_and_not_rebuilt_from_scratch() -> None:
    assert DEFAULT_TEMPLATE_PATH.is_file()


def test_review_builder_requires_project_facts_ssot(tmp_path) -> None:
    with pytest.raises(TypeError):
        build_review_workbook(tmp_path / "not-project-facts.json", tmp_path / "review.xlsx")
