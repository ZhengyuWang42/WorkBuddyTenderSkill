from __future__ import annotations

from pathlib import Path

import pymupdf
from openpyxl import load_workbook

from tender_basic.document_parser import parse_document
from tender_basic.document_models import PdfTableCell, PdfTextSpan
from tender_basic.models import ProjectFacts
from tender_basic.review_builder import build_review_workbook
from tender_basic.review_evidence import (
    EVIDENCE_STATES,
    ReviewEvidenceItem,
    ReviewSemanticClassification,
    _looks_like_navigation,
    _source_table_cell_text,
    retrieve_review_evidence,
    review_evidence_qa,
    validate_review_evidence,
    validate_semantic_classifications,
)
from tender_basic.models import PdfTableLocator
from test_review_builder import make_project_facts


def _make_review_fixture(path: Path) -> None:
    pdf = pymupdf.open()
    page = pdf.new_page(width=620, height=800)
    page.insert_text(
        (50, 80),
        "第二章 投标人须知\n"
        "资格审查：投标人资格条件和信用要求。\n"
        "本项目不接受联合体，不得转包、违法分包。\n"
        "投标截止时间：2026年12月1日10时00分；递交地点：电子交易平台。\n"
        "本项目使用全流程电子投标，不见面开标，CA解密，在线上传。\n"
        "最高投标限价：人民币100万元。\n"
        "投标保证金金额：人民币2万元；形式：银行转账；递交截止时间同投标截止时间。\n"
        "★技术参数必须逐条响应。\n"
        "第三章 评标办法 价格评分按基准价计算。",
        fontsize=11,
        fontname="china-s",
    )
    pdf.save(path)
    pdf.close()


def test_review_evidence_retrieves_generic_high_value_clauses(tmp_path: Path) -> None:
    source = tmp_path / "review-fixture.pdf"
    _make_review_fixture(source)
    document = parse_document(source)

    items, packet = retrieve_review_evidence(document)
    validation = validate_review_evidence(document, items)
    qa = review_evidence_qa(items, validation)
    by_id = {item.review_item_id: item for item in items}

    assert qa["total_review_items"] == 64
    assert qa["unclassified"] == 0
    assert qa["invalid_locators"] == 0
    assert packet
    assert by_id["R001"].state != "NOT_FOUND"
    assert by_id["R002"].state != "NOT_FOUND"
    assert by_id["R003"].state != "NOT_FOUND"
    assert by_id["R008"].state != "NOT_FOUND"
    assert by_id["R010"].state != "NOT_FOUND"
    assert by_id["R039"].state != "NOT_FOUND"
    assert by_id["R054"].state != "NOT_FOUND"


def test_fully_electronic_signal_can_be_not_applicable_candidate(tmp_path: Path) -> None:
    source = tmp_path / "electronic.pdf"
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text(
        (50, 80),
        "本项目全流程电子投标，不见面开标，在线上传，CA解密。",
        fontsize=11,
        fontname="china-s",
    )
    pdf.save(source)
    pdf.close()

    items, _packet = retrieve_review_evidence(parse_document(source))
    by_id = {item.review_item_id: item for item in items}

    assert by_id["R061"].state == "NOT_APPLICABLE_CANDIDATE"
    assert by_id["R062"].state == "NOT_APPLICABLE_CANDIDATE"


def test_review_evidence_pruning_is_bounded_and_explicit(tmp_path: Path) -> None:
    source = tmp_path / "pruning.pdf"
    _make_review_fixture(source)
    document = parse_document(source)

    items, _packet = retrieve_review_evidence(document)
    validation = validate_review_evidence(document, items)
    qa = review_evidence_qa(items, validation)

    assert qa["average_candidates_retrieved"] > qa["average_candidates_after_pruning"]
    assert qa["max_evidence_snippets_written"] <= 3
    assert qa["rows_with_more_than_3_written_snippets"] == 0
    assert qa["unclassified"] == 0
    assert all(item.state in EVIDENCE_STATES for item in items)
    assert sum(item.state == "MULTIPLE_FOUND" for item in items) < 64


def test_one_strong_clause_is_found_not_multiple_found(tmp_path: Path) -> None:
    source = tmp_path / "one-clause.pdf"
    pdf = pymupdf.open()
    page = pdf.new_page(width=620, height=800)
    page.insert_text(
        (50, 80),
        "投标截止时间：2026年12月1日10时00分，在电子交易平台上传加密的电子投标文件。",
        fontsize=11,
        fontname="china-s",
    )
    pdf.save(source)
    pdf.close()

    items, _packet = retrieve_review_evidence(parse_document(source))
    by_id = {item.review_item_id: item for item in items}

    assert by_id["R003"].state == "FOUND"
    assert len(by_id["R003"].candidates) == 1


def test_table_cell_text_ignores_tall_artwork_span() -> None:
    cell = PdfTableCell(
        row_index=0,
        column_index=0,
        text="污染后的扁平文本",
        bbox=(0, 0, 180, 1000),
        locator=PdfTableLocator(page=1, table_index=0, row_index=0, column_index=0),
        spans=[
            PdfTextSpan(text="项目名称", bbox=(2, 10, 60, 22), font_size=11),
            PdfTextSpan(text="二维码水印", bbox=(5, 100, 170, 220), font_size=10),
        ],
    )

    text = _source_table_cell_text(cell)

    assert text == "项目名称"
    assert "二维码" not in text


def test_long_contents_lines_are_not_review_evidence() -> None:
    contents = (
        "(一)响应函................................................................................38\n"
        "(二)响应函附录............................................................................39\n"
        "三、授权委托书..............................................................................40"
    )

    assert _looks_like_navigation(contents) is True


def test_scoring_row_does_not_accept_non_scoring_qualification_phrase(tmp_path: Path) -> None:
    source = tmp_path / "non-scoring.pdf"
    pdf = pymupdf.open()
    page = pdf.new_page(width=620, height=800)
    page.insert_text(
        (50, 80),
        "第三章 评标办法\n有效的主体资格证明材料。",
        fontsize=11,
        fontname="china-s",
    )
    pdf.save(source)
    pdf.close()

    items, _packet = retrieve_review_evidence(parse_document(source))
    by_id = {item.review_item_id: item for item in items}

    assert by_id["R049"].state == "NOT_FOUND"
    assert "有效的主体资格证明材料" not in by_id["R049"].evidence_text


def test_semantic_classification_must_reference_existing_candidate() -> None:
    packet = [
        {
            "review_item_id": "R003",
            "candidate_index": 0,
            "page": 1,
            "section": "投标人须知",
            "locator": {"locator_type": "pdf_block", "page": 1, "block_index": 0},
            "evidence_text": "投标截止时间",
            "retrieval_method": "keyword_family",
        }
    ]
    accepted = validate_semantic_classifications(
        packet,
        [
            ReviewSemanticClassification(
                review_item_id="R003",
                candidate_index=0,
                classification="SUPPORTS_REVIEW_ITEM",
            ),
            {
                "review_item_id": "R003",
                "candidate_index": 9,
                "classification": "SUPPORTS_REVIEW_ITEM",
            },
        ],
    )

    assert len(accepted) == 1
    assert accepted[0].candidate_index == 0


def test_review_workbook_writes_states_into_existing_64_rows_without_stray_row(tmp_path: Path) -> None:
    output = tmp_path / "review.xlsx"
    facts: ProjectFacts = make_project_facts()
    items = [
        ReviewEvidenceItem(
            review_item_id=f"R{number:03d}",
            review_intent="测试复核意图",
            state="NOT_FOUND",
            candidates=[],
            evidence_text="NOT_FOUND：未检出明确招标文件依据，需人工复核。",
        )
        for number in range(1, 65)
    ]

    build_review_workbook(facts, output, review_evidence=items)
    workbook = load_workbook(output, data_only=False)
    worksheet = workbook["投标项目复核表"]
    try:
        numbers = [
            worksheet.cell(row=row, column=1).value
            for row in range(9, 74)
            if worksheet.cell(row=row, column=1).value is not None
        ]
        assert numbers == list(range(1, 65))
        assert worksheet.cell(row=73, column=1).value is None
        assert worksheet.cell(row=9, column=5).value.startswith("NOT_FOUND")
        assert worksheet.cell(row=63, column=1).value == 55
        assert worksheet.cell(row=64, column=1).value == 56
        assert all(
            "投标保证金" not in str(worksheet.cell(row=row, column=2).value or "")
            for row in range(9, 74)
            if worksheet.cell(row=row, column=1).value is None
        )
        assert all(
            "自动填充" not in str(cell.value)
            for row in worksheet.iter_rows()
            for cell in row
            if cell.value is not None
        )
    finally:
        workbook.close()
