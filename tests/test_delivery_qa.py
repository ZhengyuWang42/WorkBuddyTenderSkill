from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from scripts.qa_delivery import main as qa_delivery_main
from tender_basic.bid_document_builder import build_bid_document
from tender_basic.delivery_qa import run_delivery_qa
from tender_basic.document_parser import parse_document
from tender_basic.format_extractor import extract_bid_format
from tender_basic.models import (
    CandidateFact,
    DocxParagraphLocator,
    FactStatus,
    FieldName,
    ProjectFacts,
    ProjectFields,
    ResolvedFact,
    SourceDocument,
    SourceType,
)
from tender_basic.output_helpers import load_output_fields
from tender_basic.review_builder import TEMPLATE_SHEET_NAME, build_review_workbook


FIELD_VALUES = {
    "project_name": "测试工程",
    "project_number": "PRJ-123",
    "tender_number": "BID-456",
    "lot_name": "一标段",
    "lot_number": "LOT-001",
    "purchaser": "采购人甲",
    "tender_agency": "代理机构乙",
    "project_location": "上海市",
    "procurement_scope": "平台建设",
    "budget": "80万元",
    "max_price": "100万元",
    "duration": "180天",
    "quality_target": "合格",
    "bid_deadline": "2026-09-10 09:30",
    "bid_open_time": "2026-09-10 09:30",
    "bid_open_location": "线上开标",
    "bid_bond_amount": "10万元",
    "bid_bond_form": "银行保函",
    "consortium_allowed": False,
    "procurement_method": "公开招标",
    "bid_validity": "90日历天",
    "submission_method": "电子",
    "electronic_platform": "测试交易平台",
}


def _candidate(value: object, index: int, *, method: str = "label_value_same_line") -> CandidateFact:
    return CandidateFact(
        value=value,
        normalized_value=value,
        confidence=0.95 if method != "keyword_window" else 0.70,
        method=method,
        source_file="qa-sample.docx",
        source_type=SourceType.DOCX,
        locator=DocxParagraphLocator(paragraph_index=index),
        section="招标公告",
        evidence_text=f"字段证据：{value}",
    )


def make_facts(
    *,
    review: tuple[str, ...] = (),
    not_found: tuple[str, ...] = (),
) -> ProjectFacts:
    review_set = set(review)
    not_found_set = set(not_found)
    values: dict[str, ResolvedFact] = {}
    for index, field in enumerate(FieldName):
        value = FIELD_VALUES[field.value]
        if field.value in not_found_set:
            values[field.value] = ResolvedFact(
                field=field,
                resolved_value=None,
                status=FactStatus.NOT_FOUND,
                confidence=0.0,
                candidates=[],
                resolution_reason="No trusted candidate was found.",
            )
        elif field.value in review_set:
            values[field.value] = ResolvedFact(
                field=field,
                resolved_value=None,
                status=FactStatus.NEEDS_REVIEW,
                confidence=0.70,
                candidates=[
                    _candidate(f"{value}候选A", index * 2),
                    _candidate(f"{value}候选B", index * 2 + 1),
                ],
                resolution_reason="Conflicting candidates require review.",
            )
        else:
            values[field.value] = ResolvedFact(
                field=field,
                resolved_value=value,
                status=FactStatus.RESOLVED,
                confidence=0.95,
                candidates=[_candidate(value, index)],
                resolution_reason="One deterministic candidate.",
            )
    return ProjectFacts.from_fields(
        source_document=SourceDocument(
            source_file="qa-sample.docx",
            source_type=SourceType.DOCX,
        ),
        fields=ProjectFields(**values),
    )


def _write_artifacts(tmp_path: Path, facts: ProjectFacts) -> tuple[Path, Path, Path]:
    tmp_path.mkdir(parents=True, exist_ok=True)
    facts_path = tmp_path / "project_facts.json"
    xlsx_path = tmp_path / "投标项目复核表.xlsx"
    docx_path = tmp_path / "基础投标文件.docx"
    facts_path.write_text(facts.model_dump_json(indent=2), encoding="utf-8")
    build_review_workbook(facts, xlsx_path)
    build_bid_document(facts, docx_path)
    return facts_path, xlsx_path, docx_path


def _qa(tmp_path: Path, facts: ProjectFacts) -> dict:
    return run_delivery_qa(*_write_artifacts(tmp_path, facts))


def _docx_info_table(document: Document):
    labels = {spec.label for spec in load_output_fields()}
    best = None
    best_count = -1
    for table in document.tables:
        count = sum(
            1
            for row in table.rows
            if len(row.cells) >= 2 and row.cells[0].text.strip() in labels
        )
        if count > best_count:
            best = table
            best_count = count
    assert best is not None
    return best


def _set_docx_info_value(path: Path, label: str, value: str) -> None:
    document = Document(path)
    table = _docx_info_table(document)
    for row in table.rows:
        if row.cells[0].text.strip() == label:
            row.cells[1].text = value
            document.save(path)
            return
    raise AssertionError(f"DOCX field row not found: {label}")


def test_complete_consistent_output_is_pass(tmp_path: Path) -> None:
    report = _qa(tmp_path, make_facts())

    assert report["overall_status"] == "PASS"
    assert report["errors"] == []
    assert report["format_source"] == "GENERIC_FALLBACK"
    assert report["summary"]["candidate_count"] == 23


def test_review_and_not_found_are_pass_with_review(tmp_path: Path) -> None:
    review_report = _qa(tmp_path / "review", make_facts(review=("project_name",)))
    missing_report = _qa(tmp_path / "missing", make_facts(not_found=("bid_bond_amount",)))

    assert review_report["overall_status"] == "PASS_WITH_REVIEW"
    assert missing_report["overall_status"] == "PASS_WITH_REVIEW"
    assert review_report["errors"] == []
    assert missing_report["errors"] == []


def test_xlsx_primary_template_sheet_is_required(tmp_path: Path) -> None:
    facts_path, xlsx_path, docx_path = _write_artifacts(tmp_path, make_facts())
    workbook = load_workbook(xlsx_path)
    workbook.create_sheet("其他")
    workbook.remove(workbook[TEMPLATE_SHEET_NAME])
    workbook.save(xlsx_path)

    report = run_delivery_qa(facts_path, xlsx_path, docx_path)

    assert report["overall_status"] == "FAIL"
    assert any(check["id"] == "xlsx_template_sheet" for check in report["errors"])


def test_xlsx_template_merged_cells_are_required(tmp_path: Path) -> None:
    facts_path, xlsx_path, docx_path = _write_artifacts(tmp_path, make_facts())
    workbook = load_workbook(xlsx_path)
    workbook[TEMPLATE_SHEET_NAME].unmerge_cells("A1:N1")
    workbook.save(xlsx_path)

    report = run_delivery_qa(facts_path, xlsx_path, docx_path)

    assert report["overall_status"] == "FAIL"
    assert any(check["id"] == "xlsx_template_structure" for check in report["errors"])


def test_top_number_and_amount_fields_cannot_be_exchanged(tmp_path: Path) -> None:
    facts_path, xlsx_path, docx_path = _write_artifacts(tmp_path, make_facts())
    workbook = load_workbook(xlsx_path)
    worksheet = workbook[TEMPLATE_SHEET_NAME]
    worksheet["F2"] = "项目编号：BID-456；招标编号：PRJ-123"
    worksheet["B4"] = "预算：100万元；最高限价：80万元"
    workbook.save(xlsx_path)

    report = run_delivery_qa(facts_path, xlsx_path, docx_path)
    ids = {check["id"] for check in report["errors"]}

    assert report["overall_status"] == "FAIL"
    assert "xlsx_top_project_number_and_tender_number" in ids
    assert "xlsx_top_budget_and_max_price" in ids


def test_procurement_method_does_not_fill_submission_method(tmp_path: Path) -> None:
    facts_path, xlsx_path, docx_path = _write_artifacts(tmp_path, make_facts())
    workbook = load_workbook(xlsx_path)
    workbook[TEMPLATE_SHEET_NAME]["L4"] = "公开招标"
    workbook.save(xlsx_path)

    report = run_delivery_qa(facts_path, xlsx_path, docx_path)

    assert report["overall_status"] == "FAIL"
    assert any(
        check["id"] == "xlsx_submission_method_not_procurement_method"
        for check in report["errors"]
    )


def test_manual_template_fields_cannot_be_guessed(tmp_path: Path) -> None:
    facts_path, xlsx_path, docx_path = _write_artifacts(tmp_path, make_facts())
    workbook = load_workbook(xlsx_path)
    workbook[TEMPLATE_SHEET_NAME]["B5"] = "电子交易平台A"
    workbook.save(xlsx_path)

    report = run_delivery_qa(facts_path, xlsx_path, docx_path)

    assert report["overall_status"] == "FAIL"
    assert any(check["id"] == "xlsx_manual_fields_not_guessed" for check in report["errors"])


def test_corrupt_docx_is_fail(tmp_path: Path) -> None:
    facts_path, xlsx_path, docx_path = _write_artifacts(tmp_path, make_facts())
    docx_path.write_bytes(b"not a valid docx")

    report = run_delivery_qa(facts_path, xlsx_path, docx_path)

    assert report["overall_status"] == "FAIL"
    assert any(check["id"] == "docx_reopen" for check in report["errors"])


def test_invalid_project_facts_contract_is_fail(tmp_path: Path) -> None:
    facts_path, xlsx_path, docx_path = _write_artifacts(tmp_path, make_facts())
    facts_path.write_text('{"fields": {}}', encoding="utf-8")

    report = run_delivery_qa(facts_path, xlsx_path, docx_path)

    assert report["overall_status"] == "FAIL"
    assert any(check["id"] == "project_facts_contract" for check in report["errors"])


def test_docx_resolved_value_corruption_is_fail(tmp_path: Path) -> None:
    facts_path, xlsx_path, docx_path = _write_artifacts(tmp_path, make_facts())
    _set_docx_info_value(docx_path, "项目名称", "篡改后的值")

    report = run_delivery_qa(facts_path, xlsx_path, docx_path)

    assert report["overall_status"] == "FAIL"
    assert any(check["id"] == "docx_field_project_name" for check in report["errors"])


def test_source_format_metadata_and_heading_order_are_qa_checked(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    source_document = Document()
    source_document.add_paragraph("第六章 投标文件格式", style="Heading 1")
    source_document.add_paragraph("一、投标函", style="Heading 2")
    source_document.add_paragraph("二、授权委托书", style="Heading 2")
    source_document.save(source)
    normalized = parse_document(source)
    template = extract_bid_format(normalized)
    facts = make_facts()
    facts_path, xlsx_path, docx_path = _write_artifacts(tmp_path / "source-output", facts)
    build_bid_document(
        facts,
        docx_path,
        normalized_document=normalized,
        format_template=template,
    )

    report = run_delivery_qa(
        facts_path,
        xlsx_path,
        docx_path,
        format_source=template.format_source,
        format_metadata=template.metadata(),
    )

    assert report["overall_status"] == "PASS"
    assert report["format_source"] == "SOURCE_DOCUMENT"
    assert any(
        check["id"] == "docx_source_format_heading_order" and check["status"] == "PASS"
        for check in report["checks"]
    )


def test_fabricated_bidder_and_submission_ready_status_are_failures(tmp_path: Path) -> None:
    facts_path, xlsx_path, docx_path = _write_artifacts(tmp_path, make_facts())
    document = Document(docx_path)
    for row in document.tables[0].rows:
        if row.cells[0].text.strip() == "投标人":
            row.cells[1].text = "虚构公司"
            break
    document.add_paragraph("READY_FOR_SUBMISSION")
    document.save(docx_path)

    report = run_delivery_qa(facts_path, xlsx_path, docx_path)

    assert report["overall_status"] == "FAIL"
    assert any(check["id"] == "docx_bidder_placeholder" for check in report["errors"])
    assert any(check["id"] == "no_submission_ready_status" for check in report["errors"])


def test_qa_cli_writes_structured_report(tmp_path: Path) -> None:
    facts_path, xlsx_path, docx_path = _write_artifacts(
        tmp_path, make_facts(review=("project_name",))
    )
    report_path = tmp_path / "qa_report.json"

    code = qa_delivery_main(
        [
            "--facts",
            str(facts_path),
            "--xlsx",
            str(xlsx_path),
            "--docx",
            str(docx_path),
            "--output",
            str(report_path),
        ]
    )

    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert code == 0
    assert report["overall_status"] == "PASS_WITH_REVIEW"
    assert isinstance(report["checks"], list)
    assert isinstance(report["errors"], list)
    assert isinstance(report["warnings"], list)
