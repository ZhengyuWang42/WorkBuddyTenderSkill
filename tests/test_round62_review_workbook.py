"""The review workbook is a review surface, and these tests hold it to that.

Every assertion here is about the *contract* between the workbook and its
authorities: ``ProjectFacts`` owns values and status, the source owns structure
and price cells, the evidence packet owns locators, and a human owns the manual
columns.  Nothing in these tests asserts a case-specific number, name or page.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from tender_basic.document_models import (
    DocumentStatus,
    NormalizedDocument,
    PdfTable,
    PdfTableCell,
    PdfTableLocator,
    PdfTableRow,
)
from tender_basic.dynamic_review import DynamicReviewItem, DynamicReviewPlan
from tender_basic.format_extractor import BidFormatTemplate, FormatElement
from tender_basic.models import (
    CandidateFact,
    FactStatus,
    FieldName,
    PdfLocator,
    ProjectFacts,
    ProjectFields,
    ResolvedFact,
    SourceDocument,
    SourceType,
)
from tender_basic.review_builder import build_review_workbook
from tender_basic.review_evidence import ReviewEvidenceCandidate, ReviewEvidenceItem
from tender_basic.review_workbook_views import (
    MANUAL_CONCLUSION,
    MANUAL_NOTE,
    MANUAL_OPTIONS,
    SHEET_TITLES,
)

DELIVERED_SHEET = "投标项目复核表"

FIELD_VALUES = {
    "project_name": "示例一体化泵站项目",
    "project_number": "ZB-2026-001",
    "tender_number": "XJ-2026-001",
    "lot_name": "一标段",
    "lot_number": "LOT-1",
    "purchaser": "示例采购人",
    "tender_agency": "示例代理机构",
    "project_location": "示例地点",
    "procurement_scope": "泵站设备与安装",
    "budget": "800万元",
    "max_price": "750万元",
    "duration": "180天",
    "quality_target": "合格",
    "bid_deadline": "2026-09-10 09:30",
    "bid_open_time": "2026-09-10 09:30",
    "bid_open_location": "线上开标",
    "bid_bond_amount": "10万元",
    "bid_bond_form": "银行保函",
    "consortium_allowed": "不允许",
    "procurement_method": "询比",
    "bid_validity": "90日历天",
    "submission_method": "电子递交",
    "electronic_platform": "示例交易平台",
}


def _candidate(value: object, page: int) -> CandidateFact:
    return CandidateFact(
        value=value,
        normalized_value=value,
        confidence=0.9,
        method="label_value_same_line",
        source_file="sample.pdf",
        source_type=SourceType.PDF,
        locator=PdfLocator(locator_type="pdf_block", page=page, block_index=0),
        section="供应商须知前附表",
        evidence_text=f"字段证据：{value}",
    )


def build_facts(*, not_found: tuple[str, ...] = ("lot_name",), review: tuple[str, ...] = ("budget",)) -> ProjectFacts:
    values: dict[str, ResolvedFact] = {}
    for index, field in enumerate(FieldName, start=1):
        value = FIELD_VALUES[field.value]
        if field.value in not_found:
            values[field.value] = ResolvedFact(
                field=field,
                resolved_value=None,
                status=FactStatus.NOT_FOUND,
                confidence=0.0,
                candidates=[],
                resolution_reason="No trusted candidate was found.",
            )
        elif field.value in review:
            values[field.value] = ResolvedFact(
                field=field,
                resolved_value=None,
                status=FactStatus.NEEDS_REVIEW,
                confidence=0.6,
                candidates=[_candidate(f"{value}A", index), _candidate(f"{value}B", index + 1)],
                resolution_reason="Conflicting candidates require review.",
            )
        else:
            values[field.value] = ResolvedFact(
                field=field,
                resolved_value=value,
                status=FactStatus.RESOLVED,
                confidence=0.9,
                candidates=[_candidate(value, index)],
                resolution_reason="One deterministic candidate.",
            )
    return ProjectFacts.from_fields(
        source_document=SourceDocument(source_file="sample.pdf", source_type=SourceType.PDF),
        fields=ProjectFields(**values),
    )


def build_item(
    item_id: str,
    *,
    module: str = "一、废标红线",
    risk: str = "一票否决",
    requirement_type: str = "REJECTION",
    topic: str = "最高限价与控制价",
    page: int = 5,
    text: str = "投标报价超过最高限价的，按无效投标处理。",
) -> DynamicReviewItem:
    action = "核对报价是否超过最高限价。"
    criteria = "报价不高于最高限价。"
    return DynamicReviewItem(
        item_id=item_id,
        module=module,
        submodule="报价",
        risk_level=risk,
        requirement_type=requirement_type,
        topic=topic,
        source_requirement=text,
        verification_action=action,
        pass_criteria=criteria,
        consequence_if_failed="否决投标。",
        source_locator=f"第{page}页 / {topic} / （pdf_block）",
        source_page=page,
        source_section="供应商须知前附表",
        source_evidence=f"源条款：{text}",
        source_requirement_ids=[f"R{page:03d}"],
        related_project_fact="max_price",
        related_fact_value="750万元",
        confidence=0.8,
        cell_text=f"【要求正文】{text}\n【复核动作】{action}\n【核验标准】{criteria}",
    )


def build_plan(extra: int = 0) -> DynamicReviewPlan:
    items = [
        build_item("DR001"),
        build_item(
            "DR002",
            module="四、报价与合同商务",
            risk="高",
            requirement_type="PRICING",
            topic="报价口径",
            page=6,
            text="投标报价应为完成本项目全部内容的含税总价。",
        ),
        build_item(
            "DR003",
            module="五、技术响应",
            risk="中",
            requirement_type="TECHNICAL",
            topic="技术参数",
            page=7,
            text="水泵参数应满足询比文件技术要求。",
        ),
    ]
    for index in range(extra):
        items.append(
            build_item(
                f"DR{index + 4:03d}",
                risk="高",
                requirement_type="QUALIFICATION",
                topic=f"资格条件{index}",
                page=8 + (index % 5),
                text=f"资格要求第 {index} 条：" + "供应商应提供有效证明材料。" * 3,
            )
        )
    return DynamicReviewPlan(
        items=tuple(items),
        index=None,
        source_requirement_count=len(items),
        dropped_duplicate_count=0,
    )


def _pdf_row(table_index: int, row_index: int, page: int, texts: list[str]) -> PdfTableRow:
    return PdfTableRow(
        row_index=row_index,
        cells=[
            PdfTableCell(
                row_index=row_index,
                column_index=column_index,
                text=text,
                locator=PdfTableLocator(
                    locator_type="pdf_table_cell",
                    page=page,
                    table_index=table_index,
                    row_index=row_index,
                    column_index=column_index,
                ),
            )
            for column_index, text in enumerate(texts)
        ],
    )


def build_document(
    *, price_rows: int = 2, blank_form: bool = True, filled: bool = True
) -> NormalizedDocument:
    tables = []
    if filled:
        tables.append(
            PdfTable(
                page=2,
                table_index=0,
                bbox=(0.0, 0.0, 500.0, 200.0),
                rows=[
                    _pdf_row(0, 0, 2, ["序号", "内容", "单位", "数量", "单价\n（元）", "小计\n（元）"]),
                    _pdf_row(0, 1, 2, ["1", "一体化综合柜", "套", "2", "50,000.00", "100,000.00"]),
                    _pdf_row(0, 2, 2, ["2", "信创服务器", "台", "5", "118,000.00", "590,000.00"]),
                ][: price_rows + 1],
            )
        )
    if blank_form:
        tables.append(
            PdfTable(
                page=3,
                table_index=0,
                bbox=(0.0, 0.0, 500.0, 200.0),
                rows=[
                    _pdf_row(0, 0, 3, ["序号", "名称", "单位", "数量", "单价（元）", "合价（元）"]),
                    _pdf_row(0, 1, 3, ["1", "", "", "", "", ""]),
                    _pdf_row(0, 2, 3, ["2", "", "", "", "", ""]),
                ],
            )
        )
    return NormalizedDocument(
        source_file="sample.pdf",
        source_type=SourceType.PDF,
        status=DocumentStatus.PARSED,
        page_count=12,
        text_length=1000,
        pages=[],
        tables=tables,
    )


def build_format_template() -> BidFormatTemplate:
    return BidFormatTemplate(
        source_heading="响应文件格式",
        confidence=0.9,
        elements=[
            FormatElement(type="heading", text="一、投标函", heading_level=1),
            FormatElement(
                type="paragraph",
                text="投标人法定代表人签字并加盖公章，填写日期。",
            ),
            FormatElement(type="heading", text="二、报价表", heading_level=1),
            FormatElement(type="paragraph", text="附营业执照复印件等证明材料。"),
        ],
    )


def build_evidence() -> list[ReviewEvidenceItem]:
    def item(review_id: str, page: int, text: str, state: str = "FOUND") -> ReviewEvidenceItem:
        candidate = ReviewEvidenceCandidate(
            review_item_id=review_id,
            candidate_index=0,
            page=page,
            section="供应商须知前附表",
            locator=PdfLocator(locator_type="pdf_block", page=page, block_index=0),
            evidence_text=text,
            retrieval_method="broad_keyword",
            retrieval_score=5.0,
            relevance_score=10.0,
            source_priority=3.0,
            final_score=15.0,
            score=15.0,
        )
        return ReviewEvidenceItem(
            review_item_id=review_id,
            review_intent="复核最高限价",
            state=state.upper(),  # type: ignore[arg-type]
            candidates=[candidate],
            evidence_text=f"{state}：第{page}页 {text}",
        )

    return [
        item("R001", 5, "最高限价：7500000.00 元。"),
        item("R001", 5, "最高限价：7500000.00 元。"),
        item("R002", 6, "报价应为含税总价。"),
    ]


def build_workbook(
    tmp_path: Path, *, extra_items: int = 0, blank_form: bool = True, filled: bool = True
) -> Path:
    path = tmp_path / "投标项目复核表.xlsx"
    build_review_workbook(
        build_facts(),
        path,
        normalized_document=build_document(blank_form=blank_form, filled=filled),
        format_template=build_format_template(),
        review_evidence=build_evidence(),
        dynamic_plan=build_plan(extra_items),
    )
    return path


def rows_of(path: Path, title: str) -> list[list[object]]:
    workbook = load_workbook(path, data_only=False)
    try:
        return [
            list(row)
            for row in workbook[title].iter_rows(min_row=2, values_only=True)
            if any(cell not in (None, "") for cell in row)
        ]
    finally:
        workbook.close()


def test_delivered_sheet_stays_first_and_views_are_appended(tmp_path: Path) -> None:
    workbook = load_workbook(build_workbook(tmp_path), data_only=False)
    try:
        assert workbook.sheetnames[0] == DELIVERED_SHEET
        assert workbook.sheetnames[1:] == list(SHEET_TITLES)
    finally:
        workbook.close()


def test_not_found_fact_shows_marker_and_never_a_value(tmp_path: Path) -> None:
    rows = {row[1]: row for row in rows_of(build_workbook(tmp_path), SHEET_TITLES[1])}
    row = rows["lot_name"]
    assert row[4] == "NOT_FOUND"
    assert row[3] == "NOT_FOUND"
    assert row[8] in (None, "")  # no evidence excerpt is invented either


def test_needs_review_fact_keeps_its_candidates_visible(tmp_path: Path) -> None:
    path = build_workbook(tmp_path)
    rows = {row[1]: row for row in rows_of(path, SHEET_TITLES[1])}
    assert rows["budget"][4] == "NEEDS_REVIEW"
    exceptions = {row[1]: row for row in rows_of(path, SHEET_TITLES[6])}
    assert "budget" in exceptions
    assert "B" in str(exceptions["budget"][3])


def test_resolved_fact_value_and_locator_equal_project_facts(tmp_path: Path) -> None:
    facts = build_facts()
    rows = {row[1]: row for row in rows_of(build_workbook(tmp_path), SHEET_TITLES[1])}
    field = facts.fields.project_name
    row = rows["project_name"]
    assert row[3] == field.resolved_value
    assert row[4] == "RESOLVED"
    assert str(row[5]) == str(field.candidates[0].locator.page)
    assert str(row[8]).startswith("字段证据")


def test_manual_columns_start_unreviewed_with_a_dropdown(tmp_path: Path) -> None:
    path = build_workbook(tmp_path)
    workbook = load_workbook(path, data_only=False)
    try:
        worksheet = workbook[SHEET_TITLES[1]]
        values = {worksheet.cell(row=row, column=13).value for row in range(2, 25)}
        assert values == {"未复核"}
        validations = [
            validation
            for validation in worksheet.data_validations.dataValidation
            if str(validation.sqref).startswith("M")
        ]
        assert validations
        assert all(option in validations[0].formula1 for option in MANUAL_OPTIONS)
    finally:
        workbook.close()


def test_no_formula_reads_a_manual_column(tmp_path: Path) -> None:
    workbook = load_workbook(build_workbook(tmp_path), data_only=False)
    try:
        manual_columns = {
            SHEET_TITLES[1]: ["M", "N"],
            SHEET_TITLES[2]: ["M"],
            SHEET_TITLES[3]: ["O"],
            SHEET_TITLES[5]: ["M"],
            SHEET_TITLES[6]: ["H"],
        }
        offenders = []
        for title, columns in manual_columns.items():
            for row in workbook[title].iter_rows():
                for cell in row:
                    if not isinstance(cell.value, str) or not cell.value.startswith("="):
                        continue
                    if any(f"'${column}$" in cell.value for column in columns):
                        offenders.append(cell.coordinate)
        assert offenders == []
    finally:
        workbook.close()


def test_dashboard_counts_are_formulas_over_the_views(tmp_path: Path) -> None:
    workbook = load_workbook(build_workbook(tmp_path), data_only=False)
    try:
        worksheet = workbook[SHEET_TITLES[0]]
        formulas = [
            cell.value
            for row in worksheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        ]
        assert len(formulas) >= 12
        assert any(SHEET_TITLES[1] in formula and "COUNTIF" in formula for formula in formulas)
        assert any(SHEET_TITLES[4] in formula for formula in formulas)
    finally:
        workbook.close()


def test_star_and_veto_are_independent_source_dimensions(tmp_path: Path) -> None:
    """★ is the source's own marker; 否决性 is the source's own consequence.

    Round 6 split the two: neither may be derived from the generator's risk
    table.  This fixture plan carries no source marker and no proven consequence
    rule, so an honest sheet shows neither -- the old behaviour (★ on the one
    ``REJECTION``-typed row) was the generator's opinion, not the source's.
    """

    rows = rows_of(build_workbook(tmp_path), SHEET_TITLES[3])
    starred = {row[1] for row in rows if str(row[7]).strip() == "★"}
    vetoed = {row[1] for row in rows if str(row[8]).strip() == "是"}
    assert starred == set()
    assert vetoed == set()
    # and the three source columns are written from three different plan fields
    # (强制性类型 J=9, 源标记 Q=16, 实质性依据 R=17, 否决依据 S=18)
    types = {str(row[9] or "").strip() for row in rows}
    markers = {str(row[16] or "").strip() for row in rows}
    bases = {str(row[18] or "").strip() for row in rows}
    assert all(marker == "" for marker in markers)
    assert all(basis == "" for basis in bases)
    # 强制性类型 keeps the source basis and, with no source-backed basis, falls
    # back to the requirement's own type rather than to a marker
    assert all(token in {"", "REJECTION", "PRICING"} for token in types)
    assert "SUBSTANTIVE_STARRED" not in types


def test_requirement_text_action_and_criteria_are_separate_columns(tmp_path: Path) -> None:
    rows = rows_of(build_workbook(tmp_path), SHEET_TITLES[3])
    row = rows[0]
    assert row[4] != row[5] != row[6]
    assert "复核" in str(row[5]) or "核对" in str(row[5])
    assert row[4].startswith("投标报价")


def test_price_rows_come_from_the_source_table_and_are_not_computed(tmp_path: Path) -> None:
    path = build_workbook(tmp_path)
    rows = rows_of(path, SHEET_TITLES[4])
    items = [row for row in rows if str(row[9]).startswith("源分项")]
    assert len(items) == 2
    assert items[0][3] == "一体化综合柜"
    assert items[0][7] == 50000
    assert items[1][8] == 590000
    workbook = load_workbook(path, data_only=False)
    try:
        worksheet = workbook[SHEET_TITLES[4]]
        assert all(
            not (isinstance(cell.value, str) and cell.value.startswith("="))
            for row in worksheet.iter_rows()
            for cell in row
        )
    finally:
        workbook.close()


def test_budget_and_max_price_stay_separate_rows(tmp_path: Path) -> None:
    rows = rows_of(build_workbook(tmp_path), SHEET_TITLES[4])
    keys = [row[1] for row in rows if row[1] in ("budget", "max_price")]
    assert keys == ["budget", "max_price"]
    assert rows[0][1] != rows[1][1]


def test_blank_price_form_is_reported_not_invented(tmp_path: Path) -> None:
    rows = rows_of(build_workbook(tmp_path / "blank", filled=False), SHEET_TITLES[4])
    items = [row for row in rows if str(row[9]).startswith("源分项")]
    forms = [row for row in rows if "空白" in str(row[4])]
    assert items == []
    assert len(forms) == 1
    assert "单价（元）" in str(forms[0][2])


def test_evidence_index_deduplicates_locators(tmp_path: Path) -> None:
    rows = rows_of(build_workbook(tmp_path), SHEET_TITLES[7])
    assert [row[0] for row in rows] == ["EV001", "EV002"]


def test_exceptions_list_only_open_items(tmp_path: Path) -> None:
    path = build_workbook(tmp_path)
    exceptions = {row[1] for row in rows_of(path, SHEET_TITLES[6])}
    facts = {row[1]: row for row in rows_of(path, SHEET_TITLES[1])}
    assert "budget" in exceptions
    assert "lot_name" in exceptions
    assert all(facts[key][4] != "RESOLVED" for key in exceptions if key in facts)


def test_structure_view_lists_the_source_headings(tmp_path: Path) -> None:
    rows = rows_of(build_workbook(tmp_path), SHEET_TITLES[5])
    titles = [row[2] for row in rows]
    assert titles == ["一、投标函", "二、报价表"]
    assert rows[0][4] == "是"  # signature required by the section text
    assert rows[0][5] == "是"  # seal required by the section text


def test_a_manual_value_never_changes_a_machine_column(tmp_path: Path) -> None:
    path = build_workbook(tmp_path)
    before = [row[3] for row in rows_of(path, SHEET_TITLES[1])]
    workbook = load_workbook(path, data_only=False)
    try:
        worksheet = workbook[SHEET_TITLES[1]]
        worksheet["M2"] = "通过"
        worksheet["N2"] = "人工填写的复核值"
        workbook.save(path)
    finally:
        workbook.close()
    after = [row[3] for row in rows_of(path, SHEET_TITLES[1])]
    assert before == after
    facts = build_facts()
    assert facts.fields.project_name.resolved_value == "示例一体化泵站项目"


def test_many_review_rows_keep_the_workbook_structurally_sound(tmp_path: Path) -> None:
    path = build_workbook(tmp_path, extra_items=995)
    workbook = load_workbook(path, data_only=False)
    try:
        assert workbook.sheetnames[0] == DELIVERED_SHEET
        assert workbook.sheetnames[1:] == list(SHEET_TITLES)
        worksheet = workbook[SHEET_TITLES[3]]
        assert worksheet.max_row >= 990
        assert worksheet.freeze_panes
        assert worksheet.auto_filter.ref
    finally:
        workbook.close()


@pytest.mark.parametrize("title", SHEET_TITLES[1:])
def test_every_view_row_has_a_review_conclusion_column(tmp_path: Path, title: str) -> None:
    path = build_workbook(tmp_path)
    workbook = load_workbook(path, data_only=False)
    try:
        headers = [
            cell.value for cell in next(workbook[title].iter_rows(min_row=1, max_row=1))
        ]
        if title not in (SHEET_TITLES[4], SHEET_TITLES[7]):
            assert MANUAL_CONCLUSION in headers or MANUAL_NOTE in headers
        assert headers[0] is not None
    finally:
        workbook.close()
