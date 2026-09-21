from __future__ import annotations

from pathlib import Path

from docx import Document

from tender_basic.document_models import (
    DocumentBlock,
    DocumentStatus,
    NormalizedDocument,
    PdfBlockElement,
)
from tender_basic.document_parser import parse_document
from tender_basic.format_extractor import (
    FORMAT_SECTION_FOUND,
    FORMAT_SECTION_NEEDS_REVIEW,
    FORMAT_SECTION_NOT_FOUND,
    extract_bid_format,
)
from tender_basic.models import PdfLocator, SourceType


def _create_source_format_docx(path: Path, *, with_next_chapter: bool = True) -> None:
    document = Document()
    document.add_paragraph("采购项目文件")
    document.add_paragraph("第六章 投标文件格式", style="Heading 1")
    document.add_paragraph("以下固定说明文字不得被事实抽取器当作项目事实。")
    document.add_paragraph("一、投标函", style="Heading 2")
    document.add_paragraph("投标函固定文本：投标人应按本格式填写。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "报价表"
    table.cell(0, 1).text = "金额"
    table.cell(1, 0).text = "投标人"
    table.cell(1, 1).text = "________________"
    document.add_paragraph("二、法定代表人身份证明", style="Heading 2")
    document.add_paragraph("三、授权委托书", style="Heading 2")
    document.add_paragraph("四、报价表", style="Heading 2")
    document.add_paragraph("签字盖章处：________________")
    if with_next_chapter:
        document.add_paragraph("第七章 评审方法", style="Heading 1")
        document.add_paragraph("评审内容不属于投标文件格式骨架。")
    document.save(path)


def _synthetic_pdf_document(texts: list[str]) -> NormalizedDocument:
    elements = []
    for index, text in enumerate(texts):
        block = DocumentBlock(
            block_index=index,
            text=text,
            bbox=(0.0, float(index), 100.0, float(index + 1)),
            block_type="text",
            locator=PdfLocator(page=1, block_index=index),
        )
        elements.append(PdfBlockElement(page_number=1, block=block))
    return NormalizedDocument(
        source_file="synthetic.pdf",
        source_type=SourceType.PDF,
        status=DocumentStatus.PARSED,
        page_count=1,
        text_length=sum(len(text) for text in texts),
        elements=elements,
    )


def test_source_format_section_is_found_and_stops_at_next_chapter(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    _create_source_format_docx(source)

    normalized = parse_document(source)
    template = extract_bid_format(normalized)

    assert template.section_status == FORMAT_SECTION_FOUND
    assert template.format_source == "SOURCE_DOCUMENT"
    assert template.source_heading == "第六章 投标文件格式"
    assert template.source_locator.model_dump() == {
        "locator_type": "docx_paragraph",
        "paragraph_index": 1,
    }
    headings = [element.text for element in template.elements if element.type == "heading"]
    assert headings == [
        "第六章 投标文件格式",
        "一、投标函",
        "二、法定代表人身份证明",
        "三、授权委托书",
        "四、报价表",
    ]
    assert all("第七章" not in element.text for element in template.elements)
    assert any(element.type == "table" for element in template.elements)


def test_body_phrase_does_not_start_a_format_section(tmp_path: Path) -> None:
    source = tmp_path / "body-phrase.docx"
    document = Document()
    document.add_paragraph("本段说明投标文件格式，但它不是格式章节标题。")
    document.save(source)

    template = extract_bid_format(parse_document(source))

    assert template.section_status == FORMAT_SECTION_NOT_FOUND
    assert template.format_source == "GENERIC_FALLBACK"
    assert template.elements == []


def test_pdf_chapter_heading_and_ordered_forms_are_preserved(tmp_path: Path) -> None:
    texts = [
        "第六章 投标文件格式",
        "一、投标函",
        "二、法定代表人身份证明",
        "三、授权委托书",
        "四、报价表",
        "第七章 评审方法",
    ]
    document = _synthetic_pdf_document(texts)

    template = extract_bid_format(document)

    assert template.section_status == FORMAT_SECTION_FOUND
    assert template.source_locator.page == 1
    assert [element.text for element in template.elements if element.type == "heading"] == [
        "第六章 投标文件格式",
        "一、投标函",
        "二、法定代表人身份证明",
        "三、授权委托书",
        "四、报价表",
    ]


def test_unbounded_large_pdf_format_tail_is_marked_for_review(tmp_path: Path) -> None:
    document = _synthetic_pdf_document(
        ["第六章 投标文件格式"] + [f"固定内容 {index}" for index in range(130)]
    )

    template = extract_bid_format(document)

    assert template.section_status == FORMAT_SECTION_NEEDS_REVIEW
    assert template.format_source == "GENERIC_FALLBACK"
    assert template.elements == []


def test_long_pdf_format_chapter_with_ordered_forms_can_end_at_eof() -> None:
    document = _synthetic_pdf_document(
        [
            "第六章 投标文件格式",
            "一、投标函",
            "二、授权委托书",
            "三、报价表",
        ]
        + [f"固定内容 {index}" for index in range(130)]
    )

    template = extract_bid_format(document)

    assert template.section_status == FORMAT_SECTION_FOUND
    assert template.format_source == "SOURCE_DOCUMENT"
    assert [element.text for element in template.elements if element.type == "heading"][:4] == [
        "第六章 投标文件格式",
        "一、投标函",
        "二、授权委托书",
        "三、报价表",
    ]
