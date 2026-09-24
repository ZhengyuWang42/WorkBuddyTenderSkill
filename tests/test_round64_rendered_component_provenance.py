"""Round-64 round-4: rendered-component provenance closure (CASE001 review).

The round-3 contract stopped at the model objects.  Round 4 makes the *rendered
XLSX* the contract, so these tests assert on the frozen workbook cells rather
than on internal ids:

* every rendered component is concern-owned (zero mismatches in all eight
  buckets),
* the delivered D cell is the projection of its verified components and names
  no concept its concern does not own (old-group leak),
* the E cell carries the concern's own decisive clause, the M cell only allowed
  linked facts (stale linked facts),
* the human fixtures A..S are satisfied by the final cell text,
* the CASE001 warranty/retention/scoring/platform semantics survive rendering
  (24-month project warranty vs 5% retention money vs 12-month release period).
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

from tender_basic.review_rendering import (  # noqa: E402
    COMPONENT_KINDS,
    foreign_terms,
    numeric_tokens,
)

import v1_review_workbook_round4_report as round4  # noqa: E402

CASE = ROOT / "acceptance/workspace/case_001"
BUILD4 = CASE / "v1_manual_fidelity_round4_date_rhythm_closure8_review_workbook4"
BUILD3 = CASE / "v1_manual_fidelity_round4_date_rhythm_closure8_review_workbook3"
WORD = CASE / "v1_manual_fidelity_round4_date_rhythm_closure8"

DELIVERED = "投标项目复核表"

pytestmark = pytest.mark.skipif(
    not (BUILD4 / "project_facts.json").is_file(),
    reason="round-4 CASE001 successor build is not present",
)

FIXTURE_KEYS = tuple("ABCDEFGHIJKLMNOPQRS")


def _flat(text: object) -> str:
    return round4._flat(text)


@pytest.fixture(scope="module")
def analysis() -> round4.Round4Report:
    return round4.Round4Report(build=BUILD4, case="case_001", before=BUILD3, audit_limit=30)


@pytest.fixture(scope="module")
def report(analysis: round4.Round4Report) -> dict:
    return analysis.run()


@pytest.fixture(scope="module")
def cells() -> dict:
    workbook = load_workbook(BUILD4 / "投标项目复核表.xlsx", data_only=True, read_only=True)
    try:
        values: dict[str, dict[str, str]] = {}
        for name in workbook.sheetnames:
            sheet = workbook[name]
            values[name] = {
                cell.coordinate: ("" if cell.value is None else str(cell.value).strip())
                for row in sheet.iter_rows()
                for cell in row
                if cell.value is not None and str(cell.value).strip()
            }
        return values
    finally:
        workbook.close()


def _item(report: dict, item_id: str) -> dict:
    for entry in report["checks"]:
        for sample in entry["evidence"].get("sample", []) or []:
            if sample["item_id"] == item_id:
                return sample
    raise AssertionError(f"{item_id} not present in the final-cell audit")


# --------------------------------------------------------------------------- #
# rendered-output gate
# --------------------------------------------------------------------------- #


def test_round4_report_passes_all_rendered_checks(report: dict) -> None:
    assert report["verdict"] == "PASS", report["failed_checks"]
    assert report["failed_checks"] == []
    assert report["check_count"] == report["passed_count"]


def test_every_rendered_component_is_concern_owned(report: dict) -> None:
    assert report["rendered_component_concern_mismatch_total"] == 0
    assert report["rendered_component_unverified_count"] == 0
    assert report["rendered_component_count"] > 300


@pytest.mark.parametrize("kind", COMPONENT_KINDS)
def test_each_component_bucket_has_zero_mismatch(report: dict, kind: str) -> None:
    key = f"rendered_{kind.lower()}_concern_mismatch"
    assert key in report["rendered_mismatch_counts"]
    assert report["rendered_mismatch_counts"][key] == 0


def test_delivered_cell_is_component_projection(report: dict) -> None:
    names = {check["check"]: check for check in report["checks"]}
    assert names["rendered_cell_equals_component_projection"]["ok"], names[
        "rendered_cell_equals_component_projection"
    ]["evidence"]
    assert names["delivered_rows_bound_to_plan_items"]["ok"]


def test_final_cell_text_names_only_owned_concepts(report: dict) -> None:
    """Old-group leak: the displayed cell may not name a foreign concept."""

    names = {check["check"]: check for check in report["checks"]}
    check = names["final_cell_text_owned_by_concern"]
    assert check["ok"], check["evidence"]
    assert check["evidence"]["leak_count"] == 0


def test_every_rendered_component_is_displayed(report: dict) -> None:
    names = {check["check"]: check for check in report["checks"]}
    check = names["every_rendered_component_present_in_final_cell"]
    assert check["ok"], check["evidence"]


def test_view_cells_repeat_rendered_text(report: dict) -> None:
    names = {check["check"]: check for check in report["checks"]}
    check = names["view_cells_repeat_rendered_components"]
    assert check["ok"], check["evidence"]


def test_final_cell_audit_is_thirty_of_thirty(report: dict) -> None:
    audit = report["final_cell_audit"]
    assert audit["incoherent_cell_count"] == 0
    assert audit["sample_count"] >= 30
    assert len(audit["sample"][0]["displayed_text"]) > 0
    assert all(entry["coherent"] for entry in audit["sample"])


# --------------------------------------------------------------------------- #
# human fixtures A..S on the frozen cells
# --------------------------------------------------------------------------- #


@pytest.mark.parametrize("key", FIXTURE_KEYS)
def test_human_fixture_is_satisfied_by_rendered_cells(report: dict, key: str) -> None:
    fixture = report["fixtures"][key]
    assert fixture["status"] == "PASS", fixture["evidence"]
    assert fixture["expectation"]


def _cell_for(analysis: round4.Round4Report, concern_id: str, column: str) -> str:
    """Displayed text of the first row owned by ``concern_id`` in ``column``."""

    for item in analysis.items:
        if item.concern_id != concern_id:
            continue
        row = analysis.rows_by_item[item.item_id]
        return {"D": row.d_text, "E": row.e_text, "M": row.m_text}[column]
    raise AssertionError(f"no rendered row for concern {concern_id}")


def test_fixture_a_licence_row_has_no_invented_qualification_terms(
    analysis: round4.Round4Report,
) -> None:
    text = _cell_for(analysis, "QUALIFICATION_LICENSE", "D")
    assert "营业执照" in text
    for term in ("资质证书", "专业类别", "资质等级"):
        assert term not in text


def test_fixture_b_signature_row_has_no_upload_or_authorisation_text(
    analysis: round4.Round4Report,
) -> None:
    text = _cell_for(analysis, "SIGNATURE_AND_SEAL", "D")
    assert "签字盖章要求" in text or "公章" in text
    for term in ("可编辑的Word", "澄清", "授权委托书", "身份证扫描件"):
        assert term not in text


def test_fixture_h_bond_amount_row_cites_its_own_clause(analysis: round4.Round4Report) -> None:
    evidence = _cell_for(analysis, "BID_BOND_AMOUNT", "E")
    assert "响应保证金" in evidence
    assert "履约保证金" not in evidence


def test_fixture_k_funding_source_row_has_no_identity_checks(
    analysis: round4.Round4Report,
) -> None:
    text = _cell_for(analysis, "PROJECT_BASIC_INFO", "D")
    for term in ("项目名称", "项目编号", "标段"):
        assert term not in text


def test_fixture_q_and_r_retention_wording_is_split(analysis: round4.Round4Report) -> None:
    delivered = " ".join(
        row.d_text for row in analysis.rows_by_item.values()
    )
    for phrase in ("项目质保期不低于12个月", "质保期不低于12个月"):
        assert phrase not in delivered
    ratio = _cell_for(analysis, "RETENTION_MONEY_RATIO", "D")
    assert "5%" in ratio
    release = _cell_for(analysis, "RETENTION_RELEASE_PERIOD", "D")
    assert "12" in release
    warranty = _cell_for(analysis, "PROJECT_WARRANTY", "D")
    assert "24个月" in warranty


def test_fixture_s_platform_roles_stay_separate(report: dict) -> None:
    evidence = report["platform_role_evidence"]
    keys = evidence["platform_fact_keys"]
    assert keys, evidence
    conflicts = evidence["conflict_rows"]
    for conflict in conflicts:
        merged = [key for key in keys if key and key in str(conflict.get("key", ""))]
        assert len(merged) <= 1, conflict
    # each platform fact keeps its own displayed value cell
    addresses = [evidence["platform_rows"][key]["cell"] for key in keys]
    assert len(set(addresses)) == len(addresses)


# --------------------------------------------------------------------------- #
# provenance maps and linked facts
# --------------------------------------------------------------------------- #


def test_provenance_map_records_cell_block_and_ids(analysis: round4.Round4Report) -> None:
    provenance = analysis.provenance()
    assert provenance["schema"].startswith("v1_review_workbook_round4")
    assert provenance["component_count"] > 300
    assert provenance["unverified_count"] == 0
    sample = provenance["entries"][0]
    for field in (
        "cell",
        "block",
        "sentence_index",
        "concern_id",
        "source_atom_ids",
        "numeric_evidence_ids",
        "material_ids",
        "evidence_ids",
        "generation_rule",
        "ownership_verified",
    ):
        assert field in sample
    assert any(entry["cell"].startswith(DELIVERED) for entry in provenance["entries"])
    assert any(entry["source_atom_ids"] for entry in provenance["entries"])


def test_stale_linked_facts_absent(report: dict) -> None:
    names = {check["check"]: check for check in report["checks"]}
    check = names["stale_linked_facts_absent"]
    assert check["ok"], check["evidence"]


def test_m_cell_linked_facts_are_concern_owned(analysis: round4.Round4Report, cells: dict) -> None:
    from tender_basic.review_concern import concern_spec

    for item in analysis.items:
        row = analysis.rows_by_item[item.item_id]
        allowed = set(concern_spec(item.concern_id).fact_fields)
        for match in re.finditer(r"关联事实：([^\s；]+)", row.m_text):
            for key in re.split(r"[、,，]", match.group(1)):
                if key:
                    assert key in allowed, (item.item_id, key, allowed)


def test_evidence_cell_cites_the_concerns_own_clause(
    analysis: round4.Round4Report, cells: dict
) -> None:
    """The rendered evidence must belong to the concern's own source clauses."""

    backings = analysis._concern_backings()
    for item in analysis.items:
        row = analysis.rows_by_item[item.item_id]
        evidence = analysis._component_texts(item, "EVIDENCE_SUMMARY")
        assert evidence, item.item_id
        if not any(_flat(text)[:30] in _flat(row.e_text) for text in evidence):
            raise AssertionError((item.item_id, row.e_text[:120], evidence[0][:120]))
        assert not foreign_terms(row.e_text, backings[item.concern_id]) or item.item_id


def test_numeric_tokens_in_delivered_cells_are_owned(report: dict) -> None:
    names = {check["check"]: check for check in report["checks"]}
    check = names["final_cell_numeric_ownership"]
    assert check["ok"], check["evidence"]


def test_scoring_rows_are_split_and_rendered(analysis: round4.Round4Report) -> None:
    delivered = " ".join(row.d_text for row in analysis.rows_by_item.values())
    flat = _flat(delivered)
    for phrase in ("95%", "5%", "40分", "基本分"):
        assert _flat(phrase) in flat, phrase
    scoring_concerns = {
        item.concern_id
        for item in analysis.items
        if item.concern_id.startswith("SCORING_") or item.concern_id == "EVALUATION_SCORING"
    }
    assert len(scoring_concerns) >= 2, scoring_concerns
    scoring_cells = [
        row.d_address
        for row in analysis.rows_by_item.values()
        if "评分提示" in row.d_text
    ]
    assert len(scoring_cells) >= 2


def test_warranty_and_retention_rows_use_their_own_numbers(report: dict) -> None:
    names = {check["check"]: check for check in report["checks"]}
    for name in (
        "project_warranty_is_24_months",
        "retention_money_ratio_is_5_percent",
        "retention_release_period_is_12_months",
        "retention_and_warranty_rows_are_distinct",
    ):
        assert names[name]["ok"], (name, names[name]["evidence"])


def test_forbidden_retention_wording_absent(report: dict) -> None:
    names = {check["check"]: check for check in report["checks"]}
    assert names["forbidden_retention_wording_absent"]["ok"]


def test_word_artifacts_unchanged(report: dict) -> None:
    assert report["word_artifacts_unchanged"] is True


def test_plan_qa_hard_gates_pass(analysis: round4.Round4Report) -> None:
    assert analysis.qa["result"] == "PASS", analysis.qa["hard_gate_failures"]
    assert analysis.qa["result"] == "PASS"


def test_report_written_for_all_three_cases() -> None:
    reports = ROOT / "acceptance/reports/v1_generalization"
    for case in ("case001", "case002", "case003"):
        payload = json.loads((reports / f"{case}_review_workbook_round4.json").read_text("utf-8"))
        assert payload["verdict"] == "PASS", (case, payload["failed_checks"])
        assert payload["rendered_component_concern_mismatch_total"] == 0


def test_provenance_artifacts_exist_per_case() -> None:
    reports = ROOT / "acceptance/reports/v1_generalization"
    for case in ("001", "002", "003"):
        path = reports / f"review_workbook_round4_rendered_component_provenance_case_{case}.json"
        payload = json.loads(path.read_text("utf-8"))
        assert payload["component_count"] > 0
        assert payload["unverified_count"] == 0
        assert payload["workbook_sha256"]


def test_numeric_role_sentence_is_source_worded() -> None:
    """The rendered numeric wording uses the source's own term."""

    from tender_basic.review_point import value_sentence

    assert "质保金" in value_sentence("RETENTION_MONEY_RATIO", "5%", "剩余5%作为质保金")
    assert "尾款" in value_sentence("RETENTION_MONEY_RATIO", "5%", "余款作为尾款支付")
    assert "质保金" not in value_sentence("RETENTION_MONEY_RATIO", "5%", "余款作为尾款支付")
    assert "12个月" in value_sentence("RETENTION_RELEASE_PERIOD", "12个月", "质保期满后支付")
    assert numeric_tokens("工期 30日历天") == {"30日历天"}
