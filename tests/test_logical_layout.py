"""Desensitized regression fixtures for logical layout, not real answers."""
import pytest
from docx import Document
from tender_basic.document_models import PdfTextLine,PdfTextSpan
from tender_basic.source_format import SourcePageElement,SourceFillSlot,SourceLine
from tender_basic.models import FieldName
from tender_basic.page_layout import build_page_layout,LogicalParagraph,form_parts,join_text
from tender_basic.word_safe_source_builder import build_source_format_docx
from tender_basic.source_fill_policy import slot_replacement
from tender_basic.logical_layout_qa import inspect_logical_layout
from test_source_format_fidelity import _paragraph,_paragraph_template,_resolved
from test_review_builder import make_project_facts


def template(lines):
    paragraphs=[]
    for index,(text,x,y,right) in enumerate(lines):
        p=_paragraph(text,block_index=index,bbox=(x,y,right,y+12),font_size=12)
        p.alignment='center'  # Deliberately unreliable upstream glyph heuristic.
        p.lines=[PdfTextLine(text=text,bbox=p.bbox,spans=[PdfTextSpan(text=text,bbox=p.bbox,font_name='SimSun',font_size=12)])]
        paragraphs.append(p)
    t=_paragraph_template(paragraphs[0])
    page=t.source_pages[0]
    page.paragraphs=paragraphs
    page.elements=[SourcePageElement(type='paragraph',index=i,bbox=p.bbox) for i,p in enumerate(paragraphs)]
    return t


def test_adjacent_blocks_form_one_paragraph_despite_glyph_edges(tmp_path):
    t=template([('我方已充分研究了招标文件的所有规定并确认愿意',96,80,520),
        ('以人民币响应报',72,100,500),('价，供货质量',72,120,400),('达到约定标准。',72,140,190)])
    output,g=build_source_format_docx(make_project_facts(),t,tmp_path/'body.docx')
    p=Document(output).paragraphs
    assert len(p)==1 and '愿意以人民币响应报价' in p[0].text
    assert p[0].alignment==0 and p[0].paragraph_format.right_indent.pt<60
    assert p[0].paragraph_format.first_line_indent.pt==24
    assert inspect_logical_layout(output,g,t)['mid_sentence_paragraph_breaks']==0


def test_addressee_has_container_not_short_glyph_width(tmp_path):
    t=template([('测试采购单位名称：',112,80,220),('我方确认招标文件所列的所有条款。',72,100,520)])
    output,_=build_source_format_docx(make_project_facts(),t,tmp_path/'name.docx')
    p=Document(output).paragraphs[0]
    assert p.paragraph_format.right_indent.pt==0
    assert p.alignment==0


def test_numbered_continuations_and_shared_list_geometry(tmp_path):
    t=template([('1、我方确认招标文件及其所有附件中的责任和义',96,80,520),
        ('务。',72,100,100),('2、我方在此声明所提供的资料真实且不存',96,120,520),
        ('在文件列明的限制条件。',72,140,250)])
    layout=build_page_layout(t.source_pages[0])
    output,_=build_source_format_docx(make_project_facts(),t,tmp_path/'list.docx')
    ps=Document(output).paragraphs
    assert len(ps)==2 and '责任和义务' in ps[0].text and '且不存在' in ps[1].text
    # The source prints the numbered first row at x=96 and returns every wrapped
    # row to x=72, so the paragraph's body boundary is 72 and its first row is
    # offset by 24 - a first-line indent.  Delivering the first row's own x as the
    # paragraph's whole left indent would put the wrapped row at 96 and the first
    # row at 72, which is the source's geometry inverted.
    content_x0=18.0
    for paragraph in ps:
        left=paragraph.paragraph_format.left_indent.pt
        first=paragraph.paragraph_format.first_line_indent.pt
        assert left+content_x0==72.0
        assert left+first+content_x0==96.0
    for item in layout.elements:
        indent=item.source_indent
        assert indent.classification=='FIRST_LINE_INDENT'
        assert indent.body_left_x==72.0 and indent.measured_continuation
        assert indent.continuation_x==72.0 and indent.first_line_x==96.0
    assert len(layout.list_groups)==1


def test_numeric_sentence_containing_document_is_not_heading():
    t=template([('9、我方认同并遵守招标文件中的全部内',96,80,520),('容。',72,100,100)])
    ps=build_page_layout(t.source_pages[0]).elements
    assert len(ps)==1 and ps[0].kind=='List' and ps[0].logical_text.endswith('内容。')


def test_appendix_catalog_rows_do_not_merge():
    t=template([('附表四 计划开工日期和施工进度网络图',96,80,420),
        ('附表五 施工总平面图',96,100,300),('附表六 临时用地表',96,120,260)])
    assert len(build_page_layout(t.source_pages[0]).elements)==3


def test_form_baselines_split_even_when_labels_are_span_fragments(tmp_path):
    t=template([('法定代表人：（签字）',220,80,400),('地',220,100,232),('址：',245,100,269),
        ('网址：',220,120,260),('电话：',220,140,260),('传真：',220,160,260)])
    # One source block with fragmented physical rows, as in a PDF form.
    p=t.source_pages[0].paragraphs[0]
    ps=t.source_pages[0].paragraphs
    p.text='\n'.join(v.text for v in ps)
    p.lines=[line for v in ps for line in v.lines]
    p.runs=[r for v in ps for r in v.runs]
    p.bbox=(220,80,400,172)
    t=_paragraph_template(p)
    output,g=build_source_format_docx(make_project_facts(),t,tmp_path/'rows.docx')
    doc=Document(output)
    assert len(doc.tables)==0 and len(doc.paragraphs)==5
    assert all(record.get('table_index') is None for record in g['logical_paragraph_records'] if record['kind']=='FormRow')
    assert inspect_logical_layout(output,g,t)['multi_label_form_row_count']==0
    assert form_parts('法定代表人：（签字）地址：网址：电话：传真：') is None


def test_guard_rejects_number_in_name_slot_but_preserves_number_slot():
    p=_paragraph('（项目名称、标段）')
    s=SourceFillSlot(slot_id='s',semantic_hint='项目编号',source_page=1,source_locator=p.locator,
        container_type='paragraph',original_text=p.text,text_start=0,text_end=len(p.text),
        allowed_fact_fields=[FieldName.PROJECT_NUMBER],match_kind='parenthetical')
    facts=_resolved(make_project_facts(),FieldName.PROJECT_NUMBER,'OWNER-001')
    assert slot_replacement(s,facts) is None
    correct=s.model_copy(update={'original_text':'（项目编号）'})
    assert slot_replacement(correct,facts)[0]=='OWNER-001'


def test_inline_slot_does_not_fragment_surrounding_prose(tmp_path):
    t=template([('我方已充分研究了（项目编号）项目所有要求并愿意',96,80,520),('参加本次投标。',72,100,240)])
    from tender_basic.source_format import _paragraph_slots
    t.fill_slots=_paragraph_slots(t.source_pages[0].paragraphs[0],0)
    f=_resolved(make_project_facts(),FieldName.PROJECT_NUMBER,'OWNER-001')
    output,_=build_source_format_docx(f,t,tmp_path/'slot.docx')
    ps=Document(output).paragraphs
    assert len(ps)==1 and 'OWNER-001' in ps[0].text and '愿意参加' in ps[0].text


def test_form_vector_blank_is_editable_visible_blank(tmp_path):
    t=template([('供应商：',72,80,120)])
    t.source_pages[0].lines=[SourceLine(bbox=(122,92,300,92))]
    output,_=build_source_format_docx(make_project_facts(),t,tmp_path/'blank.docx')
    doc=Document(output)
    assert len(doc.tables)==0
    assert not any('_' in r.text or '\u00a0' in r.text for p in doc.paragraphs for r in p.runs)
    assert any('\u2007' in p.text for p in doc.paragraphs)


def test_chinese_join_no_space_and_latin_boundary_retained():
    assert join_text('愿意','以人民币')=='愿意以人民币'
    assert join_text('hello','world')=='hello world'


def test_qa_catches_serialized_narrow_centered_body(tmp_path):
    from docx.shared import Pt
    t=template([('我方确认招标文件所列的所有条款。',72,80,520)])
    output,g=build_source_format_docx(make_project_facts(),t,tmp_path/'bad-width.docx')
    doc=Document(output)
    doc.paragraphs[0].paragraph_format.right_indent=Pt(370)
    doc.paragraphs[0].alignment=1
    doc.save(output)
    qa=inspect_logical_layout(output,g,t)
    assert qa['extreme_body_right_indent_count']==1
    assert qa['centered_body_paragraph_count']==1


def test_qa_does_not_equate_zero_breaks_with_paragraph_success(tmp_path):
    t=template([('我方已充分研究了招标文件的所有规定并确认愿意',96,80,520),('参加本次采购。',72,100,250)])
    output,g=build_source_format_docx(make_project_facts(),t,tmp_path/'fragmented.docx')
    doc=Document(output)
    doc.paragraphs[0].text='我方已充分研究了招标文件的所有规定并确认愿意'
    doc.add_paragraph('参加本次采购。')
    doc.save(output)
    qa=inspect_logical_layout(output,g,t)
    assert qa['word_manual_breaks']==0
    assert qa['paragraph_fragmentation_count']>0 and qa['result']=='NEEDS_REVIEW'
