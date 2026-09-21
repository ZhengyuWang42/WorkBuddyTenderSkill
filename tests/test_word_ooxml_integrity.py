from __future__ import annotations

import os
import subprocess
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from tender_basic.word_ooxml import scan_docx_package
from tender_basic.word_forensics import inspect_docx_ooxml


def _rewrite_document_xml(source: Path, target: Path, mutate) -> None:
    with ZipFile(source, "r") as source_zip, ZipFile(target, "w", ZIP_DEFLATED) as target_zip:
        for info in source_zip.infolist():
            payload = source_zip.read(info.filename)
            if info.filename == "word/document.xml":
                root = etree.fromstring(payload)
                mutate(root)
                payload = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            target_zip.writestr(info, payload)


def test_normal_python_docx_package_is_well_formed_and_reopens(tmp_path: Path) -> None:
    path = tmp_path / "normal.docx"
    document = Document()
    document.add_paragraph("正常 Word 包")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    document.save(path)

    Document(path)
    result = scan_docx_package(path)

    assert result["result"] == "PASS"
    assert result["dangling_relationship_ids"] == []
    assert result["duplicate_relationship_ids"] == []


def test_deep_forensic_inspector_passes_normal_package(tmp_path: Path) -> None:
    path = tmp_path / "normal-deep.docx"
    document = Document()
    document.add_paragraph("深度检查")
    document.save(path)

    report = inspect_docx_ooxml(path)

    assert report["result"] == "PASS"
    assert report["summary"]["error_count"] == 0
    assert report["package"]["zip_integrity"] == "PASS"
    assert report["structure"]["sectPr"]["body_sectPr"] == 1


def test_deep_forensic_inspector_catches_schema_order_and_direct_break(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    bad = tmp_path / "bad.docx"
    document = Document()
    paragraph = document.add_paragraph("schema marker")
    paragraph.paragraph_format.space_after = 2
    run = paragraph.add_run("run")
    run.bold = True
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "cell"
    document.save(source)

    def mutate(root) -> None:
        ppr = root.find(".//" + qn("w:pPr"))
        assert ppr is not None
        ppr.append(etree.Element(qn("w:jc")))
        rpr = root.find(".//" + qn("w:rPr"))
        assert rpr is not None
        rpr.append(etree.Element(qn("w:color")))
        tcpr = root.find(".//" + qn("w:tcPr"))
        assert tcpr is not None
        borders = etree.SubElement(tcpr, qn("w:tcBorders"))
        etree.SubElement(borders, qn("w:top"))
        etree.SubElement(borders, qn("w:start"))
        root.find(".//" + qn("w:p")).append(etree.Element(qn("w:br")))

    _rewrite_document_xml(source, bad, mutate)
    report = inspect_docx_ooxml(bad)
    categories = {issue["category"] for issue in report["errors"]}

    assert report["result"] == "FAIL"
    assert "schema_order" in categories
    assert "paragraph_structure" in categories


def test_deep_forensic_inspector_catches_invalid_xml_control_character(tmp_path: Path) -> None:
    source = tmp_path / "source-control.docx"
    bad = tmp_path / "bad-control.docx"
    document = Document()
    document.add_paragraph("CONTROL_MARKER")
    document.save(source)

    # lxml correctly refuses to serialize an illegal XML control character;
    # inject it into the package bytes so the forensic parser sees the real
    # malformed-part failure mode.
    with ZipFile(source, "r") as source_zip, ZipFile(bad, "w", ZIP_DEFLATED) as bad_zip:
        for info in source_zip.infolist():
            payload = source_zip.read(info.filename)
            if info.filename == "word/document.xml":
                payload = payload.replace(b"CONTROL_MARKER", b"CONTROL_\x01MARKER")
            bad_zip.writestr(info, payload)
    report = inspect_docx_ooxml(bad)

    assert report["result"] == "FAIL"
    assert any(
        issue["category"] in {"xml_control_character", "xml_parse"}
        for issue in report["errors"]
    )


def test_scanner_catches_legacy_floating_table_property_order(tmp_path: Path) -> None:
    path = tmp_path / "bad-table-order.docx"
    document = Document()
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "表格"
    # Simulate the pre-Round-4.1 builder: append tblpPr after tblW/tblLook.
    table._tbl.tblPr.append(OxmlElement("w:tblpPr"))
    document.save(path)

    result = scan_docx_package(path)

    assert result["result"] == "FAIL"
    assert result["schema_order_errors"]
    assert "tblpPr" in result["schema_order_errors"][0]["issue"]


@pytest.mark.skipif(
    os.name != "nt" or os.environ.get("WBTENDER_RUN_WORD_COM_TESTS") != "1",
    reason="desktop Word COM gate is run explicitly on the Windows acceptance host",
)
def test_word_com_normal_open_gate_when_enabled(tmp_path: Path) -> None:
    path = tmp_path / "word-com.docx"
    document = Document()
    document.add_paragraph("Word COM gate")
    document.save(path)
    script = Path(__file__).parents[1] / "scripts" / "validate_word_open.ps1"
    completed = subprocess.run(
        [
            "powershell.exe",
            "-NoProfile",
            "-STA",
            "-ExecutionPolicy",
            "Bypass",
            "-File",
            str(script),
            str(path),
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    assert completed.returncode == 0, completed.stdout + completed.stderr
    assert "WORD_NORMAL_OPEN=PASS" in completed.stdout
