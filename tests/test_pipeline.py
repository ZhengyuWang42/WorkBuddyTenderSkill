from __future__ import annotations

import json
from pathlib import Path

from docx import Document

try:
    import pymupdf as fitz
except ImportError:  # pragma: no cover
    import fitz  # type: ignore[no-redef]

from scripts.run_pipeline import _clear_downstream_artifacts, main as run_pipeline_main


ARTIFACT_NAMES = {
    "normalized_document.json",
    "document.lines.txt",
    "project_facts.json",
    "facts_review_packet.json",
    "fact_gap_packet.json",
    "semantic_candidate_results.json",
    "generation_report.json",
    "source_format_qa.json",
    "review_evidence_packet.json",
    "review_evidence_qa.json",
    "metadata.json",
    "投标项目复核表.xlsx",
    "基础投标文件.docx",
    "qa_report.json",
}


def _create_docx(path: Path) -> None:
    document = Document()
    document.add_paragraph("第一章 招标公告")
    document.add_paragraph("项目名称：Pipeline测试项目")
    document.add_paragraph("项目编号：PRJ-PIPE-001")
    document.add_paragraph("招标编号：BID-PIPE-002")
    document.add_paragraph("最高限价：100万元")
    document.save(path)


def _create_text_pdf(path: Path) -> None:
    pdf = fitz.open()
    first_page = pdf.new_page()
    first_page.insert_text(
        (72, 72),
        "Text tender notice page one with enough extractable content.\nProject reference ABC-001.",
    )
    first_page.insert_text(
        (72, 120),
        "第一章 招标公告\n"
        "项目名称：Pipeline测试项目\n"
        "最高投标限价：100万元，投标报价超过最高投标限价的，其投标无效。\n"
        "投标保证金金额：人民币2万元。",
        fontsize=11,
        fontname="china-s",
    )
    second_page = pdf.new_page()
    second_page.insert_text(
        (72, 72),
        "Text tender notice page two with enough extractable content.\nBid instructions follow.",
    )
    # The review rows are generated from source requirements, so the fixture
    # carries the mandatory clauses a real tender states.
    second_page.insert_text(
        (72, 120),
        "第二章 投标人须知\n"
        "投标人必须具备有效的营业执照，未被列入失信被执行人名单。\n"
        "投标有效期为投标截止之日起90日历天。\n"
        "供货期为合同签订后30日历天。\n"
        "评标办法采用综合评分法，技术分40分，价格分60分。\n"
        "投标人须按第六章响应文件格式编制响应文件并加盖单位公章。",
        fontsize=11,
        fontname="china-s",
    )
    pdf.save(path)
    pdf.close()


def _create_scan_pdf(path: Path) -> None:
    pdf = fitz.open()
    pdf.new_page()
    pdf.new_page()
    pdf.save(path)
    pdf.close()


def test_pipeline_cleanup_removes_every_known_artifact(tmp_path: Path) -> None:
    output = tmp_path / "stale-output"
    output.mkdir()
    for name in ARTIFACT_NAMES:
        (output / name).write_bytes(b"stale")

    _clear_downstream_artifacts(output)

    assert list(output.iterdir()) == []


def test_pipeline_docx_writes_all_delivery_artifacts(tmp_path: Path) -> None:
    source = tmp_path / "input.docx"
    output = tmp_path / "e2e-docx"
    _create_docx(source)

    code = run_pipeline_main([str(source), "--output", str(output)])

    assert code == 0
    assert {path.name for path in output.iterdir()} == ARTIFACT_NAMES
    report = json.loads((output / "qa_report.json").read_text(encoding="utf-8"))
    assert report["overall_status"] == "PASS_WITH_REVIEW"
    assert "READY_FOR_SUBMISSION" not in json.dumps(report, ensure_ascii=False)
    metadata = json.loads((output / "metadata.json").read_text(encoding="utf-8"))
    assert set(metadata["artifacts"]) == ARTIFACT_NAMES


def test_pipeline_text_pdf_writes_all_delivery_artifacts(tmp_path: Path) -> None:
    source = tmp_path / "input.pdf"
    output = tmp_path / "e2e-pdf"
    _create_text_pdf(source)

    code = run_pipeline_main([str(source), "--output", str(output)])

    assert code == 0
    assert {path.name for path in output.iterdir()} == ARTIFACT_NAMES
    report = json.loads((output / "qa_report.json").read_text(encoding="utf-8"))
    assert report["overall_status"] == "PASS_WITH_REVIEW"


def test_pipeline_stops_on_ocr_required_without_fake_delivery(tmp_path: Path) -> None:
    source = tmp_path / "scan.pdf"
    output = tmp_path / "e2e-ocr"
    _create_scan_pdf(source)

    code = run_pipeline_main([str(source), "--output", str(output)])

    assert code == 4
    assert (output / "normalized_document.json").is_file()
    assert (output / "document.lines.txt").is_file()
    assert not (output / "project_facts.json").exists()
    assert not (output / "投标项目复核表.xlsx").exists()
    assert not (output / "基础投标文件.docx").exists()
    assert not (output / "qa_report.json").exists()


def test_pipeline_corrupt_input_returns_nonzero_without_fake_delivery(tmp_path: Path) -> None:
    source = tmp_path / "broken.docx"
    output = tmp_path / "e2e-broken"
    source.write_bytes(b"not a valid docx")

    code = run_pipeline_main([str(source), "--output", str(output)])

    assert code == 3
    assert (output / "normalized_document.json").is_file()
    assert not (output / "project_facts.json").exists()
    assert not (output / "qa_report.json").exists()
