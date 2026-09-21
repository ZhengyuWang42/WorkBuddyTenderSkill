from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document

try:
    import pymupdf as fitz
except ImportError:  # pragma: no cover - compatibility with older PyMuPDF.
    import fitz  # type: ignore[no-redef]

from scripts.extract_document import main as extract_document_main
from tender_basic.document_models import (
    DocumentStatus,
    ParagraphElement,
    PdfBlockElement,
    PdfTableElement,
    TableElement,
)
from tender_basic.document_parser import (
    EXIT_OCR_REQUIRED,
    EXIT_UNSUPPORTED_FORMAT,
    parse_document,
    write_normalized_outputs,
)
from tender_basic.models import PdfTableLocator, SourceType


def create_ordered_docx(path: Path) -> None:
    document = Document()
    document.add_paragraph("第一章 招标公告")
    table = document.add_table(rows=3, cols=2)
    values = [
        ("项目名称", "测试智慧平台建设项目"),
        ("项目编号", "PRJ-001"),
        ("招标编号", "BID-002"),
    ]
    for row, values_for_row in zip(table.rows, values):
        row.cells[0].text = values_for_row[0]
        row.cells[1].text = values_for_row[1]
    document.add_paragraph("正文段落：投标人须知前附表")
    document.save(path)


def create_multiline_cell_docx(path: Path) -> None:
    document = Document()
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.paragraphs[0].text = "项目名称"
    cell.add_paragraph("测试项目")
    document.save(path)


def create_text_pdf(path: Path) -> None:
    pdf = fitz.open()
    first_page = pdf.new_page()
    first_page.insert_text((72, 72), "Test tender notice\nProject No: ABC-001")
    second_page = pdf.new_page()
    second_page.insert_text((72, 72), "Bidder instructions\nMax price: 1000000 CNY")
    pdf.save(path)
    pdf.close()


def _draw_vector_table(page, x_positions: list[float], y_positions: list[float], values: list[list[str]]) -> None:
    for x in x_positions:
        page.draw_line(
            (x, y_positions[0]),
            (x, y_positions[-1]),
            color=(0, 0, 0),
            width=1,
        )
    for y in y_positions:
        page.draw_line(
            (x_positions[0], y),
            (x_positions[-1], y),
            color=(0, 0, 0),
            width=1,
        )
    for row_index, row in enumerate(values):
        for column_index, value in enumerate(row):
            if not value:
                continue
            page.insert_text(
                (
                    x_positions[column_index] + 5,
                    y_positions[row_index] + 30,
                ),
                value,
                fontname="china-s",
                fontsize=10,
            )


def create_vector_table_pdf(path: Path) -> None:
    pdf = fitz.open()
    first_page = pdf.new_page(width=500, height=300)
    _draw_vector_table(
        first_page,
        [40, 240, 460],
        [40, 90, 140],
        [["项目名称", "测试项目"], ["项目编号", "PRJ-001"]],
    )
    second_page = pdf.new_page(width=500, height=300)
    _draw_vector_table(
        second_page,
        [40, 140, 240, 340, 460],
        [40, 90],
        [["项目名称", "测试项目", "项目编号", "PRJ-001"]],
    )
    pdf.save(path)
    pdf.close()


def create_scan_like_pdf(path: Path) -> None:
    pdf = fitz.open()
    pdf.new_page()
    pdf.new_page()
    pdf.save(path)
    pdf.close()


def test_normal_docx_is_parsed_with_paragraphs_and_status(tmp_path: Path):
    source = tmp_path / "sample.docx"
    create_ordered_docx(source)

    document = parse_document(source)

    assert document.status is DocumentStatus.PARSED
    assert document.source_type is SourceType.DOCX
    assert any(
        isinstance(element, ParagraphElement)
        and "第一章 招标公告" in element.paragraph.text
        for element in document.elements
    )


def test_docx_paragraph_table_order_is_preserved(tmp_path: Path):
    source = tmp_path / "ordered.docx"
    create_ordered_docx(source)

    document = parse_document(source)
    types = [element.type for element in document.elements]

    assert types == ["paragraph", "table", "paragraph"]
    assert isinstance(document.elements[1], TableElement)
    assert document.elements[1].table.rows[1].cells[0].text == "项目编号"


def test_docx_table_locators_keep_table_row_cell(tmp_path: Path):
    source = tmp_path / "table.docx"
    create_ordered_docx(source)

    document = parse_document(source)
    table_element = next(
        element for element in document.elements if isinstance(element, TableElement)
    )
    cell = table_element.table.rows[2].cells[1]

    assert cell.locator.table_index == 0
    assert cell.locator.row_index == 2
    assert cell.locator.cell_index == 1
    assert "page" not in cell.locator.model_dump(mode="json")


def test_docx_cell_newlines_are_preserved(tmp_path: Path):
    source = tmp_path / "multiline.docx"
    create_multiline_cell_docx(source)

    document = parse_document(source)
    table = next(element for element in document.elements if isinstance(element, TableElement))

    assert table.table.rows[0].cells[0].text == "项目名称\n测试项目"


def test_pdf_pages_are_one_based_and_have_text(tmp_path: Path):
    source = tmp_path / "sample.pdf"
    create_text_pdf(source)

    document = parse_document(source)

    assert document.status is DocumentStatus.PARSED
    assert document.source_type is SourceType.PDF
    assert [page.page_number for page in document.pages] == [1, 2]
    assert "Test tender notice" in document.pages[0].blocks[0].text
    assert "Max price" in document.pages[1].blocks[0].text


def test_pdf_tables_retain_two_and_four_column_cells_and_locators(tmp_path: Path):
    source = tmp_path / "tables.pdf"
    output = tmp_path / "normalized-tables"
    create_vector_table_pdf(source)

    document = parse_document(source)

    assert document.status is DocumentStatus.PARSED
    assert len(document.tables) == 2
    first = document.tables[0]
    assert first.page == 1
    assert [cell.text for cell in first.rows[0].cells] == ["项目名称", "测试项目"]
    assert [cell.text for cell in first.rows[1].cells] == ["项目编号", "PRJ-001"]
    assert isinstance(first.rows[0].cells[0].locator, PdfTableLocator)
    assert first.rows[0].cells[0].locator.model_dump(mode="json") == {
        "locator_type": "pdf_table_cell",
        "page": 1,
        "table_index": 0,
        "row_index": 0,
        "column_index": 0,
    }

    second = document.tables[1]
    assert second.page == 2
    assert [cell.text for cell in second.rows[0].cells] == [
        "项目名称",
        "测试项目",
        "项目编号",
        "PRJ-001",
    ]
    assert any(isinstance(element, PdfTableElement) for element in document.elements)

    _json_path, lines_path = write_normalized_outputs(document, output)
    lines = lines_path.read_text(encoding="utf-8")
    assert "[PDF:T:1:T:0:R:0:C:0] 项目名称" in lines
    assert "[PDF:T:2:T:0:R:0:C:3] PRJ-001" in lines


def test_pdf_second_page_and_bbox_are_serializable(tmp_path: Path):
    source = tmp_path / "sample.pdf"
    create_text_pdf(source)

    document = parse_document(source)
    second_page_element = next(
        element
        for element in document.elements
        if isinstance(element, PdfBlockElement) and element.page_number == 2
    )
    payload = document.model_dump(mode="json")

    assert second_page_element.block.locator.page == 2
    assert len(second_page_element.block.bbox) == 4
    json.dumps(payload, ensure_ascii=False)


def test_pdf_lines_locator_maps_back_to_json(tmp_path: Path):
    source = tmp_path / "sample.pdf"
    output = tmp_path / "normalized"
    create_text_pdf(source)

    document = parse_document(source)
    _, lines_path = write_normalized_outputs(document, output)
    lines = lines_path.read_text(encoding="utf-8")

    assert "[PDF:P:1:B:0]" in lines
    assert "[PDF:P:2:B:0]" in lines
    matching_block = document.pages[1].blocks[0]
    assert matching_block.locator.page == 2
    assert matching_block.locator.block_index == 0


def test_docx_lines_cell_locator_maps_back_to_json(tmp_path: Path):
    source = tmp_path / "sample.docx"
    output = tmp_path / "normalized"
    create_ordered_docx(source)

    document = parse_document(source)
    _, lines_path = write_normalized_outputs(document, output)
    lines = lines_path.read_text(encoding="utf-8")

    assert "[DOCX:T:0:R:2:C:1] BID-002" in lines
    table_element = next(element for element in document.elements if isinstance(element, TableElement))
    assert table_element.table.rows[2].cells[1].locator.table_index == 0


def test_scan_like_pdf_returns_ocr_required(tmp_path: Path):
    source = tmp_path / "scan.pdf"
    output = tmp_path / "ocr-required"
    create_scan_like_pdf(source)

    document = parse_document(source)
    code = extract_document_main([str(source), "--output", str(output)])

    assert document.status is DocumentStatus.OCR_REQUIRED
    assert document.warnings == [
        "PDF appears image-based or has insufficient extractable text."
    ]
    assert code == EXIT_OCR_REQUIRED
    assert (output / "normalized_document.json").exists()
    assert (output / "document.lines.txt").exists()


def test_unsupported_extension_returns_status_and_exit_code(tmp_path: Path):
    source = tmp_path / "input.txt"
    output = tmp_path / "unsupported"
    source.write_text("not a tender document", encoding="utf-8")

    document = parse_document(source)
    code = extract_document_main([str(source), "--output", str(output)])

    assert document.status is DocumentStatus.UNSUPPORTED_FORMAT
    assert code == EXIT_UNSUPPORTED_FORMAT
    assert document.source_type is None


@pytest.mark.parametrize("suffix", [".pdf", ".docx"])
def test_corrupt_supported_file_returns_parse_error_without_traceback(
    tmp_path: Path, capsys: pytest.CaptureFixture[str], suffix: str
):
    source = tmp_path / f"broken{suffix}"
    output = tmp_path / "broken-output"
    source.write_bytes(b"not a valid document")

    code = extract_document_main([str(source), "--output", str(output)])
    captured = capsys.readouterr()

    assert code == 3
    assert "Status: PARSE_ERROR" in captured.out + captured.err
    assert "Traceback" not in captured.out + captured.err


def test_unicode_amount_and_identifier_text_survives(tmp_path: Path):
    source = tmp_path / "unicode.docx"
    document = Document()
    document.add_paragraph("项目名称：测试智慧平台建设项目")
    document.add_paragraph("最高限价：￥1,000,000.00元")
    document.add_paragraph("项目编号：项目-2026-甲-001")
    document.save(source)

    normalized = parse_document(source)
    _, lines_path = write_normalized_outputs(normalized, tmp_path / "unicode-output")
    payload = json.loads(
        (tmp_path / "unicode-output" / "normalized_document.json").read_text(
            encoding="utf-8"
        )
    )
    lines = lines_path.read_text(encoding="utf-8")

    assert "测试智慧平台建设项目" in lines
    assert "￥1,000,000.00元" in lines
    assert "项目-2026-甲-001" in json.dumps(payload, ensure_ascii=False)
