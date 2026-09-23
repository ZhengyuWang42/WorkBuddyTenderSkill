"""Round 5.6 targeted regressions: source visual typography, fill patterns,
source page geometry and blank-kind fidelity in the production builder."""

from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace

import pytest

from tender_basic.source_fill_patterns import (
    APPEND_TO_LABEL,
    COMPOSITE_FIELDS,
    FILL_BLANK,
    LITERAL_UNDERSCORE,
    PRESERVE_EMPTY_BLANK,
    REPLACE_BLANK_KEEP_PARENS,
    TAB_LEADER,
    VECTOR_RULE,
    WHITESPACE_GAP,
    SourceFillPattern,
    blank_kind_for,
    pattern_from_blank,
    pattern_index,
    replacement_mode_for,
)
from tender_basic.source_page_geometry import section_geometry
from tender_basic.source_visual_typography import (
    BOLD_RATIO,
    bold_overrides,
    cjk_prefix,
    infer_visual_typography,
)

ROOT = Path(__file__).resolve().parents[1]
CASE_001_RUN = ROOT / "acceptance/workspace/case_001/run_5_6_visual_fidelity"
CASE_003_RUN = ROOT / "acceptance/workspace/case_003/run_5_6_visual_fidelity"
CASE_001_PDF = ROOT / "acceptance/private/7.28引江济淮郸城配套项目一体化泵站询比文件.pdf"


def _blank(representation: str, *, x0=100.0, x1=200.0, placeholder="", slot="项目名称"):
    return SimpleNamespace(
        representation_kind=representation,
        raw_placeholder=placeholder,
        semantic_slot=slot,
        width_pt=x1 - x0,
        source_x0=x0,
        source_x1=x1,
        source_locator=SimpleNamespace(page=40),
    )


def _layout(elements, *, content_top=79.7):
    return SimpleNamespace(elements=elements, content_box=(0.0, content_top, 595.3, 800.0))


def _page(*, tables=(), number=40):
    return SimpleNamespace(width=595.3, height=841.9, page=number, tables=list(tables))


def test_rendered_weight_decides_bold_and_regular_roles() -> None:
    """A measured role is bold when its rendered ink clearly exceeds the body."""

    profile = {"roles": {"TOC_TITLE": {"bold": True}, "BODY": {"bold": False}, "TOC_1": {"bold": None}}}
    overrides = bold_overrides(profile)
    assert overrides == {"TOC_TITLE": True, "BODY": False}
    assert BOLD_RATIO > 1.0


def test_visual_typography_is_unavailable_without_a_rendered_source(tmp_path: Path) -> None:
    """No source render means no invented visual evidence."""

    result = infer_visual_typography(None, SimpleNamespace(source_pages=[]))
    assert result["status"] == "UNAVAILABLE"
    assert result["roles"] == {}
    assert "no source document path" in result["reason"]


@pytest.mark.skipif(not CASE_001_PDF.is_file(), reason="case001 source PDF is not available")
def test_case001_source_roles_are_measured_from_rendered_glyphs() -> None:
    from tender_basic.document_parser import parse_document
    from tender_basic.format_extractor import extract_bid_format
    from tender_basic.source_format import build_source_format_model

    document = parse_document(CASE_001_PDF)
    template = extract_bid_format(document)
    model = build_source_format_model(document, template)
    result = infer_visual_typography(CASE_001_PDF, model)
    assert result["status"] == "MEASURED"
    roles = result["roles"]
    assert roles["TOC_TITLE"]["bold"] is True
    assert roles["HEADING_1"]["bold"] is True
    assert roles["HEADING_2"]["bold"] is True
    assert roles["BODY"]["bold"] is False
    assert roles["TOC_1"]["bold"] is False
    assert roles["HEADING_1"]["ratio"] > roles["BODY"]["ratio"]


# ``source_page_geometry`` is retained as the *legacy* per-page derivation and is
# used only as the root-cause witness in scripts/v1_page_frame_audit.py.  The
# production builder derives its section margins from ``source_page_frame``
# instead, which is what the Stage B tests below cover.
def test_source_page_geometry_derives_margins_from_the_source_body() -> None:
    body = SimpleNamespace(bbox=(70.8, 100.0, 524.4, 120.0), kind="PARAGRAPH")
    geometry = section_geometry(_layout([body]), _page())
    assert geometry.left_margin == pytest.approx(70.8, abs=0.1)
    assert geometry.right_margin == pytest.approx(70.9, abs=0.1)
    assert geometry.usable_text_width == pytest.approx(453.6, abs=0.2)
    assert geometry.top_margin == pytest.approx(79.7, abs=0.1)


def test_source_page_geometry_ignores_running_header_and_footer() -> None:
    body = SimpleNamespace(bbox=(70.8, 100.0, 524.4, 120.0), kind="PARAGRAPH")
    header = SimpleNamespace(bbox=(70.8, 52.0, 384.0, 60.0), kind="PARAGRAPH")
    footer = SimpleNamespace(bbox=(60.0, 800.0, 90.0, 812.0), kind="PAGE_NUMBER")
    with_chrome = section_geometry(_layout([header, body, footer]), _page())
    without_chrome = section_geometry(_layout([body]), _page())
    assert with_chrome.left_margin == without_chrome.left_margin
    assert with_chrome.right_margin == without_chrome.right_margin
    assert with_chrome.usable_text_width == without_chrome.usable_text_width
    assert with_chrome.top_margin == without_chrome.top_margin


# --------------------------------------------------------------------------- #
# Stage B: the section frame is stable, and no single page's content can move it
# --------------------------------------------------------------------------- #

PAGE_W = 595.3
PAGE_H = 841.9


def _span(text, x0, y0, x1, *, height=12.3):
    return SimpleNamespace(
        text=text, bbox=(x0, y0, x1, y0 + height), font_size=11.5, font_name="SimSun"
    )


def _source_page(number, paragraphs):
    """A source page carrying lines at explicit x extents, in source order."""

    return SimpleNamespace(
        page=number,
        width=PAGE_W,
        height=PAGE_H,
        tables=[],
        paragraphs=[
            SimpleNamespace(
                text="".join(span.text for span in spans),
                bbox=(min(s.bbox[0] for s in spans), min(s.bbox[1] for s in spans),
                      max(s.bbox[2] for s in spans), max(s.bbox[3] for s in spans)),
                lines=[SimpleNamespace(text=span.text, bbox=span.bbox,
                                       spans=[span], font_size=span.font_size)
                       for span in spans],
            )
            for spans in paragraphs
        ],
    )


def _elements(extents):
    return [
        SimpleNamespace(bbox=(x0, 100.0 + 20.0 * index, x1, 112.0 + 20.0 * index),
                        kind="PARAGRAPH")
        for index, (x0, x1) in enumerate(extents)
    ]


def _frame_group():
    """Four page flavours of one source section, plus deliberate outliers.

    ``dense`` is an ordinary body page, ``sparse`` carries only a short centred
    title, ``wide_a``/``wide_b`` carry a table that overhangs the body frame on
    both sides, and ``form`` carries a form row that starts at the body edge and
    ends at the body edge.  Every one of them belongs to the *same* source
    section, so all five must resolve to the same frame.
    """

    body_left, body_right = 70.8, 524.4
    wide_table = (65.33, 541.77)

    def page(number, extents, *, table=None):
        source = _source_page(number, [[_span("正文", x0, 100.0, x1)] for x0, x1 in extents])
        if table is not None:
            source.tables = [
                SimpleNamespace(bbox=(table[0], 200.0, table[1], 600.0),
                                column_widths=[table[1] - table[0]], row_heights=[])
            ]
        layout = SimpleNamespace(elements=_elements(extents), horizontal_rules=[])
        return source, layout

    dense = page(1, [(body_left, body_right)] * 4)
    sparse = page(2, [(245.0, 350.3)])
    wide_a = page(3, [(body_left, 300.0)] * 3, table=wide_table)
    wide_b = page(4, [(body_left, 300.0)] * 3, table=wide_table)
    form = page(5, [(body_left, 250.0)] * 2 + [(body_left, body_right)])
    return [dense, sparse, wide_a, wide_b, form]


def test_every_page_flavour_of_one_source_section_resolves_to_one_frame() -> None:
    from tender_basic.source_page_frame import derive_source_section_page_frames

    pages = [source for source, _layout in _frame_group()]
    layouts = [layout for _source, layout in _frame_group()]
    frames = derive_source_section_page_frames(pages, layouts)

    assert len(frames) == 5
    assert len({frame.frame_id for frame in frames.values()}) == 1
    lefts = {round(frame.left_margin, 2) for frame in frames.values()}
    rights = {round(frame.right_margin, 2) for frame in frames.values()}
    usable = {round(frame.usable_text_width, 2) for frame in frames.values()}
    assert lefts == {70.8}, "a sparse or wide-table page must not move the body frame"
    assert rights == {round(PAGE_W - 524.4, 2)}
    assert usable == {round(524.4 - 70.8, 2)}
    for frame in frames.values():
        assert frame.left_anchor.kind == "text_left"
        assert frame.left_anchor.support == 4
        assert frame.right_anchor.kind == "text_right"
        assert frame.right_anchor.support == 2


def test_sparse_title_only_page_does_not_change_the_section_frame() -> None:
    from tender_basic.source_page_frame import derive_source_section_page_frames

    group = _frame_group()
    pages = [source for source, _layout in group]
    layouts = [layout for _source, layout in group]
    with_sparse = derive_source_section_page_frames(pages, layouts)[2]
    keep = [index for index, page in enumerate(pages) if page.page != 2]
    without_sparse = derive_source_section_page_frames(
        [pages[index] for index in keep], [layouts[index] for index in keep]
    )[1]
    assert with_sparse.frame_id == without_sparse.frame_id
    assert with_sparse.left_margin == without_sparse.left_margin
    assert with_sparse.right_margin == without_sparse.right_margin
    assert with_sparse.usable_text_width == without_sparse.usable_text_width


def test_wide_table_page_cannot_contract_the_section_frame() -> None:
    from tender_basic.source_page_frame import derive_source_section_page_frames

    pages = [source for source, _layout in _frame_group()]
    layouts = [layout for _source, layout in _frame_group()]
    frames = derive_source_section_page_frames(pages, layouts)
    for number in (3, 4):
        frame = frames[number]
        assert frame.body_x0 == pytest.approx(70.8, abs=0.01)
        assert frame.body_x1 == pytest.approx(524.4, abs=0.01)
        # The overhanging table is evidence, but never the body frame.
        assert frame.left_anchor.x > 65.33
        assert frame.right_anchor.x < 541.77


def test_single_visual_row_element_keeps_one_emission_row() -> None:
    """One physical source row drawn as several PDF text objects stays one row."""

    from tender_basic.page_layout import join_single_visual_row

    same_row = [
        SimpleNamespace(bbox=(72.12, 296.82, 84.12, 309.1), text="系"),
        SimpleNamespace(bbox=(156.48, 296.82, 230.04, 309.1), text="(供应商名称)"),
        SimpleNamespace(bbox=(324.72, 296.82, 410.52, 309.1), text="的法定代表人。"),
    ]
    assert join_single_visual_row(
        "系\n(供应商名称)\n的法定代表人。", same_row
    ) == "系(供应商名称)的法定代表人。"
    # A genuinely multi-row element keeps every source line break.
    multi_row = [
        SimpleNamespace(bbox=(70.8, 100.0, 300.0, 112.0), text="第一行"),
        SimpleNamespace(bbox=(70.8, 126.0, 300.0, 138.0), text="第二行"),
    ]
    assert join_single_visual_row("第一行\n第二行", multi_row) == "第一行\n第二行"


def test_table_width_keeps_the_source_width_when_only_the_frame_is_narrower() -> None:
    from tender_basic.word_safe_source_builder import plan_table_width

    # Source page 9's table: 65.33..541.77 inside a frame of 70.8..524.4.
    total, available, clamped = plan_table_width(
        541.77 - 65.33, table_left=65.33, usable=453.6, page_width=595.3
    )
    assert clamped is False
    assert total == pytest.approx(476.44, abs=0.01)
    assert 65.33 + total == pytest.approx(541.77, abs=0.01)
    # A table that genuinely does not fit the physical page is still bounded.
    total, _available, clamped = plan_table_width(
        700.0, table_left=32.0, usable=453.6, page_width=595.3
    )
    assert clamped is True
    assert total < 700.0
    assert 32.0 + total <= 595.3


def test_blank_kind_is_derived_from_the_source_representation() -> None:
    assert blank_kind_for("VECTOR_LINE") == VECTOR_RULE
    assert blank_kind_for("UNDERLINED_WHITESPACE") == "UNDERLINED_WHITESPACE"
    assert blank_kind_for("TAB_LEADER") == TAB_LEADER
    assert blank_kind_for("LITERAL_UNDERSCORES") == LITERAL_UNDERSCORE
    assert blank_kind_for("SOURCE_WHITESPACE_GAP") == WHITESPACE_GAP
    # A whitespace gap that still carries underscore glyphs is a literal blank.
    assert blank_kind_for("SOURCE_WHITESPACE_GAP", "___") == LITERAL_UNDERSCORE


def test_fill_pattern_keeps_the_source_line_for_a_drawn_rule() -> None:
    pattern = pattern_from_blank(
        _blank("VECTOR_LINE", x0=205.6, x1=389.6), semantic_field="供应商", has_value=True
    )
    assert pattern.blank_kind == VECTOR_RULE
    assert pattern.replacement_mode == FILL_BLANK
    assert pattern.value_is_underlined is True
    assert pattern.source_blank_width == pytest.approx(184.0, abs=0.1)
    assert pattern.style_payload()["underline"] is True


def test_empty_source_blank_stays_fillable_without_inventing_a_value() -> None:
    pattern = pattern_from_blank(_blank("VECTOR_LINE"), semantic_field="供应商", has_value=False)
    assert pattern.replacement_mode == PRESERVE_EMPTY_BLANK
    assert pattern.value_is_underlined is True


def test_parenthetical_placeholder_keeps_its_shape_and_removes_the_value_underline() -> None:
    pattern = pattern_from_blank(
        _blank("SOURCE_WHITESPACE_GAP", placeholder="（项目名称、标段）"),
        semantic_field="项目名称",
        has_value=True,
    )
    assert pattern.preserve_parentheses is True
    assert pattern.preserve_placeholder is True
    assert pattern.replacement_mode == REPLACE_BLANK_KEEP_PARENS
    assert pattern.value_is_underlined is False


def test_composite_and_label_append_modes_are_distinct() -> None:
    assert (
        replacement_mode_for(
            blank_kind=VECTOR_RULE,
            has_prefix_label=True,
            preserve_parentheses=False,
            has_value=True,
            composite_fields=["project_name", "lot_name"],
        )
        == COMPOSITE_FIELDS
    )
    assert (
        replacement_mode_for(
            blank_kind=WHITESPACE_GAP,
            has_prefix_label=True,
            preserve_parentheses=False,
            has_value=True,
            composite_fields=["project_name"],
        )
        == APPEND_TO_LABEL
    )


def test_pattern_index_prefers_the_most_confident_source_pattern() -> None:
    weak = SourceFillPattern(semantic_field="供应商", blank_kind=WHITESPACE_GAP, confidence=0.2)
    strong = SourceFillPattern(
        semantic_field="供应商", blank_kind=VECTOR_RULE, confidence=0.9, source_blank_width=184.0
    )
    index = pattern_index([weak, strong])
    assert index["供应商"] is strong


@pytest.mark.skipif(not CASE_001_RUN.is_dir(), reason="Round 5.6 case001 run is not available")
def test_production_builder_consumes_the_source_visual_models() -> None:
    report = json.loads((CASE_001_RUN / "generation_report.json").read_text(encoding="utf-8"))
    typography = report["source_visual_typography"]
    assert typography["status"] == "MEASURED"
    assert typography["roles"]["HEADING_1"]["bold"] is True
    assert report["source_fill_patterns"], "the production builder must publish fill patterns"
    assert report["source_fill_pattern_index"], "the patterns must be indexed by source field"
    assert report["fill_pattern_usage"]["filled_slots"] >= 1
    assert report["section_geometry"], "the production builder must publish section geometry"
    assert report["blank_representation_usage"]["solid_rule_blanks"] >= 1


@pytest.mark.skipif(not CASE_001_RUN.is_dir(), reason="Round 5.6 case001 run is not available")
def test_generated_cover_blanks_are_continuous_underlined_runs() -> None:
    from docx import Document

    report = json.loads((CASE_001_RUN / "generation_report.json").read_text(encoding="utf-8"))
    usage = report["blank_representation_usage"]
    assert usage.get("solid_rule_blanks", 0) > 0
    document = Document(str(CASE_001_RUN / "基础投标文件.docx"))
    underlined = [
        run
        for paragraph in document.paragraphs
        for run in paragraph.runs
        if run.underline and "\u2007" in (run.text or "")
    ]
    assert underlined, "source drawn rules must become underlined Word runs"
    literal = [
        run.text
        for paragraph in document.paragraphs
        for run in paragraph.runs
        if set((run.text or "").strip()) == {"_"}
    ]
    assert all(len(text) < 40 for text in literal), "no invented underscore placeholder runs"


@pytest.mark.skipif(not CASE_003_RUN.is_dir(), reason="Round 5.6 case003 run is not available")
def test_case003_source_facts_are_unchanged() -> None:
    facts = json.loads((CASE_003_RUN / "project_facts.json").read_text(encoding="utf-8"))
    fields = facts["fields"]
    assert fields["lot_name"]["status"] == "RESOLVED"
    assert fields["lot_name"]["resolved_value"] == "三标段"
    assert fields["project_name"]["resolved_value"] == "肇源县城市供水管网漏损治理项目"


def test_cjk_sampling_window_skips_latin_and_spaces() -> None:
    assert cjk_prefix("  我方已充分研究了(项目名称)") == "我方已充分研究了"
    assert cjk_prefix("ABC") == ""
    assert len(cjk_prefix("河南省水利第二工程局集团有限公司引江济淮")) == 8
