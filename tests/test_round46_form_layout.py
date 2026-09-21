"""Synthetic Round 4.6 regressions for forms, lists, and table artifacts."""

from __future__ import annotations

from docx import Document

from tender_basic.document_models import PdfTableCell, PdfTextLine, PdfTextSpan
from tender_basic.models import PdfLocator, PdfTableLocator, SourceType
from tender_basic.page_layout import (
    FormBlock,
    build_page_layout,
    inline_blank_token,
    restore_inline_rule_blanks,
)
from tender_basic.source_format import (
    SourceCell,
    SourceFillSlot,
    SourceFormatTemplate,
    SourceLine,
    SourcePage,
    SourcePageElement,
    SourceParagraph,
    SourceRow,
    SourceRun,
    SourceTable,
    _clean_table_cell_text,
)
from tender_basic.word_safe_source_builder import build_source_format_docx
from tender_basic.layout_qa import docx_layout_counts

from test_review_builder import make_project_facts


def _paragraph(
    text: str,
    *,
    page: int = 1,
    block: int = 0,
    bbox=(40.0, 40.0, 520.0, 56.0),
    lines=None,
    runs=None,
):
    if runs is None:
        runs = [SourceRun(text=text, bbox=bbox, font_name="SimSun", font_size=10.5)]
    if lines is None:
        lines = [PdfTextLine(
            text=text,
            bbox=bbox,
            spans=[PdfTextSpan(text=text, bbox=bbox, font_name="SimSun", font_size=10.5)],
        )]
    return SourceParagraph(
        paragraph_id=f"round46-{page}-{block}",
        page=page,
        bbox=bbox,
        text=text,
        lines=lines,
        runs=runs,
        locator=PdfLocator(page=page, block_index=block),
    )


def _template(paragraphs, *, lines=(), tables=()):
    elements = [
        SourcePageElement(type="paragraph", index=i, bbox=paragraph.bbox)
        for i, paragraph in enumerate(paragraphs)
    ]
    offset = len(elements)
    elements.extend(
        SourcePageElement(type="table", index=i, bbox=table.bbox)
        for i, table in enumerate(tables)
    )
    page = SourcePage(
        page=1,
        width=595.0,
        height=842.0,
        paragraphs=list(paragraphs),
        tables=list(tables),
        elements=elements,
        lines=list(lines),
    )
    return SourceFormatTemplate(
        source_heading="Round 4.6 synthetic source",
        source_locator=PdfLocator(page=1, block_index=0),
        source_type=SourceType.PDF,
        source_pages=[page],
    )


def _bottom_border_count(document):
    seen_tables=set()
    seen_cells=set()
    def count_table(table):
        table_key=table._tbl.getroottree().getpath(table._tbl)
        if table_key in seen_tables:
            return 0
        seen_tables.add(table_key)
        total=0
        for row in table.rows:
            for cell in row.cells:
                cell_key=cell._tc.getroottree().getpath(cell._tc)
                if cell_key in seen_cells:
                    continue
                seen_cells.add(cell_key)
                total += int(bool(cell._tc.xpath('./w:tcPr/w:tcBorders/w:bottom')))
                total += sum(count_table(nested) for nested in cell.tables)
        return total
    return sum(count_table(table) for table in document.tables)


def test_visible_pdf_rule_becomes_word_bottom_rule_and_plain_empty_stays_plain(tmp_path):
    ruled = _paragraph("供应商：", bbox=(72, 80, 140, 92))
    plain = _paragraph("投标人：", block=1, bbox=(72, 120, 140, 132))
    template = _template(
        [ruled, plain],
        lines=[SourceLine(bbox=(145, 97, 300, 97))],
    )
    output, report = build_source_format_docx(make_project_facts(), template, tmp_path / "blanks.docx")
    document = Document(output)
    assert len(document.tables) == 0
    assert len(document.paragraphs) == 2
    assert "_" not in "".join(paragraph.text for paragraph in document.paragraphs)
    assert "\u00a0" not in "".join(paragraph.text for paragraph in document.paragraphs)
    assert any("\u2007" in paragraph.text for paragraph in document.paragraphs)
    # Round 5.6: a source *drawn* rule must render as one continuous underlined
    # Word run.  An underscore tab leader paints separate underscore glyphs and
    # is a different blank mechanism, so it is no longer emitted for vector
    # rules (see tests/test_round56_source_visual_fidelity.py).
    assert docx_layout_counts(output)["layout_tab_count"] == 0
    underlined = [
        run
        for paragraph in document.paragraphs
        for run in paragraph.runs
        if run.underline and "\u2007" in (run.text or "")
    ]
    assert underlined, "the drawn rule must survive as an underlined run"
    assert report["visible_rule_blanks"] == 1
    assert report["plain_empty_blanks"] == 1
    assert report["nbsp_only_placeholder_count"] == 0


def test_project_number_and_signature_rows_use_editable_lines(tmp_path):
    rows = [
        _paragraph("项目编号：", bbox=(72, 80, 150, 92)),
        _paragraph("法定代表人或其委托代理人：", block=1, bbox=(72, 110, 260, 122)),
    ]
    template = _template(rows, lines=[
        SourceLine(bbox=(265, 124, 430, 124)),
    ])
    output, report = build_source_format_docx(make_project_facts(), template, tmp_path / "signature.docx")
    document = Document(output)
    assert len(document.tables) == 0
    assert len(document.paragraphs) == 2
    # Round 5.3 onward: a *measured vector rule* is a vector line, not literal
    # source underscore text. V0.9 renders those accepted blanks through the
    # Word-native positioned tab path, so the assertion is that the row survives
    # as an editable blank with its own recorded representation - never that a
    # particular filler glyph was invented for it.
    assert "_" not in "".join(paragraph.text for paragraph in document.paragraphs)
    assert document.paragraphs[0].text.startswith("项目编号")
    assert all(
        blank["representation_kind"] in {"VECTOR_LINE", "TAB_LEADER", "SOURCE_WHITESPACE_GAP"}
        for blank in report["editable_blanks"]
    )
    assert report["invisible_required_blank_count"] == 0
    assert report["editable_blanks"]


def test_date_row_has_three_stable_editable_cells(tmp_path):
    paragraph = _paragraph("____年____月____日", bbox=(72, 80, 280, 92))
    template = _template([paragraph])
    output, report = build_source_format_docx(make_project_facts(), template, tmp_path / "date.docx")
    document = Document(output)
    assert len(document.tables) == 0
    assert len(document.paragraphs) == 1
    assert document.paragraphs[0].text.count("____") == 3
    assert document.paragraphs[0].text.count("____") == 3
    assert docx_layout_counts(output)["layout_tab_count"] == 0
    assert report["visible_rule_blanks"] == 0
    assert all(blank["representation_kind"] == "LITERAL_UNDERSCORES" for blank in report["editable_blanks"])
    assert report["nbsp_only_placeholder_count"] == 0


def test_inline_sentence_blank_stays_in_one_paragraph_and_is_underlined():
    text = "人民币（大写）日"
    left = SourceRun(text="人民币（大写）", bbox=(40, 80, 120, 92), font_name="SimSun", font_size=10.5)
    right = SourceRun(text="日", bbox=(190, 80, 200, 92), font_name="SimSun", font_size=10.5)
    rule = SourceLine(bbox=(125, 94, 184, 94))
    restored, _, count = restore_inline_rule_blanks(text, [left, right], [rule], [])
    # The token carries the rule's own measured span, so the blank's right edge
    # comes from the source rule and never from a repeated filler glyph count.
    assert restored == "人民币（大写）" + inline_blank_token(59, 125, 184) + "日"
    assert count == 1


def test_inline_sentence_blank_renders_as_visible_underlined_glyphs(tmp_path):
    text = "人民币（大写）日"
    runs = [
        SourceRun(text="人民币（大写）", bbox=(40, 80, 120, 92), font_name="SimSun", font_size=10.5),
        SourceRun(text="日", bbox=(190, 80, 200, 92), font_name="SimSun", font_size=10.5),
    ]
    lines = [PdfTextLine(text=text, bbox=(40, 80, 200, 92), spans=[
        PdfTextSpan(text=runs[0].text, bbox=runs[0].bbox, font_name="SimSun", font_size=10.5),
        PdfTextSpan(text=runs[1].text, bbox=runs[1].bbox, font_name="SimSun", font_size=10.5),
    ])]
    template = _template([_paragraph(text, runs=runs, lines=lines)], lines=[SourceLine(bbox=(125, 94, 184, 94))])
    output, report = build_source_format_docx(make_project_facts(), template, tmp_path / "inline.docx")
    document = Document(output)
    assert len(document.paragraphs) == 1
    # The accepted V0.9 rendering of an inline source gap is the Word-native
    # positioned blank at the rule's own span; the source sentence stays one
    # paragraph and the blank stays visible and editable.
    assert "\t" in document.paragraphs[0].text
    assert document.paragraphs[0].text.startswith("人民币（大写）")
    assert report["inline_blank_count"] == 1
    assert report["nbsp_only_placeholder_count"] == 0
    assert report["positioned_blank_count"] >= 1


def test_numbered_and_nested_lists_use_true_hanging_indent(tmp_path):
    main = _paragraph("1、主条款内容很长", bbox=(96, 80, 500, 92))
    nested = _paragraph("（1）子条款", block=1, bbox=(82, 110, 360, 122))
    nested_two = _paragraph("（2）子条款", block=2, bbox=(82, 130, 360, 142))
    final = _paragraph("2、另一主条款", block=3, bbox=(96, 160, 360, 172))
    template = _template([main, nested, nested_two, final])
    output, generation = build_source_format_docx(make_project_facts(), template, tmp_path / "lists.docx")
    document = Document(output)
    list_paragraphs = [p for p in document.paragraphs if p.text.startswith(("1、", "2、", "（1）", "（2）"))]
    assert len(list_paragraphs) == 4
    assert all((p.paragraph_format.first_line_indent.pt or 0) < 0 for p in list_paragraphs)
    records = [record for record in generation["logical_paragraph_records"] if record["kind"] == "List"]
    assert {record["list_level"] for record in records} == {0, 1}
    assert len({tuple(record["list_group_key"]) for record in records}) == 3


def test_form_block_has_one_shared_grid_for_related_rows(tmp_path):
    rows = [
        _paragraph("供应商名称：____", bbox=(72, 80, 260, 92)),
        _paragraph("单位性质：____", block=1, bbox=(72, 105, 260, 117)),
        _paragraph("地址：____", block=2, bbox=(72, 130, 260, 142)),
    ]
    template = _template(rows, lines=[
        SourceLine(bbox=(160, 94, 300, 94)),
        SourceLine(bbox=(160, 119, 300, 119)),
        SourceLine(bbox=(160, 144, 300, 144)),
    ])
    output, generation = build_source_format_docx(make_project_facts(), template, tmp_path / "grid.docx")
    document = Document(output)
    assert len(document.tables) == 0
    assert len(document.paragraphs) == 3
    form_records = [record for record in generation["logical_paragraph_records"] if record["kind"] == "FormRow"]
    assert len({record["form_block_index"] for record in form_records}) == 1
    assert len({record["label_x"] for record in form_records}) == 1
    assert len({record["value_x"] for record in form_records}) == 1
    assert generation["synthetic_layout_tables"] == 0
    assert generation["tab_stop_usage"]["form_tab_leader_count"] == 0
    assert all(record["representation_kind"] == "LITERAL_UNDERSCORES" for record in generation["editable_blanks"])


def _artifact_cell(text, bbox, spans):
    return PdfTableCell(
        row_index=0,
        column_index=0,
        text=text,
        bbox=bbox,
        locator=PdfTableLocator(page=1, table_index=0, row_index=0, column_index=0),
        spans=spans,
    )


def test_artifact_span_outside_cell_bbox_is_rejected():
    cell = _artifact_cell(
        "西",
        (40, 40, 100, 70),
        [PdfTextSpan(text="西", bbox=(180, 20, 190, 30), font_size=10.5)],
    )
    report = {}
    text, runs = _clean_table_cell_text(cell, (), page_size=(595, 842), artifact_report=report)
    assert text == ""
    assert runs == []
    assert report["artifact_spans_removed"] == 1


def test_rotated_watermark_span_is_rejected_but_valid_short_business_cell_survives():
    watermark = _artifact_cell(
        "西",
        (40, 40, 120, 70),
        [PdfTextSpan(text="西", bbox=(80, 0, 95, 220), font_size=10.5, flags=16)],
    )
    valid = _artifact_cell(
        "西",
        (40, 40, 120, 70),
        [PdfTextSpan(text="西", bbox=(50, 48, 60, 60), font_size=10.5)],
    )
    rejected, _ = _clean_table_cell_text(watermark, (), page_size=(595, 842), artifact_report={})
    preserved, runs = _clean_table_cell_text(valid, (), page_size=(595, 842), artifact_report={})
    assert rejected == ""
    assert preserved == "西"
    assert len(runs) == 1
