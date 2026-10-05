"""Bid response vs contract risk: the review-stage separation contract.

The human product decision this suite pins:

    报价/商务响应 = 投标响应项
    合同条款     = 投标前风险识别项

Three rules are asserted, each from the *delivered* artifact or from the
source-derived classifier, never from a case id:

* a clause is staged by **when the obligation applies** and **who must do what**,
  so a contract-chapter *price* clause is still a bid response and a payment
  condition used as an evaluation factor is still a scoring response;
* a pure post-award contract condition is a **pre-bid notice**: it never decides
  compliance, rejection or score, and it never asks the bidder to declare
  anything in its response;
* the delivered workbook shows the two stages separately -- different module, a
  non-rejection risk label, no mandatory-sheet listing -- and the dashboard counts
  them apart.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

from tender_basic.document_models import NormalizedDocument  # noqa: E402
from tender_basic.dynamic_review import build_dynamic_review_plan  # noqa: E402
from tender_basic.models import ProjectFacts  # noqa: E402
from tender_basic.review_stage import (  # noqa: E402
    BID_RESPONSE,
    CONTRACT_RISK_NOTICE,
    INFORMATIONAL,
    SCORING_RESPONSE,
    review_stage_for,
    source_is_post_award,
    source_requires_response,
)
from tender_basic.stage_invariants import (  # noqa: E402
    contract_risk_contamination,
    check_stage_separation,
)

CASE = ROOT / "acceptance/workspace/case_001"
#: Round-10 successor: the review-stage separation round
BUILD10 = CASE / "v1_manual_fidelity_round4_date_rhythm_closure9_review_workbook10"
BUILD9R2 = CASE / "v1_manual_fidelity_round4_date_rhythm_closure9_review_workbook9r2"

CASES = {
    "case_001": ("v1_manual_fidelity_round4_date_rhythm_closure9_review_workbook10"),
    "case_002": ("v1_round4_closure9_review_workbook10"),
    "case_003": ("v1_round4_closure9_review_workbook10"),
}

pytestmark = pytest.mark.skipif(
    not (BUILD10 / "project_facts.json").is_file(),
    reason="round-10 CASE001 successor build is not present",
)


# --------------------------------------------------------------------------- #
# the classifier: stage by effect, not by chapter
# --------------------------------------------------------------------------- #


def test_pure_contract_clause_is_a_pre_bid_risk_notice() -> None:
    assert (
        review_stage_for(
            None,
            concern_id="RETENTION_MONEY_RATIO",
            contract_kind="CONTRACT_ONLY",
            source_text="剩余 5%作为质保金，质保期 12 个月，质保期满后无息付清余款",
        )
        == CONTRACT_RISK_NOTICE
    )


def test_contract_chapter_price_clause_stays_a_bid_response() -> None:
    """A cost clause inside the contract chapter still shapes the quotation."""

    assert (
        review_stage_for(
            None,
            concern_id="PRICE_INCLUDED_COST_SCOPE",
            contract_kind="MANDATORY",
            source_text="设备单价中含运输费、装卸费、安装费、损耗和税金等费用",
        )
        == BID_RESPONSE
    )


def test_contract_like_scoring_clause_is_a_scoring_response() -> None:
    assert (
        review_stage_for(
            None,
            concern_id="SCORING_PAYMENT_CONDITION",
            contract_kind="SCORING",
            source_text="付款条件评分：95% 的得 12 分",
        )
        == SCORING_RESPONSE
    )


def test_contract_risk_with_response_requirement_is_not_demoted() -> None:
    """Independent source evidence keeps a contract-kind clause on the response."""

    assert (
        review_stage_for(
            None,
            concern_id="CONTRACT_DELIVERY",
            contract_kind="CONTRACT_ONLY",
            source_text="供应商应在响应文件中载明合同交付安排，签订合同后按约定履行",
        )
        == BID_RESPONSE
    )


def test_informational_concern_is_context_only() -> None:
    assert (
        review_stage_for(None, concern_id="PRE_BID_MEETING", contract_kind="INFORMATIONAL")
        == INFORMATIONAL
    )


def test_source_stage_wording_helpers() -> None:
    assert source_is_post_award("质保期满后无息付清余款")
    assert source_is_post_award("在签订合同前提交履约保证金")
    assert not source_is_post_award("投标有效期 90 日历天")
    assert source_requires_response("供应商应在响应文件中载明该比例")
    assert source_requires_response("报价应包含运输费和安装费")
    assert not source_requires_response("采购人按合同约定支付货款")


# --------------------------------------------------------------------------- #
# the delivered plan
# --------------------------------------------------------------------------- #


def _build_dir(case: str) -> Path:
    return ROOT / "acceptance/workspace" / case / CASES[case]


@pytest.fixture(scope="module")
def case001() -> dict:
    build = _build_dir("case_001")
    document = NormalizedDocument.model_validate(
        json.loads((build / "normalized_document.json").read_text(encoding="utf-8"))
    )
    facts = ProjectFacts.model_validate(
        json.loads((build / "project_facts.json").read_text(encoding="utf-8"))
    )
    plan = build_dynamic_review_plan(document, facts)
    return {"plan": plan, "build": build, "items": list(plan.items)}


def _item(case001: dict, concern_id: str) -> dict:
    return next(
        (item for item in case001["items"] if item.concern_id == concern_id),
        None,
    )


@pytest.mark.parametrize(
    "concern_id",
    (
        "CONTRACT_RISK",
        "PERFORMANCE_BOND",
        "CONTRACT_PAYMENT",
        "RETENTION_RELEASE_PERIOD",
        "DELIVERY_ACCEPTANCE_COMPLETION",
        "CONTRACT_TERMINATION_REFUND",
        "RETENTION_MONEY_RATIO",
    ),
)
def test_case001_contract_risk_concerns_are_notices(case001: dict, concern_id: str) -> None:
    item = _item(case001, concern_id)
    assert item is not None, concern_id
    assert item.review_stage == CONTRACT_RISK_NOTICE, concern_id
    assert item.module.startswith("九、"), (concern_id, item.module)
    assert item.risk_level == "风险提示", (concern_id, item.risk_level)


@pytest.mark.parametrize(
    "concern_id",
    (
        "PRICE_CEILING",
        "PRICING_COMPLETENESS",
        "PRICE_TAX_BASIS",
        "PRICE_INCLUDED_COST_SCOPE",
        "DELIVERY_PERIOD",
        "DELIVERY_LOCATION",
        "QUALITY_TARGET",
        "PROJECT_WARRANTY",
        "BID_VALIDITY",
        "BID_BOND_EVIDENCE",
    ),
)
def test_case001_response_concerns_stay_responses(case001: dict, concern_id: str) -> None:
    item = _item(case001, concern_id)
    assert item is not None, concern_id
    assert item.review_stage == BID_RESPONSE, (concern_id, item.review_stage)
    assert not item.module.startswith("九、"), (concern_id, item.module)


@pytest.mark.parametrize(
    "concern_id",
    ("SCORING_BANK_ACCEPTANCE", "SCORING_PAYMENT_CONDITION", "SCORING_PRICE_FORMULA"),
)
def test_case001_scoring_concerns_stay_scoring(case001: dict, concern_id: str) -> None:
    item = _item(case001, concern_id)
    assert item is not None, concern_id
    assert item.review_stage == SCORING_RESPONSE, (concern_id, item.review_stage)
    assert not item.module.startswith("九、"), (concern_id, item.module)


def test_project_warranty_and_retention_are_not_merged(case001: dict) -> None:
    """24-month warranty is a response duty; 12-month retention is a contract risk."""

    warranty = _item(case001, "PROJECT_WARRANTY")
    release = _item(case001, "RETENTION_RELEASE_PERIOD")
    ratio = _item(case001, "RETENTION_MONEY_RATIO")
    assert warranty is not None and release is not None and ratio is not None
    assert warranty.review_stage == BID_RESPONSE
    assert release.review_stage == CONTRACT_RISK_NOTICE
    assert ratio.review_stage == CONTRACT_RISK_NOTICE
    assert warranty.item_id != release.item_id != ratio.item_id


def test_performance_bond_after_award_is_a_contract_risk(case001: dict) -> None:
    detail = _item(case001, "PERFORMANCE_BOND")
    assert detail is not None
    assert detail.review_stage == CONTRACT_RISK_NOTICE
    assert not detail.rejection_consequence
    assert not detail.substantive_requirement


def test_bid_bond_stays_a_bid_response(case001: dict) -> None:
    for concern_id in ("BID_BOND_AMOUNT", "BID_BOND_FORM", "BID_BOND_TRANSFER", "BID_BOND_EVIDENCE"):
        item = _item(case001, concern_id)
        assert item is not None, concern_id
        assert item.review_stage == BID_RESPONSE, (concern_id, item.review_stage)


def test_contract_risk_rows_generate_no_rejection_or_substantive_status(case001: dict) -> None:
    for item in case001["items"]:
        if item.review_stage != CONTRACT_RISK_NOTICE:
            continue
        assert not item.rejection_consequence, item.item_id
        assert not item.substantive_requirement, item.item_id
        assert not [
            value
            for value in item.mandatory_types
            if str(value).startswith("SUBSTANTIVE") or str(value) == "REJECTION"
        ], (item.item_id, item.mandatory_types)


def test_contract_risk_rows_require_no_response_declaration(case001: dict) -> None:
    for item in case001["items"]:
        if item.review_stage != CONTRACT_RISK_NOTICE:
            continue
        delivered = " ".join(
            [
                str(item.source_requirement or ""),
                str(item.verification_action or ""),
                str(item.pass_criteria or ""),
                str(item.consequence_if_failed or ""),
            ]
        )
        for cue in (
            "核对响应文件已载明",
            "确认响应文件接受",
            "响应文件中被接受",
            "与响应文件一致",
        ):
            assert cue not in delivered, (item.item_id, cue)


def test_case001_stage_separation_invariants_pass(case001: dict) -> None:
    rows = [
        {
            "item_id": item.item_id,
            "concern_id": item.concern_id,
            "stage": item.review_stage,
            "expected_stage": item.review_stage,
            "risk": item.risk_level,
            "veto": "",
            "requirement": item.source_requirement,
            "action": item.verification_action,
            "criteria": item.pass_criteria,
            "consequence": item.consequence_if_failed,
            "module": item.module,
            "rejection": item.rejection_consequence,
            "substantive": item.substantive_requirement,
            "mandatory_types": list(item.mandatory_types),
        }
        for item in case001["items"]
    ]
    separation = check_stage_separation(rows)
    assert separation["contract_risk_as_bid_blocker_count"] == 0, separation[
        "contract_risk_as_bid_blockers"
    ]
    assert separation["contract_risk_row_checked"] >= 7
    assert separation["per_stage"].get(SCORING_RESPONSE, 0) >= 3


# --------------------------------------------------------------------------- #
# the saved workbook
# --------------------------------------------------------------------------- #


@pytest.fixture(scope="module")
def saved() -> dict:
    workbook = load_workbook(BUILD10 / "投标项目复核表.xlsx", data_only=True, read_only=True)
    try:
        clause = workbook["02_关键条款"]
        headers = [str(cell.value or "").strip() for cell in next(clause.iter_rows(min_row=1, max_row=1))]
        clause_rows = {
            str(row[headers.index("requirement_id")]): {
                name: row[index] for index, name in enumerate(headers) if index < len(row)
            }
            for row in clause.iter_rows(min_row=2, values_only=True)
            if row[headers.index("requirement_id")]
        }
        mandatory = workbook["03_资格否决与强制项"]
        mandatory_headers = [
            str(cell.value or "").strip() for cell in next(mandatory.iter_rows(min_row=1, max_row=1))
        ]
        mandatory_ids = {
            str(row[mandatory_headers.index("requirement_id")])
            for row in mandatory.iter_rows(min_row=2, values_only=True)
            if row[mandatory_headers.index("requirement_id")]
        }
        dashboard = workbook["00_复核总览"]
        dashboard_rows = [
            (str(row[0] or ""), row[1]) for row in dashboard.iter_rows(values_only=True) if row
        ]
        return {
            "clause": clause_rows,
            "mandatory_ids": mandatory_ids,
            "dashboard_labels": [label for label, _value in dashboard_rows],
        }
    finally:
        workbook.close()


def test_saved_workbook_shows_contract_risk_in_its_own_module(saved: dict, case001: dict) -> None:
    for item in case001["items"]:
        if item.review_stage != CONTRACT_RISK_NOTICE:
            continue
        row = saved["clause"].get(item.item_id)
        assert row is not None, item.item_id
        assert str(row.get("类别") or "").startswith("九、"), (item.item_id, row.get("类别"))
        assert str(row.get("风险级别") or "") == "风险提示", (item.item_id, row.get("风险级别"))
        # 是否强制 answers "can this gate my bid?" -> not applicable for a notice
        assert str(row.get("是否强制") or "") == "", (item.item_id, row.get("是否强制"))


def test_saved_workbook_keeps_contract_risk_off_the_mandatory_sheet(
    saved: dict, case001: dict
) -> None:
    for item in case001["items"]:
        if item.review_stage != CONTRACT_RISK_NOTICE:
            continue
        assert item.item_id not in saved["mandatory_ids"], item.item_id


def test_saved_dashboard_counts_response_and_risk_separately(saved: dict) -> None:
    labels = saved["dashboard_labels"]
    assert any("投标响应复核项" in label for label in labels), labels
    assert any("合同风险提示项" in label for label in labels), labels


def test_saved_contract_risk_rows_carry_no_response_wording(saved: dict, case001: dict) -> None:
    for item in case001["items"]:
        if item.review_stage != CONTRACT_RISK_NOTICE:
            continue
        row = saved["clause"].get(item.item_id)
        assert row is not None, item.item_id
        delivered = " ".join(
            str(row.get(field) or "")
            for field in ("抽取结果", "复核动作", "核验标准")
        )
        assert "核对响应文件已载明" not in delivered, item.item_id
        assert "确认响应文件接受" not in delivered, item.item_id


def test_saved_response_rows_are_not_in_the_risk_module(saved: dict, case001: dict) -> None:
    for item in case001["items"]:
        if item.review_stage not in (BID_RESPONSE, SCORING_RESPONSE):
            continue
        row = saved["clause"].get(item.item_id)
        if row is None:
            continue
        assert not str(row.get("类别") or "").startswith("九、"), item.item_id


# --------------------------------------------------------------------------- #
# contamination helper
# --------------------------------------------------------------------------- #


def test_contamination_helper_flags_a_rejection_presentation() -> None:
    finding = contract_risk_contamination(
        {
            "item_id": "DR999",
            "stage": CONTRACT_RISK_NOTICE,
            "risk": "一票否决",
            "veto": "是",
            "rejection": True,
            "substantive": True,
            "mandatory_types": ["REJECTION"],
            "requirement": "核对响应文件已载明该比例",
            "action": "",
            "criteria": "",
        }
    )
    assert "risk_claims_rejection" in finding["findings"]
    assert "response_file_wording:核对响应文件已载明" in finding["findings"]
    assert "rejection_consequence" in finding["findings"]


def test_contamination_helper_accepts_a_clean_notice() -> None:
    finding = contract_risk_contamination(
        {
            "item_id": "DR998",
            "stage": CONTRACT_RISK_NOTICE,
            "risk": "风险提示",
            "veto": "",
            "rejection": False,
            "substantive": False,
            "mandatory_types": [],
            "requirement": "剩余 5%作为质保金",
            "action": "确认项目团队已知悉该条款",
            "criteria": "本项用于投标前合同风险识别",
        }
    )
    assert finding["findings"] == []


# --------------------------------------------------------------------------- #
# three-case generalization
# --------------------------------------------------------------------------- #


def test_three_case_stage_separation() -> None:
    """The stage model is generic: no case id decides a row's stage."""

    seen: set[str] = set()
    for case, build_name in CASES.items():
        build = ROOT / "acceptance/workspace" / case / build_name
        if not (build / "normalized_document.json").is_file():
            pytest.skip(f"{case} round-10 build is not present")
        document = NormalizedDocument.model_validate(
            json.loads((build / "normalized_document.json").read_text(encoding="utf-8"))
        )
        facts = ProjectFacts.model_validate(
            json.loads((build / "project_facts.json").read_text(encoding="utf-8"))
        )
        plan = build_dynamic_review_plan(document, facts)
        rows = [
            {
                "item_id": item.item_id,
                "concern_id": item.concern_id,
                "stage": item.review_stage,
                "risk": item.risk_level,
                "veto": "",
                "requirement": item.source_requirement,
                "action": item.verification_action,
                "criteria": item.pass_criteria,
                "consequence": item.consequence_if_failed,
                "rejection": item.rejection_consequence,
                "substantive": item.substantive_requirement,
                "mandatory_types": list(item.mandatory_types),
            }
            for item in plan.items
        ]
        separation = check_stage_separation(rows)
        assert separation["contract_risk_as_bid_blocker_count"] == 0, (
            case,
            separation["contract_risk_as_bid_blockers"],
        )
        for item in plan.items:
            if item.review_stage == CONTRACT_RISK_NOTICE:
                assert item.module.startswith("九、"), (case, item.item_id, item.module)
                assert item.risk_level == "风险提示", (case, item.item_id, item.risk_level)
        seen.add(case)
    assert seen == set(CASES)
