"""V1 manual-word-review follow-up: composite slot underline and source indent.

Two defect classes from the CASE001 manual Microsoft Word review, each pinned to
the production rule that closes it.  Every input here is synthetic and
source-derived: no case name, page number, rule id or resolved value literal
appears in the assertions as an authority.

* Defect A - a composite slot value must inherit the underline the source drew
  on that slot, and nothing outside the slot may be underlined with it;
* Defect B - a paragraph's Word indent must be the indent its own source rows
  show, classified per paragraph instead of assumed from a list style.
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from tender_basic.document_models import PdfTextLine, PdfTextSpan  # noqa: E402
from tender_basic.page_layout import (  # noqa: E402
    LogicalParagraph,
    build_page_layout,
    source_rule_replaces_placeholder_with_value,
    underline_owner_runs,
)
from tender_basic.source_format import SourceLine, SourcePageElement  # noqa: E402
from tender_basic.source_paragraph_indent import (  # noqa: E402
    CLASSIFICATION_FIRST_LINE,
    CLASSIFICATION_HANGING,
    CLASSIFICATION_NONE,
    classify_source_paragraph_indent,
)
from tender_basic.word_safe_source_builder import build_source_format_docx  # noqa: E402

from test_review_builder import make_project_facts  # noqa: E402
from test_round2_manual_fidelity_defects import _line_level_slot  # noqa: E402
from test_source_format_fidelity import _paragraph, _paragraph_template  # noqa: E402


# --------------------------------------------------------------------------- #
# Defect A: the slot's decoration is inherited by the value that replaces it
# --------------------------------------------------------------------------- #


def _span(text, x0, x1, y=156.8):
    return PdfTextSpan(
        text=text, bbox=(x0, y, x1, y + 12.0), font_name="SimSun", font_size=10.5
    )


def _rule(x0, x1, y=156.8):
    return SourceLine(bbox=(x0, y, x1, y), orientation="horizontal")


def _composite_line():
    """The source's own shape: prose, a parenthetical placeholder, then prose.

    The rule spans the placeholder *and* the blank on either side of it, so its
    extent is only partly covered by glyphs - which is why its occupancy label is
    a form-layout rule and not an underline.
    """

    spans = [
        _span("在", 99.6, 111.6),
        _span("（项目名称、标段）", 142.68, 251.52),
        _span("询比活动中，我公司保证做到：", 324.72, 494.03),
    ]
    rule = _rule(111.8, 318.6)
    return rule, spans


def test_a_slot_bound_rule_decorates_the_value_that_replaces_its_placeholder():
    """A rule bound to a value-replacing slot is decoration evidence.

    Its occupancy label classifies how much of the rule the source's glyphs
    cover; it says nothing about whether the rule is a line.  The binding names
    the slot, so the rule is eligible as a decoration owner even though the
    label-based fallback would skip it.
    """

    rule, spans = _composite_line()
    assert underline_owner_runs([rule], spans) == []
    owners = underline_owner_runs(
        [rule], spans, bound_decoration_rule_ids={id(rule)}
    )
    assert len(owners) == 1
    owner_span, relation, owner_rule = owners[0]
    assert owner_rule is rule
    assert owner_span.text == "（项目名称、标段）"
    # The rule is not an underline relation, so the caller's decoration-ownership
    # decision - not the label - is what authorises the inheritance.
    assert relation["relation_type"] == "FORM_LAYOUT_RULE"


def test_only_a_value_replacing_binding_authorises_the_inheritance():
    """A policy alone is not enough, and a non-replacing policy never is."""

    def entry(policy, resolved):
        return {"transformation_policy": policy, "resolved_fact_fields": resolved}

    assert source_rule_replaces_placeholder_with_value(
        entry("PLACEHOLDER_REPLACED_BY_VALUE", ["PROJECT_NAME"])
    )
    assert source_rule_replaces_placeholder_with_value(
        entry("RESOLVED_VALUE_IN_FIXED_SLOT", ["PROJECT_NAME"])
    )
    # A value that will not be written decorates nothing.
    assert not source_rule_replaces_placeholder_with_value(
        entry("PLACEHOLDER_REPLACED_BY_VALUE", [])
    )
    # A placeholder the document keeps, or a rule production could not resolve,
    # must not hand its decoration to a value that is not there.
    for policy in ("PRESERVE_SOURCE_PLACEHOLDER", "FIXED_EMPTY_SLOT", "UNRESOLVED"):
        assert not source_rule_replaces_placeholder_with_value(
            entry(policy, ["PROJECT_NAME"])
        ), policy
    assert not source_rule_replaces_placeholder_with_value(None)


def test_prose_outside_the_slot_is_never_an_underline_owner():
    """Ownership is per run, so no neighbouring prose is painted with the slot."""

    rule, spans = _composite_line()
    owner_span, _relation, _rule = underline_owner_runs(
        [rule], spans, bound_decoration_rule_ids={id(rule)}
    )[0]
    for not_owned in ("在", "询比活动中，我公司保证做到："):
        assert not_owned != owner_span.text
        other = next(span for span in spans if span.text == not_owned)
        assert float(other.bbox[2]) <= 111.8 + 1.5 or float(other.bbox[0]) >= 318.6 - 1.5


def test_the_absent_component_marker_shares_the_slot_without_becoming_a_fact():
    """Every component occupying the slot shares its one decoration intent.

    The marker names the field the source form leaves empty.  It is a
    presentation component, so it cannot become a fact or change that field's
    status - and because all three components occupy the *same* source slot, the
    slot's single underline intent covers each of them.
    """

    from tender_basic.source_fill_policy import (
        COMPONENT_FACT_VALUE,
        COMPONENT_SOURCE_FORM_NOT_APPLICABLE_MARKER,
        COMPONENT_SOURCE_TEMPLATE_LITERAL,
        SOURCE_FORM_NOT_APPLICABLE_MARKER,
        recorded_slot_presentation,
        reset_recorded_slot_presentations,
        slot_replacement,
    )

    reset_recorded_slot_presentations()
    slot = _line_level_slot()
    facts = make_project_facts()
    value, _method = slot_replacement(slot, facts)
    presentation = recorded_slot_presentation(slot)
    assert presentation is not None
    assert [item["kind"] for item in presentation["components"]] == [
        COMPONENT_FACT_VALUE,
        COMPONENT_SOURCE_TEMPLATE_LITERAL,
        COMPONENT_SOURCE_FORM_NOT_APPLICABLE_MARKER,
    ]
    # One slot, one decoration intent, three components - so the underline covers
    # the resolved fact, the source's own separator and the marker alike.
    assert len(presentation["components"]) == 3
    assert SOURCE_FORM_NOT_APPLICABLE_MARKER in value
    assert presentation["marker_components"] == ["lot_name"]
    assert presentation["fact_components"] == ["project_name"]
    assert facts.fields.lot_name.status.value == "NOT_FOUND"
    assert facts.fields.lot_name.resolved_value is None
    assert facts.fields.lot_name.candidates == []


# --------------------------------------------------------------------------- #
# Defect B: the source's own row geometry decides the paragraph's Word indent
# --------------------------------------------------------------------------- #


def test_first_row_right_of_the_wrapped_rows_is_a_first_line_indent():
    indent = classify_source_paragraph_indent([96.0, 72.0, 72.0])
    assert indent.classification == CLASSIFICATION_FIRST_LINE
    assert indent.body_left_x == 72.0
    assert indent.first_line_x == 96.0
    assert indent.continuation_x == 72.0
    assert indent.first_line_indent_pt == 24.0
    assert indent.hanging_indent_pt == 0.0
    assert indent.word_first_line_indent_pt == 24.0


def test_first_row_left_of_the_wrapped_rows_is_a_hanging_indent():
    indent = classify_source_paragraph_indent([72.0, 96.0, 96.0])
    assert indent.classification == CLASSIFICATION_HANGING
    assert indent.body_left_x == 96.0
    assert indent.hanging_indent_pt == 24.0
    # Word expresses a hanging indent as a negative first-line offset.
    assert indent.word_first_line_indent_pt == -24.0


def test_rows_in_one_column_claim_no_special_first_line_indent():
    for rows in ([72.0], [80.0, 80.0], [72.0, 73.0, 71.5]):
        indent = classify_source_paragraph_indent(rows)
        assert indent.classification == CLASSIFICATION_NONE, rows
        assert indent.word_first_line_indent_pt == 0.0, rows


def test_a_centred_line_is_not_read_as_a_first_line_indent():
    """A centred line is positioned from the container centre, not a boundary.

    Its own left origin is far right of the body boundary, which is exactly the
    shape a first-line indent has, so the classification must refuse it.
    """

    template = _page([("十四、反商业贿赂承诺书", 221.4)])
    template.source_pages[0].paragraphs[0].alignment = "center"
    items = [
        item
        for item in build_page_layout(template.source_pages[0]).elements
        if isinstance(item, LogicalParagraph)
    ]
    assert len(items) == 1
    indent = items[0].source_indent
    assert indent.classification == CLASSIFICATION_NONE
    assert indent.body_left_x == 221.4
    assert indent.word_first_line_indent_pt == 0.0
    assert "centred" in indent.evidence


# --------------------------------------------------------------------------- #
# Defect B, end to end: the emitted paragraph carries that classification
# --------------------------------------------------------------------------- #


def _page(*blocks):
    """A source page of blocks, each block a list of ``(text, x)`` source rows.

    Every row is its own source paragraph object, exactly as the PDF extractor
    delivers it; which rows belong to one reconstructed paragraph is decided by
    the production page layout from the rows' own geometry and markers.
    """

    rows = []
    top = 80.0
    for block in blocks:
        for text, x in block:
            rows.append((text, x, top))
            top += 24.0
    paragraphs = []
    for index, (text, x, y) in enumerate(rows):
        paragraph = _paragraph(
            text, block_index=index, bbox=(x, y, 520.0, y + 12.0), font_size=12
        )
        paragraph.lines = [
            PdfTextLine(
                text=text,
                bbox=(x, y, 520.0, y + 12.0),
                spans=[
                    PdfTextSpan(
                        text=text,
                        bbox=(x, y, 520.0, y + 12.0),
                        font_name="SimSun",
                        font_size=12,
                    )
                ],
            )
        ]
        paragraphs.append(paragraph)
    template = _paragraph_template(paragraphs[0])
    page = template.source_pages[0]
    page.paragraphs = paragraphs
    page.elements = [
        SourcePageElement(type="paragraph", index=index, bbox=paragraph.bbox)
        for index, paragraph in enumerate(paragraphs)
    ]
    return template


def _indents(template, tmp_path, name):
    output, report = build_source_format_docx(
        make_project_facts(), template, tmp_path / name
    )
    document = Document(output)
    records = [
        record
        for record in report["logical_paragraph_records"]
        if record.get("source_indent") is not None
    ]
    return document, records


def test_wrapped_numbered_paragraph_keeps_the_source_continuation_column(tmp_path):
    """The first row's x is not the paragraph's left edge.

    The source prints the numbered first row at 96 and returns every wrapped row
    to 72, so the delivered paragraph's body boundary is 72 and its first line is
    offset by +24.  Writing the first row's x as a whole-paragraph left indent -
    with a first-line offset of zero - moves every wrapped row onto the marker
    column and inverts the source's geometry.
    """

    document, records = _indents(
        _page(
            [
                ("1、我方确认招标文件及其所有附件中的责任和义", 96.0),
                ("务。", 72.0),
            ]
        ),
        tmp_path,
        "first-line.docx",
    )
    assert len(document.paragraphs) == 1, [p.text for p in document.paragraphs]
    formatting = document.paragraphs[0].paragraph_format
    left = formatting.left_indent.pt or 0.0
    first = formatting.first_line_indent.pt or 0.0
    assert abs(left - (72.0 - 18.0)) <= 0.05
    assert abs(first - 24.0) <= 0.05
    assert abs(left + first + 18.0 - 96.0) <= 0.05
    assert records
    classification = records[0]["source_indent"]
    assert classification["classification"] == CLASSIFICATION_FIRST_LINE
    assert classification["body_left_x"] == 72.0
    assert classification["word_first_line_indent_pt"] == 24.0


def test_a_single_row_paragraph_uses_its_container_measured_body_boundary(tmp_path):
    """A paragraph that never wraps reads the boundary its container measured.

    The page carries one paragraph that does wrap and returns to the container's
    body boundary, so the boundary is measured evidence rather than an
    assumption; the single-row paragraph right of it is therefore a first-line
    indent and not a whole-paragraph indent.
    """

    document, records = _indents(
        _page(
            [
                ("我方已充分研究了招标文件的所有规定并确认愿意", 96.0),
                ("以人民币响应报价。", 72.0),
            ],
            [("在（项目名称、标段）询比活动中，我公司保证做到：", 99.6)],
        ),
        tmp_path,
        "container.docx",
    )
    assert len(document.paragraphs) == 2
    single = document.paragraphs[1].paragraph_format
    assert abs((single.left_indent.pt or 0.0) - (72.0 - 18.0)) <= 0.05
    assert abs((single.first_line_indent.pt or 0.0) - (99.6 - 72.0)) <= 0.05
    classification = records[1]["source_indent"]
    assert classification["classification"] == CLASSIFICATION_FIRST_LINE
    assert classification["measured_continuation"] is False
    assert "container" in classification["evidence"]


def test_a_hanging_source_paragraph_is_delivered_as_a_hanging_indent(tmp_path):
    """The three indent kinds are three deliveries, not one."""

    document, records = _indents(
        _page(
            [
                ("1、第一项内容在此结束并不换", 72.0),
                ("行而是继续写完。", 96.0),
            ]
        ),
        tmp_path,
        "hanging.docx",
    )
    assert len(document.paragraphs) == 1, [p.text for p in document.paragraphs]
    formatting = document.paragraphs[0].paragraph_format
    assert abs((formatting.left_indent.pt or 0.0) - (96.0 - 18.0)) <= 0.05
    assert abs((formatting.first_line_indent.pt or 0.0) + 24.0) <= 0.05
    assert records[0]["source_indent"]["classification"] == CLASSIFICATION_HANGING


def test_no_paragraph_is_given_an_indent_its_own_rows_contradict(tmp_path):
    """Every emitted paragraph's body boundary is a boundary the source used."""

    template = _page(
        [
            ("二、杜绝任何形式的商业贿赂行为。不向国家工作人员提供礼品礼金、有", 96.0),
            ("价证券、购物券、回扣、佣金、咨询费、劳务费、资助费、宣传费、宴请。", 72.0),
        ],
        [
            ("三、若出现上述行为，我公司愿意接受按照国", 96.0),
            ("家法律法规等有关规定给予的处罚。", 72.0),
        ],
        [("四、这一项从不换行。", 72.0)],
    )
    document, records = _indents(template, tmp_path, "rows.docx")
    layout = build_page_layout(template.source_pages[0])
    page_origins = {
        round(float(line.bbox[0]), 4)
        for item in layout.elements
        if isinstance(item, LogicalParagraph)
        for line in item.source_lines
    }
    subjects = [
        item
        for item in layout.elements
        if isinstance(item, LogicalParagraph) and item.source_lines
    ]
    assert subjects
    for item in subjects:
        indent = item.source_indent
        assert indent is not None
        origins = [round(float(line.bbox[0]), 4) for line in item.source_lines]
        # The body boundary is always a boundary some source row on this page
        # actually starts at: never a first row's x standing in for a body edge.
        assert indent.body_left_x in page_origins
        if indent.classification == CLASSIFICATION_NONE:
            assert indent.word_first_line_indent_pt == 0.0
            assert indent.body_left_x == min(origins)
        assert abs(
            (indent.body_left_x + indent.word_first_line_indent_pt) - indent.first_line_x
        ) <= 1e-6
    # No paragraph is delivered with a negative left indent or an unnamed
    # first-line offset.
    for paragraph in document.paragraphs:
        formatting = paragraph.paragraph_format
        assert (formatting.left_indent.pt or 0.0) >= 0.0
        assert formatting.first_line_indent is not None
    assert records
