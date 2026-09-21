from __future__ import annotations

from docx import Document
import pytest

from tender_basic.bid_document_builder import build_bid_document
from tender_basic.document_models import (
    DocumentBlock,
    DocumentStatus,
    NormalizedDocument,
    PdfBlockElement,
    PdfTable,
    PdfTableCell,
    PdfTableElement,
    PdfTableRow,
)
from tender_basic.document_parser import parse_document
from tender_basic.format_extractor import extract_bid_format
from tender_basic.output_helpers import load_output_fields
from tender_basic.models import PdfLocator, PdfTableLocator, SourceType

from test_review_builder import make_project_facts


def _document_text(document: Document) -> str:
    paragraph_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    table_text = "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    return f"{paragraph_text}\n{table_text}"


def _synthetic_pdf_source_format() -> NormalizedDocument:
    blocks = [
        DocumentBlock(
            block_index=0,
            text="第六章 投标文件格式",
            bbox=(0.0, 0.0, 200.0, 20.0),
            locator=PdfLocator(page=1, block_index=0),
        ),
        DocumentBlock(
            block_index=1,
            text="一、投标函",
            bbox=(0.0, 30.0, 200.0, 50.0),
            locator=PdfLocator(page=1, block_index=1),
        ),
        DocumentBlock(
            block_index=2,
            text="项目名称 测试工程（PDF表格文本块）",
            bbox=(10.0, 80.0, 190.0, 100.0),
            locator=PdfLocator(page=1, block_index=2),
        ),
        DocumentBlock(
            block_index=3,
            text="第七章 评审方法",
            bbox=(0.0, 300.0, 200.0, 320.0),
            locator=PdfLocator(page=1, block_index=3),
        ),
    ]
    table_rows = []
    for row_index, values in enumerate(
        [["项目名称", ""], ["项目编号", ""], ["投标人", "________________"]]
    ):
        table_rows.append(
            PdfTableRow(
                row_index=row_index,
                cells=[
                    PdfTableCell(
                        row_index=row_index,
                        column_index=column_index,
                        text=value,
                        bbox=(0.0, 60.0 + row_index * 30, 100.0, 80.0),
                        locator=PdfTableLocator(
                            page=1,
                            table_index=0,
                            row_index=row_index,
                            column_index=column_index,
                        ),
                    )
                    for column_index, value in enumerate(values)
                ],
            )
        )
    table = PdfTable(
        page=1,
        table_index=0,
        bbox=(0.0, 60.0, 200.0, 160.0),
        rows=table_rows,
    )
    elements = [
        PdfBlockElement(page_number=1, block=blocks[0]),
        PdfBlockElement(page_number=1, block=blocks[1]),
        PdfBlockElement(page_number=1, block=blocks[2]),
        PdfTableElement(table=table),
        PdfBlockElement(page_number=1, block=blocks[3]),
    ]
    return NormalizedDocument(
        source_file="synthetic.pdf",
        source_type=SourceType.PDF,
        status=DocumentStatus.PARSED,
        page_count=1,
        text_length=sum(len(block.text) for block in blocks),
        pages=[],
        tables=[table],
        elements=elements,
    )


def test_bid_document_is_readable_and_contains_basic_information_table(tmp_path) -> None:
    output = tmp_path / "基础投标文件.docx"
    build_bid_document(make_project_facts(), output)

    document = Document(output)
    assert output.stat().st_size > 0
    assert document.paragraphs
    assert document.tables
    info_tables = [table for table in document.tables if len(table.rows) == len(load_output_fields()) + 1]
    assert len(info_tables) == 1
    info_text = "\n".join(cell.text for row in info_tables[0].rows for cell in row.cells)
    for field in load_output_fields():
        assert field.label in info_text


def test_bid_document_preserves_status_placeholders_and_separate_numbers(tmp_path) -> None:
    output = tmp_path / "bid.docx"
    build_bid_document(make_project_facts(), output)
    text = _document_text(Document(output))

    assert "项目编号" in text and "PRJ-123" in text
    assert "招标编号" in text and "BID-456" in text
    assert "PRJ-123" != "BID-456"
    assert "地点A" not in text
    assert "地点B" not in text
    assert "【待人工确认】" in text
    assert "【待补充】" in text
    assert "投标人：________________" in text
    assert "投标人：采购人甲" not in text


def test_bid_document_contains_only_the_basic_skeleton_sections(tmp_path) -> None:
    output = tmp_path / "bid.docx"
    build_bid_document(make_project_facts(), output)
    text = _document_text(Document(output))

    for heading in (
        "投 标 文 件",
        "项目基本信息",
        "一、投标函",
        "二、资格审查文件",
        "三、商务响应文件",
        "四、技术响应文件",
        "五、报价文件",
        "六、其他材料",
        "签字盖章复核提示",
    ):
        assert heading in text
    assert "【提示：本目录为基础工作骨架" in text
    assert "确保中标" not in text
    assert "完全满足全部要求" not in text


def test_bid_builder_requires_project_facts_ssot(tmp_path) -> None:
    with pytest.raises(TypeError):
        build_bid_document(tmp_path / "not-project-facts.json", tmp_path / "bid.docx")


def test_source_format_becomes_the_primary_bid_document_skeleton(tmp_path) -> None:
    source = tmp_path / "source.docx"
    source_document = Document()
    source_document.add_paragraph("招标文件")
    source_document.add_paragraph("第六章 投标文件格式", style="Heading 1")
    source_document.add_paragraph("一、投标函", style="Heading 2")
    source_document.add_paragraph("固定说明：本段来自招标文件格式。")
    source_document.add_paragraph("二、法定代表人身份证明", style="Heading 2")
    source_document.add_paragraph("三、授权委托书", style="Heading 2")
    source_document.add_paragraph("四、报价表", style="Heading 2")
    source_document.add_paragraph("第七章 评审方法", style="Heading 1")
    source_document.save(source)

    normalized = parse_document(source)
    template = extract_bid_format(normalized)
    output = tmp_path / "基础投标文件.docx"
    build_bid_document(
        make_project_facts(),
        output,
        normalized_document=normalized,
        format_template=template,
    )

    document = Document(output)
    headings = [paragraph.text for paragraph in document.paragraphs]
    expected = [
        "第六章 投标文件格式",
        "一、投标函",
        "二、法定代表人身份证明",
        "三、授权委托书",
        "四、报价表",
    ]
    positions = [headings.index(title) for title in expected]
    assert positions == sorted(positions)
    assert "二、资格审查文件" not in headings
    assert document.core_properties.subject == "format_source=SOURCE_DOCUMENT"

    assert not any(
        len(table.rows) == len(load_output_fields()) + 1
        for table in document.tables
    )
    assert "投 标 文 件" not in headings
    assert "项目基本信息" not in headings


def test_source_format_without_fact_values_does_not_override_project_facts(tmp_path) -> None:
    source = tmp_path / "source.docx"
    source_document = Document()
    source_document.add_paragraph("第六章 投标文件格式", style="Heading 1")
    source_document.add_paragraph("项目名称：格式章节中的诱导值")
    source_document.add_paragraph("一、投标函", style="Heading 2")
    source_document.save(source)

    normalized = parse_document(source)
    output = tmp_path / "基础投标文件.docx"
    build_bid_document(make_project_facts(), output, normalized_document=normalized)

    document = Document(output)
    assert not any(
        len(table.rows) == len(load_output_fields()) + 1
        for table in document.tables
    )
    assert "格式章节中的诱导值" in "\n".join(paragraph.text for paragraph in document.paragraphs)


def test_pdf_source_format_rebuilds_tables_and_fills_only_tender_placeholders(tmp_path) -> None:
    normalized = _synthetic_pdf_source_format()
    template = extract_bid_format(normalized)
    output = tmp_path / "pdf-source-format.docx"

    assert template.format_source == "SOURCE_DOCUMENT"
    assert template.metadata()["table_count"] == 1
    build_bid_document(
        make_project_facts(),
        output,
        normalized_document=normalized,
        format_template=template,
    )

    document = Document(output)
    assert len(document.tables) == 1
    rows = [[cell.text for cell in row.cells] for row in document.tables[0].rows]
    assert rows[0] == ["项目名称", "测试工程"]
    assert rows[1] == ["项目编号", "PRJ-123"]
    assert rows[2] == ["投标人", "________________"]
    assert document.paragraphs[0].text == "第六章 投标文件格式"
    visible_text = _document_text(document)
    assert "投 标 文 件" not in visible_text
    assert "项目基本信息" not in visible_text
    assert "PDF表格文本块" not in visible_text
