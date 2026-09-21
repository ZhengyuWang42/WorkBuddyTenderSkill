from __future__ import annotations

from pathlib import Path

import pymupdf
from openpyxl import load_workbook

from tender_basic.document_parser import parse_document
from tender_basic.dynamic_review import (
    DynamicReviewPlan,
    _guard_concepts,
    build_dynamic_review_plan,
    dynamic_review_qa,
    order_review_items,
)
from tender_basic.models import ProjectFacts
from tender_basic.review_builder import (
    SIGNATURE_LABELS,
    build_review_workbook,
    count_stray_review_rows,
    read_dynamic_review_rows,
)
from test_review_builder import make_project_facts

FIXTURE_TEXT = (
    "第三章 投标人须知\n"
    "3.1 投标人必须是在中华人民共和国境内注册的独立法人，具备有效的营业执照。\n"
    "3.2 投标人不得被列入失信被执行人名单，须提供信用中国查询截图。\n"
    "4.1 最高投标限价：人民币310万元（不含税），投标报价超过最高投标限价的，其投标无效。\n"
    "4.2 投标保证金金额：人民币20000元，须在投标截止时间前到账。\n"
    "5.1 投标有效期为投标截止之日起90日历天。\n"
    "6.1 供货期为合同签订后30日历天。\n"
    "7.1 投标人须按第六章响应文件格式编制响应文件，并加盖单位公章。\n"
    "8.1 评标办法采用综合评分法，技术分40分，价格分60分。\n"
)


def _make_pdf(tmp_path: Path, text: str = FIXTURE_TEXT) -> Path:
    path = tmp_path / "dynamic-review.pdf"
    pdf = pymupdf.open()
    page = pdf.new_page(width=620, height=800)
    page.insert_text((50, 70), text, fontsize=10, fontname="china-s")
    pdf.save(path)
    pdf.close()
    return path


def _plan(tmp_path: Path) -> DynamicReviewPlan:
    document = parse_document(_make_pdf(tmp_path))
    return build_dynamic_review_plan(document, make_project_facts())


def test_plan_items_are_generated_from_the_source(tmp_path: Path) -> None:
    plan = _plan(tmp_path)

    assert plan.items, "a tender document must generate review items"
    assert plan.source_requirement_count >= len(plan.items)
    for item in plan.items:
        assert item.item_id.startswith("DR")
        assert item.module
        assert item.risk_level in {"一票否决", "高", "中", "低"}
        assert item.source_requirement.strip()
        assert item.verification_action.strip()
        assert item.pass_criteria.strip()
        assert item.source_evidence.strip()
        assert item.source_requirement_ids
        assert item.source_page is None or item.source_page >= 1
        assert item.cell_text == item.cell_text.strip()

    text = " ".join(item.cell_text for item in plan.items)
    assert "营业执照" in text
    assert "最高投标限价" in text
    assert "供货期" in text


def test_plan_does_not_invent_project_irrelevant_items(tmp_path: Path) -> None:
    plan = _plan(tmp_path)
    text = " ".join(item.cell_text for item in plan.items)

    for concept in ("数据库", "中间件", "国产化", "培训"):
        assert concept not in text, f"{concept} is not in the source and must not become a review row"


def test_hard_gates_pass_on_the_fixture(tmp_path: Path) -> None:
    document = parse_document(_make_pdf(tmp_path))
    plan = build_dynamic_review_plan(document, make_project_facts())
    qa = dynamic_review_qa(plan, document, make_project_facts())

    for name, value in qa["hard_gate_values"].items():
        if name == "dynamic_review_item_count":
            assert value > 0
            continue
        assert value == 0, f"{name} must be 0, got {value}"
    assert qa["result"] == "PASS"


def test_order_review_items_is_module_then_risk(tmp_path: Path) -> None:
    plan = _plan(tmp_path)
    ordered = order_review_items(plan.items)
    modules = [item.module for item in ordered]
    module_order = []
    for module in modules:
        if not module_order or module_order[-1] != module:
            module_order.append(module)
    assert len(module_order) == len(set(module_order)), "modules must not interleave"
    risk_rank = {"一票否决": 0, "高": 1, "中": 2, "低": 3}
    for module in module_order:
        ranks = [risk_rank[item.risk_level] for item in ordered if item.module == module]
        assert ranks == sorted(ranks), "risk must be non-decreasing inside a module"


def test_concept_guard_falls_back_instead_of_naming_absent_concepts() -> None:
    assert _guard_concepts("核对数据库与中间件清单。", "响应文件应加盖公章。") == ""
    assert _guard_concepts("核对营业执照。", "具备有效的营业执照。") == "核对营业执照。"


def test_dynamic_workbook_replaces_the_fixed_row_model(tmp_path: Path) -> None:
    document = parse_document(_make_pdf(tmp_path))
    facts: ProjectFacts = make_project_facts()
    plan = build_dynamic_review_plan(document, facts)
    output = tmp_path / "review.xlsx"

    build_review_workbook(
        facts,
        output,
        normalized_document=document,
        dynamic_plan=plan,
    )

    rows = read_dynamic_review_rows(output)
    assert len(rows) == len(plan.items)
    assert count_stray_review_rows(output) == 0

    workbook = load_workbook(output, data_only=False)
    worksheet = workbook["投标项目复核表"]
    try:
        numbers = [worksheet.cell(row=8 + index, column=1).value for index in range(1, len(rows) + 1)]
        assert numbers == list(range(1, len(rows) + 1))
        assert worksheet.cell(row=9, column=4).value == order_review_items(plan.items)[0].cell_text
        assert all(worksheet.cell(row=8 + index, column=6).value == "待核对" for index in range(1, len(rows) + 1))
        assert all(
            str(worksheet.cell(row=8 + index, column=5).value or "").strip()
            for index in range(1, len(rows) + 1)
        )
        signature_row = 10 + len(rows)
        assert [
            worksheet.cell(row=signature_row + offset, column=1).value for offset in range(3)
        ] == list(SIGNATURE_LABELS)
        merges = {str(value) for value in worksheet.merged_cells.ranges}
        assert f"A{signature_row}:C{signature_row}" in merges
        assert all(
            "自动填充" not in str(cell.value)
            for row in worksheet.iter_rows()
            for cell in row
            if cell.value is not None
        )
    finally:
        workbook.close()


def test_dynamic_workbook_qa_detects_a_mismatched_row(tmp_path: Path) -> None:
    document = parse_document(_make_pdf(tmp_path))
    facts = make_project_facts()
    plan = build_dynamic_review_plan(document, facts)
    output = tmp_path / "review.xlsx"
    build_review_workbook(facts, output, normalized_document=document, dynamic_plan=plan)

    rows = read_dynamic_review_rows(output)
    tampered = [list(row) for row in rows]
    tampered[0][3] = "被改写的复核要点"
    qa = dynamic_review_qa(plan, document, facts, workbook_rows=tampered)

    assert qa["workbook_rows_match_dynamic_items"] is False
    assert 1 in qa["workbook_mismatch_rows"]


def test_legacy_fixed_row_path_is_still_available(tmp_path: Path) -> None:
    """Callers that do not pass a plan keep the historical 64-row behaviour."""

    output = tmp_path / "legacy.xlsx"
    build_review_workbook(make_project_facts(), output)
    workbook = load_workbook(output, data_only=False)
    worksheet = workbook["投标项目复核表"]
    try:
        numbers = [
            worksheet.cell(row=row, column=1).value
            for row in range(9, 74)
            if worksheet.cell(row=row, column=1).value is not None
        ]
        assert numbers == list(range(1, 65))
    finally:
        workbook.close()
