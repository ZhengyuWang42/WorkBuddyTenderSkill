from __future__ import annotations

from types import SimpleNamespace
from zipfile import ZipFile

from docx import Document

from tender_basic.numbering import NumberingOwnership, parse_source_number_token
from tender_basic.round51_qa import _style_numbering_names
from tender_basic.style_architecture import FixtureStylePack, TenderStyleProfile, classify_paragraph


def _source_item(text: str, *, kind: str = "Heading", center: bool = False, size: float = 10.5):
    source = SimpleNamespace(
        locator=None,
        alignment="center" if center else "left",
        typography_role="BODY",
        runs=[SimpleNamespace(text=text, bold=False, font_weight_role="regular")],
    )
    return SimpleNamespace(
        logical_text=text,
        kind=kind,
        source=source,
        font_size_pt=size,
        alignment_hint="center" if center else "left",
        first_line_indent_pt=0.0,
    )


def test_source_number_token_preserves_prefix_punctuation_and_spacing():
    token = parse_source_number_token("一、　标题")
    assert token is not None
    assert token.raw_prefix == "一、　"
    assert token.separator_after_prefix == "　"
    assert token.numbering_family == "CHINESE_CLAUSE"
    assert token.ownership is NumberingOwnership.SOURCE_LITERAL

    token = parse_source_number_token("1） 内容")
    assert token is not None
    assert token.raw_prefix == "1） "
    assert token.numbering_family == "DECIMAL_BODY"


def test_heading_classification_does_not_use_chinese_prefix_alone():
    ambiguous = classify_paragraph(
        _source_item("七、表"),
        SimpleNamespace(page=3),
        layout_classification="FLOW_TEXT_PAGE",
        toc_pages=set(),
        toc_entries=set(),
        neighboring_items=[],
    )
    assert ambiguous.role != "HEADING"
    assert ambiguous.classification_status == "HEADING_CLASSIFICATION_NEEDS_REVIEW"
    assert ambiguous.numbering_ownership is NumberingOwnership.SOURCE_LITERAL
    assert ambiguous.source_number_token is not None

    strong = classify_paragraph(
        _source_item("一、响应函", center=True, size=14.0),
        SimpleNamespace(page=3),
        layout_classification="FLOW_TEXT_PAGE",
        toc_pages=set(),
        toc_entries={"响应函"},
        neighboring_items=[],
    )
    assert strong.role == "HEADING"
    assert strong.numbering_ownership is NumberingOwnership.SOURCE_LITERAL


def test_source_styles_have_no_auto_numbering_but_generated_styles_do(tmp_path):
    document = Document()
    pack = FixtureStylePack(fixture_path=tmp_path / "missing-fixture.docx")
    pack.apply_profile(document, TenderStyleProfile())
    source_heading = document.add_paragraph("一、 响应函", style="Heading 1")
    source_list = document.add_paragraph("1、 第一项", style="Tender Body List 1")
    generated = document.add_paragraph("系统生成小节", style="Tender Generated Heading 1")
    output = tmp_path / "numbering.docx"
    document.save(output)

    assert source_heading.text == "一、 响应函"
    assert source_list.text == "1、 第一项"
    assert generated.text == "系统生成小节"
    numbered_styles = _style_numbering_names(output)
    assert "Heading 1" not in numbered_styles
    assert "Tender Body List 1" not in numbered_styles
    assert "Tender Generated Heading 1" in numbered_styles


def test_missing_source_body_item_cannot_renumber_later_source_item(tmp_path):
    document = Document()
    pack = FixtureStylePack(fixture_path=tmp_path / "missing-fixture.docx")
    pack.apply_profile(document, TenderStyleProfile())
    document.add_paragraph("1、 第一项", style="Tender Body List 1")
    document.add_paragraph("3、 第三项", style="Tender Body List 1")
    output = tmp_path / "gapped.docx"
    document.save(output)
    reopened = Document(output)
    assert [paragraph.text for paragraph in reopened.paragraphs] == ["1、 第一项", "3、 第三项"]
    assert all(paragraph.style.name == "Tender Body List 1" for paragraph in reopened.paragraphs)

    with ZipFile(output) as archive:
        document_xml = archive.read("word/document.xml")
    assert document_xml.count(b"numPr") == 0
