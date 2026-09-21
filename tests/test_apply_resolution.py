from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook
from pydantic import ValidationError

from scripts.apply_resolution import main as apply_resolution_main
from tender_basic.bid_document_builder import build_bid_document
from tender_basic.delivery_qa import run_delivery_qa
from tender_basic.models import (
    CandidateFact,
    DocxParagraphLocator,
    FactStatus,
    FieldName,
    ProjectFacts,
    ProjectFields,
    ResolutionAction,
    ResolutionOverride,
    ResolutionOverrides,
    ResolvedFact,
    SourceDocument,
    SourceType,
)
from tender_basic.resolution import apply_resolution_overrides
from tender_basic.review_builder import build_review_workbook


def _candidate(value: str, index: int) -> CandidateFact:
    return CandidateFact(
        value=value,
        normalized_value=value,
        confidence=0.82,
        method="label_value_same_line",
        source_file="review.docx",
        source_type=SourceType.DOCX,
        locator=DocxParagraphLocator(paragraph_index=index),
        section="投标人须知前附表",
        evidence_text=f"项目名称：{value}",
    )


def _resolved(field: FieldName, value: str | bool) -> ResolvedFact:
    candidate_value = value if isinstance(value, str) else ("是" if value else "否")
    candidate = _candidate(candidate_value, list(FieldName).index(field) + 10)
    return ResolvedFact(
        field=field,
        resolved_value=value,
        status=FactStatus.RESOLVED,
        confidence=candidate.confidence,
        candidates=[candidate],
        resolution_reason="Deterministic test fact.",
    )


def _review(field: FieldName, values: tuple[str, str] = ("项目A", "项目B")) -> ResolvedFact:
    candidates = [_candidate(value, index) for index, value in enumerate(values)]
    return ResolvedFact(
        field=field,
        resolved_value=None,
        status=FactStatus.NEEDS_REVIEW,
        confidence=0.82,
        candidates=candidates,
        resolution_reason="Two candidates require review.",
    )


def _not_found(field: FieldName) -> ResolvedFact:
    return ResolvedFact(
        field=field,
        resolved_value=None,
        status=FactStatus.NOT_FOUND,
        confidence=0.0,
        candidates=[],
        resolution_reason="No candidate.",
    )


def make_facts() -> ProjectFacts:
    typed_values: dict[str, str | bool] = {
        FieldName.BUDGET.value: "800000",
        FieldName.MAX_PRICE.value: "1000000",
        FieldName.DURATION.value: "180日历天",
        FieldName.BID_DEADLINE.value: "2026-09-10 09:30",
        FieldName.BID_OPEN_TIME.value: "2026-09-10 09:30",
        FieldName.BID_BOND_AMOUNT.value: "100000",
        FieldName.CONSORTIUM_ALLOWED.value: False,
        FieldName.BID_VALIDITY.value: "90日历天",
    }
    fields: dict[str, ResolvedFact] = {}
    for field in FieldName:
        fields[field.value] = _resolved(
            field,
            typed_values.get(field.value, f"{field.value}-value"),
        )
    fields[FieldName.PROJECT_NAME.value] = _review(FieldName.PROJECT_NAME)
    return ProjectFacts.from_fields(
        source_document=SourceDocument(
            source_file="review.docx",
            source_type=SourceType.DOCX,
        ),
        fields=ProjectFields(**fields),
    )


def _override(
    field: str = "project_name",
    action: ResolutionAction = ResolutionAction.SELECT_CANDIDATE,
    candidate_index: int | None = 1,
    reason: str = "前附表标签和值明确对应。",
) -> ResolutionOverrides:
    return ResolutionOverrides(
        resolutions=[
            ResolutionOverride(
                field=field,
                action=action,
                candidate_index=candidate_index,
                reason=reason,
            )
        ]
    )


def test_select_existing_candidate_and_keep_original_unchanged() -> None:
    original = make_facts()
    reviewed, audit = apply_resolution_overrides(original, _override())

    assert original.fields.project_name.status is FactStatus.NEEDS_REVIEW
    assert original.fields.project_name.resolved_value is None
    assert reviewed.fields.project_name.status is FactStatus.RESOLVED
    assert reviewed.fields.project_name.resolved_value == "项目B"
    assert reviewed.fields.project_name.candidates == original.fields.project_name.candidates
    assert audit[0]["selected_candidate"]["candidate_index"] == 1
    assert audit[0]["previous_status"] == "NEEDS_REVIEW"
    assert audit[0]["new_status"] == "RESOLVED"


def test_keep_unresolved_does_not_change_field() -> None:
    original = make_facts()
    reviewed, audit = apply_resolution_overrides(
        original,
        _override(
            action=ResolutionAction.KEEP_UNRESOLVED,
            candidate_index=None,
            reason="两处证据都合理，交由用户确认。",
        ),
    )

    assert reviewed.fields.project_name == original.fields.project_name
    assert audit[0]["selected_candidate"] is None
    assert audit[0]["new_status"] == "NEEDS_REVIEW"


def test_invalid_candidate_index_is_rejected() -> None:
    with pytest.raises(ValueError, match="out of range"):
        apply_resolution_overrides(
            make_facts(),
            _override(candidate_index=2),
        )


@pytest.mark.parametrize("status", [FactStatus.RESOLVED, FactStatus.NOT_FOUND])
def test_only_needs_review_can_be_modified(status: FactStatus) -> None:
    facts = make_facts()
    field = FieldName.PROJECT_NUMBER
    if status is FactStatus.NOT_FOUND:
        facts.fields.project_number = _not_found(field)

    overrides = ResolutionOverrides(
        resolutions=[
            ResolutionOverride(
                field=field,
                action=ResolutionAction.SELECT_CANDIDATE,
                candidate_index=0,
                reason="不应修改非复核字段。",
            )
        ]
    )
    with pytest.raises(ValueError, match="NEEDS_REVIEW"):
        apply_resolution_overrides(facts, overrides)


def test_free_text_value_and_unknown_action_are_rejected() -> None:
    with pytest.raises(ValidationError):
        ResolutionOverride.model_validate(
            {
                "field": "project_name",
                "action": "FREE_TEXT_VALUE",
                "value": "模型生成值",
                "reason": "不允许自由生成。",
            }
        )


def test_resolution_contract_rejects_missing_or_inconsistent_payload() -> None:
    with pytest.raises(ValidationError):
        ResolutionOverride(
            field=FieldName.PROJECT_NAME,
            action=ResolutionAction.SELECT_CANDIDATE,
            reason="缺少索引。",
        )
    with pytest.raises(ValidationError):
        ResolutionOverride(
            field=FieldName.PROJECT_NAME,
            action=ResolutionAction.KEEP_UNRESOLVED,
            candidate_index=0,
            reason="保持未解决。",
        )


def test_other_fields_and_candidate_evidence_are_not_changed() -> None:
    original = make_facts()
    original_dump = original.model_dump(mode="json")
    reviewed, _ = apply_resolution_overrides(original, _override())

    assert reviewed.fields.project_number == original.fields.project_number
    assert reviewed.fields.tender_number == original.fields.tender_number
    assert reviewed.fields.project_name.candidates == original.fields.project_name.candidates
    assert original.model_dump(mode="json") == original_dump


def test_cli_writes_reviewed_facts_and_audit_without_overwriting_original(tmp_path: Path) -> None:
    original_path = tmp_path / "project_facts.json"
    overrides_path = tmp_path / "resolution_overrides.json"
    reviewed_path = tmp_path / "project_facts.reviewed.json"
    audit_path = tmp_path / "resolution_audit.json"
    original_text = make_facts().model_dump_json(indent=2)
    original_path.write_text(original_text, encoding="utf-8")
    overrides_path.write_text(
        json.dumps(_override().model_dump(mode="json"), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    code = apply_resolution_main(
        [
            str(original_path),
            str(overrides_path),
            "--output",
            str(reviewed_path),
            "--audit",
            str(audit_path),
        ]
    )

    assert code == 0
    assert original_path.read_text(encoding="utf-8") == original_text
    reviewed = ProjectFacts.model_validate_json(reviewed_path.read_text(encoding="utf-8"))
    assert reviewed.fields.project_name.resolved_value == "项目B"
    audit = json.loads(audit_path.read_text(encoding="utf-8"))
    assert audit["resolutions"][0]["field"] == "project_name"


def test_reviewed_project_facts_schema_and_outputs_are_valid(tmp_path: Path) -> None:
    original = make_facts()
    reviewed, _ = apply_resolution_overrides(original, _override())
    reviewed = ProjectFacts.model_validate(reviewed.model_dump(mode="json"))
    facts_path = tmp_path / "project_facts.reviewed.json"
    xlsx_path = tmp_path / "review.xlsx"
    docx_path = tmp_path / "review.docx"
    facts_path.write_text(reviewed.model_dump_json(indent=2), encoding="utf-8")
    build_review_workbook(reviewed, xlsx_path)
    build_bid_document(reviewed, docx_path)

    report = run_delivery_qa(facts_path, xlsx_path, docx_path)
    assert report["overall_status"] == "PASS"

    workbook = load_workbook(xlsx_path, read_only=True)
    try:
        worksheet = workbook["投标项目复核表"]
        assert worksheet["B2"].value == "项目B"
        assert worksheet["B2"].value != "项目A"
    finally:
        workbook.close()

    document = Document(docx_path)
    info_table = next(table for table in document.tables if len(table.rows) == 24)
    info_rows = {row.cells[0].text: row.cells[1].text for row in info_table.rows[1:]}
    assert info_rows["项目名称"] == "项目B"
    assert info_rows["项目名称"] != "项目A"


def test_generated_resolution_schema_matches_pydantic_contract() -> None:
    root = Path(__file__).resolve().parents[1]
    schema_path = root / "schemas" / "resolution_overrides.schema.json"
    assert json.loads(schema_path.read_text(encoding="utf-8")) == ResolutionOverrides.model_json_schema()
