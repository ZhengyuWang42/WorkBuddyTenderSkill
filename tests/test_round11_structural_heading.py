"""Round-11: canonical EvidenceUnit structural-heading fidelity.

The exact-formatter gate proves only that a delivered locator equals the
formatter's output for the canonical ``EvidenceUnit``.  It cannot prove the unit
is the source unit that *owns* the requirement.  These tests close that gap from
the **source document** side:

    SOURCE_ATOM -> actual PDF structural container -> canonical EvidenceUnit
                -> formatter -> saved XLSX

plus the delivered-text defects found in the same human review (duplicated
extraction wording and a source identifier cut in half).
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

from tender_basic.delivery_text import (  # noqa: E402
    GENERAL_ROLE_LABEL,
    PROJECT_ROLE_LABEL,
    source_role_labels,
)
from tender_basic.document_models import NormalizedDocument  # noqa: E402
from tender_basic.dynamic_review import (  # noqa: E402
    _collapse_wrapped_token_repeat,
    build_dynamic_review_plan,
)
from tender_basic.evidence_unit import (  # noqa: E402
    _article_heading,
    _carries_forward,
    _is_heading,
    build_evidence_units,
)
from tender_basic.models import ProjectFacts  # noqa: E402
from tender_basic.review_workbook_views import _clip  # noqa: E402
from tender_basic.structural_heading import (  # noqa: E402
    check_structural_heading_fidelity,
    container_is_structural,
    structural_container_for,
)

CASES = {
    "case_001": "v1_manual_fidelity_round4_date_rhythm_closure9_review_workbook11",
    "case_002": "v1_round4_closure9_review_workbook11",
    "case_003": "v1_round4_closure9_review_workbook11",
}
CASE = ROOT / "acceptance/workspace/case_001"
BUILD11 = CASE / CASES["case_001"]

pytestmark = pytest.mark.skipif(
    not (BUILD11 / "project_facts.json").is_file(),
    reason="round-11 CASE001 successor build is not present",
)


def _load(case: str) -> tuple[NormalizedDocument, list]:
    build = ROOT / "acceptance/workspace" / case / CASES[case]
    document = NormalizedDocument.model_validate(
        json.loads((build / "normalized_document.json").read_text(encoding="utf-8"))
    )
    facts = ProjectFacts.model_validate(
        json.loads((build / "project_facts.json").read_text(encoding="utf-8"))
    )
    return document, list(build_dynamic_review_plan(document, facts).items)


def _item(items: list, concern_id: str):
    return next((item for item in items if item.concern_id == concern_id), None)


# --------------------------------------------------------------------------- #
# the heading classifier
# --------------------------------------------------------------------------- #


def test_contract_article_titles_are_structural_headings() -> None:
    for title in (
        "第二条 合同价款及结算",
        "第三条验收方法及技术要求",
        "第四条 设备的数量和计量单位、计量方法",
        "第五条 交（提）货地点、方式及费用",
        "第六条 对设备提出异议的时间和方法",
        "第七条 其他约定",
        "第八条 附则",
    ):
        assert _article_heading(title), title
        assert _is_heading(title), title
        assert _carries_forward(title), title


def test_body_prose_is_never_a_structural_heading() -> None:
    for text in (
        "合同解除或乙方应当退换的，乙方应当在接到甲方通知后7个工作日内自费拉回",
        "作为质保金，质保期12个月，质保期满后无息付清余款",
    ):
        assert not container_is_structural(text), text
        assert not _article_heading(text), text


def test_structural_container_walks_the_documents_own_blocks() -> None:
    """The container is the article whose clauses actually follow it."""

    document, _items = _load("case_001")
    units = {unit.unit_id: unit for unit in build_evidence_units(document)}
    for unit_id, expected in (
        # clause 4.1/4.2 are printed under 第四条
        ("EU-P34-B12", "第四条"),
        ("EU-P34-B14", "第四条"),
        # clause 5.1/5.2 are printed under 第五条, not the earlier 第四条
        ("EU-P34-B16", "第五条"),
        ("EU-P34-B18", "第五条"),
        # the termination prose sits under 第七条
        ("EU-P35-B15", "第七条"),
    ):
        unit = units[unit_id]
        container = structural_container_for(document, unit.page, unit.block_index)
        assert expected in container, (unit_id, container)


# --------------------------------------------------------------------------- #
# fixture 1: AUTHORIZATION
# --------------------------------------------------------------------------- #


def test_authorization_row_cites_the_clause_that_carries_it() -> None:
    """The authority chain is stated in 3.7.3; 3.3 only shares a few words."""

    _document, items = _load("case_001")
    item = _item(items, "AUTHORIZATION")
    assert item is not None
    locator = item.source_locator
    assert "3.7" in locator, locator
    assert "响应文件的编制" in locator, locator
    assert "澄清和补正" not in locator, locator


def test_authorization_row_keeps_the_minority_source_as_a_link() -> None:
    _document, items = _load("case_001")
    item = _item(items, "AUTHORIZATION")
    linked = [
        unit
        for unit in (item.evidence_units or ())
        if str(unit.get("role")) == "LINKED"
    ]
    assert linked, "the second source must still be reachable as a linked unit"


# --------------------------------------------------------------------------- #
# fixture 2: CONTRACT TERMINATION
# --------------------------------------------------------------------------- #


def test_contract_termination_uses_an_article_not_body_prose() -> None:
    _document, items = _load("case_001")
    item = _item(items, "CONTRACT_TERMINATION_REFUND")
    assert item is not None
    locator = item.source_locator
    assert "第七条" in locator, locator
    assert "合同解除或乙方应当退换" not in locator, locator


# --------------------------------------------------------------------------- #
# fixtures 3 / 4: duplicated wording and mid-token clipping
# --------------------------------------------------------------------------- #


def test_wrapped_token_repeat_is_collapsed() -> None:
    assert (
        _collapse_wrapped_token_repeat("并符合询比公告中供应商的资 格要求 格要求")
        == "并符合询比公告中供应商的资 格要求"
    )
    assert (
        _collapse_wrapped_token_repeat("具备有效的营业 执照 执照")
        == "具备有效的营业 执照"
    )


def test_label_then_value_layout_is_not_collapsed() -> None:
    """A source label followed by its own value is not an extraction artifact."""

    for text in (
        "*3.4.1 响应保证金 响应保证金的金额：人民币贰万元整",
        "履约保证金 履约保证金的金额：拾万元整",
        "7.3 履约保证金 履约保证金的金额：拾万元整。履约保证金的形式：银行转账或银行保函。",
    ):
        assert _collapse_wrapped_token_repeat(text) == text, text


def test_clip_never_ends_inside_a_standard_identifier() -> None:
    standards = (
        "3.1 供应商提供设备应遵照国家和部颁发的下列标准、规程和规范： "
        "（1）《二次供水工程技术规程》CJJ140-2018 （2）《室外给水设计规范》GB50013-2018 "
        "（3）《建筑给水排水设计规范》GB50015-2019 （4）《城镇供水管网加压泵站无负压供水设备》CJ/T 415-2013 "
        "（5）《泵站设计标准》GB50265-2022 （6）《室外给水设计规范》GB50013-2006 "
        "（7）《建筑给水排水及采暖工程施工质量验收规范》GB50242-2002 "
        "（8）《压缩机、风机、泵安装工程施工及验收规范》GB50275-2010 "
        "（9）《低压配电设计规范》GB 50054-95 （10）《低压配电装置及线路设计规范》GBJ 54-83"
    )
    clipped = _clip(standards, 300)
    assert clipped.endswith("…"), clipped
    head = clipped[:-1].rstrip()
    assert not re.search(r"[A-Za-z]{1,6}\s*/?\s*T?\s*\d[\d.\-/]*$", head[-24:]), head[-40:]
    assert not re.search(r"\d[\d.\-/]{3,}$", head[-24:]), head[-40:]


def test_saved_cells_carry_no_duplicated_or_cut_source_text() -> None:
    workbook = load_workbook(
        BUILD11 / "投标项目复核表.xlsx", data_only=True, read_only=True
    )
    try:
        delivered = workbook["投标项目复核表"]
        duplicated = []
        for row in delivered.iter_rows(min_row=9, values_only=True):
            for value in (row[3], row[4]):
                text = str(value or "")
                if re.search(r"([\u4e00-\u9fa5]{2,8})\1(?![\u4e00-\u9fa5])", text):
                    duplicated.append(text[:80])
        assert not duplicated, duplicated[:3]
    finally:
        workbook.close()


# --------------------------------------------------------------------------- #
# fixture 5: multi-source roles
# --------------------------------------------------------------------------- #


def test_multi_source_row_names_its_second_source() -> None:
    """DR037 combines a front-table value with a post-award consequence clause.

    The row's PRIMARY evidence is the front-table value (page 11); the generic
    consequence clause lives in a different clause of the same concern (page 21 /
    7.3 履约担保).  The action must *name* that second source rather than cite a
    bare page number the row's own evidence does not show.  Round 11 named the
    second source as an explicit clause; round 12 names both source roles, so both
    delivered shapes are accepted -- but never a bare page, and never the response
    bond (page 18 / 3.4) as this row's source.
    """

    _document, items = _load("case_001")
    item = _item(items, "PERFORMANCE_BOND")
    assert item is not None
    own_page = re.search(r"第(\d+)页", item.source_locator)
    own = own_page.group(1) if own_page else ""
    action = item.verification_action
    other_pages = [
        page for page in re.findall(r"依据第(\d+)页", action) if page != own
    ]

    roles = source_role_labels(action)
    project = str(roles.get(PROJECT_ROLE_LABEL) or "")
    general = str(roles.get(GENERAL_ROLE_LABEL) or "")
    if project and general:
        # round-12 delivered shape: each source is named by the role it plays
        assert "第11页" in project and "7.3" in project, (project, action)
        assert "第21页" in general and "7.3" in general, (general, action)
        assert "履约担保" in general, (general, action)
    else:
        # round-11 delivered shape: the second source is an explicit clause
        assert f"依据第{own}页" in action if own else True
        for page in other_pages:
            assert f"依据第{page}页核对另一来源条款" in action, (page, action)
    # the row's own locator is the front-table value it displays, and a different
    # concern's bond clause is never presented as this row's source
    assert "供应商须知前附表" in item.source_locator, item.source_locator
    assert "响应保证金" not in action, action


# --------------------------------------------------------------------------- #
# fixture 6: DR048 / DR052 source heading
# --------------------------------------------------------------------------- #


@pytest.mark.parametrize(
    "concern_id",
    ("DELIVERY_ACCEPTANCE_COMPLETION", "PRICE_INCLUDED_COST_SCOPE"),
)
def test_article_five_clauses_cite_article_five(concern_id: str) -> None:
    _document, items = _load("case_001")
    item = _item(items, concern_id)
    assert item is not None
    locator = item.source_locator
    assert "第五条" in locator, (concern_id, locator)
    assert "第六条" not in locator, (concern_id, locator)


# --------------------------------------------------------------------------- #
# the independent structural gate, over all three cases
# --------------------------------------------------------------------------- #


def test_no_canonical_unit_carries_body_prose_or_a_foreign_container() -> None:
    for case in CASES:
        document, items = _load(case)
        units = {unit.unit_id: unit for unit in build_evidence_units(document)}
        records = []
        for item in items:
            primary = next(
                (
                    unit
                    for unit in (getattr(item, "evidence_units", ()) or ())
                    if str(unit.get("role")) == "PRIMARY"
                ),
                None,
            )
            if primary is None:
                continue
            unit = units.get(str(primary.get("unit_id") or ""))
            if unit is None:
                continue
            records.append(
                {
                    "item_id": item.item_id,
                    "concern_id": item.concern_id,
                    "page": unit.page,
                    "block_index": unit.block_index,
                    "unit_heading": str(unit.heading or unit.semantic_heading or ""),
                    "unit_clause": str(unit.clause_number or ""),
                    "requirement": str(item.source_requirement or ""),
                    "locator": str(item.source_locator or ""),
                }
            )
        result = check_structural_heading_fidelity(document, records)
        assert result["body_prose_heading_count"] == 0, (case, result["heading_not_structural"])
        assert result["foreign_heading_count"] == 0, (case, result["heading_foreign"])
