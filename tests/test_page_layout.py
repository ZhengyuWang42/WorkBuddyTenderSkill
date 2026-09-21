"""Geometry layer regressions without real project answers or pixel identity."""
import pytest
from docx import Document
from tender_basic.document_models import PdfTextLine, PdfTextSpan
from tender_basic.page_layout import logical_paragraphs, build_page_layout, bounded_scale, form_parts, underline_for_form
from tender_basic.layout_qa import docx_layout_counts
from tender_basic.source_font_policy import font_name
from tender_basic.word_safe_source_builder import build_source_format_docx
from tender_basic.word_safe_scan import scan_word_safe_docx
from test_source_format_fidelity import _paragraph, _paragraph_template
from test_review_builder import make_project_facts


def geometry(lines, xs=None):
    xs=xs or [72]*len(lines)
    p=_paragraph('\n'.join(lines),bbox=(min(xs),80,520,80+24*len(lines)),font_size=12)
    p.lines=[PdfTextLine(text=text,bbox=(x,80+24*i,520,92+24*i),
        spans=[PdfTextSpan(text=text,bbox=(x,80+24*i,520,92+24*i),font_name='SimSun',font_size=12)])
        for i,(text,x) in enumerate(zip(lines,xs))]
    return p


def test_visual_wrap_one_logical_paragraph(tmp_path):
    p=geometry(['本单位决定参','加本次采购活动。'],[96,72])
    layout=logical_paragraphs(p)
    assert len(layout)==1 and layout[0].logical_text=='本单位决定参加本次采购活动。'
    assert layout[0].first_line_indent_pt==24 and layout[0].line_pitch_pt==24
    output,_=build_source_format_docx(make_project_facts(),_paragraph_template(p),tmp_path/'wrapped.docx')
    doc=Document(output)
    assert len(doc.paragraphs)==1
    assert doc.paragraphs[0].paragraph_format.first_line_indent.pt==24
    assert docx_layout_counts(output)['word_manual_breaks']==0
    assert docx_layout_counts(output)['layout_tab_count']==0
    assert scan_word_safe_docx(output)['result']=='PASS'


def test_numbered_boundaries_not_merged():
    assert len(logical_paragraphs(geometry(['1、第一项。','2、第二项。'])))==2


def test_same_baseline_fragment_is_not_tab():
    p=geometry(['我单位收到贵公司','      项目文件。'])
    p.lines[1].bbox=(180,80,520,92)
    result=logical_paragraphs(p)
    assert len(result)==1 and '\t' not in result[0].logical_text


def test_incomplete_geometry_does_not_guess():
    p=geometry(['第一行','第二行'])
    p.text+='缺失字符'
    assert logical_paragraphs(p)[0].flow is False


def test_cover_and_divider_distinct():
    divider=_paragraph_template(_paragraph('第八章 投标文件格式')).source_pages[0]
    cover=_paragraph_template(_paragraph('（项目名称）施工招标项目\n投标文件\n投标人：（电子印章）')).source_pages[0]
    assert build_page_layout(divider).classification=='CHAPTER_DIVIDER'
    assert build_page_layout(cover).classification=='COVER_PAGE'


def test_center_and_vertical_band(tmp_path):
    p=geometry(['响应文件'])
    p.alignment='center'
    p.bbox=(200,350,395,375)
    p.lines[0].bbox=p.bbox
    output,_=build_source_format_docx(make_project_facts(),_paragraph_template(p),tmp_path/'cover.docx')
    paragraph=Document(output).paragraphs[0]
    assert paragraph.paragraph_format.space_before.pt==0
    assert Document(output).sections[0].top_margin.pt==350
    assert paragraph.alignment==1
    assert paragraph.paragraph_format.left_indent.pt==0
    assert '\t' not in paragraph.text


@pytest.mark.parametrize('name,expected',[('SimSun','宋体'),('FangSong_GB2312','仿宋'),('SimHei','黑体'),('KaiTi','楷体'),('TimesNewRomanPSMT','Times New Roman')])
def test_font_mapping(name,expected):
    assert font_name(name)[0]==expected


@pytest.mark.parametrize('text,count',[('供应商：________（盖单位公章）',3),('____年____月____日',6)])
def test_form_row_is_one_paragraph_with_native_tabs_or_underlines(tmp_path,text,count):
    assert len(form_parts(text))==count
    p=geometry([text])
    output,report=build_source_format_docx(make_project_facts(),_paragraph_template(p),tmp_path/'form.docx')
    doc=Document(output)
    assert len(doc.tables)==0
    assert len(doc.paragraphs)==1
    compact=lambda value: ''.join(value.replace('_','').replace('\u00a0','').replace('\u2007','').replace('\t','').split())
    assert compact(doc.paragraphs[0].text)==compact(text)
    assert report['form_layout_table_count']==0
    assert report['synthetic_layout_tables']==0
    assert docx_layout_counts(output)['layout_tab_hack']==0
    assert docx_layout_counts(output)['layout_tab_count'] >= (1 if '（' in text else 0)
    assert scan_word_safe_docx(output)['result']=='PASS'


def test_fixed_content_not_form():
    assert form_parts('供应商：固定内容') is None


def test_calibration_bounded():
    assert bounded_scale(.94)==.94
    with pytest.raises(ValueError,match='PAGE_LAYOUT_NEEDS_REVIEW'):
        bounded_scale(.93)


def test_body_is_not_centered_merely_because_margins_symmetric():
    p=geometry(['正文说明'*20,'后续内容'*10])
    p.alignment='center'
    layout=build_page_layout(_paragraph_template(p).source_pages[0])
    assert layout.elements[0].alignment_hint=='left'


def test_vector_blank_requires_field_geometry():
    from tender_basic.source_format import SourceLine
    item=logical_paragraphs(geometry(['投标人：']))[0]
    line=SourceLine(bbox=(520,92,600,92))
    assert underline_for_form(item,[line]) is line
    assert underline_for_form(logical_paragraphs(geometry(['正文说明']))[0],[line]) is None


def test_blank_pdf_lines_keep_geometry_not_printed_rows():
    p=geometry([' ',' ','投标文件',' '])
    p.text='投标文件'
    items=logical_paragraphs(p)
    assert len(items)==1 and items[0].bbox[1]==128
    assert len(items[0].source_lines)==1


def test_visible_span_font_not_blank_span_font(tmp_path):
    from tender_basic.source_format import SourceRun
    from tender_basic.word_safe_source_builder import WordSafeSourceDocumentBuilder,new_word_safe_document
    template=_paragraph_template(_paragraph('投标文件'))
    builder=WordSafeSourceDocumentBuilder(make_project_facts(),template)
    p=new_word_safe_document().add_paragraph()
    runs=[SourceRun(text=' ',bbox=(0,0,10,10),font_size=10.5),
          SourceRun(text='投标文件 ',bbox=(0,30,100,50),font_size=24)]
    builder.text(p,'投标文件',runs,[])
    assert p.runs[0].font.size.pt==24


def test_inline_blank_needs_vector_evidence():
    from tender_basic.page_layout import restore_inline_rule_blanks
    from tender_basic.source_format import SourceRun,SourceLine
    runs=[SourceRun(text='工期',bbox=(0,0,20,12),font_size=12),SourceRun(text='日',bbox=(80,0,92,12),font_size=12)]
    assert restore_inline_rule_blanks('工期日',runs,[],[])[0]=='工期日'
    text,_,count=restore_inline_rule_blanks('工期日',runs,[SourceLine(bbox=(22,12,78,12))],[])
    assert text.startswith('工期\ue000BLANK:') and text.endswith('\ue001日') and count==1
